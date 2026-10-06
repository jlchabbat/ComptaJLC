"""Recette de la saisie à distance par recalcul LibreOffice.

    python tests/recette_lot1_distance.py ComptaJLC_lot1.xlsm [dossier_de_travail]

1. Classeur de saisie : une cotisation réglée doit produire 4 lignes, Mvt 1.
2. Classeur maître, onglet Transmission : des lignes reçues (Mvt 1 et 2)
   doivent recevoir les numéros définitifs 422 et 423 et être prêtes à
   reporter ; un Mvt déséquilibré ou une ligne déjà présente dans Écritures
   doivent être signalés.
"""

import datetime as dt
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import classeur_saisie  # noqa: E402
import lot1_distance  # noqa: E402
from recette_lot1 import recalculer  # noqa: E402


def d(a, m, j):
    return (dt.date(a, m, j) - dt.date(1899, 12, 30)).days


COTISATION = [
    [d(2026, 10, 15), "B3", 1, 1, "411TAIEB001", None, "COTISATION - TAIEB JEANNE", 400, 0, 0, None, "GEN.002", None],
    [d(2026, 10, 15), "B3", 1, 1, "700000", None, "COTISATION - TAIEB JEANNE", 0, 400, 0, None, "GEN.002", None],
    [d(2026, 10, 15), "B3", 1, 1, "512200", None, "COTISATION - TAIEB JEANNE", 400, 0, 0, None, "GEN.002", None],
    [d(2026, 10, 15), "B3", 1, 1, "411TAIEB001", None, "COTISATION - TAIEB JEANNE", 0, 400, 0, None, "GEN.002", None],
]
FRAIS = [
    [d(2026, 10, 16), "B1", 2, 2, "600100", None, "FRAIS BANCAIRES", 25, 0, 0, None, "GEN.004", None],
    [d(2026, 10, 16), "B1", 2, 2, "512000", None, "FRAIS BANCAIRES", 0, 25, 0, None, "GEN.004", None],
]
DESEQUILIBRE = [FRAIS[0], [*FRAIS[1][:8], 20, *FRAIS[1][9:]]]
# ligne identique à une écriture existante (Mvt 1 du grand livre)
DOUBLON = [[d(2026, 9, 24), "B1", 1, 1, "710000", None, "FACTURE 40013 - TAIEB JEANNE", 400, 0, 0, None, "MAN.001", None],
           [d(2026, 9, 24), "B1", 1, 1, "411TAIEB001", None, "FACTURE 40013 - TAIEB JEANNE", 0, 400, 0, None, "MAN.001", None]]


def main(lot1, dossier=None):
    dossier = Path(dossier or tempfile.mkdtemp())
    saisie = dossier / "saisie.xlsx"
    classeur_saisie.main(lot1, saisie, {"Saisie!C4": d(2026, 10, 15), "Saisie!C5": "Cotisation membre",
                                        "Saisie!C6": "TAIEB JEANNE (411TAIEB001)", "Saisie!C7": 400, "Saisie!C8": "BIT",
                                        "Saisie!C10": "GEN.002"})
    cas = {"recu_ok": COTISATION + FRAIS, "recu_desequilibre": DESEQUILIBRE, "recu_doublon": DOUBLON}
    fichiers = [saisie]
    for nom, lignes in cas.items():
        f = dossier / f"{nom}.xlsm"
        lot1_distance.main(lot1, f, recu=lignes)
        fichiers.append(f)
    sortie = recalculer(fichiers, dossier)
    echecs = 0

    wb = openpyxl.load_workbook(sortie / "saisie.xlsx", data_only=True)
    s = wb["Saisie"]
    zone = [[s.cell(r, c).value for c in range(2, 15)] for r in range(42, 46)]
    ok = [(z[1], z[2], z[3], str(z[4])) for z in zone] == [("B3", 1, 1, "411TAIEB001"), ("B3", 1, 1, "700000"),
                                                        ("B3", 1, 1, "512200"), ("B3", 1, 1, "411TAIEB001")]
    print(("OK   " if ok else "ÉCHEC"), "classeur de saisie :", s["B16"].value, "|", s["B40"].value)
    echecs += not ok

    attendu = {"recu_ok": ("PRÊT", [422] * 4 + [423] * 2, 6), "recu_desequilibre": ("À CORRIGER", None, 0),
               "recu_doublon": ("À CORRIGER", None, 0)}
    for nom, (debut, mvts, n_zone) in attendu.items():
        wb = openpyxl.load_workbook(sortie / f"{nom}.xlsx", data_only=True)
        t = wb["Transmission"]
        statut = t["A3"].value or ""
        mvt_def = [t.cell(r, 14).value for r in range(7, 7 + len(cas[nom]))]
        controles = [t.cell(r, 16).value for r in range(7, 7 + len(cas[nom]))]
        zone = [[t.cell(r, c).value for c in range(20, 33)] for r in range(7, 7 + 60)]
        zone = [z for z in zone if z[0] not in (None, "")]
        ok = statut.startswith(debut) and len(zone) == n_zone and (mvts is None or mvt_def == mvts)
        if nom == "recu_ok":
            ok = ok and [z[2] for z in zone] == mvts and [z[3] for z in zone] == [833] * 4 + [834] * 2
        print(("OK   " if ok else "ÉCHEC"), nom, "|", statut, "|", sorted({c for c in controles if c}))
        echecs += not ok
    print(f"{4 - echecs}/4 cas conformes")
    return 1 if echecs else 0


if __name__ == "__main__":
    if not shutil.which("soffice"):
        sys.exit("LibreOffice (soffice) introuvable")
    sys.exit(main(*sys.argv[1:]))
