import unittest

from comptajlc import create_app
from comptajlc.models import Axe, CodeAnalytique, Compte, db
from comptajlc.vues import en_centimes, fmt_montant


class Base(unittest.TestCase):
    def setUp(self):
        self.app = create_app({"TESTING": True, "SQLALCHEMY_DATABASE_URI": "sqlite://"})
        self.c = self.app.test_client()
        self.connecter(self.c)

    def jeton(self, client):
        with client.session_transaction() as s:
            s.setdefault("csrf", "jeton-test")
            return s["csrf"]

    def connecter(self, client):
        j = self.jeton(client)
        client.post("/premier-demarrage", data={"csrf_token": j, "nom": "tresorier",
                                                "mdp": "motdepasse123", "mdp2": "motdepasse123"})

    def post(self, url, **kw):
        kw.setdefault("data", {})["csrf_token"] = self.jeton(self.c)
        return self.c.post(url, **kw)

    def ecriture(self, **extra):
        d = {"date": "2026-01-15", "journal": "OD", "libelle": "Test", "piece": "",
             "compte0": "606", "debit0": "100,50", "compte1": "512", "credit1": "100,50"}
        d.update(extra)
        return self.post("/saisie", data=d)


class TestMontants(unittest.TestCase):
    def test_conversion(self):
        self.assertEqual(en_centimes("1 234,56"), 123456)
        self.assertEqual(fmt_montant(123456), "1 234,56")
        with self.assertRaises(ValueError):
            en_centimes("abc")


class TestSaisie(Base):
    def test_equilibree(self):
        r = self.ecriture()
        self.assertEqual(r.status_code, 302)
        b = self.c.get("/balance").get_data(as_text=True)
        self.assertIn("100,50", b)

    def test_desequilibree_refusee(self):
        r = self.ecriture(credit1="99")
        self.assertEqual(r.status_code, 400)
        self.assertIn("non équilibrée", r.get_data(as_text=True))

    def test_compte_inconnu(self):
        self.assertEqual(self.ecriture(compte1="999").status_code, 400)

    def test_analytique(self):
        with self.app.app_context():
            a = Axe.query.first()
            code = CodeAnalytique(axe_id=a.id, code="ACT.1", libelle="Activité")
            db.session.add(code)
            db.session.commit()
            aid, cid = a.id, code.id
        self.assertEqual(self.ecriture(**{f"axe{aid}_0": str(cid)}).status_code, 302)
        self.assertIn("ACT.1", self.c.get(f"/analytique/{aid}").get_data(as_text=True))


class TestPlan(Base):
    def test_ajout_import_suppression(self):
        self.post("/plan", data={"numero": "6061", "libelle": "Fournitures"})
        with self.app.app_context():
            self.assertIsNotNone(db.session.get(Compte, "6061"))
        import io
        self.post("/plan/importer", data={"fichier": (io.BytesIO("7061;Dons\n".encode()), "p.csv")},
                  content_type="multipart/form-data")
        with self.app.app_context():
            self.assertEqual(db.session.get(Compte, "7061").libelle, "Dons")
        self.ecriture(compte0="6061")
        self.post("/plan/6061/supprimer")  # utilisé : refusé
        with self.app.app_context():
            self.assertIsNotNone(db.session.get(Compte, "6061"))


class TestCorrection(Base):
    def historique(self):
        from comptajlc.models import Historique
        with self.app.app_context():
            return [(h.action, h.mvt, h.motif, h.utilisateur) for h in Historique.query.order_by(Historique.id)]

    def donnees(self, **extra):
        d = {"date": "2026-02-01", "journal": "OD", "libelle": "Corrigée", "piece": "P1",
             "compte0": "606", "debit0": "50", "compte1": "512", "credit1": "50", "motif": "erreur de montant"}
        d.update(extra)
        return d

    def test_creation_tracee_et_modification(self):
        self.ecriture()
        self.assertEqual(self.historique(), [("création", 1, "", "tresorier")])
        self.assertIn("Modifier l'écriture n° 1", self.c.get("/ecriture/1/modifier").get_data(as_text=True))
        self.assertEqual(self.post("/ecriture/1/modifier", data=self.donnees()).status_code, 302)
        self.assertIn("50,00", self.c.get("/balance").get_data(as_text=True))
        self.assertNotIn("100,50", self.c.get("/balance").get_data(as_text=True))
        self.assertEqual(self.historique()[1], ("modification", 1, "erreur de montant", "tresorier"))
        h = self.c.get("/historique").get_data(as_text=True)
        self.assertIn("100,50", h)   # l'ancienne valeur reste visible
        self.assertIn("erreur de montant", h)

    def test_modification_sans_motif_ou_desequilibree(self):
        self.ecriture()
        self.assertEqual(self.post("/ecriture/1/modifier", data=self.donnees(motif="")).status_code, 400)
        self.assertEqual(self.post("/ecriture/1/modifier", data=self.donnees(credit1="1")).status_code, 400)
        self.assertIn("100,50", self.c.get("/balance").get_data(as_text=True))
        self.assertEqual(len(self.historique()), 1)

    def test_suppression_et_numero_non_reutilise(self):
        self.ecriture()
        self.post("/ecriture/1/supprimer", data={"motif": ""})      # refusé
        self.assertEqual(len(self.historique()), 1)
        self.post("/ecriture/1/supprimer", data={"motif": "doublon"})
        self.assertEqual(self.historique()[-1], ("suppression", 1, "doublon", "tresorier"))
        self.assertNotIn("100,50", self.c.get("/balance").get_data(as_text=True))
        self.ecriture()
        self.assertEqual(self.historique()[-1][:2], ("création", 2))   # 1 n'est pas réutilisé

    def test_inconnue(self):
        self.assertEqual(self.c.get("/ecriture/99/modifier").status_code, 404)


class TestExport(Base):
    def charger(self, url):
        import io
        from openpyxl import load_workbook
        r = self.c.get(url)
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.mimetype)
        return load_workbook(io.BytesIO(r.data))

    def test_journal_et_balance(self):
        self.ecriture(libelle="=1+1")
        ws = self.charger("/export/journal.xlsx").active
        self.assertEqual(ws["H2"].value, 100.5)
        self.assertEqual(ws["I3"].value, 100.5)
        self.assertEqual(ws["E2"].value, "=1+1")
        self.assertEqual(ws["E2"].data_type, "s")      # texte, pas une formule
        wb = self.charger("/export/balance.xlsx").active
        self.assertEqual(wb["C3"].value, 100.5)   # 606 (512 est en C2)
        self.assertEqual(wb["C4"].value, "=SUM(C2:C3)")

    def test_export_protege(self):
        anonyme = self.app.test_client()
        self.assertEqual(anonyme.get("/export/journal.xlsx").status_code, 302)


class TestAuth(unittest.TestCase):
    def setUp(self):
        self.app = create_app({"TESTING": True, "SQLALCHEMY_DATABASE_URI": "sqlite://"})
        self.c = self.app.test_client()

    def jeton(self):
        with self.c.session_transaction() as s:
            s.setdefault("csrf", "jeton-test")
            return s["csrf"]

    def test_premier_demarrage_puis_connexion_obligatoire(self):
        self.assertIn("/premier-demarrage", self.c.get("/plan").headers["Location"])
        r = self.c.post("/premier-demarrage", data={"csrf_token": self.jeton(), "nom": "a",
                                                    "mdp": "court", "mdp2": "court"})
        self.assertIn("au moins 10", r.get_data(as_text=True))
        self.c.post("/premier-demarrage", data={"csrf_token": self.jeton(), "nom": "a",
                                                "mdp": "motdepasse123", "mdp2": "motdepasse123"})
        self.assertEqual(self.c.get("/plan").status_code, 200)
        self.c.post("/deconnexion", data={"csrf_token": self.jeton()})
        self.assertIn("/connexion", self.c.get("/plan").headers["Location"])
        self.assertEqual(self.c.get("/premier-demarrage").status_code, 302)

    def test_mauvais_mot_de_passe_et_bon(self):
        self.c.post("/premier-demarrage", data={"csrf_token": self.jeton(), "nom": "a",
                                                "mdp": "motdepasse123", "mdp2": "motdepasse123"})
        self.c.post("/deconnexion", data={"csrf_token": self.jeton()})
        r = self.c.post("/connexion", data={"csrf_token": self.jeton(), "nom": "a", "mdp": "faux"})
        self.assertIn("incorrect", r.get_data(as_text=True))
        self.assertIn("/connexion", self.c.get("/saisie").headers["Location"])
        r = self.c.post("/connexion", data={"csrf_token": self.jeton(), "nom": "a",
                                            "mdp": "motdepasse123", "suite": "/saisie"})
        self.assertTrue(r.headers["Location"].endswith("/saisie"))

    def test_redirection_externe_refusee(self):
        self.c.post("/premier-demarrage", data={"csrf_token": self.jeton(), "nom": "a",
                                                "mdp": "motdepasse123", "mdp2": "motdepasse123"})
        self.c.post("/deconnexion", data={"csrf_token": self.jeton()})
        r = self.c.post("/connexion", data={"csrf_token": self.jeton(), "nom": "a",
                                            "mdp": "motdepasse123", "suite": "//evil.example"})
        self.assertNotIn("evil", r.headers["Location"])

    def test_csrf_refuse(self):
        self.c.post("/premier-demarrage", data={"csrf_token": self.jeton(), "nom": "a",
                                                "mdp": "motdepasse123", "mdp2": "motdepasse123"})
        self.assertEqual(self.c.post("/plan", data={"numero": "999", "libelle": "x"}).status_code, 400)
        self.assertEqual(self.c.post("/plan", data={"csrf_token": "faux", "numero": "999",
                                                    "libelle": "x"}).status_code, 400)


if __name__ == "__main__":
    unittest.main()
