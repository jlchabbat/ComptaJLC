"""Tests de src/controles.py sur un classeur fictif construit à la volée.

    python -m unittest discover tests
"""

import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
import controles  # noqa: E402

COLS = ["Date", "Jnl", "Mvt", "Pièce", "Compte", "Libellé", "Débit", "Crédit", "Anal1", "Anal2", "Let"]


def classeur(ecritures, cloture=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def table(nom, entetes, lignes):
        ws = wb.create_sheet(nom)
        ws.append(entetes)
        for l in lignes:
            ws.append(l)
        fin = openpyxl.utils.get_column_letter(len(entetes))
        ws.add_table(Table(displayName=nom, ref=f"A1:{fin}{len(lignes) + 1}"))

    table("T_PlanComptable", ["Compte", "Libellé compte", "Axe 1 (Anal1)"],
          [["512000", "BANQUE", "BIL.5"], ["700000", "COTISATIONS", "COT.1"], ["411001", "MEMBRE", "BIL.4"]])
    table("T_Journaux", ["Code", "Intitulé"], [["B1", "BANQUE 1"], ["VT", "VENTES"], ["AN", "A-NOUVEAUX"]])
    table("T_Axe1", ["Code", "Libellé", "Actif"], [["BIL.5", "TRESORERIE", 1], ["COT.1", "COTIS", 1], ["BIL.4", "TIERS", 1]])
    table("T_Axe2", ["Code", "Libellé", "Actif"], [["GEN.002", "COTISATIONS", 1]])
    table("T_Ecritures", COLS, ecritures)
    if cloture:
        wb["T_Journaux"]["J1"] = cloture
        wb.defined_names["P_DateCloture"] = DefinedName("P_DateCloture", attr_text="T_Journaux!$J$1")
    chemin = Path(tempfile.mkdtemp()) / "test.xlsx"
    wb.save(chemin)
    return openpyxl.load_workbook(chemin, data_only=True)


D1 = dt.datetime(2026, 1, 15)
COTISATION = [
    [D1, "VT", 1, 10, "411001", "Cotisation", 100, 0, "BIL.4", "GEN.002", None],
    [D1, "VT", 1, 10, "700000", "Cotisation", 0, 100, "COT.1", "GEN.002", None],
    [D1, "B1", 2, 11, "512000", "Règlement", 100, 0, "BIL.5", "GEN.002", None],
    [D1, "B1", 2, 11, "411001", "Règlement", 0, 100, "BIL.4", "GEN.002", None],
]


def regles(ecritures, **kw):
    wb = classeur(ecritures, **kw)
    return [r for r, _ in controles.controler(wb, controles.tables(wb))]


class Controles(unittest.TestCase):
    def test_classeur_sain(self):
        self.assertEqual(regles(COTISATION), [])

    def test_rg01_desequilibre(self):
        e = [l[:] for l in COTISATION]
        e[1][7] = 90
        self.assertIn("RG-01", regles(e))

    def test_rg02_reference_inconnue(self):
        e = [l[:] for l in COTISATION]
        e[0][9] = "MAN.999"
        e[1][9] = "MAN.999"
        self.assertEqual(regles(e), ["RG-02", "RG-02"])

    def test_rg03_debit_et_credit(self):
        e = [l[:] for l in COTISATION]
        e[0][7] = 5
        self.assertIn("RG-03", regles(e))

    def test_rg04_periode_close(self):
        self.assertEqual(regles(COTISATION, cloture=dt.datetime(2025, 12, 31)), [])
        self.assertEqual(regles(COTISATION, cloture=dt.datetime(2026, 1, 31)), ["RG-04", "RG-04"])

    def test_piece_partagee(self):
        e = [l[:] for l in COTISATION]
        e[2][3] = e[3][3] = 10
        self.assertIn("NUM", regles(e))

    def test_photo_et_comparaison(self):
        wb = classeur(COTISATION)
        p = controles.photographier(controles.tables(wb))
        self.assertEqual((p["lignes"], p["mouvements"], p["total_debit"], p["resultat"]), (4, 2, 200.0, 100.0))
        self.assertEqual(controles.comparer(p, p), [])
        autre = dict(p, total_debit=201.0)
        self.assertEqual(len(controles.comparer(p, autre)), 1)


if __name__ == "__main__":
    unittest.main()
