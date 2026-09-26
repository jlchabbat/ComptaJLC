"""Fichier de liaison pour un bénévole à distance (activités et gestion).

    python src/classeur_liaison.py ComptaBB.xlsm ComptaBB_Liaison.xlsx

Le bénévole ne fait pas de comptabilité : il remplit des listes.
- **Activité** : le trésorier fixe d'avance le code Axe2 de l'activité (C3).
  Le bénévole y inscrit les recettes (participants avec le nombre de
  personnes, dons, sponsors) et les dépenses de l'activité.
- **Gestion** : dons ou aides reçus et versés, et autres opérations. Le
  bénévole laisse le code Axe2 vide ; le trésorier l'attribue ensuite, avant
  l'export.

Chaque ligne est contrôlée (colonne Contrôle). L'onglet **Export** génère les
écritures, selon les mêmes conventions que l'onglet Saisie du maître, et les
présente dans les colonnes de T_Ecritures. Le trésorier les colle dans
l'onglet Transmission du classeur maître, qui leur donne leur numéro
définitif.

Les correspondances (natures → comptes, modes de paiement → journaux) sont
dans l'onglet Listes, modifiable par le trésorier. Toutes les formules sont
simples (sans LET ni FILTER) : elles sont vérifiées par LibreOffice
(`tests/recette_liaison.py`).
"""

import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))
from classeur_saisie import ajouter_sst, lire_maitre, reprendre_parts  # noqa: E402
from lot0_preparation import AUJOURDHUI, Feuille, c_texte  # noqa: E402
from lot1_saisie import (  # noqa: E402
    DXF_ERREUR, DXF_OK, S_DATE, S_ENTETE, S_GRAS, S_MONTANT, S_NOTE, S_SAISIE, S_SAISIE_DATE, S_SAISIE_MONTANT, S_SECTION,
    S_TITRE, Classeur, c_f, lettre, remplir, validation,
)

COLONNES = ["Date", "Jnl", "Mvt", "Pièce", "Compte", "Intitulé", "Libellé", "Débit", "Crédit", "Solde", "Anal1", "Anal2", "Let"]

MODES_REC = [("Espèces", "CA", "530000"), ("Chèque", "B1", "512000"), ("Virement", "B1", "512000"), ("Bit", "B3", "512200"),
             ("Non payé", "", "")]
MODES_DEP = [("Espèces", "CA", "530000"), ("Chèque", "B1", "512000"), ("Virement", "B1", "512000"), ("Bit", "B3", "512200"),
             ("Carte Isracard", "OD", "580000"), ("Avance d'un membre", "OD", "")]
NAT_REC_ACT = [("Participation", "710000", "PARTICIPATION"), ("Don", "725000", "DON"), ("Sponsor / subvention", "740000", "SUBVENTION"),
               ("Autre recette", "720000", "RECETTE")]
NAT_DEP_ACT = [("Manifestation (traiteur, artiste, matériel…)", "610000", "DEPENSE"), ("Location de salle", "600200", "SALLE"),
               ("Frais divers", "600000", "FRAIS DIVERS")]
NAT_REC_GES = [("Don reçu", "725000", "DON RECU"), ("Aide reçue", "740000", "AIDE RECUE"), ("Autre recette", "720000", "RECETTE")]
NAT_DEP_GES = [("Don versé", "625000", "DON VERSE"), ("Aide versée", "630000", "AIDE VERSEE"), ("Autre dépense", "600000", "DEPENSE")]
# Gestion : le trésorier choisit ensuite le mode (virement ou espèces), le compte de contrepartie et le code Axe2
MODES_GES = [("Virement Mizrahi", "B1", "512000"), ("Virement BIT", "B3", "512200"), ("Espèces", "CA", "530000")]
# Schéma | Ligne | Rôle | Sens | Journal (PAIEMENT = journal du mode de paiement)
SCHEMAS = [("RM4", 1, "TIERS", "D", "PAIEMENT"), ("RM4", 2, "CONTRE", "C", "PAIEMENT"), ("RM4", 3, "TRESO", "D", "PAIEMENT"),
           ("RM4", 4, "TIERS", "C", "PAIEMENT"), ("RMN", 1, "TIERS", "D", "VT"), ("RMN", 2, "CONTRE", "C", "VT"),
           ("RS", 1, "TRESO", "D", "PAIEMENT"), ("RS", 2, "CONTRE", "C", "PAIEMENT"),
           ("DS", 1, "CONTRE", "D", "PAIEMENT"), ("DS", 2, "TRESO", "C", "PAIEMENT"),
           ("DA", 1, "CONTRE", "D", "PAIEMENT"), ("DA", 2, "TIERS", "C", "PAIEMENT")]

# Tables de saisie : nom, onglet, coin, nombre de lignes, sorte, base des Mvt locaux
TABLES = [
    ("T_RecAct", "Activité", 12, 150, "rec", 1000),
    ("T_DepAct", "Activité", 12, 60, "dep", 2000),
    ("T_RecGes", "Gestion", 6, 60, "rec", 3000),
    ("T_DepGes", "Gestion", 6, 60, "dep", 4000),
]
NB_EXPORT = 600
NB_NOUVEAUX = 30          # lignes vides prêtes pour les nouveaux tiers
PROVISOIRE = "Tiers provisoire : compte à attribuer par le trésorier (onglet Tiers) | "


def tr(t, col):
    return f"{t}[[#This Row],[{col}]]"


def controle_date(t):
    d = tr(t, "Date")
    return (f'IF(NOT(ISNUMBER({d})),"Date invalide | ",IF({d}<=P_DateCloture,"Période close | ",'
            f'IF(OR({d}<P_DebutExercice,{d}>P_FinExercice),"Hors exercice | ","")))')


def colonnes_recettes(t, activite):
    nat = "T_NatRecAct" if activite else "T_NatRecGes"
    remplie = f'OR({tr(t, "Date")}<>"",{tr(t, "Montant")}<>"",{tr(t, "Membre")}<>"",{tr(t, "Autre payeur")}<>"")'
    entrees = ["Date", "Membre", "Autre payeur"] + (["Nb pers."] if activite else []) + ["Nature", "Montant", "Mode", "Remarque"] \
        + ([] if activite else ["Code Axe2"])
    nature = f'IF({tr(t, "Nature")}="","{NAT_REC_ACT[0][0]}",{tr(t, "Nature")})' if activite else tr(t, "Nature")
    nom = f'IF({tr(t, "Membre")}<>"",TRIM(LEFT({tr(t, "Membre")},FIND("(",{tr(t, "Membre")}&"(")-1)),{tr(t, "Autre payeur")})'
    anal2 = "LI_Code" if activite else tr(t, "Code Axe2")
    calc = {
        "Remplie": remplie,
        "NatureEff": nature,
        "CptContre": f'IFERROR(INDEX({nat}[Compte],MATCH({tr(t, "NatureEff")},{nat}[Nature],0))&"","")',
        "CptTiers": f'IFERROR(INDEX(T_Membres[Compte],MATCH({tr(t, "Membre")},T_Membres[Nom],0))&"","")',
        "JnlPai": f'IFERROR(INDEX(T_ModesRec[Journal],MATCH({tr(t, "Mode")},T_ModesRec[Mode],0))&"","")',
        "CptPai": f'IFERROR(INDEX(T_ModesRec[Compte],MATCH({tr(t, "Mode")},T_ModesRec[Mode],0))&"","")',
        "Schéma": f'IF({tr(t, "Membre")}<>"",IF({tr(t, "JnlPai")}="","RMN","RM4"),"RS")',
        "NbLignes": f'COUNTIF(T_SchemasL[Schéma],{tr(t, "Schéma")})',
        "Anal2": f'{anal2}&""',
        "Libellé": (f'UPPER(IFERROR(INDEX({nat}[Libellé court],MATCH({tr(t, "NatureEff")},{nat}[Nature],0)),"")'
                    + (' &" "&LI_Nom' if activite else "") + f'&" - "&{nom})'),
        "Contrôle": (f'IF(NOT({tr(t, "Remplie")}),"",{controle_date(t)}'
                     f'&IF(N({tr(t, "Montant")})<=0,"Montant | ","")'
                     f'&IF(AND({tr(t, "Membre")}="",{tr(t, "Autre payeur")}=""),"Membre ou autre payeur | ","")'
                     f'&IF(AND({tr(t, "Membre")}<>"",{tr(t, "Autre payeur")}<>""),"Membre OU autre payeur, pas les deux | ","")'
                     f'&IF({tr(t, "Membre")}="","",IF(COUNTIF(T_Membres[Nom],{tr(t, "Membre")})=0,"Membre inconnu (l\'ajouter dans l\'onglet Tiers) | ",'
                     f'IF({tr(t, "CptTiers")}="","{PROVISOIRE}","")))'
                     f'&IF({tr(t, "CptContre")}="","Nature | ","")'
                     f'&IF(COUNTIF(T_ModesRec[Mode],{tr(t, "Mode")})=0,"Mode de paiement | ",'
                     f'IF(AND({tr(t, "JnlPai")}="",{tr(t, "Membre")}=""),"« Non payé » seulement pour un membre | ",""))'
                     + ('&IF(LI_Code="","Code activité à définir par le trésorier (C3) | ","")' if activite else
                        f'&IF({tr(t, "Code Axe2")}="","Code Axe2 à attribuer par le trésorier | ",'
                        f'IF(COUNTIF(T_Axe2[Code],{tr(t, "Code Axe2")})=0,"Code Axe2 inconnu | ",""))') + ")"),
        "OK": f'AND({tr(t, "Remplie")},{tr(t, "Contrôle")}="")',
    }
    return entrees, calc


def colonnes_depenses(t, activite):
    nat = "T_NatDepAct" if activite else "T_NatDepGes"
    remplie = f'OR({tr(t, "Date")}<>"",{tr(t, "Montant")}<>"",{tr(t, "Bénéficiaire")}<>"",{tr(t, "Nature")}<>"")'
    entrees = ["Date", "Bénéficiaire", "Nature", "Montant", "Payé par", "Membre (si avance)", "Justificatif", "Remarque"] \
        + ([] if activite else ["Code Axe2"])
    anal2 = "LI_Code" if activite else tr(t, "Code Axe2")
    calc = {
        "Remplie": remplie,
        "CptContre": f'IFERROR(INDEX({nat}[Compte],MATCH({tr(t, "Nature")},{nat}[Nature],0))&"","")',
        "CptTiers": f'IFERROR(INDEX(T_Membres[Compte],MATCH({tr(t, "Membre (si avance)")},T_Membres[Nom],0))&"","")',
        "JnlPai": f'IFERROR(INDEX(T_ModesDep[Journal],MATCH({tr(t, "Payé par")},T_ModesDep[Mode],0))&"","")',
        "CptPai": f'IFERROR(INDEX(T_ModesDep[Compte],MATCH({tr(t, "Payé par")},T_ModesDep[Mode],0))&"","")',
        "Schéma": f'IF(AND({tr(t, "JnlPai")}<>"",{tr(t, "CptPai")}=""),"DA","DS")',
        "NbLignes": f'COUNTIF(T_SchemasL[Schéma],{tr(t, "Schéma")})',
        "Anal2": f'{anal2}&""',
        "Libellé": (f'UPPER(IFERROR(INDEX({nat}[Libellé court],MATCH({tr(t, "Nature")},{nat}[Nature],0)),"")'
                    + (' &" "&LI_Nom' if activite else "")
                    + f'&" - "&{tr(t, "Bénéficiaire")}&IF({tr(t, "Justificatif")}<>""," - PIECE "&{tr(t, "Justificatif")},""))'),
        "Contrôle": (f'IF(NOT({tr(t, "Remplie")}),"",{controle_date(t)}'
                     f'&IF(N({tr(t, "Montant")})<=0,"Montant | ","")'
                     f'&IF({tr(t, "Bénéficiaire")}="","Bénéficiaire | ","")'
                     f'&IF({tr(t, "CptContre")}="","Nature | ","")'
                     f'&IF(COUNTIF(T_ModesDep[Mode],{tr(t, "Payé par")})=0,"Payé par | ","")'
                     f'&IF({tr(t, "Schéma")}<>"DA","",IF({tr(t, "Membre (si avance)")}="","Membre qui a avancé la dépense | ",'
                     f'IF(COUNTIF(T_Membres[Nom],{tr(t, "Membre (si avance)")})=0,"Membre inconnu (l\'ajouter dans l\'onglet Tiers) | ",'
                     f'IF({tr(t, "CptTiers")}="","{PROVISOIRE}",""))))'
                     + ('&IF(LI_Code="","Code activité à définir par le trésorier (C3) | ","")' if activite else
                        f'&IF({tr(t, "Code Axe2")}="","Code Axe2 à attribuer par le trésorier | ",'
                        f'IF(COUNTIF(T_Axe2[Code],{tr(t, "Code Axe2")})=0,"Code Axe2 inconnu | ",""))') + ")"),
        "OK": f'AND({tr(t, "Remplie")},{tr(t, "Contrôle")}="")',
    }
    return entrees, calc


def colonnes_gestion(t, recette):
    """Gestion : le bénévole décrit l'opération ; le trésorier complète mode, compte de contrepartie et code Axe2."""
    nat = "T_NatRecGes" if recette else "T_NatDepGes"
    qui = "Origine (qui donne)" if recette else "Bénéficiaire"
    entrees = ["Date", qui, "Nature", "Montant"] + ([] if recette else ["Justificatif"]) + ["Remarque",
               "Mode (trésorier)", "Compte (trésorier)", "Code Axe2 (trésorier)"]
    remplie = f'OR({tr(t, "Date")}<>"",{tr(t, "Montant")}<>"",{tr(t, qui)}<>"",{tr(t, "Nature")}<>"")'
    manque = "à attribuer par le trésorier | "
    calc = {
        "Compte suggéré": f'IFERROR(INDEX({nat}[Compte],MATCH({tr(t, "Nature")},{nat}[Nature],0))&"","")',
        "Remplie": remplie,
        "CptContre": f'{tr(t, "Compte (trésorier)")}&""',
        "CptTiers": '""',
        "JnlPai": f'IFERROR(INDEX(T_ModesGes[Journal],MATCH({tr(t, "Mode (trésorier)")},T_ModesGes[Mode],0))&"","")',
        "CptPai": f'IFERROR(INDEX(T_ModesGes[Compte],MATCH({tr(t, "Mode (trésorier)")},T_ModesGes[Mode],0))&"","")',
        "Schéma": '"RS"' if recette else '"DS"',
        "NbLignes": f'COUNTIF(T_SchemasL[Schéma],{tr(t, "Schéma")})',
        "Anal2": f'{tr(t, "Code Axe2 (trésorier)")}&""',
        "Libellé": (f'UPPER(IFERROR(INDEX({nat}[Libellé court],MATCH({tr(t, "Nature")},{nat}[Nature],0)),"")&" - "&{tr(t, qui)}'
                    + ("" if recette else f'&IF({tr(t, "Justificatif")}<>""," - PIECE "&{tr(t, "Justificatif")},"")') + ")"),
        "Contrôle": (f'IF(NOT({tr(t, "Remplie")}),"",{controle_date(t)}'
                     f'&IF(N({tr(t, "Montant")})<=0,"Montant | ","")'
                     f'&IF({tr(t, qui)}="","{qui} | ","")'
                     f'&IF({tr(t, "Nature")}="","Nature | ","")'
                     f'&IF({tr(t, "Mode (trésorier)")}="","Mode {manque}",IF({tr(t, "JnlPai")}="","Mode inconnu | ",""))'
                     f'&IF({tr(t, "Compte (trésorier)")}="","Compte {manque}",IF(COUNTIF(T_Comptes[Compte],{tr(t, "Compte (trésorier)")})=0,'
                     '"Compte inconnu | ",""))'
                     f'&IF({tr(t, "Code Axe2 (trésorier)")}="","Code Axe2 {manque}",IF(COUNTIF(T_Axe2[Code],{tr(t, "Code Axe2 (trésorier)")})=0,'
                     '"Code Axe2 inconnu | ","")))'),
        "OK": f'AND({tr(t, "Remplie")},{tr(t, "Contrôle")}="")',
    }
    return entrees, calc


def base(donnees, chemin):
    wb = openpyxl.Workbook()
    acc = wb.active
    acc.title = "Accueil"
    li = wb.create_sheet("Listes")
    lignes = [
        ("B1", "ComptaBB – fichier de liaison"),
        ("B3", "Pour un bénévole qui gère une activité ou des dons et aides, sans connaissance comptable."),
        ("B5", "Onglet Activité : le trésorier a inscrit le code de l'activité en C3. Noter chaque recette (participants, dons, "
               "sponsors) et chaque dépense, une ligne par paiement."),
        ("B6", "Onglet Gestion : dons ou aides reçus ou versés. Remplir date, qui, nature, montant ; les trois colonnes « trésorier » "
               "(mode de paiement, compte de contrepartie, code Axe2) sont complétées ensuite par le trésorier."),
        ("B7", "Remplir les cases jaunes ; la colonne Contrôle indique ce qui manque. Enregistrer (OneDrive transmet) et prévenir le trésorier."),
        ("B8", "Tiers : rechercher un membre en tapant le début de son nom dans la liste ; une personne absente s'ajoute dans l'onglet "
               "Tiers (provisoire, le trésorier lui attribue ensuite un compte)."),
        ("B10", "Trésorier : onglet Export, copier les lignes indiquées et les coller dans l'onglet Transmission du classeur ComptaBB."),
        ("B12", f"Listes (membres, codes, correspondances) extraites du classeur maître le {AUJOURDHUI:%d/%m/%Y}."),
    ]
    for c, t in lignes:
        acc[c] = t

    def table(nom, col0, ligne0, entetes, valeurs):
        for j, h in enumerate(entetes):
            li.cell(ligne0, col0 + j, h)
        for i, v in enumerate(valeurs):
            for j, x_ in enumerate(v):
                li.cell(ligne0 + 1 + i, col0 + j, x_)
        a, b = openpyxl.utils.get_column_letter(col0), openpyxl.utils.get_column_letter(col0 + len(entetes) - 1)
        li.add_table(Table(displayName=nom, ref=f"{a}{ligne0}:{b}{ligne0 + max(1, len(valeurs))}",
                           tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))

    li["A1"] = "Listes et correspondances (modifiables par le trésorier)"
    table("T_ModesRec", 1, 3, ["Mode", "Journal", "Compte"], MODES_REC)
    table("T_ModesDep", 1, 11, ["Mode", "Journal", "Compte"], MODES_DEP)
    table("T_NatRecAct", 5, 3, ["Nature", "Compte", "Libellé court"], NAT_REC_ACT)
    table("T_NatDepAct", 5, 11, ["Nature", "Compte", "Libellé court"], NAT_DEP_ACT)
    table("T_ModesGes", 1, 20, ["Mode", "Journal", "Compte"], MODES_GES)
    table("T_Comptes", 26, 3, ["Compte", "Libellé compte"], [(c, lib) for c, lib, _ in donnees["plan"]])
    table("T_NatRecGes", 9, 3, ["Nature", "Compte", "Libellé court"], NAT_REC_GES)
    table("T_NatDepGes", 9, 11, ["Nature", "Compte", "Libellé court"], NAT_DEP_GES)
    table("T_SchemasL", 13, 3, ["Schéma", "Ligne", "Rôle", "Sens", "Journal", "Clé"], [s + (f"{s[0]}|{s[1]}",) for s in SCHEMAS])
    membres = sorted((f"{nom} ({c})", c) for c, nom, _ in donnees["plan"] if c.startswith("411"))
    tiers = wb.create_sheet("Tiers", 1)
    tiers["A1"] = "Tiers : membres existants et nouveaux tiers (provisoires)"
    tiers["A2"] = ("Rechercher : filtre de la colonne Nom, ou taper le début du nom dans les listes. Nouveau tiers : l'ajouter sous la "
                   "table, compte vide ; il est provisoire jusqu'à ce que le trésorier lui attribue un compte (onglet Codes du classeur ComptaBB).")
    for j, h in enumerate(["Nom", "Compte", "Type", "Téléphone", "E-mail", "Remarque"]):
        tiers.cell(4, 1 + j, h)
    for i, (nom_, compte) in enumerate(membres):
        tiers.cell(5 + i, 1, nom_)
        tiers.cell(5 + i, 2, compte)
        tiers.cell(5 + i, 3, "Membre")
    tiers.add_table(Table(displayName="T_Membres", ref=f"A4:F{4 + len(membres) + NB_NOUVEAUX}",
                          tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))
    for c, w in zip("ABCDEF", (36, 14, 12, 16, 28, 40)):
        tiers.column_dimensions[c].width = w
    table("T_Axe2", 23, 3, ["Code", "Libellé"], [(a, b) for a, b, _ in donnees["axe2"]])
    li["A26"] = "Paramètres copiés du classeur maître"
    for r, nom in enumerate(("P_DebutExercice", "P_FinExercice", "P_DateCloture"), start=27):
        li.cell(r, 1, nom)
        li.cell(r, 2, donnees["params"][nom])
        wb.defined_names[nom] = DefinedName(nom, attr_text=f"Listes!$B${r}")
    wb.save(chemin)


def construire(maitre, sortie, entrees=None):
    donnees = lire_maitre(maitre)
    with tempfile.TemporaryDirectory() as d:
        b = Path(d) / "base.xlsx"
        base(donnees, b)
        cl = Classeur(ajouter_sst(b))
    reprendre_parts(cl.p, maitre)
    ch = cl.ch

    # ------------------------------------------------ onglets Activité et Gestion
    defs = {}
    for nom, onglet, ligne0, n, sorte, base_mvt in TABLES:
        activite = onglet == "Activité"
        if activite:
            entrees_, calc = (colonnes_recettes if sorte == "rec" else colonnes_depenses)(nom, True)
        else:
            entrees_, calc = colonnes_gestion(nom, sorte == "rec")
        defs[nom] = (onglet, ligne0, n, sorte, base_mvt, entrees_, calc)

    parts = {}
    for onglet, apres in (("Activité", "Accueil"), ("Gestion", "Activité")):
        tabs = [(nom, d_) for nom, d_ in defs.items() if d_[0] == onglet]
        # recettes à gauche, dépenses à droite
        col0 = {tabs[0][0]: 1, tabs[1][0]: len(tabs[0][1][5]) + len(tabs[0][1][6]) + 2}
        dv, cf, largeurs = "", "", []
        prio = 1
        for nom, (_, ligne0, n, sorte, _, entrees_, calc) in tabs:
            c0 = col0[nom]
            r1, r2 = ligne0 + 1, ligne0 + n
            listes = {"Membre": "T_Membres", "Membre (si avance)": "T_Membres", "Code Axe2 (trésorier)": "T_Axe2",
                      "Mode": "T_ModesRec", "Payé par": "T_ModesDep", "Mode (trésorier)": "T_ModesGes", "Compte (trésorier)": "T_Comptes",
                      "Nature": {"T_RecAct": "T_NatRecAct", "T_DepAct": "T_NatDepAct", "T_RecGes": "T_NatRecGes",
                                 "T_DepGes": "T_NatDepGes"}[nom]}
            for j, h in enumerate(entrees_):
                c = lettre(c0 + j)
                if h in listes:
                    colonne = {"T_Membres": "Nom", "T_Axe2": "Code", "T_Comptes": "Compte"}.get(
                        listes[h], "Mode" if "Modes" in listes[h] else "Nature")
                    cl.nommer(f"LL_{nom[2:]}_{j}", f"{listes[h]}[{colonne}]")
                    dv += validation(f"{c}{r1}:{c}{r2}", f"LL_{nom[2:]}_{j}", "Choisir dans la liste.", h, "Choisir dans la liste.")
                elif h == "Date":
                    dv += validation(f"{c}{r1}:{c}{r2}", "P_DebutExercice", "Date du paiement (jj/mm/aaaa).", "Date",
                                     "Date hors de l'exercice ouvert.", type_="date", operateur="between", formule2="P_FinExercice")
                elif h == "Montant":
                    dv += validation(f"{c}{r1}:{c}{r2}", "0", "Montant en shekels, positif.", "Montant", "Saisir un nombre positif.",
                                     type_="decimal", operateur="greaterThan")
                largeurs.append((c0 + j, c0 + j, {"Date": 11, "Montant": 11, "Nb pers.": 8, "Mode": 12, "Payé par": 16,
                                                   "Code Axe2 (trésorier)": 12, "Compte (trésorier)": 12,
                                                   "Mode (trésorier)": 16, "Justificatif": 11}.get(h, 26), 0))
            k_ctrl = c0 + len(entrees_) + list(calc).index("Contrôle")
            cf += (f'<conditionalFormatting sqref="{lettre(k_ctrl)}{r1}:{lettre(k_ctrl)}{r2}"><cfRule type="expression" '
                   f'dxfId="{DXF_ERREUR}" priority="{prio}"><formula>{lettre(k_ctrl)}{r1}&lt;&gt;""</formula></cfRule></conditionalFormatting>')
            prio += 1
            for k, h in enumerate(calc):
                visible = h in ("Contrôle", "Compte suggéré")
                largeurs.append((c0 + len(entrees_) + k, c0 + len(entrees_) + k, 40 if h == "Contrôle" else 12, 0 if visible else 1))
        if onglet == "Activité":
            dv += validation("C3", "LL_Codes", "Code Axe2 de l'activité, choisi par le trésorier.", "Code activité", "Choisir un code.")
            cl.nommer("LL_Codes", "T_Axe2[Code]")
        largeurs = sorted({a: (a, b, w, h) for a, b, w, h in largeurs}.values())
        part = cl.nouvel_onglet(onglet, apres, largeurs, dv=dv, cf=cf, n_tables=2,
                                vue=f'<pane ySplit="{tabs[0][1][1]}" topLeftCell="A{tabs[0][1][1] + 1}" activePane="bottomLeft" state="frozen"/>')
        parts[onglet] = part
        for k, (nom, (_, ligne0, n, sorte, _, entrees_, calc)) in enumerate(tabs):
            styles = {h: S_SAISIE for h in entrees_}
            styles.update({"Date": S_SAISIE_DATE, "Montant": S_SAISIE_MONTANT})
            cl.nouvelle_table(part, f"rId{k + 1}", nom, f"{lettre(col0[nom])}{ligne0}", entrees_ + list(calc),
                              [[None] * (len(entrees_) + len(calc)) for _ in range(n)], calculees=calc, styles=styles)
        fe = Feuille(cl.p, part)
        rec, dep = tabs[0][0], tabs[1][0]
        titre = "Activité" if onglet == "Activité" else "Gestion : dons, aides et autres opérations"
        fe.poser("A1", c_texte("A1", ch, titre, S_TITRE))
        fe.poser(f"A{tabs[0][1][1] - 1}", c_texte(f"A{tabs[0][1][1] - 1}", ch, "Recettes", S_SECTION))
        fe.poser(f"{lettre(col0[dep])}{tabs[0][1][1] - 1}", c_texte(f"{lettre(col0[dep])}{tabs[0][1][1] - 1}", ch, "Dépenses", S_SECTION))
        if onglet == "Activité":
            for ref, t, s in (("B3", "Code de l'activité (trésorier)", None), ("B4", "Responsable", None), ("B5", "Dates", None),
                              ("B7", "Participants (nombre de personnes)", None), ("B8", "Recettes", None), ("B9", "Dépenses", None),
                              ("B10", "Résultat", S_GRAS)):
                fe.poser(ref, c_texte(ref, ch, t, s))
            for ref in ("C3", "C4", "C5"):
                fe.poser(ref, f'<c r="{ref}" s="{S_SAISIE}"/>')
            fe.poser("D3", c_f("D3", 'IF(LI_Code="","",IFERROR(INDEX(T_Axe2[Libellé],MATCH(LI_Code,T_Axe2[Code],0)),"code inconnu"))', S_GRAS))
            fe.poser("C7", c_f("C7", f'SUMIFS({rec}[Nb pers.],{rec}[OK],TRUE,{rec}[NatureEff],"{NAT_REC_ACT[0][0]}")', texte=False))
            fe.poser("C8", c_f("C8", f'SUMIFS({rec}[Montant],{rec}[OK],TRUE)', S_MONTANT, texte=False))
            fe.poser("C9", c_f("C9", f'SUMIFS({dep}[Montant],{dep}[OK],TRUE)', S_MONTANT, texte=False))
            fe.poser("C10", c_f("C10", "C8-C9", S_MONTANT, texte=False))
            fe.poser("E7", c_texte("E7", ch, "Totaux des lignes sans erreur. Colonne Contrôle en rouge = ligne à compléter.", S_NOTE))
            cl.nommer("LI_Code", "Activité!$C$3")
            cl.nommer("LI_Nom", "Activité!$D$3")
        else:
            fe.poser("A3", c_texte("A3", ch, "Bénévole : date, qui, nature, montant, remarque. Trésorier : mode de paiement, compte de "
                                             "contrepartie (voir « Compte suggéré ») et code Axe2, avant l'export.", S_NOTE))
        fe.enregistrer()
        cl.dimension(part)

    # ------------------------------------------------ onglet Export
    brut = []   # (table, i, k, base)
    for nom, (_, _, n, _, base_mvt, _, _) in defs.items():
        for i in range(1, n + 1):
            for k in range(1, 5):
                brut.append((nom, i, k, base_mvt))
    n_brut = len(brut)
    cf = (f'<conditionalFormatting sqref="A3"><cfRule type="expression" dxfId="{DXF_OK}" priority="1"><formula>LI_OK</formula></cfRule>'
          f'<cfRule type="expression" dxfId="{DXF_ERREUR}" priority="2"><formula>LI_Err&gt;0</formula></cfRule></conditionalFormatting>')
    largeurs = [(1, 1, 11, 0), (2, 2, 5, 0), (3, 4, 7, 0), (5, 5, 13, 0), (6, 6, 5, 0), (7, 7, 44, 0), (8, 10, 10, 0),
                (11, 11, 5, 0), (12, 12, 9, 0), (13, 13, 4, 0), (14, 14, 3, 0), (15, 15, 8, 1), (16, 30, 10, 1)]
    part = cl.nouvel_onglet("Export", "Gestion", largeurs, cf=cf,
                            vue='<pane ySplit="5" topLeftCell="A6" activePane="bottomLeft" state="frozen"/>')
    fe = Feuille(cl.p, part)
    fe.poser("A1", c_texte("A1", ch, "Export vers le classeur ComptaBB (trésorier)", S_TITRE))
    erreurs = "+".join(f'COUNTIF({t}[Contrôle],"?*")' for t in defs)
    nb_ops = "+".join(f"COUNTIF({t}[OK],TRUE)" for t in defs)
    for r, nom, f_ in ((1, "LI_Err", erreurs), (2, "LI_NbOps", nb_ops), (3, "LI_NbLignes", f"MAX(Q7:Q{6 + n_brut})"),
                       (4, "LI_OK", "AND(LI_Err=0,LI_NbOps>0)")):
        fe.poser(f"P{r}", c_f(f"P{r}", f_, texte=False))
        cl.nommer(nom, f"Export!$P${r}")
    fe.poser("A3", c_f("A3", 'IF(LI_NbOps+LI_Err=0,"Aucune opération saisie.",IF(LI_OK,"PRÊT : "&LI_NbOps&" opérations, "&LI_NbLignes'
                             '&" lignes d\'écritures.","À COMPLÉTER : "&LI_Err&" ligne(s) signalée(s) dans les onglets Activité et Gestion."))', S_GRAS))
    fe.poser("A4", c_f("A4", 'IF(NOT(LI_OK),"","Sélectionner A6:M"&(5+LI_NbLignes)&", Copier ; dans ComptaBB, onglet Transmission, clic droit '
                             'sur A7 › Collage spécial › Valeurs.")', S_NOTE))
    for j, h in enumerate(COLONNES):
        fe.poser(f"{lettre(1 + j)}5", c_texte(f"{lettre(1 + j)}5", ch, h, S_ENTETE))
    # grille brute (colonnes P:AD masquées) : une ligne par ligne d'écriture possible
    for idx, (t, i, k, b) in enumerate(brut):
        r = 7 + idx
        ok = f"AND(INDEX({t}[OK],{i}),{k}<=INDEX({t}[NbLignes],{i}))"
        cellules = {
            "P": (ok, False),
            "Q": (f"N(Q{r - 1})+IF(P{r},1,0)" if idx else "IF(P7,1,0)", False),
            "R": (f'IF(P{r},Q{r},"")', False),
            "S": (f'IF(P{r},MATCH(INDEX({t}[Schéma],{i})&"|{k}",T_SchemasL[Clé],0),0)', False),
            "T": (f'IF(S{r}=0,"",INDEX(T_SchemasL[Rôle],S{r}))', True),
            "U": (f'IF(S{r}=0,"",INDEX(T_SchemasL[Sens],S{r}))', True),
            "V": (f'IF(S{r}=0,"",IF(INDEX(T_SchemasL[Journal],S{r})="PAIEMENT",INDEX({t}[JnlPai],{i}),INDEX(T_SchemasL[Journal],S{r})))', True),
            "W": (f'IF(S{r}=0,"",IF(T{r}="TIERS",INDEX({t}[CptTiers],{i}),IF(T{r}="CONTRE",INDEX({t}[CptContre],{i}),INDEX({t}[CptPai],{i}))))', True),
            "X": (f'IF(S{r}=0,"",INDEX({t}[Date],{i}))', False),
            "Y": (f'IF(S{r}=0,"",{b}+{i})', False),
            "Z": (f'IF(S{r}=0,"",INDEX({t}[Libellé],{i}))', True),
            "AA": (f'IF(S{r}=0,"",IF(U{r}="D",ROUND(INDEX({t}[Montant],{i}),2),0))', False),
            "AB": (f'IF(S{r}=0,"",IF(U{r}="C",ROUND(INDEX({t}[Montant],{i}),2),0))', False),
            "AC": (f'IF(S{r}=0,"",INDEX({t}[Anal2],{i}))', True),
        }
        for c, (f_, texte) in cellules.items():
            fe.poser(f"{c}{r}", c_f(f"{c}{r}", f_, texte=texte))
    # grille compacte visible : colonnes de T_Ecritures ; Intitulé, Anal1, Let vides
    for n in range(1, NB_EXPORT + 1):
        r = 5 + n
        fe.poser(f"O{r}", c_f(f"O{r}", f'IF(NOT(LI_OK),0,IFERROR(MATCH({n},$R$7:$R${6 + n_brut},0)+6,0))', texte=False))
        g = f"$O{r}=0"
        for c, src, s, texte in (("A", "X", S_DATE, False), ("B", "V", None, True), ("C", "Y", None, False), ("D", "Y", None, False),
                                 ("E", "W", None, True), ("G", "Z", None, True), ("H", "AA", S_MONTANT, False),
                                 ("I", "AB", S_MONTANT, False), ("L", "AC", None, True)):
            fe.poser(f"{c}{r}", c_f(f"{c}{r}", f'IF({g},"",INDEX(${src}:${src},$O{r}))', s, texte=texte))
        fe.poser(f"J{r}", c_f(f"J{r}", f'IF({g},"",0)', S_MONTANT, texte=False))
    fe.enregistrer()
    cl.dimension(part)
    cl.enregistrer_noms()
    if entrees:
        remplir(cl, entrees)
    cl.ch.enregistrer()
    cl.p.enregistrer(sortie)
    print(f"{sortie} : fichier de liaison ({sum(d_[2] for d_ in defs.values())} lignes de saisie, {n_brut} lignes d'écritures possibles)")


if __name__ == "__main__":
    construire(sys.argv[1], sys.argv[2])

