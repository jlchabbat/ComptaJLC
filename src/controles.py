"""Contrôles et photographie des totaux du classeur ComptaJLC.

Lecture seule : le classeur n'est jamais enregistré (openpyxl perdrait
Power Query, les commentaires à thread et une partie des validations).
Les valeurs lues sont celles calculées par Excel au dernier enregistrement.

    python src/controles.py                      contrôles RG-01 à RG-04
    python src/controles.py --photo avant.json   + enregistre les totaux
    python src/controles.py --compare avant.json + compare aux totaux enregistrés

Le classeur est celui désigné par config.json (chemin relatif à la racine
du projet), sauf s'il est donné en argument.
"""

import argparse
import datetime as dt
import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

RACINE = Path(__file__).resolve().parent.parent
ARRONDI = 0.005
ERREURS = {"#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def lire_config():
    chemin = RACINE / "config.json"
    return json.loads(chemin.read_text(encoding="utf-8")) if chemin.exists() else {}


def tables(wb):
    """{nom de table: liste de dicts}, en-têtes pris dans la table."""
    res = {}
    for ws in wb.worksheets:
        for t in ws.tables.values():
            lignes = list(ws[t.ref])
            entetes = [c.value for c in lignes[0]]
            res[t.name] = [dict(zip(entetes, (c.value for c in l))) for l in lignes[1:]]
    return res


def nom_defini(wb, nom):
    d = wb.defined_names.get(nom)
    if d is None:
        return None
    for feuille, plage in d.destinations:
        return wb[feuille][plage.replace("$", "")].value
    return None


def montant(v):
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).replace(" ", "").replace(" ", "").replace(",", "."))


def texte(v):
    return "" if v is None else str(v).strip()


def controler(wb, t):
    ecr = t["T_Ecritures"]
    anomalies = []

    # RG-01 : chaque Mvt équilibré
    ecart = defaultdict(float)
    for l in ecr:
        ecart[l["Mvt"]] += montant(l["Débit"]) - montant(l["Crédit"])
    for mvt, e in sorted(ecart.items(), key=lambda x: (x[0] is None, x[0])):
        if abs(e) > ARRONDI:
            anomalies.append(("RG-01", f"Mvt {mvt} déséquilibré de {e:.2f}"))

    # RG-02 : références présentes dans les tables
    refs = {
        "Compte": ("T_PlanComptable", "Compte"),
        "Jnl": ("T_Journaux", "Code"),
        "Anal1": ("T_Axe1", "Code"),
        "Anal2": ("T_Axe2", "Code"),
    }
    for col, (table, cle) in refs.items():
        connus = {texte(r[cle]) for r in t.get(table, [])}
        for i, l in enumerate(ecr, start=2):
            v = texte(l.get(col))
            if v and v not in connus:
                anomalies.append(("RG-02", f"ligne {i} : {col} « {v} » absent de {table}"))

    # RG-03 : soit débit soit crédit
    for i, l in enumerate(ecr, start=2):
        d, c = montant(l["Débit"]), montant(l["Crédit"])
        if (d != 0) == (c != 0):
            anomalies.append(("RG-03", f"ligne {i} : débit {d} et crédit {c}"))

    # RG-04 : aucune écriture nouvelle dans une période close. Les Mvt
    # existants à la clôture (jusqu'à P_DernierMvtClos) y restent.
    cloture = nom_defini(wb, "P_DateCloture")
    dernier_clos = nom_defini(wb, "P_DernierMvtClos") or 0
    if isinstance(cloture, dt.datetime):
        dernier = defaultdict(lambda: None)
        for i, l in enumerate(ecr, start=2):
            if (texte(l["Jnl"]) != "AN" and isinstance(l["Date"], dt.datetime) and l["Date"] <= cloture
                    and isinstance(l["Mvt"], (int, float)) and l["Mvt"] > dernier_clos):
                dernier[l["Mvt"]] = i
        for mvt, i in dernier.items():
            anomalies.append(("RG-04", f"Mvt {mvt} (ligne {i}) daté dans la période close au {cloture:%d/%m/%Y}"))

    # Numérotation : une pièce ne sert qu'à un mouvement
    mvts_par_piece = defaultdict(set)
    for l in ecr:
        mvts_par_piece[l["Pièce"]].add(l["Mvt"])
    for piece, mvts in mvts_par_piece.items():
        if len(mvts) > 1:
            anomalies.append(("NUM", f"pièce {piece} partagée par les Mvt {sorted(mvts)}"))

    # Erreurs de formule en cache, tous onglets
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value in ERREURS:
                    anomalies.append(("ERR", f"{ws.title}!{c.coordinate} : {c.value}"))
    return anomalies


def photographier(t, jusqu_a=None):
    """Totaux de T_Ecritures ; jusqu_a limite aux Mvt existants à une photo antérieure."""
    ecr = [l for l in t["T_Ecritures"] if jusqu_a is None or (isinstance(l["Mvt"], (int, float)) and l["Mvt"] <= jusqu_a)]
    par_jnl = defaultdict(lambda: [0.0, 0.0])
    par_classe = defaultdict(float)
    for l in ecr:
        d, c = montant(l["Débit"]), montant(l["Crédit"])
        par_jnl[texte(l["Jnl"])][0] += d
        par_jnl[texte(l["Jnl"])][1] += c
        par_classe[texte(l["Compte"])[:1]] += d - c
    r = lambda x: round(x, 2)
    return {
        "lignes": len(ecr),
        "mouvements": len({l["Mvt"] for l in ecr}),
        "mvt_max": max((l["Mvt"] for l in ecr if isinstance(l["Mvt"], (int, float))), default=0),
        "total_debit": r(sum(v[0] for v in par_jnl.values())),
        "total_credit": r(sum(v[1] for v in par_jnl.values())),
        "resultat": r(-(par_classe.get("6", 0) + par_classe.get("7", 0))),
        "par_journal": {k: {"debit": r(v[0]), "credit": r(v[1])} for k, v in sorted(par_jnl.items())},
        "solde_par_classe": {k: r(v) for k, v in sorted(par_classe.items())},
    }


def comparer(avant, apres):
    """Différences hors écritures nouvelles : on ne compare que ce qui existait."""
    diffs = []
    for cle in ("lignes", "mouvements", "total_debit", "total_credit", "resultat"):
        if avant[cle] != apres[cle]:
            diffs.append(f"{cle} : {avant[cle]} → {apres[cle]}")
    return diffs


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("classeur", nargs="?")
    p.add_argument("--photo", help="enregistre les totaux dans ce fichier JSON")
    p.add_argument("--compare", help="compare aux totaux de ce fichier JSON")
    a = p.parse_args(argv)

    chemin = Path(a.classeur) if a.classeur else RACINE / lire_config().get("classeur", "ComptaJLC.xlsm")
    wb = openpyxl.load_workbook(chemin, data_only=True)
    t = tables(wb)

    anomalies = controler(wb, t)
    photo = photographier(t)
    print(f"{chemin.name} : {photo['lignes']} lignes, {photo['mouvements']} mouvements, "
          f"débit {photo['total_debit']:,.2f} / crédit {photo['total_credit']:,.2f}, résultat {photo['resultat']:,.2f}")
    for regle, msg in anomalies:
        print(f"  [{regle}] {msg}")
    print("Contrôles : OK" if not anomalies else f"Contrôles : {len(anomalies)} anomalie(s)")

    if a.photo:
        Path(a.photo).write_text(json.dumps(photo, ensure_ascii=False, indent=2), encoding="utf-8")
    diffs = []
    if a.compare:
        avant = json.loads(Path(a.compare).read_text(encoding="utf-8"))
        diffs = comparer(avant, photographier(t, jusqu_a=avant["mvt_max"]))
        print(f"Mvt 1 à {avant['mvt_max']} : " + ("totaux identiques à la photographie" if not diffs else "totaux modifiés :"))
        for d in diffs:
            print(f"  {d}")
        if photo["mvt_max"] > avant["mvt_max"]:
            n = photographier(t)
            print(f"Écritures nouvelles (Mvt {avant['mvt_max'] + 1} à {photo['mvt_max']}) : "
                  f"{n['lignes'] - avant['lignes']} lignes, {n['total_debit'] - avant['total_debit']:,.2f} au débit et au crédit, "
                  f"effet sur le résultat {n['resultat'] - avant['resultat']:,.2f}")
    return 1 if anomalies or diffs else 0


if __name__ == "__main__":
    sys.exit(main())
