"""Reprise d'une comptabilité Ciel Compta dans un site ComptaBB neuf (un seul axe, journaux en devise).

Entrées (exports Ciel) :
    Plan.xlsx      feuilles Plan (Compte, LibelCompte, Anal, LibelAnal, …, Journ) et Journ (Journ, Dev, Intitulé)
    RImport.txt    écritures au format d'échange Ciel (« ##Transfert », section Mvt, tabulations, Windows-1252)
    Traduction.xlsx (facultatif) : LibelleH (hébreu) → traduction française des libellés de relevés

Sorties (dossier de sortie et son ZIP), à déposer dans Administration › Imports / Exports puis « Tout importer » :
    Reglages.xlsx      un seul axe (axe 2 masqué), traductions des relevés
    Exercices.xlsx     exercices des dates des écritures
    Axe1.xlsx          codes analytiques Ciel (0BILAN, 1VERL…)
    PlanComptable.xlsx comptes du plan et des écritures ; 401/411 lettrables
    Journaux.xlsx      journaux, compte de trésorerie (colonne Journ du plan), devise (Dev : € → vide, Nis → ILS, $ → USD)
    Ecritures.xlsx     Mvt renumérotés 1, 2, 3… dans l'ordre des dates ; Pièce = n° du mouvement Ciel ; montant d'origine
                       gardé pour les journaux en devise ; lettrage repris
    Traductions.xlsx   si Traduction.xlsx est fourni
    Rapport.txt        chiffres, contrôles (équilibre, totaux par compte), points à vérifier

    python src/reprise_ciel.py Plan.xlsx RImport.txt [--traductions Traduction.xlsx] [--sortie dossier]
"""

import argparse
import datetime as dt
import re
import sys
import zipfile
from collections import Counter, OrderedDict, defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl

RACINE = Path(__file__).resolve().parent.parent
DEVISES = {"€": "", "EUR": "", "NIS": "ILS", "ILS": "ILS", "₪": "ILS", "$": "USD", "USD": "USD"}
ECRITURES = ["Date", "Jnl", "Mvt", "Pièce", "Compte", "Libellé", "Débit", "Crédit", "Anal2", "Let", "Montant devise"]


def _t(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _m(v):
    """« 3 381.71 » (espace insécable) → Decimal ; '' → None."""
    v = (v or "").replace("\xa0", "").replace(" ", "").replace(" ", "").replace(",", ".")
    return Decimal(v).quantize(Decimal("0.01")) if v else None


def lire_plan(chemin):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    comptes = OrderedDict()
    ws = wb["Plan"]
    entetes = [_t(c) for c in next(ws.iter_rows(max_row=1, values_only=True))]
    k = {e: i for i, e in enumerate(entetes)}
    for r in ws.iter_rows(min_row=2, values_only=True):
        numero = _t(r[k["Compte"]])
        if numero:
            comptes[numero] = {"libelle": _t(r[k["LibelCompte"]]), "anal": _t(r[k["Anal"]]), "libanal": _t(r[k["LibelAnal"]]),
                               "journal": _t(r[k["Journ"]]).upper() if "Journ" in k else ""}
    journaux = OrderedDict()
    for r in wb["Journ"].iter_rows(min_row=2, values_only=True):
        if _t(r[0]):
            journaux[_t(r[0]).upper()] = {"devise": DEVISES.get(_t(r[1]).upper(), _t(r[1]).upper()[:3]), "intitule": _t(r[2])}
    return comptes, journaux


def lire_rimport(chemin):
    """Lignes d'écritures (dict) et comptes décrits dans le fichier (lignes « compte, libellé, SR »)."""
    brut = Path(chemin).read_bytes().decode("cp1252", errors="replace")
    lignes, comptes = [], {}
    for texte in brut.splitlines():
        r = [x.strip().strip('"') for x in texte.split("\t")]
        if len(r) < 9:
            continue
        if not re.match(r"\d\d/\d\d/\d{4}$", r[2]):
            if r[0] and r[2] == "SR":
                comptes[r[0]] = r[1]
            continue
        j, m, a = r[2].split("/")
        lignes.append({"ciel": r[0], "journal": r[1].upper(), "date": dt.date(int(a), int(m), int(j)), "compte": r[3],
                       "libcompte": r[4], "montant": _m(r[5]), "sens": r[6], "libelle": r[8] or r[4],
                       "devise": _m(r[9]) if len(r) > 9 else None, "anal": r[13] if len(r) > 13 else "",
                       "libanal": r[14] if len(r) > 14 else "", "lettrage": r[23] if len(r) > 23 else ""})
    return lignes, comptes


def lire_traductions(chemin):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    res = OrderedDict()
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            h, f = _t(r[0] if r else ""), _t(r[1] if len(r) > 1 else "")
            h = h.replace("‫", "").replace("‬", "").replace("‪", "").strip()
            if h and f and h != "LibelleH" and not f.startswith("#"):
                res[h] = f
    return res


def reprendre(plan, journaux, lignes, comptes_rimport, traductions=None):
    fichiers, rapport, alertes = OrderedDict(), [], []

    # opérations Ciel, contrôle d'équilibre
    ops = OrderedDict()
    for l in lignes:
        ops.setdefault(l["ciel"], []).append(l)
    for n, ls in ops.items():
        solde = sum((l["montant"] if l["sens"] == "D" else -l["montant"]) for l in ls)
        if solde:
            alertes.append(f"Mouvement Ciel {n} déséquilibré de {solde} € : repris tel quel, à corriger.")
        if len({(l["journal"], l["date"]) for l in ls}) > 1:
            alertes.append(f"Mouvement Ciel {n} : plusieurs journaux ou dates.")

    # plan comptable : plan Ciel + comptes des écritures
    comptes = OrderedDict((c, dict(v)) for c, v in plan.items())
    for l in lignes:
        if l["compte"] not in comptes:
            comptes[l["compte"]] = {"libelle": comptes_rimport.get(l["compte"]) or l["libcompte"], "anal": l["anal"],
                                    "libanal": l["libanal"], "journal": ""}
            alertes.append(f"Compte {l['compte']} absent de Plan.xlsx : créé d'après les écritures ({l['libcompte']}).")
    codes = OrderedDict()
    for c in comptes.values():
        if c["anal"] and c["anal"] not in codes:
            codes[c["anal"]] = c["libanal"] or c["anal"]
    for l in lignes:
        if l["anal"] and l["anal"] not in codes:
            codes[l["anal"]] = l["libanal"] or l["anal"]
    fichiers["Reglages.xlsx"] = (["Clé", "Valeur", "Description"], [
        ["un_seul_axe", "oui", "Un seul axe analytique (axe 1) ; axe 2 masqué"],
        ["axe2_defaut", "GEN", "Code axe 2 mis d'office"],
        ["traductions_releve", "oui" if traductions else "non", "Traduire les libellés du relevé (hébreu → français)"],
    ])
    annees = sorted({l["date"].year for l in lignes})
    fichiers["Exercices.xlsx"] = (["Libellé", "Début", "Fin", "Clos"],
                                  [[str(a), dt.date(a, 1, 1), dt.date(a, 12, 31), "non"] for a in annees])
    fichiers["Axe1.xlsx"] = (["Code", "Libellé"], [[c, lib[:100]] for c, lib in codes.items()])
    fichiers["PlanComptable.xlsx"] = (["Compte", "Libellé", "Axe 1", "Lettrable", "Actif"], [
        [n, (c["libelle"] or n)[:100], c["anal"], "oui" if n.startswith(("401", "411")) else "non", "oui"]
        for n, c in comptes.items()])

    # journaux : compte de trésorerie = compte du plan rattaché au journal (le plus utilisé dans ses écritures)
    usage = defaultdict(Counter)
    for l in lignes:
        usage[l["journal"]][l["compte"]] += 1
    for l in lignes:
        if l["journal"] not in journaux:
            journaux[l["journal"]] = {"devise": "", "intitule": l["journal"]}
            alertes.append(f"Journal {l['journal']} absent de la feuille Journ : créé en euros.")
    lignes_j, sans_compte = [], []
    for code, j in journaux.items():
        rattaches = [n for n, c in comptes.items() if c["journal"] == code and n[:1] == "5" and not n.startswith("58")]
        rattaches.sort(key=lambda n: (-usage[code][n], not n.startswith("512"), n))
        compte = rattaches[0] if rattaches else ""
        if not compte and code in ("I1",):
            compte = next((n for n, c in comptes.items() if c["journal"] == code), "")
        type_ = "BQ" if compte.startswith("51") else ("CA" if compte.startswith("53") else ("CR" if compte.startswith("50")
                                                                                              else code[:10]))
        if not compte and code not in ("AN", "HA", "VT") and not code.startswith("OD"):
            sans_compte.append(code)
        lignes_j.append([code, j["intitule"][:60], type_, compte, "oui", j["devise"]])
    fichiers["Journaux.xlsx"] = (["Code", "Intitulé", "Type", "Compte de trésorerie", "Actif", "Devise"], lignes_j)
    devise_j = {j[0]: j[5] for j in lignes_j}

    # écritures : Mvt renumérotés par date, pièce = n° Ciel
    ordre = sorted(ops, key=lambda n: (ops[n][0]["date"], int(n) if n.isdigit() else 0, n))
    sortie = []
    for i, n in enumerate(ordre, 1):
        for l in ops[n]:
            d = l["montant"] if l["sens"] == "D" else None
            c = l["montant"] if l["sens"] != "D" else None
            devise = abs(l["devise"]) if devise_j.get(l["journal"]) and l["devise"] is not None else None
            if l["montant"] == 0:
                alertes.append(f"Mouvement Ciel {n}, compte {l['compte']} : ligne à 0 ignorée.")
                continue
            sortie.append([l["date"], l["journal"], i, int(n) if n.isdigit() else i, l["compte"], l["libelle"][:200],
                           float(d) if d else None, float(c) if c else None, "", l["lettrage"][:10],
                           float(devise) if devise is not None else None])
    fichiers["Ecritures.xlsx"] = (ECRITURES, sortie)
    if traductions:
        fichiers["Traductions.xlsx"] = (["Opération (hébreu)", "Traduction"], [[h, f] for h, f in traductions.items()])

    # totaux par compte (contrôle avec la balance Ciel)
    soldes = defaultdict(Decimal)
    for l in lignes:
        soldes[l["compte"]] += l["montant"] if l["sens"] == "D" else -l["montant"]
    rapport += [
        f"Écritures Ciel : {len(ops)} mouvements, {len(sortie)} lignes, du {min(l['date'] for l in lignes):%d/%m/%Y} "
        f"au {max(l['date'] for l in lignes):%d/%m/%Y} ; renumérotées Mvt 1 à {len(ordre)} (pièce = n° Ciel).",
        f"Comptes : {len(comptes)} ; codes axe 1 : {len(codes)} ; journaux : {len(lignes_j)} "
        f"(en devise : {', '.join(f'{j[0]} {j[5]}' for j in lignes_j if j[5]) or 'aucun'}).",
        f"Total des débits : {sum(l['montant'] for l in lignes if l['sens'] == 'D')} € ; "
        f"des crédits : {sum(l['montant'] for l in lignes if l['sens'] != 'D')} €.",
    ]
    if sans_compte:
        rapport.append("Journaux sans compte de trésorerie (aucun compte 5… rattaché dans le plan) : " + ", ".join(sans_compte) + ".")
    rapport += alertes or ["Aucune anomalie : chaque mouvement est équilibré et tous ses comptes existent."]
    rapport.append("")
    rapport.append("Soldes par compte (débit − crédit), à comparer à la balance Ciel :")
    rapport += [f"  {n:<12} {comptes.get(n, {}).get('libelle', '')[:40]:<40} {soldes[n]:>14}" for n in sorted(soldes)]
    return fichiers, rapport


def ecrire(fichiers, rapport, sortie):
    sortie.mkdir(parents=True, exist_ok=True)
    for vieux in sortie.glob("*.xlsx"):
        vieux.unlink()
    for nom, (entetes, lignes) in fichiers.items():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = nom.rsplit(".", 1)[0]
        ws.append(entetes)
        for r in lignes:
            ws.append(r)
        for k, e in enumerate(entetes, 1):
            if e in ("Date", "Début", "Fin"):
                for (c,) in ws.iter_rows(min_row=2, min_col=k, max_col=k):
                    c.number_format = "dd/mm/yyyy"
        wb.save(sortie / nom)
    (sortie / "Rapport.txt").write_text("\n".join(rapport) + "\n", encoding="utf-8")
    with zipfile.ZipFile(sortie.with_suffix(".zip"), "w", zipfile.ZIP_DEFLATED) as z:
        for nom in list(fichiers) + ["Rapport.txt"]:
            z.write(sortie / nom, nom)


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("plan")
    p.add_argument("rimport")
    p.add_argument("--traductions")
    p.add_argument("--sortie", default=str(RACINE / "Imports" / "reprise_ciel"))
    a = p.parse_args(argv)
    plan, journaux = lire_plan(a.plan)
    lignes, comptes = lire_rimport(a.rimport)
    fichiers, rapport = reprendre(plan, journaux, lignes, comptes, lire_traductions(a.traductions) if a.traductions else None)
    ecrire(fichiers, rapport, Path(a.sortie))
    print("\n".join(r for r in rapport if not r.startswith("  ")))
    print(f"Fichiers écrits dans {a.sortie} (et {Path(a.sortie).with_suffix('.zip').name})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
