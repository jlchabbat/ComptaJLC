"""Reprise de l'extrait SUMIT (classeur EXTRACT_SUMIT_*.xlsx) dans l'application web.

    python src/reprise_sumit.py EXTRACT_SUMIT.xlsx [dossier_de_sortie]

Le classeur SUMIT est adapté aux structures de la base, sans rien changer à
l'application ni à ses formats de fichiers :

- codes analytiques au format de la base : ACT1 -> ACT.1 (axe 1), MAN7 -> MAN.007
  (axe 2) ; le code BIL1 est réparti selon la classe du compte (BIL.1 classes 1-2,
  BIL.4 tiers, BIL.5 trésorerie) ; ligne sans axe 2 -> GEN.001 ;
- un Mvt par opération (pièce SUMIT « B3-001 »), numéroté dans l'ordre des dates,
  Pièce = N° de la Base SUMIT (repris aussi dans le commentaire) ;
- opérations de montant nul écartées (une ligne porte un débit OU un crédit) ;
- journaux AN, B1, B2, B3, CA, OD du plan SUMIT, plus VT et HA (saisie guidée) ;
- réglages de la Loge, relevé Mizrahi CC (B1) de l'onglet « Rappro Mizrahi CC »
  pointé comme dans SUMIT, lexique hébreu -> français en traductions.

Le tout est chargé dans une base neuve et temporaire, contrôlé contre la balance
SUMIT, puis écrit en « export complet » (Export_complet_*.zip). Sur le site :
Base de données › Remettre à zéro et recharger, avec ce fichier (case « Fichiers
modifiés » décochée). Les comptes utilisateurs du site sont gardés ; tout le reste
est remplacé (une sauvegarde est faite juste avant).
"""

import datetime as dt
import os
import shutil
import sys
import tempfile
import warnings
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent
EXTRAIT = "27/09/2026"


def code_axe(code, axe):
    """ACT1 -> ACT.1 ; MAN7 -> MAN.007."""
    lettres = code.rstrip("0123456789")
    return f"{lettres}.{int(code[len(lettres):]):0{3 if axe == 2 else 1}d}"


def axe1_du_compte(compte, code):
    if code != "BIL1":
        return code_axe(code, 1)
    return {"4": "BIL.4", "5": "BIL.5"}.get(compte[:1], "BIL.1")


def montant(v):
    return Decimal(str(v or 0)).quantize(Decimal("0.01"))


def jour(v):
    return v.date() if isinstance(v, dt.datetime) else v


def lignes_tableau(ws, entete_ligne):
    """Lignes (dict par en-tête) sous la ligne d'en-têtes, jusqu'à la première ligne vide en colonne A."""
    rangees = ws.iter_rows(min_row=entete_ligne, values_only=True)
    entetes = [str(e).strip() if e is not None else "" for e in next(rangees)]
    for r in rangees:
        if r[0] in (None, ""):
            break
        yield dict(zip(entetes, r))


def lire_sumit(chemin):
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(chemin, data_only=True)
    d = {}
    plan = wb["Plan"]
    rangees = list(plan.iter_rows(min_row=4, values_only=True))
    d["comptes"] = [(str(r[0]), r[1], r[2]) for r in rangees if r[0]]
    d["journaux"] = [(r[4], r[5], str(r[6] or "")) for r in rangees if r[4]]
    d["axe1"] = [(r[8], r[9]) for r in rangees if r[8] and r[9] and not str(r[8]).startswith("Contrôle")]
    d["axe2"] = [(r[12], r[13]) for r in rangees if r[12] and r[13]]
    d["ecritures"] = list(lignes_tableau(wb["Ecritures"], 4))
    d["base"] = {r["N°"]: r for r in lignes_tableau(wb["Base"], 1)}
    d["balance"] = {str(r["Compte"]): montant(r["Total débit"]) - montant(r["Total crédit"])
                    for r in lignes_tableau(wb["Balance"], 3)}
    d["lexique"] = [(r["LibelleH"], r["Libellé FR"]) for r in lignes_tableau(wb["Lexique HE-FR"], 4) if r["Libellé FR"]]
    d["releve_b1"] = list(lignes_tableau(wb["Rappro Mizrahi CC"], 12))
    return d


def installer_django(donnees):
    os.environ["COMPTABB_DATA"] = str(donnees)
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "comptabb.settings")
    sys.path.insert(0, str(RACINE / "appli"))
    import django
    django.setup()
    from django.core.management import call_command
    call_command("migrate", verbosity=0)


def charger(d):
    """Remplit la base neuve ; renvoie le compte rendu (liste de lignes)."""
    from django.db import transaction
    from compta import releves
    from compta.export_complet import vider
    from compta.models import (CodeAnalytique, Compte, Exercice, Journal, LigneReleve, Ligne, Modification, Mouvement,
                               ParametreReleve, Prefixe, Reglage, Traduction)
    from compta.reglages import REGLAGES
    from compta.saisie import initialiser_parametres

    rapport = []
    with transaction.atomic():
        vider()
        for cle, description, _, loge in REGLAGES:
            Reglage.objects.create(cle=cle, valeur=loge, description=description)
        for cle, valeur, description in (
                ("compte_virement", "580000", "Virements internes et paiements carte Isracard (décision Q2)"),
                ("compte_attente", "470000", "Opérations en attente d'éclaircissement (décision Q3)"),
                ("compte_report_a_nouveau", "110000", "Compte de report à nouveau (clôture)"),
                ("journal_a_nouveaux", "AN", "Journal des à-nouveaux (clôture)")):
            Reglage.objects.create(cle=cle, valeur=valeur, description=description)

        # axes et préfixes
        bil = {"BIL.1": "BILAN - FONDS ASSOCIATIFS ET IMMOBILISATIONS", "BIL.4": "BILAN - TIERS",
               "BIL.5": "BILAN - TRESORERIE"}
        for code, lib in bil.items():
            CodeAnalytique.objects.create(code=code, axe=1, libelle=lib)
        for code, lib in d["axe1"]:
            if code != "BIL1":
                CodeAnalytique.objects.create(code=code_axe(code, 1), axe=1, libelle=lib.upper()[:100])
        for code, lib in d["axe2"]:
            CodeAnalytique.objects.create(code=code_axe(code, 2), axe=2, libelle=lib[:100], statut=1)
        for axe in (1, 2):
            for p in sorted({c.split(".")[0] for c in CodeAnalytique.objects.filter(axe=axe).values_list("code", flat=True)}):
                Prefixe.objects.create(prefixe=p + ".", axe=axe, libelle="")

        # plan comptable et journaux
        for numero, lib, a1 in d["comptes"]:
            Compte.objects.create(numero=numero, libelle=lib[:100], anal1_id=axe1_du_compte(numero, a1),
                                  lettrable=numero.startswith(("401", "411")))
        types = {"AN": "AN", "B1": "BQ", "B2": "BQ", "B3": "BQ", "CA": "CA", "OD": "OD"}
        for code, intitule, compte in d["journaux"]:
            Journal.objects.create(code=code, intitule=intitule, type=types.get(code, ""), compte_id=compte or None)
        for code, intitule in (("VT", "VENTES ET RECETTES"), ("HA", "ACHATS")):
            Journal.objects.create(code=code, intitule=intitule, type=code)

        # écritures : un Mvt par pièce SUMIT
        par_piece = defaultdict(list)
        for r in d["ecritures"]:
            par_piece[r["Pièce"]].append(r)
        ecartes, mvts = [], []
        for piece, ls in par_piece.items():
            if all(montant(r["Débit"]) == 0 and montant(r["Crédit"]) == 0 for r in ls):
                ecartes.append(ls[0])
                continue
            mvts.append(ls)
        mvts.sort(key=lambda ls: (jour(ls[0]["Date"]), ls[0]["Journ"] != "AN", ls[0]["N° base"]))
        for numero, ls in enumerate(mvts, 1):
            n = ls[0]["N° base"]
            b = d["base"].get(n, {})
            commentaire = (f"Reprise SUMIT (extrait du {EXTRAIT}) : N° base {n}, pièce {ls[0]['Pièce']}, "
                           f"source {b.get('Source') or '?'}, statut {b.get('Statut') or '?'}, catégorie {b.get('Catégorie') or '?'}")
            m = Mouvement.objects.create(numero=numero, date=jour(ls[0]["Date"]), journal_id=ls[0]["Journ"], piece=n,
                                         origine="import", commentaire=commentaire)
            Ligne.objects.bulk_create([
                Ligne(mouvement=m, ordre=i, compte_id=str(r["Compte"]), libelle=str(r["Libellé écriture"])[:200],
                      debit=montant(r["Débit"]), credit=montant(r["Crédit"]),
                      anal2_id=code_axe(r["AnalAxe2"], 2) if r["AnalAxe2"] else "GEN.001")
                for i, r in enumerate(ls)])
        rapport.append(f"{len(mvts)} mouvements, {Ligne.objects.count()} lignes repris")
        rapport += [f"écarté (montant nul) : N° base {r['N° base']} {r['Date']:%d/%m/%Y} {r['Journ']} {r['Libellé écriture']}"
                    for r in ecartes]

        premiere = min(m.date for m in Mouvement.objects.all())
        Exercice.objects.create(libelle="2026", debut=min(premiere, dt.date(2026, 1, 1)), fin=dt.date(2026, 12, 31))
        if premiere < dt.date(2026, 1, 1):
            n = Mouvement.objects.filter(date__lt=dt.date(2026, 1, 1)).count()
            rapport.append(f"exercice 2026 ouvert le {premiere:%d/%m/%Y} : {n} Mvt SUMIT datés de ce jour (dates gardées)")

        for hebreu, fr in d["lexique"]:
            Traduction.objects.update_or_create(cle=Traduction.cle_de(hebreu), defaults={"hebreu": hebreu, "traduction": fr})

        initialiser_parametres()

        # relevé Mizrahi CC (B1) et pointage de SUMIT
        b1 = Journal.objects.get(code="B1")
        ParametreReleve.objects.create(journal=b1, date_reprise=dt.date(2026, 1, 1),
                                       libelle="Mizrahi-Tefahot 732-182029 (compte courant)")
        lignes = [{"date": jour(r["Date"]), "reference": str(r["Réf. banque"] or "")[:40],
                   "operation": str(r["LibelleH"] or r["Opération (traduite)"])[:200], "montant": montant(r["Montant"]),
                   "solde": montant(r["Solde progressif"])} for r in d["releve_b1"]]
        solde_depart = lignes[0]["solde"] - lignes[0]["montant"]
        releves.importer(b1, lignes, source=f"Extrait SUMIT {EXTRAIT}", solde_ouverture=solde_depart)
        rapport.append(f"relevé B1 : {len(lignes)} lignes, solde d'ouverture {solde_depart}, "
                       f"{len(releves.ecarts_solde(b1))} écart(s) de solde")
        _pointer_comme_sumit(d, b1, solde_depart, rapport)

        Modification.objects.create(auteur="Reprise", lot="Reprise SUMIT", action="Remise à zéro et reprise",
                                    objet=f"Extrait SUMIT du {EXTRAIT}",
                                    apres="; ".join(rapport)[:300])
    return rapport


def _pointer_comme_sumit(d, b1, solde_depart, rapport):
    """Pointe le solde d'ouverture sur l'à-nouveau, puis les lignes du relevé sur les écritures désignées par SUMIT
    (colonne « N° Base SUMIT ») ; les lignes hors SUMIT de somme nulle le même jour sont reliées entre elles."""
    from compta import releves
    from compta.models import Ligne, LigneReleve

    ouverture = LigneReleve.objects.get(journal=b1, ouverture=True)
    an = Ligne.objects.filter(compte=b1.compte, mouvement__journal_id="AN", debit=solde_depart).first()
    if an:
        releves.pointer(b1, [ouverture], [an], mode="auto")

    rangs = defaultdict(int)
    groupes = defaultdict(list)                    # N° base -> lignes du relevé
    nulles = defaultdict(list)                     # lignes hors SUMIT, par date
    for r in d["releve_b1"]:
        date = jour(r["Date"])
        rangs[date] += 1
        l = LigneReleve.objects.get(journal=b1, date=date, rang=rangs[date], ouverture=False)
        if l.rapprochement_id:
            continue
        numeros = frozenset(int(x) for x in str(r["N° Base SUMIT"] or "").split("+") if x.strip())
        (groupes[numeros] if numeros else nulles[date]).append(l)
    ok, refus = 0, []
    for numeros, ls in groupes.items():
        ecr = Ligne.objects.filter(compte=b1.compte, mouvement__piece__in=numeros, rapprochement__isnull=True)
        try:
            releves.pointer(b1, ls, ecr, mode="auto")
            ok += 1
        except ValueError as e:
            refus.append(f"N° {'+'.join(map(str, sorted(numeros)))} : {e}")
    for date, ls in nulles.items():
        try:
            releves.relier_nulles(ls)
            ok += 1
        except ValueError as e:
            refus.append(f"lignes hors SUMIT du {date:%d/%m/%Y} : {e}")
    reste_r = LigneReleve.objects.filter(journal=b1, rapprochement__isnull=True).count()
    reste_e = Ligne.objects.filter(compte=b1.compte, rapprochement__isnull=True).count()
    rapport.append(f"pointage B1 : {ok} groupes pointés ; restent {reste_r} ligne(s) du relevé et {reste_e} écriture(s)")
    rapport += refus


def verifier(d):
    """Contrôles de l'application et comparaison à la balance SUMIT ; renvoie les anomalies."""
    from django.db.models import Sum
    from compta import controles
    from compta.models import Ligne
    anomalies = [f"{r.regle} {r.libelle} : {r.valeur}" for r in controles.executer() if r.statut == "Anomalie"]
    base = {c: montant(s["d"]) - montant(s["c"]) for c, s in
            ((x["compte_id"], x) for x in Ligne.objects.values("compte_id").annotate(d=Sum("debit"), c=Sum("credit")))}
    for compte in sorted(set(base) | set(d["balance"])):
        a, b = base.get(compte, Decimal("0.00")), d["balance"].get(compte, Decimal("0.00"))
        if a != b:
            anomalies.append(f"compte {compte} : base {a}, balance SUMIT {b}")
    return anomalies


def a_verifier():
    from compta import controles
    return [f"{r.libelle} : {r.valeur}" for r in controles.executer() if r.statut == "À vérifier"]


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    source = Path(argv[1])
    sortie = Path(argv[2]) if len(argv) > 2 else Path.cwd()
    sortie.mkdir(parents=True, exist_ok=True)
    d = lire_sumit(source)
    travail = Path(tempfile.mkdtemp(prefix="reprise_sumit_"))
    installer_django(travail)
    rapport = charger(d)
    anomalies, remarques = verifier(d), a_verifier()
    from compta.export_complet import exporter
    zip_ = exporter(auteur="Reprise SUMIT")
    cible = sortie / zip_.name.replace("Export_complet_", "Export_complet_SUMIT_")
    shutil.copy(zip_, cible)
    print("\n".join(rapport))
    print("Contrôles :", "aucune anomalie" if not anomalies else "\n  " + "\n  ".join(anomalies))
    for r in remarques:
        print("À vérifier :", r)
    print(f"Fichier à recharger sur le site : {cible}")
    shutil.rmtree(travail, ignore_errors=True)
    return 1 if anomalies else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
