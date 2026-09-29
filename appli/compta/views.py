"""Pages de consultation : tableau de bord, écritures, grand livre, balance, analytique, contrôles."""

import datetime as dt

from django.contrib.auth.decorators import login_required, permission_required
from django.core.paginator import Paginator
from django.db.models import Case, F, Q, Sum, Value, When
from django.core.exceptions import PermissionDenied
from django.shortcuts import get_object_or_404, redirect, render

from . import controles as ctrl
from . import etats
from .models import ZERO, arrondi, CodeAnalytique, Compte, Exercice, Journal, Ligne, Modification, Mouvement, soldes

consulter = permission_required("compta.view_mouvement", raise_exception=True)


def periode(request):
    """Période choisie (?du=AAAA-MM-JJ&au=…), par défaut l'exercice ouvert."""
    ex = Exercice.ouvert()
    def lire(cle, defaut):
        try:
            return dt.date.fromisoformat(request.GET.get(cle, ""))
        except ValueError:
            return defaut
    debut = lire("du", ex.debut if ex else dt.date(2000, 1, 1))
    fin = lire("au", ex.fin if ex else dt.date(2100, 12, 31))
    return debut, fin


def par_page(request):
    """100 lignes par page ; toutes pour l'export Excel ou l'impression (?tout=1)."""
    return 10 ** 7 if request.GET.get("tout") else 100


def lignes_periode(debut, fin):
    return Ligne.objects.filter(mouvement__date__range=(debut, fin))


def resultat_par_axe(lignes, axe):
    """Produits (7), charges (6) et résultat par code d'axe 1 ou 2, en une requête."""
    champ = "compte__anal1" if axe == 1 else "anal2"
    classe7 = Q(compte__numero__startswith="7")
    classe6 = Q(compte__numero__startswith="6")
    rows = (lignes.filter(classe6 | classe7).values(champ, f"{champ}__libelle").order_by(champ)
            .annotate(produits=Sum(Case(When(classe7, then=F("credit") - F("debit")), default=Value(ZERO))),
                      charges=Sum(Case(When(classe6, then=F("debit") - F("credit")), default=Value(ZERO)))))
    return [{"code": r[champ], "libelle": r[f"{champ}__libelle"], "produits": arrondi(r["produits"]),
             "charges": arrondi(r["charges"]), "resultat": arrondi(r["produits"] - r["charges"])} for r in rows]


@login_required
def tableau_de_bord(request):
    if not request.user.has_perm("compta.view_mouvement"):
        if request.user.has_perm("compta.view_fiche"):
            return redirect("fiches")     # bénévole : ses fiches
        raise PermissionDenied
    debut, fin = periode(request)
    ls = lignes_periode(debut, fin)
    _, _, charges = soldes(ls.filter(compte__numero__startswith="6"))
    _, _, produits = soldes(ls.filter(compte__numero__startswith="7"))
    tresorerie = []
    for j in Journal.objects.filter(compte__isnull=False).select_related("compte"):
        s = etats.solde_cumule(j.compte, fin)
        tresorerie.append({"journal": j, "solde": s})
    res = ctrl.executer()
    etat, a_verifier = ctrl.etat_general(res)
    return render(request, "compta/tableau_de_bord.html", {
        "debut": debut, "fin": fin, "produits": -produits, "charges": charges, "resultat": -(produits + charges),
        "tresorerie": tresorerie, "total_tresorerie": sum((t["solde"] for t in tresorerie), ZERO),
        "axe1": resultat_par_axe(ls, 1), "axe2": resultat_par_axe(ls, 2), "etat": etat, "a_verifier": a_verifier,
        "nb_mouvements": Mouvement.objects.filter(date__range=(debut, fin)).count(),
    })


TRIS_ECRITURES = {
    "date": ("Date", "mouvement__date"), "jnl": ("Jnl", "mouvement__journal_id"), "mvt": ("Mvt", "mouvement__numero"),
    "piece": ("Pièce", "mouvement__piece"), "compte": ("Compte", "compte_id"), "libelle": ("Libellé", "libelle"),
    "debit": ("Débit", "debit"), "credit": ("Crédit", "credit"), "anal1": ("Axe 1", "compte__anal1_id"),
    "anal2": ("Axe 2", "anal2_id"),
}


COLONNES_CIEL = ["Mvt", "Journ", "Date", "Compte", "LibelCompte", "Debit", "Credit", "Npiece", "Anal", "LibelAnal", "Lettr"]


def export_ciel(lignes, debut, fin):
    """Écritures (période et filtres de la page) au format de contrôle de Ciel Compta : une ligne par ligne d'écriture ;
    Npiece = code axe 2 (texte), Anal / LibelAnal = code axe 1 du compte et son libellé, Lettr = lettrage."""
    import io

    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ecritures"
    ws.append(COLONNES_CIEL)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    for l in lignes.order_by("mouvement__date", "mouvement__numero", "ordre"):
        a1 = l.compte.anal1
        ws.append([l.mouvement.numero, l.mouvement.journal_id, l.mouvement.date, l.compte_id, l.compte.libelle,
                   float(l.debit), float(l.credit), l.anal2_id or "", a1.code if a1 else "", a1.libelle if a1 else "",
                   l.lettrage or ""])
        r = ws.max_row
        ws.cell(r, 3).number_format = "DD/MM/YYYY"
        ws.cell(r, 6).number_format = ws.cell(r, 7).number_format = "0.00"
        for col in (4, 8, 9):                             # codes en texte (Npiece = code axe 2)
            ws.cell(r, col).number_format = "@"
    for col, largeur in zip("ABCDEFGHIJK", (8, 7, 11, 14, 34, 12, 12, 12, 10, 28, 7)):
        ws.column_dimensions[col].width = largeur
    ws.freeze_panes = "A2"
    tampon = io.BytesIO()
    wb.save(tampon)
    r = HttpResponse(tampon.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r["Content-Disposition"] = f'attachment; filename="Ecritures_Ciel_{debut:%Y-%m-%d}_{fin:%Y-%m-%d}.xlsx"'
    return r


@login_required
@consulter
def ecritures(request):
    debut, fin = periode(request)
    qs = (lignes_periode(debut, fin).select_related("mouvement", "mouvement__journal", "compte", "compte__anal1", "anal2")
          .order_by("-mouvement__date", "-mouvement__numero", "ordre"))
    f = {k: request.GET.get(k, "").strip() for k in ("journal", "compte", "anal2", "q", "just")}
    if f["journal"]:
        qs = qs.filter(mouvement__journal_id=f["journal"])
    if f["compte"]:
        qs = qs.filter(compte__numero__startswith=f["compte"])
    if f["anal2"]:
        qs = qs.filter(anal2_id=f["anal2"])
    if f["q"]:
        qs = qs.filter(Q(libelle__icontains=f["q"]) | Q(compte__libelle__icontains=f["q"]))
    if f["just"] in ("avec", "sans"):                   # mouvements avec / sans justificatif joint
        qs = qs.filter(mouvement__justificatifs__isnull=(f["just"] == "sans")).distinct()
    if request.GET.get("format") == "ciel":
        return export_ciel(qs, debut, fin)
    d, c, _ = soldes(qs)
    # tri par colonne (sur toutes les pages) : ?tri=<colonne>&ordre=asc|desc
    tri, ordre = request.GET.get("tri", ""), request.GET.get("ordre", "asc")
    if tri in TRIS_ECRITURES:
        champ = TRIS_ECRITURES[tri][1]
        qs = qs.order_by(("-" if ordre == "desc" else "") + champ, "-mouvement__numero", "ordre")
    entetes = []
    for cle, (titre, _) in TRIS_ECRITURES.items():
        params = request.GET.copy()
        params.pop("page", None)
        params["tri"], params["ordre"] = cle, "desc" if (tri == cle and ordre == "asc") else "asc"
        entetes.append({"titre": titre, "url": "?" + params.urlencode(), "sens": ordre if tri == cle else "",
                        "n": cle in ("debit", "credit")})
    page = Paginator(qs, par_page(request)).get_page(request.GET.get("page"))
    from .models import Justificatif
    pieces = {}
    for mvt in Justificatif.objects.filter(mouvement__in={l.mouvement_id for l in page}).values_list("mouvement", flat=True):
        pieces[mvt] = pieces.get(mvt, 0) + 1
    return render(request, "compta/ecritures.html", {
        "page": page, "pieces": pieces, "filtres": f, "entetes": entetes, "debut": debut, "fin": fin, "total_debit": d, "total_credit": c,
        "journaux": Journal.objects.all(), "codes2": CodeAnalytique.objects.filter(axe=2), "comptes": Compte.objects.all(),
    })


@login_required
@consulter
def mouvement(request, numero):
    m = get_object_or_404(Mouvement.objects.select_related("journal"), numero=numero)
    from .corrections import refus_suppression as refus_mvt, verrou
    from .justificatifs import liste_a_classer, refus_suppression
    from .vues_justificatifs import peut_ajouter
    pieces = list(m.justificatifs.all())
    joindre = peut_ajouter(request.user)
    return render(request, "compta/mouvement.html", {"m": m, "lignes": m.lignes.select_related("compte", "anal2"),
                                                     "verrou": verrou(m), "justificatifs": pieces,
                                                     "peut_joindre": joindre,
                                                     "a_classer": liste_a_classer() if joindre else [],
                                                     "refus_suppression": refus_suppression(pieces[0]) if pieces else "",
                                                     "refus_suppression_mvt": refus_mvt(m)})


@login_required
@consulter
def grand_livre(request):
    debut, fin = periode(request)
    numero = request.GET.get("compte", "")
    compte = Compte.objects.filter(numero=numero).first()
    lignes, ouverture = [], ZERO
    if compte:
        o = etats.origine(debut)
        avant = Ligne.objects.filter(compte=compte, mouvement__date__lt=debut)
        if compte.numero[0] in "67":         # gestion : depuis le début de l'exercice de la période
            ex = Exercice.objects.filter(debut__lte=debut, fin__gte=debut).first()
            avant = avant.filter(mouvement__date__gte=ex.debut) if ex else avant
        elif o:
            avant = avant.filter(mouvement__date__gte=o)
        _, _, ouverture = soldes(avant)
        cumul = ouverture
        periode_ls = Ligne.objects.filter(compte=compte, mouvement__date__range=(debut, fin))
        o_fin = etats.origine(fin)
        if o_fin and o_fin > debut:          # la période enjambe une clôture : l'historique continue, sans les à-nouveaux
            periode_ls = periode_ls.exclude(mouvement__origine="cloture")
        for l in (periode_ls
                  .select_related("mouvement", "anal2").order_by("mouvement__date", "mouvement__numero", "ordre")):
            cumul += l.debit - l.credit
            lignes.append((l, cumul))
    return render(request, "compta/grand_livre.html", {
        "comptes": Compte.objects.all(), "compte": compte, "lignes": lignes, "ouverture": ouverture,
        "debut": debut, "fin": fin, "cloture": lignes[-1][1] if lignes else ouverture,
    })


@login_required
@consulter
def balance(request):
    debut, fin = periode(request)
    rows = (lignes_periode(debut, fin).values("compte__numero", "compte__libelle")
            .annotate(d=Sum("debit"), c=Sum("credit")).order_by("compte__numero"))
    lignes = [{"numero": r["compte__numero"], "libelle": r["compte__libelle"], "debit": arrondi(r["d"]),
               "credit": arrondi(r["c"]), "solde": arrondi(r["d"] - r["c"])} for r in rows]
    return render(request, "compta/balance.html", {
        "lignes": lignes, "debut": debut, "fin": fin,
        "total_debit": sum((l["debit"] for l in lignes), ZERO), "total_credit": sum((l["credit"] for l in lignes), ZERO),
    })


@login_required
@consulter
def controles(request):
    res = ctrl.executer()
    etat, a_verifier = ctrl.etat_general(res)
    return render(request, "compta/controles.html", {"resultats": res, "etat": etat, "a_verifier": a_verifier})


@login_required
@consulter
def modifications(request):
    q = request.GET.get("q", "").strip()
    qs = Modification.objects.all()
    if q:
        qs = qs.filter(Q(auteur__icontains=q) | Q(action__icontains=q) | Q(objet__icontains=q) | Q(lot__icontains=q))
    return render(request, "compta/modifications.html", {"page": Paginator(qs, par_page(request)).get_page(request.GET.get("page")), "q": q,
                                                         "total": qs.count()})


# ---------------------------------------------------------------- saisie guidée et codes (W1)

from django.contrib import messages  # noqa: E402
from django.db import transaction  # noqa: E402

from . import saisie as moteur  # noqa: E402
from .forms import CodeForm, MembreForm, SaisieForm, StatutForm  # noqa: E402
from .models import ModeleOperation, Prefixe  # noqa: E402


@login_required
@permission_required("compta.add_mouvement", raise_exception=True)
def saisie(request):
    initial = {"date": dt.date.today()}
    initial.update({k: request.GET[k] for k in ("date", "montant", "paiement", "libelle") if request.GET.get(k)})
    releve = request.POST.get("releve") or request.GET.get("releve") or ""
    form = SaisieForm(request.POST or None, initial=initial)
    resultat = op = None
    if request.method == "POST" and form.is_valid():
        c = form.cleaned_data
        op = moteur.Operation(date=c["date"], modele=c["modele"], tiers=c["tiers"], montant=c["montant"], paiement=c["paiement"],
                              vers=c["vers"], anal2=c["anal2"], compte=c["compte"], remboursement=c["remboursement"],
                              libelle=c["libelle"])
        resultat = moteur.controler(op)
        doublon_seul = list(resultat.erreurs) == ["Déjà enregistrée ?"] and c["forcer"]
        if "enregistrer" in request.POST and (resultat.ok or doublon_seul):
            crees = moteur.enregistrer(op, request.user, forcer_doublon=doublon_seul)
            messages.success(request, "Enregistré : " + ", ".join(f"Mvt {m.numero}" for m in crees) + f" · {resultat.libelle}")
            if releve.isdigit():
                from .vues_rapprochement import pointer_apres_saisie
                r = pointer_apres_saisie(int(releve), crees, request.user)
                if r:
                    messages.success(request, f"Ligne du relevé pointée ({r}).")
                else:
                    messages.warning(request, "Aucune écriture de banque de même montant : ligne du relevé non pointée.")
            return redirect("mouvement", crees[0].numero)
    numero, piece = Mouvement.prochain_numero(), Mouvement.prochaine_piece()
    modeles = {m.pk: {"aide": m.aide, "prefixe": m.tiers.prefixe if m.tiers else "", "vi": m.schema == "VI"}
               for m in ModeleOperation.objects.select_related("tiers")}
    return render(request, "compta/saisie.html", {"form": form, "resultat": resultat, "numero": numero, "piece": piece,
                                                  "modeles": modeles, "releve": releve})


def _journaliser(request, action, objet, avant="", apres=""):
    Modification.objects.create(auteur=request.user.get_username(), action=action, objet=objet, avant=avant, apres=apres)


@login_required
@permission_required("compta.add_codeanalytique", raise_exception=True)
def codes(request):
    code_form = CodeForm(request.POST if "creer_code" in request.POST else None, prefix="code",
                         axe1=request.user.has_perm("compta.parametrer"))
    membre_form = MembreForm(request.POST if "creer_membre" in request.POST else None, prefix="membre")
    statut_form = StatutForm(request.POST if "changer_statut" in request.POST else None, prefix="statut")
    if request.method == "POST":
        with transaction.atomic():
            if "creer_code" in request.POST and code_form.is_valid():
                c = code_form.cleaned_data
                p, lib = c["prefixe"], c["libelle"].strip().upper()
                if CodeAnalytique.objects.filter(axe=p.axe, libelle=lib).exists():
                    code_form.add_error("libelle", "Ce libellé existe déjà dans cet axe.")
                else:
                    nouveau = CodeAnalytique.objects.create(code=p.code_suivant(), axe=p.axe, libelle=lib,
                                                            statut=c["statut"] if p.axe == 2 else 1)
                    _journaliser(request, "Création", f"code axe {p.axe} {nouveau.code}", apres=lib)
                    messages.success(request, f"Code {nouveau.code} créé : {lib}.")
                    return redirect("codes")
            if "creer_membre" in request.POST and membre_form.is_valid():
                c = membre_form.cleaned_data
                libelle = f"{c['nom'].strip().upper()} {c['prenom'].strip().upper()}".strip()
                existant = Compte.objects.filter(libelle=libelle, numero__startswith=c["type"].prefixe).first()
                if existant:
                    membre_form.add_error("nom", f"Un compte existe déjà à ce nom : {existant.numero}.")
                else:
                    from .membres import creer_tiers
                    compte = creer_tiers(c["type"], c["nom"], c["prenom"], **{k: c[k] for k in (
                        "adresse", "code_postal", "ville", "telephone", "email")})
                    _journaliser(request, "Création", f"compte {compte.numero}", apres=f"{c['type']} {libelle}")
                    messages.success(request, f"{c['type']} créé : {compte.numero} – {libelle}.")
                    return redirect("codes")
            if "changer_statut" in request.POST and statut_form.is_valid():
                c = statut_form.cleaned_data
                code, nouveau = c["code"], c["statut"]
                utilise = code.lignes.count()
                if nouveau == code.statut:
                    statut_form.add_error("statut", "C'est déjà le statut de ce code.")
                elif utilise and not c["confirmation"]:
                    statut_form.add_error("confirmation", f"Code utilisé dans {utilise} ligne(s) : cocher la confirmation.")
                else:
                    avant = code.get_statut_display()
                    code.statut = nouveau
                    code.save()
                    _journaliser(request, "Statut", f"code {code.code}", avant=avant, apres=code.get_statut_display())
                    messages.success(request, f"{code.code} : {avant} → {code.get_statut_display()}.")
                    return redirect("codes")
    prefixes = [(p, p.code_suivant()) for p in Prefixe.objects.all()]
    return render(request, "compta/codes.html", {"code_form": code_form, "membre_form": membre_form, "statut_form": statut_form,
                                                 "prefixes": prefixes})
