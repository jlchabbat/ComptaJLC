"""« Excel » sur toutes les pages : ?export=xlsx renvoie les tableaux de la page dans un classeur (un onglet par tableau,
nommé d'après le titre qui le précède). Les montants « 1 234,50 » / « −50,00 » deviennent des nombres, les dates
jj/mm/aaaa des dates. Une copie est écrite dans le dossier Exports (comme les autres exports)."""

import datetime as dt
import re
from html.parser import HTMLParser

import openpyxl

from .export import ENTETE, FOND, MONTANT

MONTANT_TEXTE = re.compile(r"^[−-]?\d{1,3}(?:[\s  ]\d{3})*,\d{2}$")
DATE_TEXTE = re.compile(r"^(\d{2})/(\d{2})/(\d{4})$")
IGNORES = {"script", "style", "select", "option", "button", "textarea"}


class Tableaux(HTMLParser):
    """Tableaux de la page : [(titre, lignes [[(texte, est_entete)]])] ; titre = dernier h1/h2/h3 vu avant le tableau."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tableaux, self.titre, self.page = [], "", ""
        self._pile, self._table, self._ligne, self._cellule, self._titre, self._ignore, self._span = [], None, None, None, None, 0, 1
        self._dans_title = False

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in IGNORES:
            self._ignore += 1
        elif tag == "title":
            self._dans_title = True
        elif tag in ("h1", "h2", "h3") and self._table is None:
            self._titre = []
        elif tag == "table":
            if self._table is not None:                   # tableau imbriqué : ses cellules comptent dans le parent
                self._pile.append(None)
                return
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._ligne = []
        elif tag in ("td", "th") and self._ligne is not None:
            self._cellule, self._entete = [], tag == "th"
            try:
                self._span = max(1, int(a.get("colspan") or 1))
            except ValueError:
                self._span = 1
        elif tag == "br" and self._cellule is not None:
            self._cellule.append(" ")

    def handle_endtag(self, tag):
        if tag in IGNORES:
            self._ignore = max(0, self._ignore - 1)
        elif tag == "title":
            self._dans_title = False
        elif tag in ("h1", "h2", "h3") and self._titre is not None:
            self.titre = " ".join("".join(self._titre).split())
            self._titre = None
        elif tag in ("td", "th") and self._cellule is not None:
            texte = " ".join("".join(self._cellule).split())
            self._ligne.append((texte, self._entete))
            self._ligne.extend([("", self._entete)] * (self._span - 1))
            self._cellule = None
        elif tag == "tr" and self._ligne is not None:
            if any(t for t, _ in self._ligne):
                self._table.append(self._ligne)
            self._ligne = None
        elif tag == "table":
            if self._pile:
                self._pile.pop()
                return
            if self._table:
                self.tableaux.append((self.titre, self._table))
            self._table = None

    def handle_data(self, data):
        if self._dans_title:
            self.page += data
        if self._ignore:
            return
        if self._cellule is not None:
            self._cellule.append(data)
        elif self._titre is not None:
            self._titre.append(data)


def valeur(texte):
    t = texte.strip()
    if MONTANT_TEXTE.match(t):
        return float(re.sub(r"[\s  ]", "", t).replace("−", "-").replace(",", ".")), True
    d = DATE_TEXTE.match(t)
    if d:
        try:
            return dt.date(int(d[3]), int(d[2]), int(d[1])), False
        except ValueError:
            pass
    return t, False


def nom_onglet(titre, pris):
    base = re.sub(r"[\[\]:*?/\\]", "-", titre or "Tableau")[:28].strip() or "Tableau"
    nom, i = base, 2
    while nom.lower() in pris:
        nom, i = f"{base[:25]} ({i})", i + 1
    pris.add(nom.lower())
    return nom


def classeur(html):
    """Classeur des tableaux d'une page HTML ; None si la page n'a pas de tableau."""
    p = Tableaux()
    p.feed(html)
    if not p.tableaux:
        return None, ""
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    pris = set()
    for titre, lignes in p.tableaux:
        ws = wb.create_sheet(nom_onglet(titre, pris))
        largeurs = {}
        for r, ligne in enumerate(lignes, 1):
            for c, (texte, entete) in enumerate(ligne, 1):
                v, montant = valeur(texte)
                cell = ws.cell(r, c, v if v != "" else None)
                if entete:
                    cell.font, cell.fill = ENTETE, FOND
                elif montant:
                    cell.number_format = MONTANT
                elif isinstance(v, dt.date):
                    cell.number_format = "DD/MM/YYYY"
                largeurs[c] = max(largeurs.get(c, 8), min(50, len(texte) + 2))
        for c, l in largeurs.items():
            ws.column_dimensions[get_column_letter(c)].width = l
        if lignes and all(e for _, e in lignes[0]):
            ws.freeze_panes = "A2"
    page = p.page.split("–")[0].strip() or "ComptaBB"
    return wb, page


class ExportExcelMiddleware:
    """?export=xlsx sur une page HTML : ses tableaux en Excel (réservé aux utilisateurs connectés)."""

    def __init__(self, get_response):
        self.get_response = get_response

    def __call__(self, request):
        response = self.get_response(request)
        if (request.method != "GET" or request.GET.get("export") != "xlsx" or response.status_code != 200
                or not request.user.is_authenticated or getattr(response, "streaming", False)
                or not response.get("Content-Type", "").startswith("text/html")):
            return response
        wb, page = classeur(response.content.decode(response.charset or "utf-8"))
        if wb is None:
            return response
        from .vues_journaux import reponse_excel
        nom = re.sub(r"[^\w-]+", "_", page).strip("_") or "ComptaBB"
        return reponse_excel(wb, f"{nom}_{dt.datetime.now():%Y-%m-%d_%H%M%S}.xlsx")
