"""Édition des journaux : journal de trésorerie (banque, caisse) avec solde progressif, autres journaux par
mouvement ; impression et export Excel. Historique des opérations exportable."""

import datetime as dt
import io

from django.contrib.auth.decorators import login_required, permission_required
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render

from . import etats
from .export import feuille
from .models import ZERO, Journal, Ligne, Modification, Mouvement
from .views import periode

consulter = permission_required("compta.view_mouvement", raise_exception=True)


def reponse_excel(wb, nom):
    from .dossiers import copier_export
    tampon = io.BytesIO()
    wb.save(tampon)
    copier_export(nom, tampon.getvalue())
    return HttpResponse(tampon.getvalue(), headers={
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "Content-Disposition": f'attachment; filename="{nom}"'})


def journal_tresorerie(journal, debut, fin):
    """Lignes du compte de trésorerie du journal : recette, dépense, contrepartie, solde progressif."""
    ouverture = etats.solde_cumule(journal.compte, debut - dt.timedelta(days=1))
    lignes = (Ligne.objects.filter(compte=journal.compte, mouvement__date__range=(debut, fin))
              .select_related("mouvement", "mouvement__journal", "anal2")
              .prefetch_related("mouvement__lignes__compte")
              .order_by("mouvement__date", "mouvement__numero", "ordre"))
    o_fin = etats.origine(fin)
    if o_fin and o_fin > debut:
        lignes = lignes.exclude(mouvement__origine="cloture")
    res, solde = [], ouverture
    for l in lignes:
        solde += l.debit - l.credit
        contre = sorted({x.compte.numero + " " + x.compte.libelle for x in l.mouvement.lignes.all() if x.compte_id != l.compte_id})
        res.append({"l": l, "contrepartie": " · ".join(contre)[:120], "solde": solde})
    return ouverture, res


@login_required
@consulter
def journaux(request):
    debut, fin = periode(request)
    tous = list(Journal.objects.filter(actif=True))
    code = request.GET.get("journal") or (tous[0].code if tous else "")
    journal = next((j for j in tous if j.code == code), None)
    ctx = {"journaux": tous, "journal": journal, "debut": debut, "fin": fin}
    if journal and journal.compte_id:
        ouverture, lignes = journal_tresorerie(journal, debut, fin)
        ctx.update(tresorerie=True, ouverture=ouverture, lignes=lignes,
                   recettes=sum((x["l"].debit for x in lignes), ZERO), depenses=sum((x["l"].credit for x in lignes), ZERO),
                   cloture=lignes[-1]["solde"] if lignes else ouverture)
    elif journal:
        mvts = (Mouvement.objects.filter(journal=journal, date__range=(debut, fin)).order_by("date", "numero")
                .prefetch_related("lignes__compte", "lignes__anal2"))
        ctx.update(tresorerie=False, mouvements=mvts,
                   total_debit=sum((l.debit for m in mvts for l in m.lignes.all()), ZERO),
                   total_credit=sum((l.credit for m in mvts for l in m.lignes.all()), ZERO))
    if request.GET.get("format") == "xlsx" and journal:
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        titre = f"Journal {journal.code}"
        if ctx["tresorerie"]:
            rangs = [["", "", "", f"Solde au {debut:%d/%m/%Y}", "", "", ctx["ouverture"]]]
            rangs += [[x["l"].mouvement.date, x["l"].mouvement.numero, x["l"].libelle, x["contrepartie"],
                       x["l"].debit, x["l"].credit, x["solde"]] for x in ctx["lignes"]]
            rangs.append(["", "", "", "Totaux et solde final", ctx["recettes"], ctx["depenses"], ctx["cloture"]])
            feuille(wb, titre, ["Date", "Mvt", "Libellé", "Contrepartie", "Recette", "Dépense", "Solde"], rangs, (5, 6, 7))
        else:
            rangs = [[m.date, m.numero, l.compte_id, l.compte.libelle, l.libelle, l.debit, l.credit, l.anal2_id]
                     for m in ctx["mouvements"] for l in m.lignes.all()]
            rangs.append(["", "", "", "", "Totaux", ctx["total_debit"], ctx["total_credit"], ""])
            feuille(wb, titre, ["Date", "Mvt", "Compte", "Intitulé", "Libellé", "Débit", "Crédit", "Axe 2"], rangs, (6, 7))
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=2):
                if hasattr(row[0].value, "year"):
                    row[0].number_format = "DD/MM/YYYY"
        return reponse_excel(wb, f"Journal_{journal.code}_{debut:%Y-%m-%d}_{fin:%Y-%m-%d}.xlsx")
    return render(request, "compta/journaux.html", ctx)


@login_required
@consulter
def historique_excel(request):
    import openpyxl
    q = request.GET.get("q", "").strip()
    qs = Modification.objects.all()
    if q:
        qs = qs.filter(Q(auteur__icontains=q) | Q(action__icontains=q) | Q(objet__icontains=q) | Q(lot__icontains=q))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    feuille(wb, "Historique", ["Date", "Auteur", "Lot", "Action", "Objet", "Avant", "Après"],
            [[m.date.replace(tzinfo=None), m.auteur, m.lot, m.action, m.objet, m.avant, m.apres] for m in qs])
    for row in wb["Historique"].iter_rows(min_row=2):
        row[0].number_format = "DD/MM/YYYY HH:MM"
    return reponse_excel(wb, "ComptaJLC_historique.xlsx")


@login_required
def effacer_historique(request):
    """Efface l'historique jusqu'à une date (administrateur) ; copie Excel et sauvegarde de la base faites avant."""
    from django.contrib import messages
    from django.core.exceptions import PermissionDenied
    from django.shortcuts import redirect

    from . import base_donnees as bd
    if not request.user.is_superuser:
        raise PermissionDenied
    if request.method != "POST":
        return redirect("modifications")
    if request.POST.get("confirmation", "").strip().upper() != "EFFACER":
        messages.error(request, "Historique non effacé : tapez EFFACER pour confirmer.")
        return redirect("modifications")
    try:
        jusquau = dt.date.fromisoformat(request.POST.get("jusquau", ""))
    except ValueError:
        jusquau = dt.date.today()
    qs = Modification.objects.filter(date__date__lte=jusquau)
    n = qs.count()
    if not n:
        messages.info(request, "Aucune opération à effacer jusqu'à cette date.")
        return redirect("modifications")
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    feuille(wb, "Historique", ["Date", "Auteur", "Lot", "Action", "Objet", "Avant", "Après"],
            [[m.date.replace(tzinfo=None), m.auteur, m.lot, m.action, m.objet, m.avant, m.apres] for m in qs])
    from .dossiers import archives as dossier_archives
    archives = dossier_archives()
    copie = archives / f"Historique_jusqu_au_{jusquau:%Y-%m-%d}_efface_le_{dt.date.today():%Y-%m-%d}.xlsx"
    wb.save(copie)
    sauvegarde = bd.sauvegarder("avant-effacement-historique")
    qs.delete()
    Modification.objects.create(auteur=request.user.get_username(), lot="Base de données", action="Historique effacé",
                                objet=f"{n} opération(s) jusqu'au {jusquau:%d/%m/%Y}",
                                apres=f"copie {copie.name} ; sauvegarde {sauvegarde.name}"[:300])
    messages.success(request, f"{n} opération(s) effacée(s). Copie Excel gardée : {copie.name} (Administration › Base de données › Archives).")
    return redirect("modifications")
