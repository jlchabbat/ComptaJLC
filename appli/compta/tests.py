"""Tests du socle : modèle, contrôles, droits d'accès, reprise d'un classeur.

    python manage.py test compta
"""

import datetime as dt
import io
import tempfile
import zipfile
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.db import IntegrityError
from django.test import TestCase, TransactionTestCase, override_settings
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from . import controles, reglages
from .models import CodeAnalytique, Compte, Exercice, Journal, Ligne, Mouvement, Prefixe, Reglage, soldes

D = Decimal
# Les tests décrivent le site de la Loge : un réglage absent y prend la valeur de la Loge (Mizrahi, Bit, Isracard…).
# La classe ReglagesAssociation vérifie les valeurs neutres d'une nouvelle association.
NEUTRES = dict(reglages.NEUTRES)
reglages.NEUTRES.update({cle: loge for cle, _, _, loge in reglages.REGLAGES})


def referentiels():
    CodeAnalytique.objects.create(code="BIL.5", axe=1, libelle="TRESORERIE")
    CodeAnalytique.objects.create(code="COT.2", axe=1, libelle="COTISATIONS")
    CodeAnalytique.objects.create(code="BIL.4", axe=1, libelle="TIERS")
    CodeAnalytique.objects.create(code="GEN.002", axe=2, libelle="COTISATIONS")
    CodeAnalytique.objects.create(code="MAN.001", axe=2, libelle="RALLYE")
    CodeAnalytique.objects.create(code="MAN.007", axe=2, libelle="")
    Compte.objects.create(numero="512000", libelle="BANQUE", anal1_id="BIL.5")
    Compte.objects.create(numero="700000", libelle="COTISATIONS", anal1_id="COT.2")
    Compte.objects.create(numero="411TEST001", libelle="MEMBRE TEST", anal1_id="BIL.4")
    Journal.objects.create(code="B1", intitule="BANQUE 1", compte_id="512000")
    Exercice.objects.create(libelle="2025", debut=dt.date(2025, 10, 1), fin=dt.date(2025, 12, 31), clos=True)
    Exercice.objects.create(libelle="2026", debut=dt.date(2026, 1, 1), fin=dt.date(2026, 12, 31))


def mouvement(numero, date, lignes, origine="saisie"):
    m = Mouvement.objects.create(numero=numero, date=date, journal_id="B1", piece=numero, origine=origine)
    for i, (compte, d, c) in enumerate(lignes):
        Ligne.objects.create(mouvement=m, ordre=i, compte_id=compte, libelle="TEST", debit=D(d), credit=D(c), anal2_id="GEN.002")
    return m


class Modele(TestCase):
    def setUp(self):
        referentiels()

    def test_soldes_et_resultat(self):
        mouvement(1, dt.date(2026, 2, 1), [("411TEST001", 400, 0), ("700000", 0, 400), ("512000", 400, 0), ("411TEST001", 0, 400)])
        self.assertEqual(soldes(Ligne.objects.all())[:2], (D("800.00"), D("800.00")))
        self.assertEqual(soldes(Ligne.objects.filter(compte_id="512000"))[2], D("400.00"))
        self.assertEqual(soldes(Ligne.objects.filter(compte__numero__startswith="7"))[2], D("-400.00"))

    def test_debit_et_credit_refuses(self):
        m = mouvement(1, dt.date(2026, 2, 1), [])
        with self.assertRaises(IntegrityError):
            Ligne.objects.create(mouvement=m, compte_id="512000", libelle="X", debit=D(1), credit=D(1), anal2_id="GEN.002")

    def test_code_suivant(self):
        Prefixe.objects.create(prefixe="MAN.", axe=2)
        Prefixe.objects.create(prefixe="COT.", axe=1)
        Prefixe.objects.create(prefixe="PJT.", axe=2)
        self.assertEqual(Prefixe.objects.get(prefixe="MAN.").code_suivant(), "MAN.008")
        self.assertEqual(Prefixe.objects.get(prefixe="COT.").code_suivant(), "COT.3")
        self.assertEqual(Prefixe.objects.get(prefixe="PJT.").code_suivant(), "PJT.001")

    def test_numerotation(self):
        mouvement(41, dt.date(2026, 2, 1), [("512000", 1, 0), ("700000", 0, 1)])
        self.assertEqual((Mouvement.prochain_numero(), Mouvement.prochaine_piece()), (42, 42))


class Controles(TestCase):
    def setUp(self):
        referentiels()
        Reglage.objects.create(cle="compte_virement", valeur="512000")

    def statut(self, libelle):
        return next(r for r in controles.executer() if r.libelle.startswith(libelle)).statut

    def test_equilibre(self):
        mouvement(1, dt.date(2026, 2, 1), [("512000", 100, 0), ("700000", 0, 100)])
        self.assertEqual(self.statut("Mouvements déséquilibrés"), "OK")
        mouvement(2, dt.date(2026, 2, 2), [("512000", 100, 0), ("700000", 0, 90)])
        self.assertEqual(self.statut("Mouvements déséquilibrés"), "Anomalie")
        self.assertEqual(controles.etat_general()[0], "2 anomalie(s)")

    def test_periode_close(self):
        mouvement(1, dt.date(2025, 11, 1), [("512000", 100, 0), ("700000", 0, 100)], origine="import")
        self.assertEqual(self.statut("Écritures nouvelles dans un exercice clos"), "OK")
        mouvement(2, dt.date(2025, 11, 2), [("512000", 100, 0), ("700000", 0, 100)])
        self.assertEqual(self.statut("Écritures nouvelles dans un exercice clos"), "Anomalie")

    def test_compte_de_liaison(self):
        mouvement(1, dt.date(2026, 2, 1), [("512000", 100, 0), ("700000", 0, 100)])
        self.assertEqual(self.statut("Compte de virement interne"), "À vérifier")


class Droits(TestCase):
    def setUp(self):
        referentiels()
        call_command("migrate", verbosity=0)   # crée les rôles (post_migrate)

    def client_pour(self, role):
        u = User.objects.create_user(role, password="motdepasse-test")
        u.groups.add(Group.objects.get(name=role))
        self.client.force_login(u)

    def test_pages_bureau(self):
        mouvement(1, dt.date(2026, 2, 1), [("512000", 100, 0), ("700000", 0, 100)])
        self.client_pour("Bureau")
        for url in ("/", "/ecritures/", "/mouvement/1/", "/grand-livre/?compte=512000", "/balance/", "/analytique/",
                    "/controles/", "/modifications/"):
            self.assertEqual(self.client.get(url).status_code, 200, url)

    def test_benevole_refuse(self):
        self.client_pour("Bénévole")
        self.assertRedirects(self.client.get("/"), "/fiches/")
        self.assertEqual(self.client.get("/ecritures/").status_code, 403)

    def test_anonyme_redirige(self):
        self.assertEqual(self.client.get("/").status_code, 302)


class Reprise(TestCase):
    def classeur(self):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        def table(nom, entetes, lignes):
            ws = wb.create_sheet(nom)
            ws.append(entetes)
            for l in lignes:
                ws.append(l)
            ws.add_table(Table(displayName=nom, ref=f"A1:{openpyxl.utils.get_column_letter(len(entetes))}{len(lignes) + 1}"))
            return ws

        table("T_Axe1", ["Code", "Libellé", "Actif"], [["BIL.5", "TRESORERIE", 1], ["COT.2", "COTIS", 1]])
        table("T_Axe2", ["Code", "Libellé", "Actif"], [["GEN.002", "COTISATIONS", 1]])
        table("T_Prefixes", ["Préfixe", "Axe", "Libellé", "Code suivant"], [["GEN.", 2, "General", "GEN.003"]])
        table("T_PlanComptable", ["Compte", "Libellé compte", "Axe 1 (Anal1)", "Lettrable"],
              [["512000", "BANQUE", "BIL.5", ""], ["700000", "COTISATIONS", "COT.2", ""]])
        table("T_Journaux", ["Code", "Intitulé", "Type", "N° de compte", "Actif"], [["B1", "BANQUE 1", "BQ", 512000, 1]])
        e = ["Date", "Jnl", "Mvt", "Pièce", "Compte", "Libellé", "Débit", "Crédit", "Anal2", "Let"]
        table("T_Ecritures", e, [[dt.datetime(2026, 2, 1), "B1", 1, 1, "512000", "COT", 50.5, 0, "GEN.002", None],
                                 [dt.datetime(2026, 2, 1), "B1", 1, 1, "700000", "COT", 0, 50.5, "GEN.002", None],
                                 [dt.datetime(2025, 11, 1), "B1", 2, 2, "512000", "COT", 10, 0, "GEN.002", None],
                                 [dt.datetime(2025, 11, 1), "B1", 2, 2, "700000", "COT", 0, 10, "GEN.002", None]])
        p = wb.create_sheet("Paramètres")
        for r, (n, v) in enumerate((("P_DebutExercice", dt.datetime(2026, 1, 1)), ("P_FinExercice", dt.datetime(2026, 12, 31)),
                                    ("P_DateCloture", dt.datetime(2025, 12, 31)), ("P_CompteVirement", "580000")), start=1):
            p.cell(r, 2, v)
            wb.defined_names[n] = DefinedName(n, attr_text=f"Paramètres!$B${r}")
        chemin = Path(tempfile.mkdtemp()) / "test.xlsx"
        wb.save(chemin)
        return chemin

    def test_reprise(self):
        call_command("importer_classeur", str(self.classeur()), stdout=open("/dev/null", "w"))
        self.assertEqual(Mouvement.objects.count(), 2)
        self.assertEqual(soldes(Ligne.objects.all())[:2], (D("60.50"), D("60.50")))
        self.assertEqual(Journal.objects.get(code="B1").compte_id, "512000")
        self.assertEqual(list(Exercice.objects.values_list("clos", flat=True)), [True, False])
        self.assertEqual(Reglage.lire("compte_virement"), "580000")
        self.assertEqual(controles.etat_general()[0], "OK")


# ---------------------------------------------------------------- W1 : saisie guidée (scénarios de la recette Excel du Lot 1)

from . import saisie as moteur  # noqa: E402
from .models import ModeleOperation, MoyenPaiement  # noqa: E402


def referentiels_saisie():
    for code, lib in (("BIL.5", "TRESORERIE"), ("BIL.4", "TIERS"), ("COT.2", "COTISATIONS"), ("FON.1", "FONCTIONNEMENT")):
        CodeAnalytique.objects.create(code=code, axe=1, libelle=lib)
    for code, lib in (("GEN.001", "GENERAL"), ("GEN.002", "COTISATIONS"), ("GEN.004", "BANQUE"), ("MAN.001", "RALLYE"),
                      ("SOC.006", "BOURSES")):
        CodeAnalytique.objects.create(code=code, axe=2, libelle=lib)
    comptes = [("512000", "MIZRAHI", "BIL.5"), ("512100", "MIZRAHI EPARGNE", "BIL.5"), ("512200", "BIT", "BIL.5"),
               ("530000", "CAISSE", "BIL.5"), ("580000", "VIREMENTS INTERNES", "BIL.5"), ("411TAIEB001", "TAIEB JEANNE", "BIL.4"),
               ("401000", "FOURNIS DIVERS", "BIL.4"), ("600000", "ACHATS", "FON.1"), ("600100", "FRAIS BANCAIRES", "FON.1"),
               ("610000", "SERVICES", "FON.1")]
    comptes += [(n, "PRODUIT", "COT.2") for n in ("700000", "710000", "720000", "725000", "740000", "750000")]
    for n, lib, a1 in comptes:
        Compte.objects.create(numero=n, libelle=lib, anal1_id=a1)
    for code, compte in (("B1", "512000"), ("B2", "512100"), ("B3", "512200"), ("CA", "530000"), ("VT", None), ("HA", None), ("OD", None)):
        Journal.objects.create(code=code, intitule=code, compte_id=compte)
    Exercice.objects.create(libelle="court", debut=dt.date(2025, 10, 25), fin=dt.date(2025, 12, 31), clos=True)
    Exercice.objects.create(libelle="2026", debut=dt.date(2026, 1, 1), fin=dt.date(2026, 12, 31))
    Reglage.objects.create(cle="compte_virement", valeur="580000")
    moteur.initialiser_parametres()
    m = Mouvement.objects.create(numero=421, date=dt.date(2026, 9, 24), journal_id="B1", piece=421, origine="import")
    Ligne.objects.create(mouvement=m, ordre=0, compte_id="411TAIEB001", libelle="FACTURE 40013 - TAIEB JEANNE", debit=D(400),
                         anal2_id="MAN.001")
    Ligne.objects.create(mouvement=m, ordre=1, compte_id="710000", libelle="FACTURE 40013 - TAIEB JEANNE", credit=D(400),
                         anal2_id="MAN.001")


def operation(date=dt.date(2026, 10, 15), modele=None, tiers=None, montant=400, paiement=None, vers=None, anal2=None, compte=None,
              remboursement=False, libelle=""):
    mp = lambda lib: MoyenPaiement.objects.get(libelle=lib) if lib else None  # noqa: E731
    return moteur.Operation(date=date, modele=ModeleOperation.objects.get(type=modele) if modele else None,
                            tiers=Compte.objects.get(numero=tiers) if tiers else None, montant=D(montant), paiement=mp(paiement),
                            vers=mp(vers), anal2=CodeAnalytique.objects.get(code=anal2) if anal2 else None,
                            compte=Compte.objects.get(numero=compte) if compte else None, remboursement=remboursement,
                            libelle=libelle)


SCENARIOS = {
    "cotisation_reglee": (dict(modele="Cotisation membre", tiers="411TAIEB001", paiement="BIT", anal2="GEN.002"),
                          [("B3", 1, "411TAIEB001", 400, 0), ("B3", 1, "700000", 0, 400), ("B3", 1, "512200", 400, 0),
                           ("B3", 1, "411TAIEB001", 0, 400)], set()),
    "cotisation_non_reglee": (dict(modele="Cotisation membre", tiers="411TAIEB001", paiement="Non réglé", anal2="GEN.002"),
                              [("VT", 1, "411TAIEB001", 400, 0), ("VT", 1, "700000", 0, 400)], set()),
    "frais_bancaires": (dict(modele="Frais bancaires", montant=25, paiement="Mizrahi compte courant", anal2="GEN.004"),
                        [("B1", 1, "600100", 25, 0), ("B1", 1, "512000", 0, 25)], set()),
    "virement_interne": (dict(modele="Virement interne", montant=1000, paiement="Mizrahi compte courant", vers="BIT", anal2="GEN.001"),
                         [("B1", 1, "580000", 1000, 0), ("B1", 1, "512000", 0, 1000), ("B3", 2, "512200", 1000, 0),
                          ("B3", 2, "580000", 0, 1000)], set()),
    "isracard": (dict(modele="Paiement carte Isracard", montant=500, paiement="Mizrahi compte courant", anal2="GEN.004"),
                 [("OD", 1, "600000", 500, 0), ("OD", 1, "580000", 0, 500), ("B1", 2, "580000", 500, 0),
                  ("B1", 2, "512000", 0, 500)], set()),
    "remboursement": (dict(modele="Cotisation membre", tiers="411TAIEB001", paiement="BIT", anal2="GEN.002", remboursement=True),
                      [("B3", 1, "411TAIEB001", 0, 400), ("B3", 1, "700000", 400, 0), ("B3", 1, "512200", 0, 400),
                       ("B3", 1, "411TAIEB001", 400, 0)], set()),
    "facture_fournisseur_caisse": (dict(modele="Facture fournisseur", tiers="401000", paiement="Caisse (espèces)", anal2="MAN.001",
                                        compte="610000"),
                                   [("CA", 1, "610000", 400, 0), ("CA", 1, "401000", 0, 400), ("CA", 1, "401000", 400, 0),
                                    ("CA", 1, "530000", 0, 400)], set()),
    "fournisseur_sans_compte": (dict(modele="Facture fournisseur", tiers="401000", paiement="Caisse (espèces)", anal2="MAN.001"),
                                None, {"Compte"}),
    "date_close_et_mauvais_tiers": (dict(date=dt.date(2025, 12, 15), modele="Cotisation membre", tiers="401000", paiement="BIT",
                                         anal2="GEN.002"), None, {"Date", "Tiers"}),
    "depense_classe_7": (dict(modele="Dépense directe", paiement="BIT", anal2="GEN.001", compte="700000"), None, {"Compte"}),
    "paiement_manquant": (dict(modele="Frais bancaires", anal2="GEN.004"), None, {"Moyen de paiement"}),
    "virement_meme_compte": (dict(modele="Virement interne", paiement="BIT", vers="BIT", anal2="GEN.001"), None, {"Virement"}),
    "deja_enregistree": (dict(date=dt.date(2026, 9, 24), modele="Facture manifestation (membre)", tiers="411TAIEB001",
                              paiement="Mizrahi compte courant", anal2="MAN.001", libelle="FACTURE 40013 - TAIEB JEANNE"),
                         None, {"Déjà enregistrée ?"}),
}


class Saisie(TestCase):
    def setUp(self):
        referentiels_saisie()

    def test_scenarios(self):
        for nom, (entrees, attendu, erreurs) in SCENARIOS.items():
            with self.subTest(nom):
                r = moteur.controler(operation(**entrees))
                self.assertEqual(set(r.erreurs), erreurs, r.erreurs)
                if attendu:
                    obtenu = [(l.journal.code, l.mvt, l.compte.numero, l.debit, l.credit) for l in r.lignes]
                    self.assertEqual(obtenu, [(j, m, c, D(d), D(cr)) for j, m, c, d, cr in attendu])

    def test_enregistrer(self):
        u = User.objects.create_user("t")
        crees = moteur.enregistrer(operation(**SCENARIOS["virement_interne"][0]), u)
        self.assertEqual([(m.numero, m.piece, m.journal_id) for m in crees], [(422, 422, "B1"), (423, 423, "B3")])
        self.assertEqual(soldes(Ligne.objects.filter(mouvement__numero__gte=422))[:2], (D(2000), D(2000)))
        self.assertEqual(Ligne.objects.filter(mouvement__numero=422).first().libelle, "VIREMENT INTERNE")
        with self.assertRaises(ValueError):   # même opération : doublon
            moteur.enregistrer(operation(**SCENARIOS["virement_interne"][0]), u)
        self.assertEqual(len(moteur.enregistrer(operation(**SCENARIOS["virement_interne"][0]), u, forcer_doublon=True)), 2)

    def test_initialiser_idempotent(self):
        n = ModeleOperation.objects.count()
        moteur.initialiser_parametres()
        self.assertEqual((n, ModeleOperation.objects.count()), (14, 14))



class ReglagesAssociation(TestCase):
    """Un site par association (étape 1) : réglages propres à l'association, neutres sur une base neuve."""

    def test_migration_garde_les_valeurs_du_site_existant(self):
        from django.apps import apps
        creer = __import__("importlib").import_module("compta.migrations.0016_reglages_association").creer
        creer(apps, None)                                             # base neuve : rien d'enregistré
        self.assertFalse(Reglage.objects.filter(cle__in=[c for c, *_ in reglages.REGLAGES]).exists())
        Journal.objects.create(code="B1", intitule="Mizrahi")
        Reglage.objects.create(cle="association", valeur="Loge de Jérusalem")
        creer(apps, None)                                             # site existant : ses valeurs actuelles
        self.assertEqual([Reglage.lire(c) for c in ("nom_association", "devise", "releves_mizrahi", "releve_bit",
                                                    "carte_bancaire", "hebergeur_liens", "traductions_releve")],
                         ["Loge de Jérusalem", "₪", "B1,B2", "B3", "Isracard", "SUMIT", "oui"])

    def test_nouvelle_association_sans_mizrahi_bit_ni_isracard(self):
        from unittest import mock

        from .models import LigneReleve, ModeFiche
        with mock.patch.dict(reglages.NEUTRES, NEUTRES):
            referentiels_saisie()
            self.assertFalse(ModeleOperation.objects.filter(type__icontains="carte").exists())
            self.assertEqual(ModeleOperation.objects.count(), 13)
            moyens = set(MoyenPaiement.objects.values_list("libelle", flat=True))
            self.assertEqual(moyens, {"B1", "B2", "Caisse (espèces)", "Non réglé"})      # nom du journal, pas « Mizrahi »
            modes = set(ModeFiche.objects.values_list("libelle", flat=True))
            self.assertIn("Virement bancaire", modes)
            self.assertFalse({"Bit", "Virement BIT", "Virement Mizrahi"} & modes or any("Carte" in m for m in modes))
            self.assertEqual(reglages.montant(D("1234.5")), "1 234,50 €")
            with self.assertRaises(ech.Refus):
                ech.imp_bit([], "Bit.xlsx")
            l = LigneReleve.objects.create(journal_id="B1", date=dt.date(2026, 1, 5), rang=1, operation="VIREMENT RECU",
                                           montant=D(10), solde=D(10))
            self.assertEqual(l.traduction, "VIREMENT RECU")                       # pas de traduction : libellé tel quel
            Reglage.objects.create(cle="devise", valeur="$")
            self.assertEqual(reglages.montant(D(-3)), "-3,00 $")

    def test_nom_et_devise_affiches(self):
        referentiels_saisie()
        call_command("migrate", verbosity=0)
        Reglage.objects.update_or_create(cle="nom_association", defaults={"valeur": "Loge de Jérusalem"})
        u = User.objects.create_user("tresorier", password="x")
        u.groups.add(Group.objects.get(name="Trésorier"))
        self.client.force_login(u)
        r = self.client.get("/?du=2026-01-01&au=2026-12-31")
        self.assertContains(r, "Loge de Jérusalem")
        self.assertContains(r, " ₪</div>")
        self.assertContains(r, "Traductions du relevé")
        Reglage.objects.create(cle="traductions_releve", valeur="non")
        self.assertNotContains(self.client.get("/"), "Traductions du relevé")



class PremierDemarrage(TestCase):
    """Site neuf : premier administrateur (code d'installation), puis assistant de démarrage."""

    def setUp(self):
        call_command("migrate", verbosity=0)                   # rôles
        from . import demarrage
        demarrage.chemin_code().unlink(missing_ok=True)

    def administrateur(self):
        from . import demarrage
        self.assertRedirects(self.client.get("/ecritures/"), "/demarrage/compte/")
        code = demarrage.code_installation()
        r = self.client.post("/demarrage/compte/", {"code": "FAUX", "identifiant": "tresorerie@asso.org",
                                                    "mot_de_passe": "une phrase longue 2026", "confirmation": "une phrase longue 2026"})
        self.assertContains(r, "installation incorrect")
        r = self.client.post("/demarrage/compte/", {"code": code.lower(), "identifiant": "tresorerie@asso.org",
                                                    "mot_de_passe": "une phrase longue 2026", "confirmation": "une phrase longue 2026"})
        self.assertRedirects(r, "/demarrage/", fetch_redirect_response=False)
        u = User.objects.get(username="tresorerie@asso.org")
        self.assertTrue(u.is_superuser and u.groups.filter(name="Administrateur").exists())
        self.assertFalse(demarrage.chemin_code().exists())
        self.assertEqual(demarrage.code_installation(), "")         # plus de code une fois l'administrateur créé
        self.assertRedirects(self.client.get("/"), "/demarrage/")
        self.assertEqual(self.client.get("/base/").status_code, 200)   # recharger une sauvegarde reste possible

    def test_plan_de_base(self):
        from .models import ModeFiche, MoyenPaiement
        self.administrateur()
        r = self.client.post("/demarrage/", {"nom": "Amis du musée", "devise": "€", "plan": "base", "debut": "2027-01-01",
                                             "fin": "2027-12-31", "banque1": "Banque Leumi", "format1": "excel",
                                             "banque2": "Bit", "format2": "bit", "banque3": "", "format3": "excel",
                                             "caisse": "on", "carte": "", "hebergeur": ""})
        self.assertRedirects(r, "/", fetch_redirect_response=False)
        self.assertEqual(sorted(Journal.objects.values_list("code", flat=True)), ["AN", "B1", "B2", "CA", "HA", "OD", "VT"])
        self.assertEqual(Journal.objects.get(code="B1").compte_id, "512000")
        self.assertEqual([reglages.lire(c) for c in ("nom_association", "devise", "releves_mizrahi", "releve_bit",
                                                     "carte_bancaire", "traductions_releve")], ["Amis du musée", "€", "", "B2", "", "non"])
        self.assertEqual(Exercice.objects.get().libelle, "2027")
        self.assertEqual(set(MoyenPaiement.objects.values_list("libelle", flat=True)),
                         {"Banque Leumi", "BIT", "Caisse (espèces)", "Non réglé"})
        self.assertEqual(ModeleOperation.objects.count(), 13)                      # pas de carte
        self.assertEqual(ModeFiche.objects.get(libelle="Bit").compte_id, "512100")  # compte du journal Bit
        r = self.client.get("/")
        self.assertContains(r, "Amis du musée")
        self.assertContains(r, " €</div>")
        self.assertRedirects(self.client.get("/demarrage/"), "/", fetch_redirect_response=False)   # une seule fois
        self.assertEqual(controles.etat_general(controles.executer())[0], "OK")

    def test_mes_propres_fichiers(self):
        self.administrateur()
        r = self.client.post("/demarrage/", {"nom": "Club", "devise": "$", "plan": "importer", "debut": "2027-01-01",
                                             "fin": "2027-12-31", "format1": "excel", "format2": "excel", "format3": "excel"})
        self.assertRedirects(r, "/echanges/", fetch_redirect_response=False)
        self.assertFalse(Journal.objects.exists() or Compte.objects.exists())
        self.assertEqual((reglages.lire("devise"), Reglage.lire("demarrage")), ("$", "fait"))
        self.assertEqual(self.client.get("/").status_code, 200)



class Licences(TestCase):
    """Étape 3 : sites créés par l'assistant ; essai, lecture seule, licence signée ; le site de la Loge est exempté."""

    def setUp(self):
        from unittest import mock

        from cryptography.hazmat.primitives import serialization
        from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey

        from . import licence
        self.licence = licence
        self.cle = Ed25519PrivateKey.generate()
        publique = licence._b64(self.cle.public_key().public_bytes(serialization.Encoding.Raw, serialization.PublicFormat.Raw))
        patch = mock.patch.object(licence, "CLE_PUBLIQUE", publique)
        patch.start()
        self.addCleanup(patch.stop)
        referentiels_saisie()
        call_command("migrate", verbosity=0)
        self.u = User.objects.create_user("admin", password="x")
        from .vues_utilisateurs import donner_role
        donner_role(self.u, "Administrateur")
        self.client.force_login(self.u)

    def site_neuf(self, il_y_a):
        Reglage.objects.create(cle="demarrage", valeur="fait")
        Reglage.objects.create(cle="demarrage_le", valeur=(dt.date.today() - dt.timedelta(days=il_y_a)).isoformat())

    def test_site_de_la_loge_exempte(self):
        self.assertEqual(self.licence.etat()["code"], "exempt")
        self.assertNotContains(self.client.get("/"), "Licence")

    def test_essai_puis_lecture_seule_puis_licence(self):
        self.site_neuf(10)
        e = self.licence.etat()
        self.assertEqual((e["code"], e["lecture_seule"]), ("essai", False))
        self.assertContains(self.client.get("/"), "50 jour(s) restant(s)")
        Reglage.objects.filter(cle="demarrage_le").update(valeur=(dt.date.today() - dt.timedelta(days=61)).isoformat())
        self.assertTrue(self.licence.etat()["lecture_seule"])
        n = Mouvement.objects.count()
        r = self.client.post("/codes/", {"nouveau_code": "1"}, follow=True)
        self.assertContains(r, "Période d&#x27;essai terminée")
        self.assertEqual(Mouvement.objects.count(), n)
        self.assertEqual(self.client.get("/ecritures/").status_code, 200)            # consultation
        self.assertEqual(self.client.post("/echanges/", {"exporter": "Tiers"}).status_code, 200)   # export permis
        fin = dt.date.today() + dt.timedelta(days=365)
        texte = self.licence.signer(self.cle, "Amis du musée", "testserver", fin)
        self.client.post("/licence/", {"licence": texte[:40] + "\n" + texte[40:]})          # coupée par un retour à la ligne
        e = self.licence.etat()
        self.assertEqual((e["code"], e["licence"]["association"]), ("valide", "Amis du musée"))
        self.assertTrue(Modification.objects.filter(action="Licence enregistrée").exists())

    def test_licence_falsifiee_ou_d_un_autre_site(self):
        self.site_neuf(0)
        fin = dt.date.today() + dt.timedelta(days=30)
        autre = self.licence.signer(self.cle, "Club", "club.pythonanywhere.com", fin)
        Reglage.objects.create(cle="licence", valeur=autre)
        self.assertEqual(self.licence.etat()["code"], "invalide")
        bonne = self.licence.signer(self.cle, "Club", "*", fin)
        contenu = bonne.split(".")[1]
        fausse = bonne.replace(contenu, self.licence._b64(b'{"a": "Club", "s": "*", "f": "2099-12-31"}'))
        with self.assertRaises(self.licence.LicenceInvalide):
            self.licence.lire(fausse)
        Reglage.objects.filter(cle="licence").update(valeur=bonne)
        self.assertEqual(self.licence.etat()["code"], "valide")
        Reglage.objects.filter(cle="licence").update(valeur=self.licence.signer(self.cle, "Club", "*", dt.date.today() - dt.timedelta(days=1)))
        self.assertEqual(self.licence.etat()["code"], "expiree")

    def test_commande_de_l_editeur(self):
        d = Path(tempfile.mkdtemp())
        out = io.StringIO()
        call_command("licence", "cles", dossier=str(d), stdout=out)
        self.assertTrue((d / "privee.pem").exists())
        with self.assertRaises(Exception):
            call_command("licence", "cles", dossier=str(d), stdout=io.StringIO())      # jamais remplacée
        out = io.StringIO()
        call_command("licence", "creer", association="Club", site="club.pythonanywhere.com", fin="2027-12-31", dossier=str(d), stdout=out)
        texte = out.getvalue().strip()
        out = io.StringIO()
        call_command("licence", "verifier", texte, dossier=str(d), stdout=out)
        self.assertIn("Club · site : club.pythonanywhere.com · fin : 31/12/2027", out.getvalue())


class Ecrans(TestCase):
    def setUp(self):
        referentiels_saisie()
        call_command("migrate", verbosity=0)
        self.u = User.objects.create_user("tresorier", password="x")
        self.u.groups.add(Group.objects.get(name="Trésorier"))
        self.client.force_login(self.u)

    def formulaire(self, **k):
        e = SCENARIOS["cotisation_reglee"][0]
        data = {"date": "2026-10-15", "modele": ModeleOperation.objects.get(type=e["modele"]).pk, "tiers": e["tiers"], "montant": "400",
                "paiement": MoyenPaiement.objects.get(libelle=e["paiement"]).pk, "anal2": e["anal2"]}
        data.update(k)
        return data

    def test_pages_et_listes_cherchables(self):
        for url in ("/saisie/", "/codes/", "/ecritures/", "/grand-livre/"):
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200, url)
            self.assertContains(r, "data-cherchable", msg_prefix=url)
            self.assertContains(r, "compta/liste.js", msg_prefix=url)
        self.assertContains(self.client.get("/saisie/"), "411TAIEB001 – TAIEB JEANNE")
        self.assertContains(self.client.get("/codes/"), "SOC.006 – BOURSES")

    def test_apercu_puis_enregistrement(self):
        r = self.client.post("/saisie/", self.formulaire(apercu="1"))
        self.assertContains(r, "Tous les contrôles sont OK")
        self.assertFalse(Mouvement.objects.filter(numero=422).exists())
        r = self.client.post("/saisie/", self.formulaire(enregistrer="1"))
        self.assertRedirects(r, "/mouvement/422/")
        self.assertEqual(Mouvement.objects.get(numero=422).cree_par, self.u)
        r = self.client.post("/saisie/", self.formulaire(enregistrer="1"))
        self.assertContains(r, "Déjà enregistrée ?")
        self.assertRedirects(self.client.post("/saisie/", self.formulaire(enregistrer="1", forcer="on")), "/mouvement/423/")

    def test_saisie_refusee_au_bureau(self):
        b = User.objects.create_user("bureau")
        b.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(b)
        self.assertEqual(self.client.get("/saisie/").status_code, 403)
        self.assertEqual(self.client.get("/codes/").status_code, 403)

    def test_codes(self):
        Prefixe.objects.create(prefixe="SOC.", axe=2, libelle="Social")
        r = self.client.post("/codes/", {"code-prefixe": "SOC.", "code-libelle": "aide aux familles", "code-statut": "1",
                                         "creer_code": "1"})
        self.assertRedirects(r, "/codes/")
        self.assertEqual(CodeAnalytique.objects.get(code="SOC.007").libelle, "AIDE AUX FAMILLES")
        membre, fournisseur = TypeTiers.objects.get(libelle="Membre"), TypeTiers.objects.get(libelle="Fournisseur")
        self.client.post("/codes/", {"membre-type": membre.pk, "membre-nom": "Taïeb", "membre-prenom": "Paul", "membre-ville": "Netanya",
                                     "creer_membre": "1"})
        self.client.post("/codes/", {"membre-type": fournisseur.pk, "membre-nom": "Partner", "membre-adresse": "8 rue Ha-Barzel",
                                     "membre-code_postal": "6971005", "membre-ville": "Tel Aviv", "creer_membre": "1"})
        from .models import Membre
        f = Membre.objects.get(compte_id="401PARTN001")
        self.assertEqual((f.type, f.adresse_complete), (fournisseur, "8 rue Ha-Barzel, 6971005 Tel Aviv"))
        self.assertEqual(Membre.objects.get(compte_id="411TAIEB002").ville, "Netanya")
        self.assertTrue(Compte.objects.filter(numero="411TAIEB002", libelle="TAÏEB PAUL").exists())
        r = self.client.post("/codes/", {"statut-code": "MAN.001", "statut-statut": "2", "changer_statut": "1"})
        self.assertContains(r, "cocher la confirmation")
        self.client.post("/codes/", {"statut-code": "MAN.001", "statut-statut": "2", "statut-confirmation": "on", "changer_statut": "1"})
        self.assertEqual(CodeAnalytique.objects.get(code="MAN.001").statut, 2)


# ---------------------------------------------------------------- W2 : fiches bénévoles (cas de la recette du fichier de liaison)

from . import fiches as fiches_moteur  # noqa: E402
from .forms import BenevoleForm  # noqa: E402
from .models import Fiche, LigneFiche, ModeFiche, NatureFiche, TiersProvisoire, TypeTiers  # noqa: E402


def referentiels_fiches():
    referentiels_saisie()
    for n, lib in (("600200", "LOCATION DE SALLE"), ("625000", "DONS EMIS"), ("630000", "DEPENSES SOCIALES")):
        Compte.objects.create(numero=n, libelle=lib, anal1_id="FON.1")
    moteur.initialiser_parametres()


def ligne_fiche(fiche, sens, nature, montant, mode=None, tiers=None, autre="", personnes=None, date=dt.date(2026, 10, 15), **k):
    return LigneFiche.objects.create(
        fiche=fiche, sens=sens, date=date, nature=NatureFiche.objects.get(type_fiche=fiche.type, sens=sens, libelle=nature),
        montant=D(montant), mode=ModeFiche.objects.get(type_fiche=fiche.type, libelle=mode) if mode else None,
        tiers=Compte.objects.get(numero=tiers) if tiers else None, autre=autre, personnes=personnes, **k)


def ecritures(l):
    journal, lignes = fiches_moteur.generer(l)
    return journal.code, [(x.compte.numero, x.debit, x.credit) for x in lignes]


class Fiches(TestCase):
    def setUp(self):
        referentiels_fiches()
        self.act = Fiche.objects.create(type="activite", titre="Rallye", anal2_id="MAN.001")
        self.ges = Fiche.objects.create(type="gestion", titre="Dons 2026")

    def remplir(self):
        a = self.act
        return [ligne_fiche(a, "R", "Participation", 200, "Bit", tiers="411TAIEB001", personnes=2),
                ligne_fiche(a, "R", "Don", 50, "Espèces", autre="M. Levy", personnes=1),
                ligne_fiche(a, "R", "Participation", 200, "Non payé", tiers="411TAIEB001"),
                ligne_fiche(a, "D", "Location de salle", 150, "Carte Isracard", autre="Salle Beit Ha'am"),
                ligne_fiche(a, "D", "Frais divers", 140, "Avance d'un membre", tiers="411TAIEB001"),
                ligne_fiche(self.ges, "R", "Don reçu", 1000, "Virement BIT", autre="Fondation X", anal2_id="SOC.006"),
                ligne_fiche(self.ges, "D", "Aide versée", 300, "Espèces", autre="Famille Y", anal2_id="SOC.007"
                            if CodeAnalytique.objects.filter(code="SOC.007").exists() else "SOC.006")]

    def test_parametres(self):
        self.assertEqual((NatureFiche.objects.count(), ModeFiche.objects.count()), (13, 10))
        self.assertEqual(ModeFiche.objects.get(libelle="Carte Isracard").compte_id, "580000")

    def test_ecritures(self):
        ls = self.remplir()
        attendu = [
            ("B3", [("411TAIEB001", 200, 0), ("710000", 0, 200), ("512200", 200, 0), ("411TAIEB001", 0, 200)]),
            ("CA", [("530000", 50, 0), ("725000", 0, 50)]),
            ("VT", [("411TAIEB001", 200, 0), ("710000", 0, 200)]),
            ("OD", [("600200", 150, 0), ("580000", 0, 150)]),
            ("OD", [("600000", 140, 0), ("411TAIEB001", 0, 140)]),
            ("B3", [("512200", 1000, 0), ("725000", 0, 1000)]),
            ("CA", [("630000", 300, 0), ("530000", 0, 300)]),
        ]
        for l, (j, e) in zip(ls, attendu):
            with self.subTest(l.nature.libelle):
                self.assertEqual(fiches_moteur.controler(l).erreurs, {})
                self.assertEqual(ecritures(l), (j, [(c, D(d), D(cr)) for c, d, cr in e]))
        self.assertEqual(fiches_moteur.totaux(self.act),
                         {"participants": 3, "recettes": D(450), "depenses": D(290), "resultat": D(160)})
        self.assertEqual(fiches_moteur.controler(ls[0]).libelle, "PARTICIPATION - TAIEB JEANNE")

    def test_controles(self):
        l = ligne_fiche(self.ges, "R", "Don reçu", 100, autre="Z")
        self.assertEqual(set(fiches_moteur.controler(l).erreurs), {"Mode de paiement", "Axe 2"})
        l = ligne_fiche(self.act, "R", "Participation", 100, "Non payé", autre="Inconnu")
        self.assertIn("demande un membre", fiches_moteur.controler(l).erreurs["Mode de paiement"])
        p = TiersProvisoire.objects.create(nom="Nouveau")
        l = ligne_fiche(self.act, "R", "Participation", 100, "Bit", provisoire=p)
        self.assertIn("Tiers", fiches_moteur.controler(l).erreurs)
        sans_code = Fiche.objects.create(type="activite", titre="Sans code")
        l = ligne_fiche(sans_code, "D", "Frais divers", 10, "Espèces", date=dt.date(2025, 11, 2))
        self.assertEqual(set(fiches_moteur.controler(l).erreurs), {"Axe 2", "Date"})

    def test_report(self):
        self.remplir()
        u = User.objects.create_user("t")
        crees = fiches_moteur.reporter(self.act, u)
        self.assertEqual([m.numero for m in crees], [422, 423, 424, 425, 426])
        self.assertEqual(Ligne.objects.filter(mouvement__origine="liaison").count(), 12)
        self.assertEqual(set(Ligne.objects.filter(mouvement__origine="liaison").values_list("anal2_id", flat=True)), {"MAN.001"})
        self.assertEqual(self.act.statut, "reportee")
        self.assertEqual(fiches_moteur.reporter(self.act, u), [])     # rien de nouveau
        ligne_fiche(self.ges, "R", "Don reçu", 100, autre="Z")         # gestion : colonnes du trésorier vides
        with self.assertRaises(ValueError):
            fiches_moteur.reporter(self.ges, u)
        self.assertFalse(Mouvement.objects.filter(numero=427).exists())


class DocumentsFichesExport(TransactionTestCase):
    """Export complet puis réinjection avec des reçus de bénévoles (sauvegarde de la base : hors transaction de test)."""

    def setUp(self):
        EcransFiches.setUp(self)

    def ligne_post(self, **k):
        return EcransFiches.ligne_post(self, **k)

    @override_settings(DATA_DIR=Path(tempfile.mkdtemp()), EXPORTS_DIR=Path(tempfile.mkdtemp()))
    def test_documents_de_fiche_dans_l_export_complet(self):
        from django.core.files.uploadedfile import SimpleUploadedFile as F
        from . import export_complet
        self.client.force_login(self.benevole)
        self.client.post(f"/fiches/{self.act.pk}/", self.ligne_post(qui="c:411TAIEB001"))
        self.client.post(f"/fiches/{self.act.pk}/", {"joindre_document": "1", "ligne": self.act.lignes.get().pk,
                                                     "documents": [F("recu.pdf", b"%PDF-1.4 recu")]})
        self.client.post(f"/fiches/{self.act.pk}/", {"joindre_document": "1", "ligne": "", "documents": [F("tout.pdf", b"%PDF t")]})
        chemin = export_complet.exporter()
        export_complet.reinjecter(chemin)                                            # contrôle : base identique à l'export
        f = Fiche.objects.get(titre="Rallye")
        docs = {d.nom: d for d in f.documents.all()}
        self.assertEqual(set(docs), {"recu.pdf", "tout.pdf"})
        self.assertEqual(docs["recu.pdf"].ligne, f.lignes.get())
        self.assertIsNone(docs["tout.pdf"].ligne)
        from . import justificatifs as just
        self.assertEqual((just.dossier() / docs["recu.pdf"].chemin).read_bytes(), b"%PDF-1.4 recu")


class EcransFiches(TestCase):
    def setUp(self):
        referentiels_fiches()
        call_command("migrate", verbosity=0)
        self.tresorier = User.objects.create_user("tresorier")
        self.tresorier.groups.add(Group.objects.get(name="Trésorier"))
        self.benevole = User.objects.create_user("rachel")
        self.benevole.groups.add(Group.objects.get(name="Bénévole"))
        self.act = Fiche.objects.create(type="activite", titre="Rallye", anal2_id="MAN.001")
        self.act.benevoles.add(self.benevole)
        self.autre = Fiche.objects.create(type="gestion", titre="Pas à elle")

    def ligne_post(self, **k):
        data = {"ajouter": "1", "sens": "R", "date": "2026-10-15", "nature": NatureFiche.objects.get(type_fiche="activite", libelle="Participation").pk,
                "montant": "120", "mode": ModeFiche.objects.get(type_fiche="activite", libelle="Bit").pk, "personnes": "2"}
        data.update(k)
        return data

    def test_benevole(self):
        self.client.force_login(self.benevole)
        self.assertRedirects(self.client.get("/"), "/fiches/")
        r = self.client.get("/fiches/")
        self.assertContains(r, "Rallye")
        self.assertNotContains(r, "Pas à elle")
        self.assertNotContains(r, "Nouvelle fiche")
        self.assertEqual(self.client.get(f"/fiches/{self.autre.pk}/").status_code, 403)
        for url in ("/ecritures/", "/saisie/", "/tiers-provisoires/"):
            self.assertEqual(self.client.get(url).status_code, 403, url)
        r = self.client.get(f"/fiches/{self.act.pk}/")
        self.assertContains(r, "411TAIEB001")         # recherche des membres existants
        self.assertNotContains(r, "Compte de contrepartie")
        # membre existant, puis nouveau tiers provisoire
        self.client.post(f"/fiches/{self.act.pk}/", self.ligne_post(qui="c:411TAIEB001"))
        self.client.post(f"/fiches/{self.act.pk}/", self.ligne_post(nouveau_nom="Cohen", nouveau_prenom="Dan"))
        self.assertEqual(self.act.lignes.count(), 2)
        p = TiersProvisoire.objects.get()
        self.assertEqual((str(p), p.cree_par), ("COHEN DAN (provisoire)", self.benevole))
        self.assertContains(self.client.get(f"/fiches/{self.act.pk}/"), "compte à attribuer par le trésorier")
        # transmission : plus de saisie
        self.client.post(f"/fiches/{self.act.pk}/", {"transmettre": "1"})
        self.act.refresh_from_db()
        self.assertEqual(self.act.statut, "transmise")
        self.client.post(f"/fiches/{self.act.pk}/", self.ligne_post(qui="c:411TAIEB001"))
        self.assertEqual(self.act.lignes.count(), 2)
        self.assertEqual(self.client.get(f"/fiches/ligne/{self.act.lignes.first().pk}/").status_code, 403)

    @override_settings(DATA_DIR=Path(tempfile.mkdtemp()))
    def test_documents_du_benevole_deviennent_justificatifs(self):
        from django.core.files.uploadedfile import SimpleUploadedFile as F
        self.client.force_login(self.benevole)
        self.client.post(f"/fiches/{self.act.pk}/", self.ligne_post(qui="c:411TAIEB001"))
        l = self.act.lignes.get()
        page = self.client.get(f"/fiches/{self.act.pk}/")
        self.assertContains(page, "Documents (reçus, factures)")
        self.client.post(f"/fiches/{self.act.pk}/", {"joindre_document": "1", "ligne": l.pk, "description": "reçu Bit",
                                                     "documents": [F("recu.pdf", b"%PDF-1.4 recu")]})
        self.client.post(f"/fiches/{self.act.pk}/", {"joindre_document": "1", "ligne": "",
                                                     "documents": [F("affiche.jpg", b"\xff\xd8x"), F("virus.exe", b"MZ")]})
        self.assertEqual(sorted(self.act.documents.values_list("nom", flat=True)), ["affiche.jpg", "recu.pdf"])
        d = self.act.documents.get(nom="recu.pdf")
        self.assertEqual(b"".join(self.client.get(f"/fiches/document/{d.pk}/").streaming_content), b"%PDF-1.4 recu")
        autre = User.objects.create_user("autre")
        autre.groups.add(Group.objects.get(name="Bénévole"))
        self.client.force_login(autre)                                           # pas sa fiche
        self.assertEqual(self.client.get(f"/fiches/document/{d.pk}/").status_code, 403)
        self.client.force_login(self.tresorier)
        self.client.post(f"/fiches/{self.act.pk}/", {"reporter": "1"})
        m = Mouvement.objects.get(origine="liaison")
        self.assertEqual(sorted(m.justificatifs.values_list("nom", flat=True)), ["affiche.jpg", "recu.pdf"])
        self.assertEqual(m.justificatifs.get(nom="recu.pdf").description, "reçu Bit")

    def test_tresorier(self):
        self.client.force_login(self.benevole)
        self.client.post(f"/fiches/{self.act.pk}/", self.ligne_post(nouveau_nom="Cohen", nouveau_prenom="Dan"))
        self.client.force_login(self.tresorier)
        for url in ("/fiches/", f"/fiches/{self.act.pk}/", "/tiers-provisoires/", f"/fiches/ligne/{self.act.lignes.get().pk}/"):
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200, url)
            self.assertContains(r, "data-cherchable", msg_prefix=url)
        # report refusé tant que le tiers est provisoire
        self.client.post(f"/fiches/{self.act.pk}/", {"reporter": "1"})
        self.assertFalse(Mouvement.objects.filter(origine="liaison").exists())
        p = TiersProvisoire.objects.get()
        self.assertRedirects(self.client.post("/tiers-provisoires/", {"provisoire": p.pk, f"p{p.pk}-compte": "",
                                                                  f"p{p.pk}-type": TypeTiers.objects.get(libelle="Membre").pk}), "/tiers-provisoires/")
        p.refresh_from_db()
        self.assertEqual((p.compte_id, p.compte.libelle), ("411COHEN001", "COHEN DAN"))
        self.client.post(f"/fiches/{self.act.pk}/", {"reporter": "1"})
        m = Mouvement.objects.get(origine="liaison")
        self.assertEqual(list(m.lignes.values_list("compte_id", flat=True)), ["411COHEN001", "710000", "512200", "411COHEN001"])
        self.assertEqual(m.cree_par, self.tresorier)
        self.assertEqual(self.client.get(f"/fiches/ligne/{self.act.lignes.get().pk}/").status_code, 403)   # verrouillée

    def test_creations_par_le_tresorier(self):
        from .membres import creer_manquants
        from .models import Membre
        creer_manquants()
        Membre.objects.filter(compte_id="411TAIEB001").update(nom="Taieb", prenom="Jeanne")
        self.client.force_login(self.tresorier)
        r = self.client.post("/fiches/", {"benevole-identifiant": "jeanne", "benevole-membre": "9999",
                                          "benevole-mot_de_passe": "motdepasse-8", "creer_benevole": "1"})
        self.assertFalse(User.objects.filter(username="jeanne").exists())          # le bénévole doit être un membre
        r = self.client.post("/fiches/", {"benevole-identifiant": "jeanne", "benevole-membre": "411TAIEB001",
                                          "benevole-mot_de_passe": "motdepasse-8", "creer_benevole": "1"})
        self.assertRedirects(r, "/fiches/")
        jeanne = User.objects.get(username="jeanne")
        self.assertTrue(jeanne.groups.filter(name="Bénévole").exists())
        self.assertEqual((jeanne.first_name, jeanne.last_name, jeanne.membre.compte_id), ("Jeanne", "Taieb", "411TAIEB001"))
        self.assertNotIn("411TAIEB001", [m.pk for m in BenevoleForm().fields["membre"].queryset])  # déjà bénévole
        r = self.client.post("/fiches/", {"fiche-type": "gestion", "fiche-titre": "Aides", "fiche-benevoles": [jeanne.pk],
                                          "creer_fiche": "1"})
        f = Fiche.objects.get(titre="Aides")
        self.assertRedirects(r, f"/fiches/{f.pk}/")
        self.assertEqual(list(f.benevoles.all()), [jeanne])
        # code axe 2 choisi dans la liste, ou créé avec la fiche
        Prefixe.objects.create(prefixe="MAN.", axe=2, libelle="Manifestations")
        self.client.post("/fiches/", {"fiche-type": "activite", "fiche-titre": "Rallye 2", "fiche-anal2": "MAN.001",
                                      "creer_fiche": "1"})
        self.assertEqual(Fiche.objects.get(titre="Rallye 2").anal2_id, "MAN.001")
        self.client.post("/fiches/", {"fiche-type": "activite", "fiche-titre": "Gala", "fiche-nouveau_prefixe": "MAN.",
                                      "fiche-nouveau_libelle": "Gala 2026", "creer_fiche": "1"})
        g = Fiche.objects.get(titre="Gala")
        self.assertEqual((g.anal2_id, g.anal2.libelle, g.anal2.axe), ("MAN.002", "GALA 2026", 2))
        r = self.client.post("/fiches/", {"fiche-type": "activite", "fiche-titre": "Double", "fiche-anal2": "MAN.001",
                                          "fiche-nouveau_prefixe": "MAN.", "fiche-nouveau_libelle": "Autre", "creer_fiche": "1"})
        self.assertContains(r, "pas les deux")
        self.assertFalse(Fiche.objects.filter(titre="Double").exists())
        self.assertFalse(CodeAnalytique.objects.filter(libelle="AUTRE").exists())


# ---------------------------------------------------------------- W3 : rapprochement bancaire

from . import releves as rap  # noqa: E402
from .models import LigneReleve, ParametreReleve, Rapprochement, Traduction  # noqa: E402

CSV_MODELE = ("﻿Date;Référence;Opération (relevé);Montant;Solde relevé\n"
              "05/01/2026;11;עמלת מסלול;-10,00;990,00\n"
              "10/01/2026;12;הפקדת שיק;1 550,00;2 540,00\n"
              "20/01/2026;14;העברה באינטרנט;-400,00;2 140,00\n").encode("utf-8")


class Rapprochements(TestCase):
    def setUp(self):
        referentiels_saisie()
        self.b1 = Journal.objects.get(code="B1")
        ParametreReleve.objects.create(journal=self.b1, date_reprise=dt.date(2026, 1, 1))
        Traduction.objects.create(cle=Traduction.cle_de("עמלת מסלול"), hebreu="עמלת מסלול", traduction="Frais de forfait")
        self.u = User.objects.create_user("t")
        # compta : à-nouveau 1 000 au 01/01, frais 10, deux chèques remis ensemble, un virement de 400 daté 3 jours plus tôt
        n = 500
        for d, lignes in ((dt.date(2026, 1, 1), [("512000", 1000, 0), ("580000", 0, 1000)]),
                          (dt.date(2026, 1, 5), [("600100", 10, 0), ("512000", 0, 10)]),
                          (dt.date(2026, 1, 9), [("512000", 1250, 0), ("411TAIEB001", 0, 1250)]),
                          (dt.date(2026, 1, 9), [("512000", 300, 0), ("411TAIEB001", 0, 300)]),
                          (dt.date(2026, 1, 17), [("600000", 400, 0), ("512000", 0, 400)]),
                          (dt.date(2026, 1, 30), [("600000", 55, 0), ("512000", 0, 55)])):
            m = Mouvement.objects.create(numero=n, date=d, journal=self.b1, piece=n, origine="saisie")
            for i, (c, db, cr) in enumerate(lignes):
                Ligne.objects.create(mouvement=m, ordre=i, compte_id=c, libelle=f"MVT {n}", debit=D(db), credit=D(cr), anal2_id="GEN.001")
            n += 1

    def importer(self, contenu=CSV_MODELE, nom="releve.csv", **k):
        return rap.importer(self.b1, rap.lire(nom, contenu), source=nom, **k)

    def test_lecture_et_import_sans_doublon(self):
        lignes = rap.lire("r.csv", CSV_MODELE)
        self.assertEqual([(l["date"].day, l["montant"]) for l in lignes], [(5, D("-10")), (10, D("1550")), (20, D("-400"))])
        self.assertEqual(self.importer(), (3, 0, 0))
        ouverture = LigneReleve.objects.get(ouverture=True)
        self.assertEqual((ouverture.montant, ouverture.date), (D("1000"), dt.date(2026, 1, 4)))   # déduit du premier solde
        self.assertEqual(self.importer(), (0, 3, 0))                                              # réimport : rien en double
        self.assertEqual(LigneReleve.objects.get(reference="11").traduction, "Frais de forfait")
        self.assertEqual(LigneReleve.objects.get(reference="14").traduction, "À traduire")

    def test_colonnes_hebreu_credit_debit_et_ordre_inverse(self):
        rangees = [["Relevé Mizrahi"], ["תאריך", "סוג תנועה", "זכות", "חובה", "יתרה", "אסמכתה"],
                   ["20/01/26", "העברה באינטרנט", "", "400.00", "2,140.00", "14"],
                   ["10/01/26", "הפקדת שיק", "1,250.00", "", "2,540.00", "12"]]
        lignes = rap.normaliser(rangees)
        self.assertEqual([(l["date"], l["montant"], l["solde"]) for l in lignes],
                         [(dt.date(2026, 1, 10), D("1250.00"), D("2540.00")), (dt.date(2026, 1, 20), D("-400.00"), D("2140.00"))])
        # PDF lu à l'envers (ordre visuel) : en-têtes et libellés retournés
        inverses = [[c[::-1] if isinstance(c, str) and rap.HEBREU.search(c) else c for c in r] for r in rangees]
        self.assertEqual(rap.normaliser(inverses)[1]["operation"], "העברה באינטרנט")

    def test_premier_import_sans_solde(self):
        with self.assertRaises(ValueError):
            rap.importer(self.b1, [{"date": dt.date(2026, 1, 5), "reference": "", "operation": "x", "montant": D(5), "solde": None}])

    def test_automatique_manuel_et_etat(self):
        self.importer()
        LigneReleve.objects.create(journal=self.b1, date=dt.date(2025, 12, 20), rang=1, reference="9", operation="avant", montant=D(0.01))
        LigneReleve.objects.filter(ouverture=True).update(montant=D("999.99"))  # ouverture + ligne avant reprise = à-nouveau
        self.assertEqual(rap.automatique(self.b1, self.u), 3)   # à-nouveau, frais, virement (± 3 jours)
        e = rap.etat(self.b1, dt.date(2026, 1, 31))
        self.assertEqual((len(e["releve_non_pointe"]), len(e["ecritures_non_pointees"])), (1, 3))
        self.assertEqual((e["ecart"], e["ecart_explique"]), (D("55.00"), D("0.00")))
        # une remise de chèques au relevé = deux règlements en compta ; totaux différents refusés
        rel = LigneReleve.objects.filter(rapprochement__isnull=True, ouverture=False)
        cheques = rap.ecritures(self.b1).filter(debit__in=[1250, 300])
        with self.assertRaises(ValueError):
            rap.pointer(self.b1, rel, cheques[:1], self.u)
        r = rap.pointer(self.b1, rel, cheques, self.u)
        self.assertEqual((r.releves.count(), r.ecritures.count(), r.total), (1, 2, D("1550")))
        e = rap.etat(self.b1, dt.date(2026, 1, 31))
        self.assertEqual([x.debit - x.credit for x in e["ecritures_non_pointees"]], [D("-55")])
        rap.depointer(r)
        self.assertEqual(Ligne.objects.filter(rapprochement__isnull=False).count(), 3)
        mois = rap.par_mois(self.b1)
        self.assertEqual((mois[0]["compta"], mois[-1]["ecart"]), (None, D("0.00")))   # arrêté au dernier jour du relevé


class EcransRapprochement(TestCase):
    def setUp(self):
        Rapprochements.setUp(self)
        call_command("migrate", verbosity=0)
        self.u.groups.add(Group.objects.get(name="Trésorier"))
        self.client.force_login(self.u)

    def test_import_affectation_et_ecriture(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        r = self.client.post("/rapprochement/B1/import/", {"fichier": SimpleUploadedFile("releve.csv", CSV_MODELE)})
        self.assertRedirects(r, "/rapprochement/B1/")
        self.client.post("/rapprochement/B1/import/", {"fichier": SimpleUploadedFile("releve.csv", CSV_MODELE)})
        self.assertEqual(LigneReleve.objects.count(), 4)                         # relevé retéléchargé : pas de doublon
        for url in ("/rapprochement/", "/rapprochement/B1/", "/rapprochement/traductions/"):
            self.assertEqual(self.client.get(url).status_code, 200, url)
        for url in ("/rapprochement/B1/pointage/", "/rapprochement/B1/automatique/"):
            self.assertEqual(self.client.get(url).status_code, 404, url)          # plus de pointage
        Mouvement.objects.all().delete()                  # compta vide : aucune écriture « déjà en compta »
        self.client.post("/rapprochement/traductions/", {"h_1": "עמלת מסלול", "t_1": "Frais de forfait"})
        frais, remise, virement = LigneReleve.objects.filter(ouverture=False).order_by("date")
        page = self.client.get("/rapprochement/B1/")
        self.assertContains(page, '<option value="600100 – FRAIS BANCAIRES">')
        self.assertContains(page, '<option value="GEN.004 – BANQUE">')
        self.assertNotContains(page, '<option value="512000')                     # pas la banque elle-même
        # frais : code choisi dans la liste ; virement : libellé tapé ; remise laissée vide ; code inconnu refusé
        r = self.client.post("/rapprochement/B1/", {
            f"compte_{frais.pk}": "600100 – FRAIS BANCAIRES", f"anal2_{frais.pk}": "GEN.004 – BANQUE",
            f"compte_{virement.pk}": "taieb jeanne", f"anal2_{virement.pk}": "GEN.002", "creer": "1"})
        self.assertRedirects(r, "/rapprochement/B1/")
        m = Mouvement.objects.get(lignes__compte_id="600100")
        self.assertEqual([(x.compte_id, x.debit, x.credit, x.anal2_id, x.libelle) for x in m.lignes.order_by("ordre")],
                         [("512000", D(0), D(10), "GEN.004", "FRAIS DE FORFAIT"), ("600100", D(10), D(0), "GEN.004", "FRAIS DE FORFAIT")])
        self.assertEqual((m.date, m.journal_id), (frais.date, "B1"))
        v = Mouvement.objects.get(lignes__compte_id="411TAIEB001", lignes__debit=400)
        self.assertEqual(v.lignes.get(compte_id="512000").credit, D(400))
        frais.refresh_from_db()
        self.assertEqual(frais.rapprochement.ecritures.get().mouvement, m)
        page = self.client.get("/rapprochement/B1/")
        self.assertNotContains(page, f'name="compte_{frais.pk}"')                   # traitée : sortie de la liste
        self.assertContains(page, f'name="compte_{remise.pk}"')
        n = Mouvement.objects.count()
        self.client.post("/rapprochement/B1/", {f"compte_{frais.pk}": "600100", f"anal2_{frais.pk}": "GEN.004", "creer": "1"})
        self.assertEqual(Mouvement.objects.count(), n)                              # double envoi : rien de plus
        r = self.client.post("/rapprochement/B1/", {f"compte_{remise.pk}": "600100", f"anal2_{remise.pk}": "ZZZ", "creer": "1"})
        self.assertContains(r, "Code axe 2 introuvable")
        self.assertContains(r, 'value="600100"')                                    # saisie gardée
        # déjà en compta (même montant, date proche) : relier plutôt que créer
        mv = Mouvement.objects.create(numero=900, date=remise.date, journal_id="B1", piece=900)
        e = Ligne.objects.create(mouvement=mv, ordre=1, compte_id="512000", libelle="REMISE", debit=D(1550), anal2_id="GEN.001")
        Ligne.objects.create(mouvement=mv, ordre=2, compte_id="411TAIEB001", libelle="REMISE", credit=D(1550), anal2_id="GEN.001")
        self.assertContains(self.client.get("/rapprochement/B1/"), f'value="{remise.pk}:{e.pk}"')
        r = self.client.post("/rapprochement/B1/", {f"compte_{remise.pk}": "411TAIEB001", f"anal2_{remise.pk}": "GEN.001", "creer": "1"})
        self.assertContains(r, "existe déjà")
        self.assertEqual(Mouvement.objects.count(), n + 1)
        self.client.post("/rapprochement/B1/", {"relier": f"{remise.pk}:{e.pk}"})
        remise.refresh_from_db()
        self.assertEqual(remise.rapprochement.ecritures.get(), e)
        self.assertEqual(Mouvement.objects.count(), n + 1)                          # reliée, pas d'écriture en plus
        self.assertContains(self.client.get("/rapprochement/B1/"), "Tout le relevé téléchargé est en comptabilité")

    def test_situation_financiere_pdf_et_excel(self):
        mv = Mouvement.objects.create(numero=907, date=dt.date(2026, 3, 1), journal_id="B1", piece=907)
        Ligne.objects.create(mouvement=mv, ordre=1, compte_id="512000", libelle="FRAIS", credit=D(50), anal2_id="GEN.004")
        Ligne.objects.create(mouvement=mv, ordre=2, compte_id="600100", libelle="FRAIS", debit=D(50), anal2_id="GEN.004")
        self.assertContains(self.client.get("/?du=2026-01-01&au=2026-12-31"), "Situation (PDF)")
        page = self.client.get("/situation/?du=2026-01-01&au=2026-12-31")
        for texte in ("Situation financière", "Période du <b>01/01/2026</b> au <b>31/12/2026</b>", "Charges par compte", "600100",
                      "Résultat par nature (axe 1)"):
            self.assertContains(page, texte)
        r = self.client.get("/situation/excel/?du=2026-01-01&au=2026-12-31")
        self.assertIn("Situation_2026-01-01_2026-12-31.xlsx", r["Content-Disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(r.content))["Situation"]
        valeurs = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertIn("Charges par compte", valeurs)
        self.assertIn(50, valeurs)

    def test_une_ligne_du_releve_pour_plusieurs_ecritures(self):
        """Prélèvement Isracard de 760 = deux écritures (160 + 600) passées séparément sur la banque."""
        l = LigneReleve.objects.create(journal=self.b1, date=dt.date(2026, 2, 10), rang=1, operation="ISRACARD", montant=D(-760))
        es = []
        for n, m in ((903, 160), (904, 600), (905, 45)):
            mv = Mouvement.objects.create(numero=n, date=dt.date(2026, 2, 10), journal_id="B1", piece=n)
            es.append(Ligne.objects.create(mouvement=mv, ordre=1, compte_id="512000", libelle="CARTE", credit=D(m), anal2_id="GEN.004"))
            Ligne.objects.create(mouvement=mv, ordre=2, compte_id="600100", libelle="CARTE", debit=D(m), anal2_id="GEN.004")
        page = self.client.get("/rapprochement/B1/")
        self.assertContains(page, "Plusieurs écritures")
        self.assertContains(page, f'value="{l.pk}:{es[0].pk},{es[1].pk}"')
        self.client.post("/rapprochement/B1/", {"relier": f"{l.pk}:{es[0].pk},{es[2].pk}"})       # 205 ≠ 760 : refusé
        l.refresh_from_db()
        self.assertIsNone(l.rapprochement_id)
        r = self.client.post("/rapprochement/B1/", {"relier": f"{l.pk}:{es[0].pk},{es[1].pk}"}, follow=True)
        self.assertContains(r, "reliée(s) à Mvt 903 + Mvt 904")
        l.refresh_from_db()
        self.assertEqual(sorted(e.pk for e in l.rapprochement.ecritures.all()), [es[0].pk, es[1].pk])
        self.assertEqual(Mouvement.objects.filter(numero__gte=903).count(), 3)                 # aucune écriture créée

    def test_frais_du_mois_en_une_ecriture_et_depot_renouvele(self):
        forfait = LigneReleve.objects.create(journal=self.b1, date=dt.date(2026, 2, 1), rang=1, operation="FORFAIT", montant=D(-10))
        guichet = LigneReleve.objects.create(journal=self.b1, date=dt.date(2026, 2, 1), rang=2, operation="GUICHET", montant=D("-13.60"))
        mv = Mouvement.objects.create(numero=906, date=dt.date(2026, 2, 1), journal_id="B1", piece=906)
        e = Ligne.objects.create(mouvement=mv, ordre=1, compte_id="512000", libelle="FRAIS", credit=D("23.60"), anal2_id="GEN.004")
        Ligne.objects.create(mouvement=mv, ordre=2, compte_id="600100", libelle="FRAIS", debit=D("23.60"), anal2_id="GEN.004")
        page = self.client.get("/rapprochement/B1/")
        self.assertContains(page, "Les relier ensemble")
        self.assertContains(page, f'value="{forfait.pk},{guichet.pk}:{e.pk}"')
        r = self.client.post("/rapprochement/B1/", {"relier": f"{forfait.pk},{guichet.pk}:{e.pk}"}, follow=True)
        self.assertContains(r, "2 lignes du relevé reliée(s) à Mvt 906")
        forfait.refresh_from_db(); guichet.refresh_from_db()
        self.assertEqual(forfait.rapprochement_id, guichet.rapprochement_id)
        # dépôt renouvelé : +3029,94 ; +0,01 ; −3029,95 le même jour : reliées entre elles, sans écriture
        d = dt.date(2026, 3, 6)
        rs = [LigneReleve.objects.create(journal=self.b1, date=d, rang=i, operation=o, montant=D(m))
              for i, (o, m) in enumerate((("DEPOT", "3029.94"), ("INTERETS", "0.01"), ("DEPOT", "-3029.95")), 10)]
        n = Mouvement.objects.count()
        page = self.client.get("/rapprochement/B1/")
        self.assertContains(page, "Relier sans écriture")
        valeur = ",".join(str(x.pk) for x in rs) + ":"
        self.assertContains(page, f'value="{valeur}"')
        self.client.post("/rapprochement/B1/", {"relier": f"{rs[0].pk},{rs[1].pk}:"})           # somme non nulle : refusé
        self.assertIsNone(LigneReleve.objects.get(pk=rs[0].pk).rapprochement_id)
        self.client.post("/rapprochement/B1/", {"relier": valeur})
        self.assertEqual(len({LigneReleve.objects.get(pk=x.pk).rapprochement_id for x in rs}), 1)
        self.assertIsNotNone(LigneReleve.objects.get(pk=rs[0].pk).rapprochement_id)
        self.assertEqual(Mouvement.objects.count(), n)

    def test_pistes_quand_rien_n_est_propose(self):
        l = LigneReleve.objects.create(journal=self.b1, date=dt.date(2026, 3, 20), rang=1, operation="x", montant=D(-71.93))
        mv = Mouvement.objects.create(numero=902, date=dt.date(2026, 2, 20), journal_id="B1", piece=902)     # 28 jours avant
        e = Ligne.objects.create(mouvement=mv, ordre=1, compte_id="512000", libelle="FRAIS", credit=D("71.93"), anal2_id="GEN.004")
        Ligne.objects.create(mouvement=mv, ordre=2, compte_id="600100", libelle="FRAIS", debit=D("71.93"), anal2_id="GEN.004")
        page = self.client.get("/rapprochement/B1/")
        self.assertContains(page, "Même montant, date plus éloignée")
        self.assertContains(page, f'value="{l.pk}:{e.pk}"')                       # reliable malgré la date
        self.client.post("/rapprochement/B1/", {"relier": f"{l.pk}:{e.pk}"})
        l.refresh_from_db()
        self.assertEqual(l.rapprochement.ecritures.get(), e)
        # même montant déjà relié à une autre ligne (relevé importé deux fois)
        double = LigneReleve.objects.create(journal=self.b1, date=dt.date(2026, 3, 20), rang=2, operation="x", montant=D("-71.93"))
        page = self.client.get("/rapprochement/B1/")
        self.assertContains(page, "relevé importé deux fois ?")
        self.assertNotContains(page, f'value="{double.pk}:{e.pk}"')
        LigneReleve.objects.create(journal=self.b1, date=dt.date(2026, 3, 22), rang=3, operation="y", montant=D(-3.21))
        self.assertContains(self.client.get("/rapprochement/B1/"), "Aucune écriture de ce montant à 60 jours près")

    def test_nouvelle_ecriture_malgre_un_montant_proche(self):
        l = LigneReleve.objects.create(journal=self.b1, date=dt.date(2026, 2, 3), rang=1, operation="x", montant=D(-10))
        mv = Mouvement.objects.create(numero=901, date=dt.date(2026, 2, 1), journal_id="B1", piece=901)
        Ligne.objects.create(mouvement=mv, ordre=1, compte_id="512000", libelle="AUTRE", credit=D(10), anal2_id="GEN.004")
        Ligne.objects.create(mouvement=mv, ordre=2, compte_id="600100", libelle="AUTRE", debit=D(10), anal2_id="GEN.004")
        self.client.post("/rapprochement/B1/", {f"compte_{l.pk}": "600100", f"anal2_{l.pk}": "GEN.004", f"nouvelle_{l.pk}": "on",
                                                "creer": "1"})
        l.refresh_from_db()
        self.assertIsNotNone(l.rapprochement_id)

    def test_droits(self):
        b = User.objects.create_user("bureau")
        b.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(b)
        l = LigneReleve.objects.create(journal=self.b1, date=dt.date(2026, 2, 3), rang=1, operation="x", montant=D(-10))
        r = self.client.get("/rapprochement/B1/")
        self.assertContains(r, "Lignes téléchargées sans écriture (1)")
        self.assertNotContains(r, f'name="compte_{l.pk}"')
        self.client.post("/rapprochement/B1/", {f"compte_{l.pk}": "600100", f"anal2_{l.pk}": "GEN.004", "creer": "1"})
        l.refresh_from_db()
        self.assertIsNone(l.rapprochement_id)


# ---------------------------------------------------------------- W4 : états annuels et clôture

from . import cloture as clot  # noqa: E402
from . import etats  # noqa: E402
from .models import Budget  # noqa: E402



def temporaire(d=None):
    """Dossiers de test : données, Imports et Exports hors du projet."""
    d = Path(d or tempfile.mkdtemp())
    return dict(DATA_DIR=d, IMPORTS_DIR=d / "Imports", EXPORTS_DIR=d / "Exports")

ARCHIVES_TEST = Path(tempfile.mkdtemp())


@override_settings(**temporaire(ARCHIVES_TEST))
class Cloture(TestCase):
    def setUp(self):
        referentiels_saisie()
        Exercice.objects.all().delete()
        self.e25 = Exercice.objects.create(libelle="Exercice 2025", debut=dt.date(2025, 1, 1), fin=dt.date(2025, 12, 31))
        self.e26 = Exercice.objects.create(libelle="Exercice 2026", debut=dt.date(2026, 1, 1), fin=dt.date(2026, 12, 31))
        Journal.objects.create(code="AN", intitule="A-NOUVEAUX", type="AN")
        Compte.objects.create(numero="110000", libelle="REPORT A NOUVEAU", anal1_id="BIL.4")
        self.u = User.objects.create_user("t")
        self.n = 1
        self.mvt(dt.date(2025, 3, 1), [("512000", 1000, 0), ("700000", 0, 1000)])      # cotisations 2025
        self.mvt(dt.date(2025, 6, 1), [("600000", 300, 0), ("512000", 0, 300)])        # achats 2025
        self.mvt(dt.date(2025, 9, 1), [("411TAIEB001", 50, 0), ("700000", 0, 50)])     # cotisation due
        self.mvt(dt.date(2026, 2, 1), [("512000", 200, 0), ("700000", 0, 200)])       # 2026

    def mvt(self, date, lignes):
        m = Mouvement.objects.create(numero=self.n, date=date, journal_id="B1", piece=self.n)
        for i, (c, d, cr) in enumerate(lignes):
            Ligne.objects.create(mouvement=m, ordre=i, compte_id=c, libelle="T", debit=D(d), credit=D(cr), anal2_id="GEN.001")
        self.n += 1

    def test_etats(self):
        cr = etats.compte_de_resultat(self.e26, self.e25)
        self.assertEqual((cr["resultat_n"], cr["resultat_n1"]), (D(600), D(750)))   # 2026 : + la facture 421 des référentiels
        b = etats.bilan(self.e25.fin, self.e25.debut)
        self.assertEqual((b["total_actif"], b["resultat"], b["equilibre"]), (D(750), D(750), True))
        Budget.objects.create(exercice=self.e25, nature="C", compte_id="600000", montant=D(500))
        Budget.objects.create(exercice=self.e25, nature="P", anal2_id="GEN.001", montant=D(1000))
        self.assertEqual([(x["realise"], x["ecart"], x["pourcent"]) for x in etats.budget(self.e25)],
                         [(D(300), D(-200), D("-40.0")), (D(1050), D(50), D("5.0"))])

    def test_ordre_des_clotures(self):
        p = clot.preparer(self.e26)
        self.assertIn("Un exercice antérieur n'est pas clos : le clôturer d'abord.", p.bloquants)

    def test_cloture(self):
        p = clot.preparer(self.e25)
        self.assertEqual(p.bloquants, [])
        self.assertEqual(p.lignes, [("110000", D(0), D(750)), ("411TAIEB001", D(50), D(0)), ("512000", D(700), D(0))])
        clot.cloturer(self.e25, self.u, CodeAnalytique.objects.get(code="GEN.001"))
        self.e25.refresh_from_db()
        an = self.e25.mouvement_an
        self.assertEqual((self.e25.clos, self.e25.resultat, an.date, an.journal_id), (True, D(750), dt.date(2026, 1, 1), "AN"))
        self.assertTrue((clot.dossier_archives() / self.e25.archive).exists())
        # soldes : l'exercice 2026 repart des à-nouveaux (pas de double compte)
        self.assertEqual(etats.solde_cumule(Compte.objects.get(numero="512000"), dt.date(2026, 6, 30)), D(900))
        b = etats.bilan(self.e26.fin, self.e26.debut)
        self.assertEqual((b["total_actif"], b["resultat"], b["resultats_anterieurs"], b["equilibre"]), (D(1350), D(600), D(0), True))
        # 2025 verrouillé pour la saisie
        op = operation(date=dt.date(2025, 12, 1), modele="Frais bancaires", montant=5, paiement="Mizrahi compte courant", anal2="GEN.004")
        self.assertIn("Date", moteur.controler(op).erreurs)
        # le rapprochement ignore les à-nouveaux de clôture
        self.assertFalse(rap.ecritures(Journal.objects.get(code="B1")).filter(mouvement=an).exists())
        # deuxième clôture impossible
        with self.assertRaises(ValueError):
            clot.cloturer(self.e25, self.u, CodeAnalytique.objects.get(code="GEN.001"))
        # clôture de 2026 : crée l'exercice 2027
        clot.cloturer(self.e26, self.u, CodeAnalytique.objects.get(code="GEN.001"))
        self.assertTrue(Exercice.objects.filter(debut=dt.date(2027, 1, 1), fin=dt.date(2027, 12, 31), clos=False).exists())
        self.assertEqual(etats.solde_cumule(Compte.objects.get(numero="110000"), dt.date(2027, 1, 1)), D(-1350))


@override_settings(**temporaire(ARCHIVES_TEST))
class EcransEtats(TestCase):
    def setUp(self):
        Cloture.setUp(self)
        call_command("migrate", verbosity=0)
        self.u.groups.add(Group.objects.get(name="Trésorier"))
        self.client.force_login(self.u)

    mvt = Cloture.mvt

    def test_pages(self):
        r = self.client.get(f"/etats/?exercice={self.e25.pk}")
        self.assertContains(r, "Bilan simplifié")
        self.assertContains(r, "équilibré")
        r = self.client.post(f"/etats/?exercice={self.e25.pk}", {"budget-nature": "C", "budget-compte": "600000", "budget-montant": "500",
                                                                 "ajouter_budget": "1"})
        self.assertEqual(Budget.objects.count(), 1)
        r = self.client.get(f"/etats/export/?exercice={self.e25.pk}")
        from . import dossiers as dos
        self.assertTrue(list(dos.exports().glob("ComptaBB_etats_*.xlsx")))                        # copie dans Exports
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb = openpyxl.load_workbook(__import__("io").BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Compte de résultat", "Résultat axe 1", "Résultat axe 2", "Bilan", "Balance", "Grand livre", "Budget",
                                         "Historique"])
        self.assertContains(self.client.get("/cloture/"), "À-nouveaux qui seront créés")
        self.client.post("/cloture/", {"anal2": "GEN.001", "confirmation": "on"})
        self.e25.refresh_from_db()
        self.assertTrue(self.e25.clos)
        self.assertEqual(self.client.get(f"/cloture/archive/{self.e25.pk}/").status_code, 200)
        self.assertEqual(Reglage.lire("code_axe2_a_nouveaux"), "GEN.001")

    def test_droits(self):
        b = User.objects.create_user("bureau")
        b.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(b)
        self.assertEqual(self.client.get("/etats/").status_code, 200)
        self.assertNotContains(self.client.get("/etats/"), "Ajouter une ligne de budget")
        self.assertEqual(self.client.get("/cloture/").status_code, 403)


# ---------------------------------------------------------------- W5 : suivi des membres

from . import membres as mbr  # noqa: E402
from .models import Membre  # noqa: E402


class Membres(TestCase):
    def setUp(self):
        referentiels_saisie()           # crée la fiche de 411TAIEB001 (et la facture 421 du 24/09/2026, 400 à 710000)
        Compte.objects.create(numero="411000", libelle="ADHERENTS", anal1_id="BIL.4")
        mbr.creer_manquants()
        self.c = Compte.objects.get(numero="411TAIEB001")
        self.u = User.objects.create_user("t")
        n = 600
        for d, lignes in ((dt.date(2026, 1, 10), [("411TAIEB001", 500, 0), ("700000", 0, 500)]),       # cotisation facturée
                          (dt.date(2026, 2, 1), [("512000", 300, 0), ("411TAIEB001", 0, 300)]),       # acompte
                          (dt.date(2026, 3, 1), [("411TAIEB001", 200, 0), ("710000", 0, 200),          # manifestation réglée
                                                 ("512000", 200, 0), ("411TAIEB001", 0, 200)])):
            m = Mouvement.objects.create(numero=n, date=d, journal_id="B1", piece=n)
            for i, (c, db, cr) in enumerate(lignes):
                Ligne.objects.create(mouvement=m, ordre=i, compte_id=c, libelle=f"MVT {n}", debit=D(db), credit=D(cr), anal2_id="GEN.001")
            n += 1

    def test_fiches_creees(self):
        self.assertEqual(list(Membre.objects.order_by("compte").values_list("compte_id", "nom", "prenom", "type__libelle")),
                         [("401000", "FOURNIS DIVERS", "", "Fournisseur"), ("411TAIEB001", "TAIEB", "Jeanne", "Membre")])
        self.assertEqual(mbr.creer_manquants(), 0)

    def test_situation(self):
        s = mbr.situation(self.c, dt.date(2026, 10, 1))
        self.assertEqual((s.facture, s.regle, s.solde), (D(1100), D(500), D(600)))
        # FIFO : les 500 réglés soldent la cotisation de janvier ; restent la manifestation de mars et la facture 421
        self.assertEqual([(l.mouvement.numero, r) for l, r in s.impayes], [(602, D(200)), (421, D(400))])

    def test_lettrage(self):
        self.assertEqual(mbr.lettrage_automatique(self.c), 1)          # manifestation de mars : même mouvement
        self.assertEqual(sorted(Ligne.objects.filter(compte=self.c, lettrage="A").values_list("debit", "credit")),
                         [(D(0), D(200)), (D(200), D(0))])
        cot, acompte = Ligne.objects.get(compte=self.c, debit=500), Ligne.objects.get(compte=self.c, credit=300)
        with self.assertRaises(ValueError):
            mbr.lettrer(self.c, [cot, acompte])                       # 500 ≠ 300
        s = mbr.situation(self.c, dt.date(2026, 10, 1))
        self.assertEqual([(l.mouvement.numero, r) for l, r in s.impayes], [(600, D(200)), (421, D(400))])
        self.assertEqual(mbr.delettrer(self.c, "A"), 2)
        self.assertEqual(mbr.code_suivant(self.c), "A")

    def test_cotisations_et_import(self):
        self.assertEqual(Reglage.lire("compte_cotisations"), "700000")      # créé par l'initialisation
        ex = Exercice.objects.get(libelle="2026")
        r = mbr.importer_tableau([["Compte", "Nom", "Prénom", "Téléphone", "E-mail", "Date d'adhésion", "Statut", "Cotisation annuelle"],
                                  ["411TAIEB001", "Taieb", "Jeanne", "050", "j@example.org", "01/09/2020", "honoraire", "500,00"],
                                  ["999XXX", "Inconnu", "", "", "", "", "actif", ""]])
        self.assertEqual((len(r.crees), len(r.mis_a_jour), len(r.erreurs)), (0, 1, 1))
        m = Membre.objects.get(compte_id="411TAIEB001")
        self.assertEqual((m.nom, m.statut, m.cotisation, m.date_adhesion), ("TAIEB", "honoraire", D(500), dt.date(2020, 9, 1)))
        lignes, tot = mbr.cotisations(ex)
        self.assertEqual((lignes[0]["attendue"], lignes[0]["facturee"], lignes[0]["recue"], lignes[0]["du"]), (D(0), D(500), D(500), D(0)))
        self.assertEqual(tot["taux"], D("100.0"))


class EcransMembres(TestCase):
    def setUp(self):
        Membres.setUp(self)
        call_command("migrate", verbosity=0)
        self.u.groups.add(Group.objects.get(name="Trésorier"))
        self.client.force_login(self.u)

    def test_pages_et_lettrage_manuel(self):
        for url in ("/membres/", "/membres/cotisations/", "/membres/411TAIEB001/", "/favicon.ico"):
            self.assertIn(self.client.get(url).status_code, (200, 301), url)
        self.assertNotContains(self.client.get("/membres/411TAIEB001/"), "relance")
        cot, acompte = Ligne.objects.get(compte=self.c, debit=500), Ligne.objects.get(compte=self.c, credit=300)
        cr = Ligne.objects.get(compte=self.c, credit=200)
        self.client.post("/membres/411TAIEB001/", {"lettrer": "1", "ligne": [cot.pk, acompte.pk]})
        self.assertEqual(Ligne.objects.filter(compte=self.c).exclude(lettrage="").count(), 0)   # refusé : déséquilibré
        self.client.post("/membres/411TAIEB001/", {"lettrer": "1", "ligne": [Ligne.objects.get(compte=self.c, debit=200).pk, cr.pk]})
        self.assertEqual(Ligne.objects.get(pk=cr.pk).lettrage, "A")
        self.client.post("/membres/411TAIEB001/", {"enregistrer": "1", "nom": "TAIEB", "prenom": "Jeanne", "statut": "actif",
                                                    "email": "jeanne@example.org", "cotisation": "450"})
        self.assertEqual(Membre.objects.get(compte_id="411TAIEB001").cotisation, D(450))

    def test_droits(self):
        b = User.objects.create_user("bureau")
        b.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(b)
        self.assertEqual(self.client.get("/membres/411TAIEB001/").status_code, 200)
        self.assertNotContains(self.client.get("/membres/411TAIEB001/"), "Lettrer la sélection")
        self.client.post("/membres/", {"lettrage_auto": "1"})
        self.assertFalse(Ligne.objects.exclude(lettrage="").exists())


# ---------------------------------------------------------------- W6 : corrections d'écritures, tri

from . import corrections as corr  # noqa: E402
from .models import Modification  # noqa: E402


class Corrections(TestCase):
    def setUp(self):
        referentiels_saisie()
        self.u = User.objects.create_user("t")
        self.m = Mouvement.objects.create(numero=500, date=dt.date(2026, 3, 1), journal_id="B1", piece=500)
        self.l1 = Ligne.objects.create(mouvement=self.m, ordre=0, compte_id="600100", libelle="FRAIS", debit=D(10), anal2_id="GEN.004")
        self.l2 = Ligne.objects.create(mouvement=self.m, ordre=1, compte_id="512000", libelle="FRAIS", credit=D(10), anal2_id="GEN.004")
        self.g = CodeAnalytique.objects.get(code="GEN.004")

    def saisie(self, l, compte=None, d=None, c=None, id=True):
        return corr.LigneSaisie(l.pk if id else None, Compte.objects.get(numero=compte or l.compte_id), l.libelle,
                                D(d if d is not None else l.debit), D(c if c is not None else l.credit), self.g)

    def test_modifier(self):
        b1 = Journal.objects.get(code="B1")
        with self.assertRaises(ValueError):                      # motif obligatoire
            corr.modifier(self.m, self.m.date, b1, [self.saisie(self.l1), self.saisie(self.l2)], "", self.u)
        with self.assertRaises(ValueError):                      # déséquilibré
            corr.modifier(self.m, self.m.date, b1, [self.saisie(self.l1, d=12), self.saisie(self.l2)], "erreur", self.u)
        r = Rapprochement.objects.create(journal=b1, mode="manuel")
        Ligne.objects.filter(pk=self.l2.pk).update(rapprochement=r)
        corr.modifier(self.m, self.m.date, b1, [self.saisie(self.l1, d=12), self.saisie(self.l2, c=12)], "montant réel 12", self.u)
        self.m.refresh_from_db()
        self.assertEqual([(l.compte_id, l.debit, l.credit, l.rapprochement_id) for l in self.m.lignes.all()],
                         [("600100", D(12), D(0), None), ("512000", D(0), D(12), None)])   # dépointé
        self.assertIn("montant réel 12", self.m.commentaire)
        self.assertTrue(Modification.objects.filter(action="Modification", objet__contains="Mvt 500").exists())
        # 3 lignes : une ajoutée, le compte d'une autre changé
        corr.modifier(self.m, self.m.date, b1, [self.saisie(self.l1, compte="600000", d=10), self.saisie(self.l1, d=2, id=False),
                                                self.saisie(self.l2, c=12)], "ventilation", self.u)
        self.assertEqual(self.m.lignes.count(), 3)

    def test_verrou_exercice_clos(self):
        Exercice.objects.filter(libelle="2026").update(clos=True)
        self.assertIn("exercice clos", corr.verrou(self.m))
        with self.assertRaises(ValueError):
            corr.modifier(self.m, self.m.date, self.m.journal, [self.saisie(self.l1), self.saisie(self.l2)], "x", self.u)

    def test_ecriture_libre(self):
        m = corr.creer(dt.date(2026, 4, 1), Journal.objects.get(code="OD"),
                       [corr.LigneSaisie(None, Compte.objects.get(numero="470000") if Compte.objects.filter(numero="470000").exists()
                                         else Compte.objects.get(numero="600000"), "RECLASSEMENT", D(5), D(0), self.g),
                        corr.LigneSaisie(None, Compte.objects.get(numero="600100"), "RECLASSEMENT", D(0), D(5), self.g)],
                       "reclassement", self.u)
        self.assertEqual((m.journal_id, m.lignes.count()), ("OD", 2))


class EcransCorrections(TestCase):
    def setUp(self):
        Corrections.setUp(self)
        call_command("migrate", verbosity=0)
        self.u.groups.add(Group.objects.get(name="Trésorier"))
        self.client.force_login(self.u)

    def test_formulaire(self):
        self.assertContains(self.client.get("/mouvement/500/"), "Modifier")
        self.assertEqual(self.client.get("/mouvement/500/modifier/").status_code, 200)
        data = {"date": "2026-03-01", "journal": "B1", "motif": "montant", "l-TOTAL_FORMS": "3", "l-INITIAL_FORMS": "2",
                "l-MIN_NUM_FORMS": "0", "l-MAX_NUM_FORMS": "1000",
                "l-0-id": self.l1.pk, "l-0-compte": "600100", "l-0-libelle": "FRAIS", "l-0-debit": "15", "l-0-anal2": "GEN.004",
                "l-1-id": self.l2.pk, "l-1-compte": "512000", "l-1-libelle": "FRAIS", "l-1-credit": "15", "l-1-anal2": "GEN.004"}
        self.assertRedirects(self.client.post("/mouvement/500/modifier/", data), "/mouvement/500/")
        self.assertEqual(Ligne.objects.get(pk=self.l1.pk).debit, D(15))
        self.assertRedirects(self.client.get("/mouvement/rappel/?numero=500"), "/mouvement/500/modifier/")
        self.assertContains(self.client.get("/mouvement/rappel/?numero=9999"), "Aucun mouvement n° 9999")
        self.assertEqual(self.client.get("/mouvement/500/contrepasser/").status_code, 404)
        self.assertEqual(self.client.get("/mouvement/nouveau/").status_code, 200)

    def test_tri_et_droits(self):
        r = self.client.get("/ecritures/?tri=debit&ordre=desc")
        self.assertEqual(r.context["page"][0].debit, D(400))          # plus gros débit en tête
        self.assertContains(r, 'data-sens="desc"')
        b = User.objects.create_user("bureau")
        b.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(b)
        self.assertNotContains(self.client.get("/mouvement/500/"), "Modifier")
        self.assertEqual(self.client.get("/mouvement/500/modifier/").status_code, 403)


class ImportTiers(TestCase):
    def setUp(self):
        referentiels_saisie()
        mbr.creer_manquants()
        Membre.objects.filter(compte_id="411TAIEB001").update(telephone="050-111", ville="Netanya")

    def classeur(self, lignes):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Feuil1"
        ws.append(["Liste des tiers 2026"])                      # une ligne de titre avant les en-têtes
        ws.append(["Nom", "Prénom", "Type", "Adresse", "CP", "Ville", "Tél", "Mail", "Cotisation"])
        for l in lignes:
            ws.append(l)
        tampon = __import__("io").BytesIO()
        wb.save(tampon)
        return tampon.getvalue()

    def test_import_xlsx(self):
        contenu = self.classeur([
            ["Taieb", "Jeanne", "Membre", "3 rue Weizmann", 4250000, None, None, "jeanne@example.org", 450],   # mise à jour
            ["Levy", "Rachel", None, None, None, "Haïfa", "054-222", None, None],                                # nouveau membre
            ["Partner", None, "Fournisseur", "8 Ha-Barzel", "6971005", "Tel Aviv", None, "pas-un-mail", None],   # nouveau fournisseur
            ["Inconnu", None, "Sponsor", None, None, None, None, None, None],                                  # type inconnu
            [None, None, None, None, None, None, None, None, None]])
        r = mbr.importer_tableau(mbr.lire_tableau("Tiers.xlsx", contenu))
        self.assertEqual((len(r.crees), len(r.mis_a_jour), len(r.erreurs)), (2, 1, 2))
        t = Membre.objects.get(compte_id="411TAIEB001")
        self.assertEqual((t.adresse, t.code_postal, t.ville, t.telephone, t.email, t.cotisation),
                         ("3 rue Weizmann", "4250000", "Netanya", "050-111", "jeanne@example.org", D(450)))   # cellules vides : rien d'effacé
        levy = Membre.objects.get(nom="LEVY")
        self.assertEqual((levy.compte_id, levy.type.libelle, levy.ville), ("411LEVY001", "Membre", "Haïfa"))
        p = Membre.objects.get(nom="PARTNER")
        self.assertEqual((p.compte_id, p.email, p.code_postal), ("401PARTN001", "", "6971005"))
        # réimport : plus de création, que des mises à jour
        r = mbr.importer_tableau(mbr.lire_tableau("Tiers.xlsx", contenu))
        self.assertEqual((len(r.crees), len(r.mis_a_jour)), (0, 3))

    def test_ecran_et_modele(self):
        call_command("migrate", verbosity=0)
        u = User.objects.create_user("t")
        u.groups.add(Group.objects.get(name="Trésorier"))
        self.client.force_login(u)
        r = self.client.get("/membres/modele-tiers.xlsx")
        self.assertEqual(r["Content-Disposition"], 'attachment; filename="Tiers.xlsx"')
        from django.core.files.uploadedfile import SimpleUploadedFile as Fichier
        b = User.objects.create_user("bureau")
        b.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(b)
        self.client.post("/membres/", {"fichier": Fichier("Tiers.xlsx", r.content)})          # rôle Bureau : pas d'import
        self.assertFalse(Membre.objects.filter(compte_id="411COHEN001").exists())
        self.client.force_login(u)
        self.client.post("/membres/", {"fichier": Fichier("Tiers.xlsx", r.content)})          # trésorier : pas d'import non plus
        self.assertFalse(Membre.objects.filter(compte_id="411COHEN001").exists())
        a = User.objects.create_user("admin")
        a.groups.add(Group.objects.get(name="Administrateur"))
        self.client.force_login(a)                                                       # import : administrateur
        from django.core.files.uploadedfile import SimpleUploadedFile
        modele = openpyxl.load_workbook(__import__("io").BytesIO(r.content))
        self.assertEqual(modele["Tiers"]["C1"].value, "Nom")
        self.client.post("/membres/", {"fichier": SimpleUploadedFile("Tiers.xlsx", r.content)})
        self.assertTrue(Membre.objects.filter(compte_id="411COHEN001", ville="Netanya").exists())
        self.assertTrue(Membre.objects.filter(compte__numero__startswith="401TRAIT", type__libelle="Fournisseur").exists())


class Journaux(TestCase):
    def setUp(self):
        Corrections.setUp(self)                   # Mvt 500 : frais 10 sur B1 (512000), + facture 421 en B1
        call_command("migrate", verbosity=0)
        self.u.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(self.u)

    def test_journal_banque(self):
        r = self.client.get("/journaux/?journal=B1&du=2026-01-01&au=2026-12-31")
        self.assertTrue(r.context["tresorerie"])
        self.assertEqual((r.context["depenses"], r.context["cloture"]), (D(10), D(-10)))
        self.assertIn("600100", r.context["lignes"][0]["contrepartie"])
        x = self.client.get("/journaux/?journal=B1&du=2026-01-01&au=2026-12-31&format=xlsx")
        from . import dossiers as dos
        self.assertTrue((dos.exports() / "Journal_B1_2026-01-01_2026-12-31.xlsx").exists())      # copie dans Exports
        ws = openpyxl.load_workbook(__import__("io").BytesIO(x.content))["Journal B1"]
        self.assertEqual([c.value for c in ws[1]], ["Date", "Mvt", "Pièce", "Libellé", "Contrepartie", "Recette", "Dépense", "Solde"])

    def test_journal_ventes_et_historique(self):
        r = self.client.get("/journaux/?journal=VT&du=2026-01-01&au=2026-12-31")
        self.assertFalse(r.context["tresorerie"])
        Modification.objects.create(auteur="t", action="Saisie", objet="Mvt 500")
        self.assertContains(self.client.get("/modifications/?q=saisie"), "Mvt 500")
        self.assertEqual(self.client.get("/modifications/excel/").status_code, 200)


# ---------------------------------------------------------------- base de données : sauvegarde, restauration, reprise



from . import base_donnees as bd  # noqa: E402


@override_settings(**temporaire())
class BaseDonnees(TransactionTestCase):
    def setUp(self):
        call_command("migrate", verbosity=0)
        self.admin = User.objects.create_superuser("admin", "", "motdepasse-long")
        self.client.force_login(self.admin)

    def test_trois_sauvegardes_gardees_et_restauration_de_la_plus_ancienne(self):
        import os
        import time
        referentiels()
        mouvement(1, dt.date(2026, 2, 1), [("512000", 100, 0), ("700000", 0, 100)])
        faites = []
        for i in range(5):
            faites.append(bd.sauvegarder(f"essai{i}"))
            os.utime(faites[-1], (time.time() + i, time.time() + i))        # ordre des dates garanti
        self.assertEqual(bd.liste(), faites[:1:-1])                           # les 3 dernières seulement
        mouvement(2, dt.date(2026, 2, 2), [("512000", 5, 0), ("700000", 0, 5)])
        plus_ancienne = bd.liste()[-1]
        n, avant = bd.restaurer(plus_ancienne)                                # sa copie « avant » ne l'efface pas
        self.assertEqual((n, Mouvement.objects.count()), (1, 1))
        self.assertTrue(plus_ancienne.exists() and avant.exists())

    def test_sauvegarde_et_restauration(self):
        referentiels()
        mouvement(1, dt.date(2026, 2, 1), [("512000", 100, 0), ("700000", 0, 100)])
        s = bd.sauvegarder("test")
        self.assertEqual(bd.verifier(s), 1)
        mouvement(2, dt.date(2026, 2, 2), [("512000", 5, 0), ("700000", 0, 5)])
        n, avant = bd.restaurer(s)
        self.assertEqual((n, Mouvement.objects.count(), bd.verifier(avant)), (1, 1, 2))
        faux = Path(tempfile.mkdtemp()) / "faux.sqlite3"
        faux.write_bytes(b"pas une base")
        with self.assertRaises(ValueError):
            bd.restaurer(faux)

    def test_page_et_reprise_du_classeur(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.assertContains(self.client.get("/base/"), "Remettre à zéro")
        self.client.post("/base/", {"sauvegarder": "1"})
        self.assertEqual(len(bd.liste()), 1)
        self.assertEqual(self.client.get("/base/telecharger/").status_code, 200)
        classeur = Reprise.classeur(self)
        refuse = self.client.post("/base/", {"remplacer": "1", "confirmation": "non",
                                             "fichier": SimpleUploadedFile("ComptaBB.xlsx", classeur.read_bytes())})
        self.assertEqual(refuse.status_code, 200)                       # confirmation manquante : rien ne change
        self.client.post("/base/", {"remplacer": "1", "confirmation": "REMPLACER",
                                    "fichier": SimpleUploadedFile("ComptaBB.xlsx", classeur.read_bytes())})
        self.assertEqual(Mouvement.objects.count(), 2)
        self.assertTrue(User.objects.filter(username="admin").exists())      # utilisateurs gardés
        self.assertTrue(Modification.objects.filter(action="Remise à zéro et reprise").exists())

    def test_reserve_a_l_administrateur(self):
        u = User.objects.create_user("tresorier")
        u.groups.add(Group.objects.get(name="Trésorier"))
        self.client.force_login(u)
        self.assertEqual(self.client.get("/base/").status_code, 403)


@override_settings(**temporaire())
class EffacerHistorique(TransactionTestCase):
    def test_effacer(self):
        call_command("migrate", verbosity=0)
        admin = User.objects.create_superuser("admin", "", "motdepasse-long")
        for i in range(3):
            Modification.objects.create(auteur="t", action="Saisie", objet=f"Mvt {i}")
        self.client.force_login(admin)
        self.client.post("/modifications/effacer/", {"jusquau": dt.date.today().isoformat(), "confirmation": "non"})
        self.assertEqual(Modification.objects.count(), 3)                       # pas confirmé
        self.client.post("/modifications/effacer/", {"jusquau": dt.date.today().isoformat(), "confirmation": "effacer"})
        self.assertEqual(list(Modification.objects.values_list("action", flat=True)), ["Historique effacé"])
        from django.conf import settings
        self.assertEqual(len(list((settings.EXPORTS_DIR / "Archives").glob("Historique_*.xlsx"))), 1)
        self.assertEqual(len(bd.liste()), 1)
        u = User.objects.create_user("tresorier")
        u.groups.add(Group.objects.get(name="Trésorier"))
        self.client.force_login(u)
        self.assertEqual(self.client.post("/modifications/effacer/", {"confirmation": "EFFACER"}).status_code, 403)


@override_settings(**temporaire())
class CommandeSauvegarder(TransactionTestCase):
    def test_commande(self):
        call_command("migrate", verbosity=0)
        call_command("sauvegarder", stdout=open("/dev/null", "w"))
        self.assertTrue(bd.liste()[0].name.endswith("_auto.sqlite3"))

    def test_restaurer(self):
        call_command("migrate", verbosity=0)
        referentiels_saisie()
        call_command("sauvegarder", stdout=open("/dev/null", "w"))
        Mouvement.objects.all().delete()
        sortie = __import__("io").StringIO()
        call_command("restaurer", stdout=sortie)
        self.assertIn("1. comptabb_", sortie.getvalue())
        call_command("restaurer", "1", oui=True, stdout=sortie)
        self.assertTrue(Mouvement.objects.filter(numero=421).exists())
        self.assertTrue(Modification.objects.filter(action="Restauration").exists())
        from django.core.management.base import CommandError
        with self.assertRaises(CommandError):
            call_command("restaurer", "99", oui=True, stdout=sortie)


from . import parametres as prm  # noqa: E402


@override_settings(**temporaire())
class ImportParametres(TestCase):
    def setUp(self):
        referentiels_saisie()
        moteur.initialiser_parametres()
        call_command("migrate", verbosity=0)

    def classeur(self):
        return openpyxl.load_workbook(__import__("io").BytesIO(prm.contenu_classeur()))

    def octets(self, wb):
        tampon = __import__("io").BytesIO()
        wb.save(tampon)
        return tampon.getvalue()

    def test_aller_retour_sans_changement(self):
        wb = self.classeur()
        self.assertIn("Plan comptable", wb.sheetnames)
        comptes = [r[0] for r in wb["Plan comptable"].iter_rows(min_row=2, values_only=True)]
        self.assertIn("512000", comptes)
        self.assertNotIn("411TAIEB001", comptes)                 # comptes de tiers : fichier Tiers.xlsx
        r = prm.importer(self.octets(wb))
        self.assertEqual((r.erreurs, r.crees, r.modifies), ([], [], []))
        self.assertGreater(r.inchanges, 30)

    def test_modification_et_creation(self):
        wb = self.classeur()
        pc = wb["Plan comptable"]
        for row in pc.iter_rows(min_row=2):
            if row[0].value == "600100":
                row[1].value, row[4].value = "FRAIS DE BANQUE", "Non"
            if row[0].value == "610000":
                row[1].value = None                               # cellule vide : rien ne change
        pc.append([622000, "HONORAIRES", "FON.1", "Non", "Oui"])    # nombre lu par Excel
        wb["Axe 2"].append(["MAN.002", "GALA", "En cours"])
        wb["Traductions"].append(["עמלה", "COMMISSION"])
        r = prm.importer(self.octets(wb), auteur="t")
        self.assertEqual(r.erreurs, [])
        self.assertEqual((len(r.crees), len(r.modifies)), (3, 1))
        c = Compte.objects.get(numero="600100")
        self.assertEqual((c.libelle, c.actif), ("FRAIS DE BANQUE", False))
        self.assertEqual(Compte.objects.get(numero="610000").libelle, "SERVICES")
        self.assertEqual(Compte.objects.get(numero="622000").anal1_id, "FON.1")
        self.assertEqual(CodeAnalytique.objects.get(code="MAN.002").axe, 2)
        self.assertEqual(Traduction.traduire("עמלה"), "COMMISSION")
        m = Modification.objects.get(action="Modification (import)")
        self.assertIn("SERVICES", Compte.objects.get(numero="610000").libelle)
        self.assertIn("FRAIS BANCAIRES", m.avant)
        self.assertTrue(Modification.objects.filter(action="Import Parametres.xlsx").exists())

    def test_erreur_rien_n_est_enregistre(self):
        wb = self.classeur()
        wb["Plan comptable"].append(["623000", "PUBLICITE", "ZZZ.9", None, None])   # code axe 1 inconnu
        wb["Axe 1"].append(["GEN.001", "DOUBLON", None])                            # code déjà sur l'axe 2
        wb["Journaux"].append(["BQ", None, None, None, None])                       # intitulé manquant
        wb["Réglages"].append(["compte_attente", "470000", "Compte d'attente"])     # ligne correcte
        r = prm.importer(self.octets(wb))
        self.assertEqual(len(r.erreurs), 3, r.erreurs)
        self.assertTrue(any("Plan comptable, ligne" in e and "ZZZ.9" in e for e in r.erreurs))
        self.assertFalse(Reglage.objects.filter(cle="compte_attente").exists())    # tout ou rien
        self.assertFalse(Compte.objects.filter(numero="623000").exists())
        self.assertEqual(CodeAnalytique.objects.get(code="GEN.001").libelle, "GENERAL")



@override_settings(**temporaire())
class EcranParametres(TransactionTestCase):
    """Imports avec sauvegarde SQLite préalable : hors transaction de test (comme BaseDonnees)."""

    def setUp(self):
        call_command("migrate", verbosity=0)
        referentiels_saisie()
        moteur.initialiser_parametres()

    octets = ImportParametres.octets

    def test_droits_et_dossier_imports(self):
        from django.conf import settings
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.assertEqual(set(Group.objects.values_list("name", flat=True)), {"Administrateur", "Trésorier", "Bureau", "Bénévole"})
        for role in ("Bureau", "Trésorier", "Bénévole"):
            self.assertFalse(Group.objects.get(name=role).permissions.filter(codename="parametrer").exists())
            u = User.objects.create_user(role)
            u.groups.add(Group.objects.get(name=role))
            self.client.force_login(u)
            self.assertEqual(self.client.get("/parametres/").status_code, 403)
            self.assertEqual(self.client.get("/parametres/Parametres.xlsx").status_code, 403)
            self.assertEqual(self.client.post("/parametres/", {"preparer": "1"}).status_code, 403)
        self.assertFalse((settings.IMPORTS_DIR / "Parametres.xlsx").exists())
        from .vues_utilisateurs import donner_role
        u = User.objects.create_user("admin")
        donner_role(u, "Administrateur")
        self.client.force_login(u)
        self.assertContains(self.client.get("/"), "Paramètres (Excel)")
        r = self.client.get("/parametres/Parametres.xlsx")
        self.assertEqual(r["Content-Disposition"], 'attachment; filename="Parametres.xlsx"')
        self.assertTrue((settings.EXPORTS_DIR / "Parametres.xlsx").exists())
        wb = openpyxl.load_workbook(__import__("io").BytesIO(r.content))
        self.assertEqual(wb["Plan comptable"]["A1"].font.size, 12)
        wb["Réglages"].append(["compte_attente", "470000", ""])
        self.client.post("/parametres/", {"fichier": SimpleUploadedFile("Parametres.xlsx", self.octets(wb))})
        self.assertEqual(Reglage.objects.get(cle="compte_attente").valeur, "470000")
        self.assertEqual(bd.dossier(), settings.EXPORTS_DIR / "Sauvegardes")
        self.assertTrue(any(p.name.endswith("_avant_import_parametres.sqlite3") for p in bd.liste()))
        # par le dossier Imports : préparer, modifier dans « Excel », importer
        chemin = settings.IMPORTS_DIR / "Parametres.xlsx"
        self.client.post("/parametres/", {"preparer": "1"})
        self.client.post("/parametres/", {"preparer": "1"})           # l'ancien fichier est gardé dans Exports
        self.assertEqual(len(list(settings.EXPORTS_DIR.glob("Parametres_*_remplace.xlsx"))), 1)
        wb = openpyxl.load_workbook(chemin)
        wb["Réglages"].append(["compte_don", "750000", "Dons"])
        wb.save(chemin)
        self.client.post("/parametres/", {"importer_dossier": "1"})
        self.assertEqual(Reglage.objects.get(cle="compte_don").valeur, "750000")

    def test_anciennes_sauvegardes_deplacees(self):
        from django.conf import settings
        ancien = settings.DATA_DIR / "sauvegardes"
        ancien.mkdir(parents=True, exist_ok=True)
        (ancien / "comptabb_ancienne.sqlite3").write_bytes(b"x")
        self.assertIn("comptabb_ancienne.sqlite3", [p.name for p in bd.liste()])
        self.assertFalse(ancien.exists())

    def test_commande(self):
        chemin = Path(tempfile.mkdtemp()) / "Parametres.xlsx"
        call_command("parametres", "exporter", str(chemin), stdout=open("/dev/null", "w"))
        wb = openpyxl.load_workbook(chemin)
        wb["Types de tiers"].append(["Donateur", "412"])
        wb.save(chemin)
        call_command("parametres", "importer", str(chemin), stdout=open("/dev/null", "w"))
        self.assertEqual(TypeTiers.objects.get(libelle="Donateur").prefixe, "412")


from . import export_complet as ec  # noqa: E402


@override_settings(**temporaire())
class ExportComplet(TransactionTestCase):
    """Tout exporter, remettre à zéro, réinjecter : la base doit redevenir identique."""

    def setUp(self):
        call_command("migrate", verbosity=0)
        referentiels_saisie()
        moteur.initialiser_parametres()
        self.u = User.objects.create_user("tresorier", is_superuser=True, is_staff=True)
        benevole = User.objects.create_user("rachel")
        ex = Exercice.objects.create(libelle="2026", debut=dt.date(2026, 1, 1), fin=dt.date(2026, 12, 31))
        m = Mouvement.objects.create(numero=1, date=dt.date(2026, 2, 1), journal_id="B1", piece=1, origine="saisie",
                                     commentaire="Cotisation", cree_par=self.u)
        p = Rapprochement.objects.create(journal_id="B1", mode="manuel", cree_par=self.u)
        Ligne.objects.create(mouvement=m, ordre=1, compte_id="411TAIEB001", libelle="COTISATION", credit=D(400), lettrage="A",
                             anal2_id="GEN.002")
        Ligne.objects.create(mouvement=m, ordre=2, compte_id="512000", libelle="COTISATION", debit=D(400), rapprochement=p,
                             anal2_id="GEN.002")
        LigneReleve.objects.create(journal_id="B1", date=dt.date(2026, 2, 2), rang=0, operation="העברה", montant=D(400),
                                   solde=D(1400), rapprochement=p)
        Budget.objects.create(exercice=ex, nature="P", compte_id="700000", montant=D(5000))
        t = TiersProvisoire.objects.create(nom="Lévy", prenom="Rachel", cree_par=benevole)
        f = Fiche.objects.create(type="activite", titre="Rallye", anal2_id="MAN.001")
        f.benevoles.add(benevole)
        nature = NatureFiche.objects.filter(type_fiche="activite", sens="R").first()
        mode = ModeFiche.objects.filter(type_fiche="activite").first()
        LigneFiche.objects.create(fiche=f, sens="R", date=dt.date(2026, 3, 1), provisoire=t, nature=nature, montant=D(80),
                                  mode=mode, cree_par=benevole)
        Membre.objects.filter(compte_id="411TAIEB001").update(ville="Netanya", cotisation=D(400), utilisateur=benevole)
        Modification.objects.create(auteur="tresorier", action="Essai", objet="avant export")

    def test_aller_retour_identique(self):
        avant = ec.empreinte()
        chemin = ec.exporter(auteur="tresorier")
        with zipfile.ZipFile(chemin) as z:
            self.assertTrue({"comptabb.sqlite3", "Parametres.xlsx", "Tiers.xlsx", "Ecritures.xlsx", "Donnees.xlsx",
                             "Etats_2026.xlsx", "controle.json"} <= set(z.namelist()))
            self.assertEqual(openpyxl.load_workbook(io.BytesIO(z.read("Ecritures.xlsx"))).sheetnames, ["Écritures"])
            self.assertNotIn("Écritures", openpyxl.load_workbook(io.BytesIO(z.read("Donnees.xlsx"))).sheetnames)
        Mouvement.objects.create(numero=2, date=dt.date(2026, 4, 1), journal_id="OD", piece=2)      # après l'export
        Compte.objects.filter(numero="600000").update(libelle="MODIFIÉ")
        message, _ = ec.reinjecter(chemin, auteur="tresorier")
        self.assertIn("identique", message)
        self.assertTrue(Modification.objects.filter(action="Réinjection d'un export complet").exists())
        Modification.objects.filter(action="Réinjection d'un export complet").delete()
        self.assertEqual(ec.ecarts(avant, ec.empreinte()), [])          # même contenu, historique compris
        self.assertFalse(Mouvement.objects.filter(numero=2).exists())
        self.assertEqual(Compte.objects.get(numero="600000").libelle, "ACHATS")
        self.assertEqual(LigneReleve.objects.get().rang, 0)
        l = Ligne.objects.get(compte_id="512000")
        self.assertEqual(l.rapprochement.releves.get().montant, D(400))
        self.assertEqual(Fiche.objects.get().benevoles.get().username, "rachel")
        self.assertEqual(LigneFiche.objects.get().provisoire.nom, "Lévy")
        self.assertEqual(Membre.objects.get(compte_id="411TAIEB001").ville, "Netanya")
        self.assertEqual(Membre.objects.get(compte_id="411TAIEB001").utilisateur.username, "rachel")
        self.assertEqual(Mouvement.objects.get(numero=1).cree_par, self.u)

    def modifier(self, chemin, nom, changer, dossier="", retirer=()):
        """Copie de l'export, le classeur nom modifié par changer(wb) ; les fichiers sont rangés dans dossier (recompression)."""
        altere = chemin.with_name(f"Export_complet_altere_{len(dossier)}.zip")
        with zipfile.ZipFile(chemin) as z, zipfile.ZipFile(altere, "w") as sortie:
            for n in z.namelist():
                if n in retirer:
                    continue
                contenu = z.read(n)
                if n == nom:
                    wb = openpyxl.load_workbook(io.BytesIO(contenu))
                    changer(wb)
                    tampon = io.BytesIO()
                    wb.save(tampon)
                    contenu = tampon.getvalue()
                sortie.writestr(dossier + n, contenu)
        return altere

    def test_export_altere_rien_n_est_modifie(self):
        chemin = ec.exporter()
        altere = self.modifier(chemin, "Ecritures.xlsx", lambda wb: wb["Écritures"].__setitem__("K2", "LIBELLÉ CHANGÉ"))
        Compte.objects.filter(numero="600000").update(libelle="GARDÉ")
        with self.assertRaisesRegex(ec.ExportInvalide, "Fichiers modifiés"):
            ec.reinjecter(altere)
        self.assertEqual(Compte.objects.get(numero="600000").libelle, "GARDÉ")              # tout ou rien
        with self.assertRaises(ec.ExportInvalide):
            ec.reinjecter(b"pas un zip")

    def test_fichiers_modifies_recharges_apres_remise_a_zero(self):
        numeros = sorted(set(Mouvement.objects.values_list("numero", flat=True)) | {5})
        chemin = ec.exporter()

        def changer(wb):
            ws = wb["Écritures"]
            ws["K2"] = ws["K3"] = "COTISATION 2026"                    # libellés corrigés
            ws["L3"], ws["M2"] = 450, 450                              # montant corrigé, Mvt toujours équilibré
            ws.append([5, dt.datetime(2026, 5, 1), "OD", 5, "", "", None, "", 1, "600000", "AJOUT", 30, None, "GEN.002"])
            ws.append([5, dt.datetime(2026, 5, 1), "OD", 5, "", "", None, "", 2, "512000", "AJOUT", None, 30, "GEN.002"])
        modifie = self.modifier(chemin, "Ecritures.xlsx", changer, dossier="Export_complet_corrige/",
                                retirer=("controle.json",))
        with self.assertRaises(ec.ExportInvalide):
            ec.reinjecter(modifie)                                      # sans la case « Fichiers modifiés »
        message, _ = ec.reinjecter(modifie, modifie=True)
        self.assertIn("cohérence", message)
        self.assertEqual(sorted(Mouvement.objects.values_list("numero", flat=True)), numeros)
        self.assertEqual(set(Ligne.objects.filter(mouvement__numero=1).values_list("libelle", "debit", "credit")),
                         {("COTISATION 2026", D(450), D(0)), ("COTISATION 2026", D(0), D(450))})
        self.assertEqual(Fiche.objects.get().titre, "Rallye")                                   # le reste est rechargé
        self.assertEqual(Ligne.objects.get(compte_id="512000", mouvement__numero=1).rapprochement.releves.count(), 1)
        # déséquilibré ou compte inconnu : refusé, rien n'est modifié
        for cellule, valeur in (("L2", 999), ("J2", "999999")):
            faux = self.modifier(chemin, "Ecritures.xlsx", lambda wb: wb["Écritures"].__setitem__(cellule, valeur))
            with self.assertRaises(ec.ExportInvalide):
                ec.reinjecter(faux, modifie=True)
            self.assertEqual(sorted(Mouvement.objects.values_list("numero", flat=True)), numeros)

    def test_export_anterieur_sans_ecritures_xlsx(self):
        """Exports d'avant Ecritures.xlsx : les écritures sont dans la feuille Écritures de Donnees.xlsx."""
        avant = ec.empreinte()
        chemin = ec.exporter()
        ancien = chemin.with_name("Export_complet_ancien.zip")
        with zipfile.ZipFile(chemin) as z, zipfile.ZipFile(ancien, "w") as sortie:
            for n in z.namelist():
                if n == "Donnees.xlsx":
                    wb = openpyxl.load_workbook(io.BytesIO(z.read(n)))
                    ws = wb.create_sheet("Écritures", 2)
                    for r in openpyxl.load_workbook(io.BytesIO(z.read("Ecritures.xlsx")))["Écritures"].iter_rows(values_only=True):
                        ws.append(r)
                    sortie.writestr(n, ec._octets(wb))
                elif n != "Ecritures.xlsx":
                    sortie.writestr(n, z.read(n))
        ec.reinjecter(ancien)
        Modification.objects.filter(action="Réinjection d'un export complet").delete()
        self.assertEqual(ec.ecarts(avant, ec.empreinte()), [])

    def test_page_base(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.u)
        self.client.post("/base/", {"export_complet": "1"})
        nom = ec.liste()[0].name
        r = self.client.get(f"/base/export/{nom}")
        contenu = b"".join(r.streaming_content)
        Mouvement.objects.create(numero=9, date=dt.date(2026, 4, 1), journal_id="OD", piece=9)
        self.client.post("/base/", {"remplacer": "1", "sauvegarde": nom, "confirmation": "REMPLACER"})
        self.assertFalse(Mouvement.objects.filter(numero=9).exists())
        Mouvement.objects.create(numero=9, date=dt.date(2026, 4, 1), journal_id="OD", piece=9)
        self.client.post("/base/", {"remplacer": "1", "fichier": SimpleUploadedFile(nom, contenu), "confirmation": "REMPLACER"})
        self.assertFalse(Mouvement.objects.filter(numero=9).exists())
        b = User.objects.create_user("bureau")
        b.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(b)
        self.assertEqual(self.client.get(f"/base/export/{nom}").status_code, 403)


# ---------------------------------------------------------------- rôles et comptes de connexion

from .vues_utilisateurs import donner_role  # noqa: E402


class Utilisateurs(TestCase):
    def setUp(self):
        referentiels_saisie()
        Prefixe.objects.create(prefixe="MAN.", axe=2)
        Prefixe.objects.create(prefixe="COT.", axe=1)
        self.admin = User.objects.create_user("admin", password="ancien-mot-de-passe-1")
        donner_role(self.admin, "Administrateur")
        self.tresorier = User.objects.create_user("tresorier", password="ancien-mot-de-passe-1")
        donner_role(self.tresorier, "Trésorier")

    def test_tresorier_sans_parametrage_de_base(self):
        self.client.force_login(self.tresorier)
        for url in ("/utilisateurs/", f"/utilisateurs/{self.admin.pk}/", "/parametres/", "/base/", "/admin/journal/"):
            self.assertIn(self.client.get(url).status_code, (302, 403), url)
        self.assertEqual(self.client.post("/rapprochement/B1/parametres/", {"date_reprise": "2026-01-01"}).status_code, 403)
        r = self.client.get("/codes/")
        self.assertNotContains(r, '<option value="COT."')                     # axe 1 : administrateur
        self.client.post("/codes/", {"code-prefixe": "COT.", "code-libelle": "X", "code-statut": 1, "creer_code": "1"})
        self.assertFalse(CodeAnalytique.objects.filter(libelle="X").exists())
        self.client.post("/codes/", {"code-prefixe": "MAN.", "code-libelle": "Gala", "code-statut": 1, "creer_code": "1"})
        self.assertTrue(CodeAnalytique.objects.filter(code="MAN.002").exists())       # axe 2 : trésorier
        self.assertEqual(self.client.get("/saisie/").status_code, 200)
        self.client.force_login(self.admin)
        self.assertContains(self.client.get("/codes/"), '<option value="COT."')
        self.assertEqual(self.client.get("/utilisateurs/").status_code, 200)

    def test_mon_compte_identifiant_email_et_mot_de_passe(self):
        self.client.force_login(self.tresorier)
        self.assertEqual(self.client.get("/mon-compte/").status_code, 200)
        r = self.client.post("/mon-compte/", {"id-identifiant": "admin", "id-email": "", "identifiant_maj": "1"})
        self.assertContains(r, "déjà utilisé")
        self.client.post("/mon-compte/", {"id-identifiant": "tresor@exemple.org", "id-email": "tresor@exemple.org",
                                          "identifiant_maj": "1"})
        self.tresorier.refresh_from_db()
        self.assertEqual(self.tresorier.username, "tresor@exemple.org")
        self.client.post("/mon-compte/", {"mdp-old_password": "ancien-mot-de-passe-1", "mdp-new_password1": "Nouveau-mdp-2026!",
                                          "mdp-new_password2": "Nouveau-mdp-2026!", "mot_de_passe_maj": "1"})
        self.client.logout()
        self.assertTrue(self.client.login(username="tresor@exemple.org", password="Nouveau-mdp-2026!"))
        self.admin.email = "chef@exemple.org"
        self.admin.save()
        self.assertTrue(self.client.login(username="chef@exemple.org", password="ancien-mot-de-passe-1"))   # connexion par e-mail

    def test_administrateur_gere_les_utilisateurs(self):
        self.client.force_login(self.admin)
        self.client.post("/utilisateurs/", {"n-identifiant": "bureau@exemple.org", "n-email": "bureau@exemple.org",
                                            "n-role": "Bureau", "n-mot_de_passe": "Bureau-mdp-2026!", "creer": "1"})
        b = User.objects.get(username="bureau@exemple.org")
        self.assertEqual(list(b.groups.values_list("name", flat=True)), ["Bureau"])
        self.assertFalse(b.has_perm("compta.add_mouvement"))
        self.assertTrue(b.has_perm("compta.view_mouvement"))
        self.client.post(f"/utilisateurs/{self.tresorier.pk}/", {"identifiant": "tresorier", "email": "", "role": "Administrateur",
                                                                 "actif": "on", "mot_de_passe": "", "prenom": "Jean-Luc",
                                                                 "nom": "CHABBAT"})
        self.tresorier.refresh_from_db()
        self.assertEqual(self.tresorier.get_full_name(), "Jean-Luc CHABBAT")                    # nom modifiable après création
        self.assertContains(self.client.get(f"/utilisateurs/{self.tresorier.pk}/"), 'value="CHABBAT"')
        self.assertTrue(self.tresorier.is_superuser)
        self.assertEqual(list(self.tresorier.groups.values_list("name", flat=True)), ["Administrateur"])
        # toujours au moins un administrateur actif
        self.client.post(f"/utilisateurs/{self.tresorier.pk}/", {"identifiant": "tresorier", "role": "Trésorier", "actif": "on"})
        r = self.client.post(f"/utilisateurs/{self.admin.pk}/", {"identifiant": "admin", "role": "Bureau", "actif": "on"})
        self.assertContains(r, "au moins un administrateur")
        self.assertTrue(User.objects.get(pk=self.admin.pk).is_superuser)


from . import echanges as ech  # noqa: E402


@override_settings(DATA_DIR=Path(tempfile.mkdtemp()))
class Echanges(TransactionTestCase):
    """Imports et exports par fichiers .xlsx des dossiers Imports et Exports."""

    def setUp(self):
        self.dossier = Path(tempfile.mkdtemp())
        self.reglage = override_settings(IMPORTS_DIR=self.dossier / "Imports", EXPORTS_DIR=self.dossier / "Exports")
        self.reglage.enable()
        call_command("migrate", verbosity=0)
        referentiels_saisie()
        mbr.creer_manquants()
        Prefixe.objects.create(prefixe="MAN.", axe=2, libelle="Manifestations")
        u = User.objects.create_user("admin")
        donner_role(u, "Administrateur")                    # page réservée : elle importe le paramétrage de base
        self.client.force_login(u)

    def tearDown(self):
        self.reglage.disable()

    def fichier(self, nom, lignes, colonnes=None):
        f = ech.format_de(nom)
        wb = ech.classeur(f, lignes)
        if colonnes:
            for i, c in enumerate(colonnes, 1):
                wb.active.cell(1, i, c)
        wb.save(ech.imports() / nom)

    def importer(self, nom):
        return self.client.post("/echanges/", {"importer": nom}, follow=True)

    def test_aller_retour_des_referentiels(self):
        self.client.post("/echanges/", {"exporter": "tout"})
        exportes = sorted(p.name for p in ech.exports().iterdir() if not p.name.startswith("Lexique"))
        self.assertEqual(len(exportes), len(ech.FORMATS))
        avant = (Compte.objects.count(), Journal.objects.count(), CodeAnalytique.objects.count(), Membre.objects.count())
        for f in ech.FORMATS:
            if f.nom in ("Ecritures", "Banque1", "Banque2", "Bit", "Caisse", "Budget", "Traductions"):
                continue                        # écritures : Mvt déjà présents ; relevés et budget vides
            source = next(p for p in ech.exports().iterdir() if f.reconnait(p))
            (ech.imports() / source.name).write_bytes(source.read_bytes())
            r = self.importer(source.name)
            self.assertContains(r, f"{source.name} importé (", msg_prefix=f.nom)
        self.assertEqual(avant, (Compte.objects.count(), Journal.objects.count(), CodeAnalytique.objects.count(), Membre.objects.count()))
        self.assertEqual(ech.a_importer(), [])                                    # rangés dans Importés
        self.assertEqual(len(list((ech.imports() / "Importés").iterdir())), 9)
        wb = openpyxl.load_workbook(next(p for p in ech.exports().iterdir() if p.name.startswith("Ecritures")))
        self.assertEqual([c.value for c in wb.active[1]], ["Date", "Jnl", "Mvt", "Pièce", "Compte", "Libellé", "Débit", "Crédit", "Anal2", "Let"])
        self.assertEqual(wb.active["C2"].value, 421)

    def test_refus_tout_ou_rien(self):
        self.fichier("PlanComptable.xlsx", [["999000", "NOUVEAU", "", "non", "oui"]], colonnes=["Compte", "Intitule"])
        r = self.importer("PlanComptable.xlsx")
        self.assertContains(r, "En-têtes de la ligne 1 non conformes")
        self.assertTrue((ech.imports() / "PlanComptable.xlsx").exists())            # reste dans Imports
        self.fichier("PlanComptable.xlsx", [["999000", "NOUVEAU", "", "non", "oui"], ["999001", "AUTRE", "XXX.9", "non", "oui"]])
        self.assertContains(self.importer("PlanComptable.xlsx"), "Ligne 3 : code axe 1 « XXX.9 » inconnu")
        self.assertFalse(Compte.objects.filter(numero="999000").exists())          # rien d'enregistré
        openpyxl.Workbook().save(ech.imports() / "Inconnu.xlsx")
        self.assertContains(self.client.get("/echanges/"), "nom non reconnu")

    def test_ecritures(self):
        d = dt.datetime(2026, 3, 1)
        lignes = [[d, "B1", 900, 900, "600100", "FRAIS", 10, None, "GEN.004", ""],
                  [d, "B1", 900, 900, "512000", "FRAIS", None, 9, "GEN.004", ""]]
        self.fichier("Ecritures.xlsx", lignes)
        self.assertContains(self.importer("Ecritures.xlsx"), "Mvt 900 déséquilibré")
        lignes[1][7] = 10
        self.fichier("Ecritures.xlsx", lignes)
        self.assertContains(self.importer("Ecritures.xlsx"), "1 mouvement(s) ajouté(s), 0 modifié(s), 0 inchangé(s)")
        m = Mouvement.objects.get(numero=900)
        self.assertEqual((m.origine, m.total_debit, m.total_credit), ("import", D(10), D(10)))

    def test_ecritures_modifiees_a_la_main(self):
        """Exporter, corriger le fichier dans Excel, le réinjecter : seuls les Mvt changés sont mis à jour."""
        chemin, _ = ech.exporter(ech.PAR_NOM["Ecritures"])
        wb = openpyxl.load_workbook(chemin)
        ws = wb.active
        ws["F2"], ws["F3"] = "FACTURE CORRIGEE", "FACTURE CORRIGEE"          # Mvt 421 : libellés
        ws["G2"], ws["H3"] = 450, 450                                          # et montant
        wb.save(ech.imports() / chemin.name)
        r = self.importer(chemin.name)
        self.assertContains(r, "0 mouvement(s) ajouté(s), 1 modifié(s), 0 inchangé(s)")
        m = Mouvement.objects.get(numero=421)
        self.assertEqual((m.total_debit, m.lignes.first().libelle), (D(450), "FACTURE CORRIGEE"))
        self.assertIn("import Ecritures_", m.commentaire)
        self.assertTrue(Modification.objects.filter(action="Modification", objet__startswith="Mvt 421").exists())
        # réinjecter le même fichier : rien ne change
        (ech.imports() / "Ecritures.xlsx").write_bytes(next(ech.importes().glob("Ecritures_*")).read_bytes())
        self.assertContains(self.importer("Ecritures.xlsx"), "0 mouvement(s) ajouté(s), 0 modifié(s), 1 inchangé(s)")
        # exercice clos : modification refusée
        Exercice.objects.filter(libelle="2026").update(clos=True)
        ws["G2"], ws["H3"] = 460, 460
        wb.save(ech.imports() / "Ecritures.xlsx")
        self.assertContains(self.importer("Ecritures.xlsx"), "non modifiable")
        self.assertEqual(Mouvement.objects.get(numero=421).total_debit, D(450))
        ws["G2"], ws["H3"] = 450, 450                                          # inchangé : accepté malgré la clôture
        wb.save(ech.imports() / "Ecritures.xlsx")
        self.assertContains(self.importer("Ecritures.xlsx"), "1 inchangé(s)")

    def test_ecritures_en_double_refusees(self):
        """Un fichier dont les n° de Mvt ne sont pas ceux du site ne crée pas de doublons."""
        chemin, _ = ech.exporter(ech.PAR_NOM["Ecritures"])
        wb = openpyxl.load_workbook(chemin)
        ws = wb.active
        ws["C2"], ws["C3"], ws["D2"], ws["D3"] = 5000, 5000, 5000, 5000      # Mvt 421 renuméroté : même contenu
        ws["F2"], ws["F3"] = "AUTRE LIBELLE", "AUTRE LIBELLE"
        wb.save(ech.imports() / chemin.name)
        n = Mouvement.objects.count()
        r = self.importer(chemin.name)
        self.assertContains(r, "Mvt 5000 absent du site, mais identique au Mvt 421")
        self.assertContains(r, "rien n&#x27;a été enregistré")
        self.assertEqual(Mouvement.objects.count(), n)
        ws["G2"], ws["H3"] = 999, 999                                          # autre montant : Mvt vraiment nouveau
        wb.save(ech.imports() / chemin.name)
        self.assertContains(self.importer(chemin.name), "1 mouvement(s) ajouté(s), 0 modifié(s)")
        self.assertEqual(Mouvement.objects.count(), n + 1)

    def test_kit_des_modeles_vierges(self):
        self.assertContains(self.client.get("/echanges/"), "Modèles vierges (ZIP)")
        r = self.client.get("/echanges/modeles.zip")
        with zipfile.ZipFile(io.BytesIO(r.content)) as z:
            noms = z.namelist()
            self.assertIn("Lexique.xlsx", noms)
            self.assertIn("LISEZMOI.txt", noms)
            self.assertEqual(len([n for n in noms if n[:2].isdigit()]), len(ech.FORMATS))
            tiers = next(n for n in noms if n.endswith("_Tiers.xlsx"))
            ws = openpyxl.load_workbook(io.BytesIO(z.read(tiers))).active
            self.assertEqual((ws.max_row, ws["C1"].value), (1, "Nom"))                      # vierge : les colonnes seulement

    def test_supprimer_un_fichier_du_dossier_imports(self):
        self.fichier("Axe1.xlsx", [["ZZZ.1", "Essai"]])
        page = self.client.get("/echanges/")
        self.assertContains(page, 'name="retirer" value="Axe1.xlsx"')
        r = self.client.post("/echanges/", {"retirer": "Axe1.xlsx"}, follow=True)
        self.assertContains(r, "supprimé du dossier Imports")
        self.assertFalse((ech.imports() / "Axe1.xlsx").exists())
        self.assertFalse(CodeAnalytique.objects.filter(code="ZZZ.1").exists())               # rien d'importé
        self.assertContains(self.client.post("/echanges/", {"retirer": "../secret.txt"}, follow=True), "pas (ou plus)")

    def test_libelles_seulement_mvt_retrouves_par_leur_contenu(self):
        """Fichier venu d'une autre base : n° décalés ; seuls les libellés sont repris, Mvt retrouvés par leur contenu."""
        chemin, _ = ech.exporter(ech.PAR_NOM["Ecritures"])
        wb = openpyxl.load_workbook(chemin)
        ws = wb.active
        premier = ws["C2"].value
        m = Mouvement.objects.get(numero=premier)
        avant = {l.pk: (l.compte_id, l.debit, l.credit, l.rapprochement_id, l.lettrage) for l in m.lignes.all()}
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 3).value == premier:
                ws.cell(r, 3).value = premier + 7000                      # autre numérotation
                ws.cell(r, 4).value = 1
                ws.cell(r, 6).value = f"LIBELLE BANQUE {r}"
        ws.append([dt.datetime(2020, 5, 5), "OD", 8888, 1, "600100", "INCONNU", 1, None, ws["I2"].value, None])
        ws.append([dt.datetime(2020, 5, 5), "OD", 8888, 1, "512000", "INCONNU", None, 1, ws["I2"].value, None])
        # Mvt dont le montant diffère sur le site : le rapport propose le Mvt du site et la différence
        m2 = m                                                            # seul Mvt de la base d'essai
        for i, l in enumerate(m2.lignes.order_by("ordre")):
            ws.append([m2.date, m2.journal_id, 9999, m2.piece, l.compte_id, "LIBELLE MODIFIE",
                       (l.debit + 1) if l.debit else None, (l.credit + 1) if l.credit else None, l.anal2_id, None])
        wb.save(ech.imports() / "Libelles_2026-09-27.xlsx")
        self.assertEqual(ech.format_de("Libellés_2026-09-27.xlsx").nom, "Libelles")          # accent et majuscule admis
        self.assertEqual(ech.format_de("LIBELLES.xlsx").nom, "Libelles")
        self.assertEqual(ech.format_de("Ecritures_2026-09-27_libelles_banque.xlsx").nom, "Ecritures")
        n = Mouvement.objects.count()
        r = self.importer("Libelles_2026-09-27.xlsx")
        self.assertContains(r, "1 mouvement(s) : libellés mis à jour")
        self.assertContains(r, "2 introuvable(s) sur le site")
        self.assertContains(r, "Mvt du fichier 8888")
        rapport = next(ech.exports().glob("Ecarts_libelles_*.xlsx"))
        feuille = openpyxl.load_workbook(rapport)["Introuvables"]
        self.assertEqual(feuille["A2"].value, 8888)
        self.assertIn("absent du site", feuille["K2"].value)
        self.assertEqual((feuille["A3"].value, feuille["F3"].value), (9999, m2.numero))
        self.assertIn("montant", feuille["K3"].value)
        self.assertEqual(Mouvement.objects.count(), n)                                  # rien de créé
        m.refresh_from_db()
        self.assertTrue(all(l.libelle.startswith("LIBELLE BANQUE") for l in m.lignes.all()))
        self.assertEqual(avant, {l.pk: (l.compte_id, l.debit, l.credit, l.rapprochement_id, l.lettrage) for l in m.lignes.all()})
        self.assertEqual(m.numero, premier)                                             # n° du site gardé
        self.assertTrue(Modification.objects.filter(action="Libellés mis à jour", objet__contains=f"fichier : Mvt {premier + 7000}").exists())
        wb.save(ech.imports() / "Libelles_2026-09-27.xlsx")                            # réimport : déjà à jour
        self.assertContains(self.importer("Libelles_2026-09-27.xlsx"), "0 mouvement(s) : libellés mis à jour")
        # exercice clos : libellés non modifiés, signalés
        Exercice.objects.filter(debut__lte=m.date, fin__gte=m.date).update(clos=True)
        ws["F2"] = "ENCORE AUTRE"
        wb.save(ech.imports() / "Libelles_2026-09-27.xlsx")
        self.assertContains(self.importer("Libelles_2026-09-27.xlsx"), "1 dans un exercice clos, non modifié(s)")

    def test_bit_remplace_le_releve(self):
        d = dt.datetime(2026, 1, 11)
        self.fichier("Bit.xlsx", [["B2", d, "TAIEB JEANNE - FETES - RACLETTE", 440, None]])
        self.assertContains(self.importer("Bit.xlsx"), "journal « B2 » : B3 attendu")
        self.fichier("Bit.xlsx", [["B3", d, "TAIEB JEANNE - FETES - RACLETTE", 440, None],
                                  ["B3", d, "BANQUE - BANQUE - VIREMENT BIT VERS BANQUE", None, 6820]])
        self.assertContains(self.importer("Bit.xlsx"), "2 ligne(s) importée(s), 0 ancienne(s) remplacée(s)")
        self.assertEqual(sorted(LigneReleve.objects.filter(journal_id="B3", ouverture=False).values_list("montant", flat=True)),
                         [D(-6820), D(440)])
        self.fichier("Bit_2026-09-27.xlsx", [["B3", d, "TAIEB JEANNE - FETES - RACLETTE", 440, None]])
        self.assertContains(self.importer("Bit_2026-09-27.xlsx"), "1 ligne(s) importée(s), 2 ancienne(s) remplacée(s)")
        self.assertEqual(LigneReleve.objects.filter(journal_id="B3").count(), 2)          # ouverture + 1 ligne
        chemin, n = ech.exporter(ech.PAR_NOM["Bit"])
        ws = openpyxl.load_workbook(chemin).active
        self.assertEqual([c.value for c in ws[2]][:3] + [ws["D2"].value], ["B3", dt.datetime(2026, 1, 11), "TAIEB JEANNE - FETES - RACLETTE", 440])

    def test_banque_sans_doublon(self):
        d = dt.datetime(2026, 1, 5)
        self.fichier("Banque1.xlsx", [[d, "11", "עמלת מסלול", -10, 990], [d, "12", "הפקדת שיק", 1550, 2540]])
        self.assertContains(self.importer("Banque1.xlsx"), "2 ligne(s) ajoutée(s), 0 déjà présente(s)")
        self.fichier("Banque1.xlsx", [[d, "11", "עמלת מסלול", -10, 990], [d, "12", "הפקדת שיק", 1550, 2540]])
        self.assertContains(self.importer("Banque1.xlsx"), "0 ligne(s) ajoutée(s), 2 déjà présente(s)")
        self.assertEqual(LigneReleve.objects.get(journal_id="B1", ouverture=True).montant, D(1000))

    def test_droits(self):
        b = User.objects.create_user("bureau")
        b.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(b)
        self.assertEqual(self.client.get("/echanges/").status_code, 403)

    def test_lexique_et_dossiers(self):
        self.client.get("/echanges/")
        for d in (ech.imports(), ech.exports()):
            wb = openpyxl.load_workbook(d / "Lexique.xlsx")
            self.assertEqual(wb.sheetnames, ["Fichiers", "Colonnes", "Règles"])
            self.assertEqual(wb["Fichiers"].max_row, len(ech.FORMATS) + 1)
        self.assertEqual(ech.a_importer(), [])                                  # le lexique ne s'importe pas
        autre = Path(tempfile.mkdtemp()) / "Mes imports"
        self.client.post("/echanges/", {"dossiers": "1", "dossier_imports": str(autre), "dossier_exports": ""})
        self.assertEqual((ech.imports(), ech.exports()), (autre, self.dossier / "Exports"))
        self.assertTrue(autre.is_dir())
        self.assertEqual(Reglage.lire("dossier_imports"), str(autre))

    def test_deposer_et_telecharger(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        r = self.client.post("/echanges/", {"exporter": "Ecritures"})
        self.assertIn("attachment", r["Content-Disposition"])                         # téléchargé aussitôt
        self.assertIn("spreadsheetml", r["Content-Type"])
        page = self.client.get("/echanges/")
        nom = next(n for n, _ in page.context["exportes"] if n.startswith("Ecritures_"))
        r = self.client.get(f"/echanges/exports/{nom}")
        contenu = b"".join(r.streaming_content)
        self.assertEqual(self.client.get("/echanges/exports/..%2Fsecret.txt").status_code, 404)
        tampon = io.BytesIO()
        with zipfile.ZipFile(tampon, "w") as z:
            z.writestr("Dossier/Budget.xlsx", contenu)
            z.writestr("Dossier/notes.txt", "ignoré")
        r = self.client.post("/echanges/", {"deposer": "1", "fichiers": [SimpleUploadedFile(nom, contenu),
                                                                         SimpleUploadedFile("lot.zip", tampon.getvalue())]},
                             follow=True)
        self.assertContains(r, "Déposé(s) dans Imports")
        self.assertEqual(sorted(p.name for p, _ in ech.a_importer()), ["Budget.xlsx", nom])
        r = self.client.post("/echanges/", {"deposer": "1", "fichiers": SimpleUploadedFile("virus.exe", b"x")}, follow=True)
        self.assertContains(r, "fichier .xlsx, .pdf ou .zip attendu")

    def test_tout_reinjecter(self):
        """Tout exporter, corriger à la main, tout réinjecter : les données de chaque fichier sont remplacées."""
        d = dt.datetime(2026, 1, 5)
        self.fichier("Banque1.xlsx", [[d, "11", "עמלת מסלול", -10, 990]])
        self.importer("Banque1.xlsx")
        b1 = Journal.objects.get(code="B1")
        m = mouvement(600, dt.date(2026, 1, 5), [("600100", 10, 0), ("512000", 0, 10)])
        rap.pointer(b1, LigneReleve.objects.filter(journal=b1, ouverture=False), m.lignes.filter(compte_id="512000"))
        Mouvement.objects.filter(numero=421).update(origine="saisie", commentaire="saisi à la main")
        self.client.post("/echanges/", {"exporter": "tout"})
        for p in ech.exports().glob("*_*.xlsx"):
            if ech.format_de(p).nom in ("Ecritures", "Banque1", "PlanComptable"):
                (ech.imports() / p.name).write_bytes(p.read_bytes())
        plan = next(ech.imports().glob("PlanComptable_*"))
        wb = openpyxl.load_workbook(plan)
        wb.active.append(["600200", "LOCATION DE SALLE", "FON.1", "non", "oui"])            # compte ajouté à la main
        wb.save(plan)
        ecr = next(ech.imports().glob("Ecritures_*"))
        wb = openpyxl.load_workbook(ecr)
        ws = wb.active
        ws.delete_rows(2, 2)                                                                     # Mvt 421 supprimé à la main
        wb.save(ecr)
        self.assertContains(self.client.post("/echanges/", {"reinjecter": "1", "confirmation": "non"}, follow=True), "Taper REMPLACER")
        r = self.client.post("/echanges/", {"reinjecter": "1", "confirmation": "remplacer"}, follow=True)
        self.assertContains(r, "Réinjection terminée")
        self.assertEqual(list(Mouvement.objects.values_list("numero", flat=True)), [600])
        self.assertTrue(Compte.objects.filter(numero="600200").exists())
        self.assertEqual(Rapprochement.objects.count(), 1)                                        # pointage recollé
        self.assertEqual(Mouvement.objects.get(numero=600).origine, "saisie")                    # origine gardée
        self.assertEqual(ech.a_importer(), [])
        # un fichier en erreur : tout est annulé
        ws.append([dt.datetime(2026, 2, 1), "B1", 700, 700, "999999", "X", 5, None, "GEN.001", ""])
        wb.save(ech.imports() / "Ecritures.xlsx")
        r = self.client.post("/echanges/", {"reinjecter": "1", "confirmation": "REMPLACER"}, follow=True)
        self.assertContains(r, "compte 999999 inconnu")
        self.assertEqual(list(Mouvement.objects.values_list("numero", flat=True)), [600])


# ---------------------------------------------------------------- dossiers paramétrables, jeu complet d'exports

from . import dossiers  # noqa: E402


@override_settings(DATA_DIR=Path(tempfile.mkdtemp()))
class DossiersEtExportComplet(TransactionTestCase):
    def setUp(self):
        self.racine = Path(tempfile.mkdtemp())
        self.reglage = override_settings(IMPORTS_DIR=self.racine / "Imports", EXPORTS_DIR=self.racine / "Exports")
        self.reglage.enable()
        call_command("migrate", verbosity=0)
        referentiels_saisie()

    def tearDown(self):
        self.reglage.disable()

    def test_chemins_windows_effaces_et_refuses(self):
        from django.apps import apps
        __import__("importlib").import_module("compta.migrations.0012_dossiers_pc").dossiers_pc(apps, None)
        self.assertEqual(Reglage.lire("dossier_exports"), r"D:\OneDrive\Applications\ComptaBB\Exports")   # ancien programme du PC
        __import__("importlib").import_module("compta.migrations.0015_sans_programme_pc").effacer_chemins_pc(apps, None)
        self.assertEqual((Reglage.lire("dossier_exports"), Reglage.lire("dossier_sauvegardes")), ("", ""))
        self.assertEqual(dossiers.sauvegardes(), self.racine / "Exports" / "Sauvegardes")
        for chemin in (r"E:\Sauvegardes", "relatif/Exports"):
            with self.assertRaises(ech.Refus):
                ech.changer_dossiers({"dossier_sauvegardes": chemin})
        self.assertEqual(Reglage.lire("dossier_sauvegardes"), "")

    def test_tout_exporter_et_sauvegarder(self):
        exp, sauv = self.racine / "Mes exports", self.racine / "Mes sauvegardes"
        ech.changer_dossiers({"dossier_exports": str(exp), "dossier_sauvegardes": str(sauv)}, auteur="t")
        u = User.objects.create_user("admin")
        donner_role(u, "Administrateur")
        self.client.force_login(u)
        page = self.client.get("/echanges/")
        self.assertContains(page, str(sauv))
        self.client.post("/echanges/", {"tout_sauvegarder": "1"})
        noms = [p.name for p in exp.iterdir()]
        for f in ech.FORMATS:
            self.assertTrue(any(n.startswith(f.nom + "_") for n in noms), f.nom)
        self.assertTrue(any(n.startswith("Parametres_") for n in noms))
        self.assertEqual([p.suffix for p in sauv.iterdir()], [".sqlite3"])
        self.assertEqual(bd.dossier(), sauv)
        call_command("exporter_tout", stdout=open("/dev/null", "w"))
        self.assertEqual(len(list(sauv.glob("*.sqlite3"))), 2)


# ---------------------------------------------------------------- exports des référentiels, pages et barre des boutons

@override_settings(DATA_DIR=Path(tempfile.mkdtemp()))
class ReferentielsEtPages(TransactionTestCase):
    def setUp(self):
        self.racine = Path(tempfile.mkdtemp())
        self.reglage = override_settings(IMPORTS_DIR=self.racine / "Imports", EXPORTS_DIR=self.racine / "Exports")
        self.reglage.enable()
        call_command("migrate", verbosity=0)
        referentiels_saisie()
        mbr.creer_manquants()

    def tearDown(self):
        self.reglage.disable()

    def test_export_de_chaque_referentiel(self):
        from .vues_referentiels import referentiels
        t = User.objects.create_user("tresorier")
        donner_role(t, "Trésorier")
        self.client.force_login(t)
        self.assertEqual(self.client.get("/referentiels/").status_code, 403)          # exports : administrateur seulement
        self.assertEqual(self.client.get("/documentation/installation.pdf").status_code, 403)
        page = self.client.get("/")
        self.assertNotContains(page, "Administration ▾")                             # onglet réservé
        self.assertNotContains(page, "Installation et mises à jour")
        t = User.objects.create_user("admin2")
        donner_role(t, "Administrateur")
        self.client.force_login(t)
        noms = [r[0] for r in referentiels(t)]
        self.assertIn("PlanComptable", noms)
        self.assertContains(self.client.get("/referentiels/"), "PlanComptable.xlsx")
        for nom in noms + ["Parametres"]:
            r = self.client.get(f"/referentiels/{nom}.xlsx")
            self.assertEqual(r.status_code, 200, nom)
            openpyxl.load_workbook(__import__("io").BytesIO(r.content))
        r = self.client.get("/referentiels/PlanComptable.xlsx")
        wb = openpyxl.load_workbook(__import__("io").BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Mode d'emploi", "Plan comptable"])
        self.assertIn("512000", [c.value for c in wb["Plan comptable"]["A"]])
        self.assertEqual(prm.importer(r.content).erreurs, [])                      # se réimporte tel quel
        self.assertTrue(list((self.racine / "Exports").glob("PlanComptable_*.xlsx")))
        self.assertIn("Ecritures", noms)                                           # les écritures s'exportent aussi
        r = self.client.get("/referentiels/Ecritures.xlsx")
        self.assertEqual(openpyxl.load_workbook(io.BytesIO(r.content)).active["C2"].value, 421)
        r = self.client.get("/referentiels/Tout.xlsx")                               # Tout exporter : .zip
        self.assertIn(".zip", r["Content-Disposition"])
        with zipfile.ZipFile(io.BytesIO(r.content)) as z:
            self.assertTrue({"Parametres.xlsx", "Ecritures.xlsx", "Tiers.xlsx", "Banque1.xlsx", "Lexique.xlsx"} <= set(z.namelist()))
        self.assertTrue(list((self.racine / "Exports").glob("Exports_*.zip")))
        self.assertEqual(self.client.get("/referentiels/Utilisateurs.xlsx").status_code, 200)             # administrateur
        b = User.objects.create_user("bureau")
        donner_role(b, "Bureau")
        self.client.force_login(b)
        self.assertEqual(self.client.get("/referentiels/").status_code, 403)

    def test_toutes_les_pages_et_leur_barre(self):
        import re as _re
        a = User.objects.create_user("admin")
        donner_role(a, "Administrateur")
        self.client.force_login(a)
        pages = ["/", "/saisie/", "/mouvement/421/", "/mouvement/421/modifier/", "/mouvement/nouveau/", "/mouvement/rappel/",
                 "/codes/", "/fiches/", "/membres/", "/membres/411TAIEB001/", "/membres/cotisations/", "/tiers-provisoires/",
                 "/ecritures/", "/journaux/", "/grand-livre/", "/balance/", "/analytique/", "/rapprochement/B1/",
                 "/rapprochement/traductions/", "/etats/", "/cloture/", "/controles/", "/modifications/", "/mon-compte/",
                 "/utilisateurs/", f"/utilisateurs/{a.pk}/", "/parametres/", "/echanges/", "/base/", "/referentiels/"]
        for url in pages:
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200, url)
            html = r.content.decode()
            self.assertNotIn('class="boutons"', html, url)                        # plus de boutons en bas de page
            ids = set(_re.findall(r'<form[^>]* id="([^"]+)"', html))
            for cible in _re.findall(r'<button[^>]* form="([^"]+)"', html):
                self.assertIn(cible, ids, f"{url} : bouton relié au formulaire absent {cible}")
        self.assertContains(self.client.get("/"), "Mode d'emploi (PDF)")
        for nom in ("presentation", "mode-emploi", "installation"):
            r = self.client.get(f"/documentation/{nom}.pdf")
            self.assertEqual((r.status_code, r["Content-Type"]), (200, "application/pdf"), nom)
            self.assertTrue(b"".join(r.streaming_content).startswith(b"%PDF"))
        self.assertEqual(self.client.get("/documentation/autre.pdf").status_code, 404)


# ---------------------------------------------------------------- rapports analytiques (axe 1, axe 2)

class RapportsAnalytiques(TestCase):
    def setUp(self):
        referentiels_saisie()                   # Mvt 421 : 400 de produit (710000, COT.2) sur MAN.001, contre 411TAIEB001
        Compte.objects.create(numero="630000", libelle="AIDES")                                   # compte sans code d'axe 1
        m = Mouvement.objects.create(numero=500, date=dt.date(2026, 3, 1), journal_id="CA", piece=500)
        Ligne.objects.create(mouvement=m, ordre=0, compte_id="600000", libelle="LOCATION", debit=D(150), anal2_id="MAN.001")
        Ligne.objects.create(mouvement=m, ordre=1, compte_id="630000", libelle="AIDE", debit=D(50), anal2_id="SOC.006")
        Ligne.objects.create(mouvement=m, ordre=2, compte_id="530000", libelle="LOCATION", credit=D(200), anal2_id="MAN.001")
        u = User.objects.create_user("bureau")
        u.groups.add(Group.objects.get(name="Bureau"))                          # consultation seule suffit
        self.client.force_login(u)

    def test_synthese_gestion_avec_totaux(self):
        r = self.client.get("/analytique/?du=2026-01-01&au=2026-12-31")
        axe1, axe2 = r.context["axes"]
        self.assertEqual({l["code"]: l["s"] for l in axe2["lignes"]}, {"MAN.001": D(250), "SOC.006": D(-50)})
        self.assertEqual(axe2["total"], {"a": D(400), "b": D(200), "s": D(200)})
        self.assertEqual([l["code"] for l in axe1["lignes"]], ["(sans code)", "COT.2", "FON.1"])
        self.assertContains(r, "Imprimer")
        self.assertContains(r, "Résultat")
        self.assertContains(r, "<h1>Analytique – Nature (axe 1)</h1>")                        # une page par axe
        self.assertNotContains(r, "RALLYE")
        r = self.client.get("/analytique/?axe=2&du=2026-01-01&au=2026-12-31")
        self.assertContains(r, "<h1>Analytique – Objet (axe 2 : événement, projet)</h1>")
        self.assertContains(r, "/analytique/detail/?comptes=gestion&amp;axe=2&amp;code=MAN.001")
        self.assertContains(r, '<option value="MAN.001">RALLYE (MAN.001)</option>')             # choix par le libellé
        x = self.client.get("/analytique/?du=2026-01-01&au=2026-12-31&format=xlsx")
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ["Axe 1", "Axe 2"])
        self.assertEqual([c.value for c in wb["Axe 2"][wb["Axe 2"].max_row]], ["TOTAL", None, 400, 200, 200])

    def test_synthese_bilan_comptes_1_a_5(self):
        r = self.client.get("/analytique/?comptes=bilan&du=2026-01-01&au=2026-12-31")
        axe1, axe2 = r.context["axes"]
        self.assertEqual({l["code"]: (l["a"], l["b"], l["s"]) for l in axe2["lignes"]},
                         {"MAN.001": (D(400), D(200), D(200))})                    # 411 au débit, 530000 au crédit
        self.assertEqual({l["code"]: l["s"] for l in axe1["lignes"]}, {"BIL.4": D(400), "BIL.5": D(-200)})
        self.assertContains(r, "Solde")
        self.assertNotContains(r, "Résultat")

    def test_detail_par_code(self):
        r = self.client.get("/analytique/detail/?axe=2&code=MAN.001&du=2026-01-01&au=2026-12-31")
        (s,) = r.context["sections"]
        self.assertEqual([(c["numero"], c["a"], c["b"]) for c in s["comptes"]], [("600000", D(0), D(150)), ("710000", D(400), D(0))])
        self.assertEqual(s["total"]["s"], D(250))
        self.assertEqual(len(s["ecritures"]), 2)                                   # classes 6 et 7 seulement
        self.assertContains(r, "Total MAN.001")
        tous = self.client.get("/analytique/detail/?axe=1&du=2026-01-01&au=2026-12-31").context["sections"]
        self.assertEqual([s["code"] for s in tous], ["(sans code)", "COT.2", "FON.1"])
        sans = self.client.get("/analytique/detail/?axe=1&code=(sans code)&du=2026-01-01&au=2026-12-31").context["sections"]
        self.assertEqual(sans[0]["comptes"][0]["numero"], "630000")
        bilan = self.client.get("/analytique/detail/?comptes=bilan&axe=2&code=MAN.001&du=2026-01-01&au=2026-12-31")
        self.assertEqual([c["numero"] for c in bilan.context["sections"][0]["comptes"]], ["411TAIEB001", "530000"])
        x = self.client.get("/analytique/detail/?axe=2&du=2026-01-01&au=2026-12-31&format=xlsx")
        ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
        valeurs = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertIn("État détaillé – MAN.001 – " + s["libelle"], valeurs)           # présentation soignée, un bloc par code
        self.assertIn("Sous-total 600000", valeurs)
        self.assertIn("Répartition par axe 1", valeurs)
        # présentation : chiffres clés, répartition par l'autre axe, détail par compte avec sous-totaux
        self.assertEqual(s["nb_mouvements"], 2)
        self.assertEqual([(x["code"], x["s"]) for x in s["autres"]], [("COT.2", D(400)), ("FON.1", D(-150))])
        self.assertContains(r, "Répartition par nature (axe 1)")
        self.assertContains(r, "Sous-total 600000")
        self.assertContains(r, 'data-cherchable')


# ---------------------------------------------------------------- Excel et PDF sur toutes les pages

class ExcelDeToutesLesPages(TestCase):
    def setUp(self):
        referentiels_saisie()
        mouvement(600, dt.date(2026, 2, 1), [("600000", 1234.5, 0), ("512000", 0, 1234.5)])
        u = User.objects.create_user("admin", is_superuser=True, is_staff=True)
        self.client.force_login(u)

    def test_pages_en_excel(self):
        for url in ("/balance/", "/grand-livre/?compte=512000", "/ecritures/", "/controles/", "/modifications/", "/membres/",
                    "/membres/cotisations/", "/", "/analytique/", "/etats/", "/rapprochement/"):
            page = self.client.get(url)
            self.assertContains(page, "Imprimer / PDF", msg_prefix=url)
            self.assertContains(page, 'id="vers-excel"', msg_prefix=url)
            r = self.client.get(url + ("&" if "?" in url else "?") + "export=xlsx&tout=1")
            self.assertEqual(r.status_code, 200, url)
            self.assertIn("spreadsheetml", r["Content-Type"], url)
            openpyxl.load_workbook(io.BytesIO(r.content))
        wb = openpyxl.load_workbook(io.BytesIO(self.client.get("/balance/?export=xlsx&tout=1").content))
        ws = wb.worksheets[0]
        valeurs = [c.value for row in ws.iter_rows() for c in row]
        self.assertIn(1234.5, valeurs)                                      # montant « 1 234,50 » devenu nombre
        self.assertIn("Compte", valeurs)
        from . import dossiers as dos
        self.assertTrue(list(dos.exports().glob("Balance_*.xlsx")))            # copie dans Exports

    def test_export_ciel(self):
        page = self.client.get("/ecritures/?du=2026-01-01&au=2026-12-31&journal=OD")
        self.assertContains(page, "Export Ciel (Excel)")
        self.assertContains(page, "journal=OD&amp;du=2026-01-01&amp;au=2026-12-31&amp;format=ciel")
        r = self.client.get("/ecritures/?du=2026-01-01&au=2026-12-31&format=ciel")
        self.assertIn("Ecritures_Ciel_2026-01-01_2026-12-31.xlsx", r["Content-Disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[1]], ["Mvt", "Journ", "Date", "Compte", "LibelCompte", "Debit", "Credit", "Npiece",
                                                    "Anal", "LibelAnal", "Lettr"])
        lignes = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
        charge = next(l for l in lignes if l[3] == "600000")
        self.assertEqual((charge[0], charge[5], charge[6]), (600, 1234.5, 0))
        self.assertEqual(ws.cell(2, 8).number_format, "@")                               # Npiece en texte

    def test_toutes_les_lignes_malgre_la_pagination(self):
        for n in range(1, 131):
            mouvement(1000 + n, dt.date(2026, 3, 1), [("600000", 1, 0), ("512000", 0, 1)])
        self.assertEqual(len(self.client.get("/ecritures/").context["page"].object_list), 100)
        self.assertGreater(len(self.client.get("/ecritures/?tout=1").context["page"].object_list), 130)

    def test_page_sans_tableau(self):
        r = self.client.get("/mon-compte/?export=xlsx")
        self.assertIn("text/html", r["Content-Type"])                          # rien à exporter : la page s'affiche
from . import justificatifs as just  # noqa: E402
from .models import Justificatif  # noqa: E402


@override_settings(**temporaire())
class Justificatifs(TransactionTestCase):
    def setUp(self):
        ExportComplet.setUp(self)
        self.m = Mouvement.objects.get(numero=1)

    def fichier(self, nom="facture.pdf", contenu=b"%PDF-1.4 facture"):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return SimpleUploadedFile(nom, contenu)

    def test_ajout_consultation_suppression(self):
        self.client.force_login(self.u)
        r = self.client.post("/mouvement/1/justificatifs/", {"fichiers": [self.fichier(), self.fichier("recu.jpg", b"\xff\xd8photo")],
                                                             "description": "Traiteur"})
        self.assertRedirects(r, "/mouvement/1/")
        self.assertEqual(self.m.justificatifs.count(), 2)
        j = self.m.justificatifs.first()
        self.assertTrue(just.chemin(j).exists())
        self.assertRegex(j.chemin, r"^2026/Mvt1_\d+_facture\.pdf$")
        page = self.client.get("/mouvement/1/")
        self.assertContains(page, "facture.pdf")
        self.assertIn("<title>Mvt 1 ", page.content.decode())            # section dans la page, pas dans le titre
        self.assertNotIn("Justificatifs", page.content.decode().split("</title>")[0])
        self.assertContains(page, "Traiteur")
        r = self.client.get(f"/justificatif/{j.pk}/")
        self.assertEqual(b"".join(r.streaming_content), b"%PDF-1.4 facture")
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertIn("attachment", self.client.get(f"/justificatif/{j.pk}/?telecharger=1")["Content-Disposition"])
        # colonne 📎 et filtre des écritures
        self.assertContains(self.client.get("/ecritures/?du=2026-01-01&au=2026-12-31"), "📎2")
        sans = self.client.get("/ecritures/?du=2026-01-01&au=2026-12-31&just=sans").context["page"]
        avec = self.client.get("/ecritures/?du=2026-01-01&au=2026-12-31&just=avec").context["page"]
        self.assertNotIn(1, {l.mouvement.numero for l in sans})
        self.assertEqual({l.mouvement.numero for l in avec}, {1})
        self.assertEqual(avec.paginator.count, 2)
        # suppression tracée, fichier effacé
        fichier = just.chemin(j)
        self.client.post(f"/justificatif/{j.pk}/supprimer/")
        self.assertFalse(fichier.exists())
        self.assertEqual(self.m.justificatifs.count(), 1)
        self.assertTrue(Modification.objects.filter(action="Suppression d'un justificatif").exists())
        self.assertTrue(Modification.objects.filter(action="Ajout d'un justificatif").exists())

    def test_refus(self):
        self.client.force_login(self.u)
        self.client.post("/mouvement/1/justificatifs/", {"fichiers": [self.fichier("virus.exe", b"MZ")]})
        self.assertEqual(Justificatif.objects.count(), 0)
        with self.assertRaises(ValueError):
            just.ajouter(self.m, self.fichier("gros.pdf", b"x" * (just.TAILLE_MAXI + 1)))
        # bureau (consultation) : consulte mais ne joint ni ne supprime
        j = just.ajouter(self.m, self.fichier(), auteur="tresorier")
        v = User.objects.create_user("verif")
        v.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(v)
        self.assertEqual(self.client.get(f"/justificatif/{j.pk}/").status_code, 200)
        self.assertEqual(self.client.post("/mouvement/1/justificatifs/", {"fichiers": [self.fichier()]}).status_code, 403)
        self.assertEqual(self.client.post(f"/justificatif/{j.pk}/supprimer/").status_code, 403)
        self.assertNotContains(self.client.get("/mouvement/1/"), "Joindre")
        # exercice clos : on peut encore joindre, plus supprimer
        Exercice.objects.filter(libelle="2026").update(clos=True)
        with self.assertRaises(ValueError):
            just.supprimer(j)
        self.assertTrue(just.chemin(j).exists())
        just.ajouter(self.m, self.fichier("complement.png", b"\x89PNG"))

    def test_export_complet_avec_justificatifs(self):
        j = just.ajouter(self.m, self.fichier(), "Traiteur", "tresorier")
        contenu = just.chemin(j).read_bytes()
        chemin = ec.exporter(auteur="tresorier")
        with zipfile.ZipFile(chemin) as z:
            self.assertEqual(z.read(f"Justificatifs/{j.chemin}"), contenu)
        just.chemin(j).unlink()                              # fichier perdu sur le serveur
        Justificatif.objects.all().delete()
        message, _ = ec.reinjecter(chemin, auteur="tresorier")
        self.assertIn("1 justificatif", message)
        j = Justificatif.objects.get()
        self.assertEqual((j.description, just.chemin(j).read_bytes()), ("Traiteur", contenu))
        # export dont le scan a été altéré : refusé, rien n'est modifié
        altere = chemin.with_name("Export_complet_scan_altere.zip")
        with zipfile.ZipFile(chemin) as z, zipfile.ZipFile(altere, "w") as sortie:
            for n in z.namelist():
                sortie.writestr(n, b"%PDF autre" if n.startswith("Justificatifs/") else z.read(n))
        with self.assertRaises(ec.ExportInvalide):
            ec.reinjecter(altere)
        self.assertEqual(just.chemin(Justificatif.objects.get()).read_bytes(), contenu)


@override_settings(**temporaire())
class JustificatifsExistants(TransactionTestCase):
    def setUp(self):
        ExportComplet.setUp(self)                         # Mvt 1 : pièce 1, 01/02/2026, 400,00
        import shutil
        shutil.rmtree(just.dossier())                     # dossier partagé par les tests de la classe
        Mouvement.objects.filter(numero=1).update(piece=739)
        m = Mouvement.objects.create(numero=5, date=dt.date(2026, 3, 15), journal_id="OD", piece=740)
        Ligne.objects.create(mouvement=m, ordre=1, compte_id="600000", libelle="TRAITEUR", debit=D("450.00"), anal2_id="MAN.001")
        Ligne.objects.create(mouvement=m, ordre=2, compte_id="512000", libelle="TRAITEUR", credit=D("450.00"), anal2_id="MAN.001")

    def test_document_depose_remplace_le_lien_du_meme_mouvement(self):
        m = Mouvement.objects.get(numero=5)
        lien = Justificatif.objects.create(mouvement=m, lien="https://documents.exemple.org/fichier/aaa/", nom="Document en ligne")
        just.deposer("5 facture.pdf", b"%PDF-1.4 test")
        j = just.rattacher(just.a_classer()[0]["nom"], m, "", "test", remplacer=True)
        self.assertEqual(list(m.justificatifs.values_list("pk", flat=True)), [j.pk])
        self.assertFalse(Justificatif.objects.filter(pk=lien.pk).exists())

    def test_tout_telecharger_en_zip(self):
        self.client.force_login(self.u)
        m = Mouvement.objects.get(numero=5)
        just.ajouter(m, __import__("django.core.files.uploadedfile", fromlist=["x"]).SimpleUploadedFile("facture.pdf", b"%PDF-1.4 test"), "", "test")
        Justificatif.objects.create(mouvement=m, lien="https://documents.exemple.org/fichier/aaa/", nom="Document en ligne")
        r = self.client.get("/justificatifs/tout.zip")
        self.assertEqual(r.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(b"".join(r.streaming_content))) as z:
            noms = z.namelist()
        self.assertTrue(any(n.endswith("facture.pdf") for n in noms))
        self.assertIn("Liens.txt", noms)

    def test_supprimer_les_lignes_cochees(self):
        self.client.force_login(self.u)
        just.deposer("a.pdf", b"%PDF-1.4 a")
        just.deposer("b.pdf", b"%PDF-1.4 b")
        noms = [l["nom"] for l in just.a_classer()]
        self.client.post("/justificatifs/a-classer/", {"ecarter_coches": "1", "nom": noms, "garder_0": "1"})
        self.assertEqual([l["nom"] for l in just.a_classer()], [noms[1]])

    def test_propositions(self):
        cas = {"Mvt 5 facture.pdf": (5, True), "mvt_1.jpg": (1, True), "Pièce 739.pdf": (1, True),
               "PJ-740 traiteur.pdf": (5, True), "scan 739.pdf": (1, False), "2026-03-15 traiteur 450,00.pdf": (5, False),
               "15.03.2026.pdf": (5, False), "5.pdf": (5, False), "photo.jpg": (None, False), "Mvt 999.pdf": (None, False),
               # n° de Mvt en tête du nom puis une espace (convention du trésorier), y compris dans un dossier du ZIP
               "5 facture traiteur.pdf": (5, True), "1 cotisation 2026-03-15 450,00.pdf": (1, True),
               "Justificatifs__5 recu.pdf": (5, True), "739 facture.pdf": (None, False)}
        self.assertIn("aucun mouvement le 16/03/2026", just.proposer("2026-03-16 recu.pdf")[1])
        for nom, (numero, sur) in cas.items():
            m, raison, s = just.proposer(nom)
            self.assertEqual((m.numero if m else None, s), (numero, sur), f"{nom} : {raison}")

    def test_rappel_depuis_le_mouvement(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.u)
        self.assertNotContains(self.client.get("/mouvement/5/"), "document déjà déposé")          # rien à classer
        self.client.post("/justificatifs/a-classer/", {"deposer": "1", "fichiers": [
            SimpleUploadedFile("scan traiteur.pdf", b"%PDF t"), SimpleUploadedFile("recu.jpg", b"\xff\xd8r")]})
        page = self.client.get("/mouvement/5/")
        self.assertContains(page, "Choisir un document déjà déposé (2 à classer)")
        noms = [nom for nom, _ in page.context["a_classer"]]
        r = self.client.post("/mouvement/5/justificatifs/", {"a_classer": noms, "description": "rappel"}, follow=True)
        self.assertContains(r, "2 document(s) déposé(s) rattaché(s) au mouvement 5")
        self.assertEqual(sorted(Mouvement.objects.get(numero=5).justificatifs.values_list("nom", flat=True)),
                         ["recu.jpg", "scan traiteur.pdf"])
        self.assertEqual(just.liste_a_classer(), [])
        b = User.objects.create_user("bureau")                 # consultation seulement : ni liste ni rattachement
        donner_role(b, "Bureau")
        self.client.force_login(b)
        self.assertEqual(self.client.post("/mouvement/5/justificatifs/", {"a_classer": ["x"]}).status_code, 403)

    def test_numero_en_tete_et_plusieurs_mouvements(self):
        import io
        tampon = io.BytesIO()
        with zipfile.ZipFile(tampon, "w") as z:
            z.writestr("Justificatifs/5 facture traiteur.pdf", b"%PDF t")
            z.writestr("Justificatifs/5+1 facture commune.pdf", b"%PDF c")
            z.writestr("Justificatifs/sans numero.pdf", b"%PDF s")
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.u)
        self.client.post("/justificatifs/a-classer/", {"deposer": "1", "fichiers": [SimpleUploadedFile("J.zip", tampon.getvalue())]})
        page = self.client.get("/justificatifs/a-classer/")
        lignes = {l["affiche"]: l for l in page.context["lignes"]}
        self.assertEqual(set(lignes), {"5 facture traiteur.pdf", "5+1 facture commune.pdf", "sans numero.pdf"})
        self.assertTrue(lignes["5 facture traiteur.pdf"]["sur"])
        self.assertEqual(lignes["5+1 facture commune.pdf"]["saisie"], "5+1")
        self.assertContains(page, 'value="5+1"')
        self.assertFalse(lignes["sans numero.pdf"]["sur"])
        donnees = {"rattacher": "1", "nom": [l["nom"] for l in page.context["lignes"]]}
        for i, l in enumerate(page.context["lignes"]):
            if l["sur"]:
                donnees.update({f"mvt_{i}": l["saisie"] or str(l["mouvement"].numero), f"garder_{i}": "1"})
        self.client.post("/justificatifs/a-classer/", donnees)
        self.assertEqual(sorted(Justificatif.objects.filter(mouvement__numero=5).values_list("nom", flat=True)),
                         ["5 facture traiteur.pdf", "5+1 facture commune.pdf"])
        j1 = Justificatif.objects.get(mouvement__numero=1)
        self.assertEqual(just.chemin(j1).read_bytes(), b"%PDF c")                  # copie pour le Mvt 1
        self.assertEqual([l["affiche"] for l in just.a_classer()], ["sans numero.pdf"])
        # depuis la page du mouvement : « Joindre aussi » au Mvt 1
        j = Justificatif.objects.get(mouvement__numero=5, nom="5 facture traiteur.pdf")
        self.assertContains(self.client.get("/mouvement/5/"), "Joindre aussi")
        r = self.client.post(f"/justificatif/{j.pk}/copier/", {"mvt": "1+999"}, follow=True)
        self.assertContains(r, "joint aussi au mouvement 1")
        self.assertContains(r, "Mvt 999 inconnu")
        self.assertEqual(Justificatif.objects.filter(mouvement__numero=1).count(), 2)
        self.client.post(f"/justificatif/{j.pk}/copier/", {"mvt": "5"})               # déjà sur ce mouvement
        self.assertEqual(Justificatif.objects.filter(mouvement__numero=5).count(), 2)

    def test_depot_zip_et_rattachement(self):
        import io
        tampon = io.BytesIO()
        with zipfile.ZipFile(tampon, "w") as z:
            z.writestr("Factures 2026/Piece 739.pdf", b"%PDF 739")
            z.writestr("Factures 2026/Piece 739.pdf.bak", b"x")                 # format refusé
            z.writestr("__MACOSX/._Piece 739.pdf", b"x")                        # ignoré
            z.writestr("Divers/photo reçu.jpg", b"\xff\xd8photo")
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.u)
        self.client.post("/justificatifs/a-classer/", {"deposer": "1", "fichiers": [
            SimpleUploadedFile("scans.zip", tampon.getvalue()), SimpleUploadedFile("Mvt 5.pdf", b"%PDF 5"),
            SimpleUploadedFile("Mvt 5.pdf", b"%PDF 5 bis")]})
        lignes = just.a_classer()
        self.assertEqual(len(lignes), 4)
        page = self.client.get("/justificatifs/a-classer/")
        self.assertContains(page, "Piece 739.pdf")
        noms = [l["nom"] for l in page.context["lignes"]]
        self.assertTrue(any("doublon" in n for n in noms))                       # deux « Mvt 5.pdf »
        donnees = {"rattacher": "1", "nom": noms}
        for i, l in enumerate(page.context["lignes"]):
            donnees[f"mvt_{i}"] = str(l["mouvement"].numero) if l["mouvement"] else "1"   # photo : saisie à la main
            donnees[f"garder_{i}"] = "1"
            donnees[f"desc_{i}"] = "reprise"
        self.client.post("/justificatifs/a-classer/", donnees)
        self.assertEqual(just.a_classer(), [])
        self.assertEqual(Justificatif.objects.filter(mouvement__numero=5).count(), 2)
        self.assertEqual(set(Justificatif.objects.filter(mouvement__numero=1).values_list("nom", flat=True)),
                         {"Piece 739.pdf", "photo recu.jpg"})
        self.assertEqual(set(Justificatif.objects.filter(mouvement__numero=5).values_list("nom", flat=True)), {"Mvt 5.pdf"})
        # écarter, droits
        just.deposer("inutile.pdf", b"%PDF")
        self.client.post("/justificatifs/a-classer/", {"ecarter": just.a_classer()[0]["nom"]})
        self.assertEqual(just.a_classer(), [])
        b = User.objects.create_user("bureau")
        b.groups.add(Group.objects.get(name="Bureau"))
        self.client.force_login(b)
        self.assertEqual(self.client.get("/justificatifs/a-classer/").status_code, 403)

    def test_page_justificatifs_ouverte_au_tresorier(self):
        t = User.objects.create_user("tresorier_j")
        t.groups.add(Group.objects.get(name="Trésorier"))
        self.client.force_login(t)
        page = self.client.get("/justificatifs/a-classer/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(self.client.get("/"), "Justificatifs (dépôt et classement)")

    def test_commande(self):
        dossier = Path(tempfile.mkdtemp()) / "Scans"
        (dossier / "2026").mkdir(parents=True)
        (dossier / "2026" / "Pièce 740.pdf").write_bytes(b"%PDF 740")
        (dossier / "inconnu.jpg").write_bytes(b"\xff\xd8")
        call_command("importer_justificatifs", str(dossier), "--rattacher", stdout=open("/dev/null", "w"))
        self.assertEqual(Justificatif.objects.get().mouvement.numero, 5)
        self.assertEqual([l["affiche"] for l in just.a_classer()], ["inconnu.jpg"])
        (dossier / "5+1 facture commune.pdf").write_bytes(b"%PDF c")               # ZIP envoyé par Files, puis console
        call_command("importer_justificatifs", str(dossier / "5+1 facture commune.pdf"), "--rattacher",
                     stdout=open("/dev/null", "w"))
        self.assertEqual(sorted(Justificatif.objects.filter(nom="5+1 facture commune.pdf")
                                .values_list("mouvement__numero", flat=True)), [1, 5])


@override_settings(**temporaire())
class LiensEnLigne(TransactionTestCase):
    def setUp(self):
        JustificatifsExistants.setUp(self)               # Mvt 1 (01/02/2026, 400) et Mvt 5 (15/03/2026, 450)
        Reglage.objects.update_or_create(cle="hebergeur_liens", defaults={"valeur": "Documents en ligne"})

    def extrait(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DEPENSES"
        ws.append(["Card name", "תאריך", "ספק/ית", "סכום", "פריט הוצאה", "סוג תשלום", "סטטוס", "קובץ מקושר"])
        for date, somme, item, lien in ((dt.datetime(2026, 3, 15), -450, "2 - FETES - RALLYE", "https://documents.exemple.org/fichier/aaa/"),
                                        (dt.datetime(2026, 2, 1), -400, "1 - LOGE - COTISATION", "https://documents.exemple.org/fichier/bbb/"),
                                        (dt.datetime(2026, 4, 1), -99, "4 - FRAIS", "https://documents.exemple.org/fichier/ccc/"),
                                        (dt.datetime(2026, 4, 2), -10, "4 - FRAIS", None)):
            ws.append(["חשבונית", date, None, somme, item, "Cash", "Draft", "https://documents.exemple.org/fichier" if lien else None])
            if lien:
                ws.cell(ws.max_row, 8).hyperlink = lien
        return self.octets(wb)

    def octets(self, wb):
        tampon = __import__("io").BytesIO()
        wb.save(tampon)
        return tampon.getvalue()

    def test_lien_ecrit_en_texte(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["תאריך", "סכום", "קובץ מקושר"])
        ws.append([dt.datetime(2026, 3, 15), -450, "https://documents.exemple.org/fichier/ddd/"])     # sans hyperlien
        self.assertEqual(len(just.deposer("EXTRAIT.xlsx", self.octets(wb))[0]), 1)
        l = just.a_classer()[0]
        self.assertEqual((l["lien"], l["mouvement"].numero), ("https://documents.exemple.org/fichier/ddd/", 5))

    def test_excel_fait_a_la_main(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Liens"
        ws.append(["Mvt", "Pièce", "Lien", "Remarque"])
        ws.append([1, None, "https://documents.exemple.org/fichier/m01/", "cotisation"])
        ws.append([None, Mouvement.objects.get(numero=5).piece, "https://documents.exemple.org/fichier/p05/", None])
        ws.append([999, None, "https://documents.exemple.org/fichier/x99/", None])
        ws.append([None, None, "https://documents.exemple.org/fichier/sans/", None])        # rien pour le rattacher
        self.assertEqual(len(just.deposer("Liens.xlsx", self.octets(wb))[0]), 3)
        lignes = {l["lien"][-4:-1]: l for l in just.a_classer()}
        self.assertEqual((lignes["m01"]["mouvement"].numero, lignes["m01"]["sur"], lignes["m01"]["affiche"]), (1, True, "cotisation"))
        self.assertEqual(lignes["p05"]["mouvement"].numero, 5)
        self.assertIsNone(lignes["x99"]["mouvement"])
        self.client.force_login(self.u)
        self.assertContains(self.client.get("/justificatifs/a-classer/"), "Mvt 999 inexistant")
        just.rattacher(lignes["m01"]["nom"], lignes["m01"]["mouvement"])
        self.assertEqual(Justificatif.objects.get(mouvement__numero=1).nom, "Liens (lien)")

    def test_documents_en_ligne_copies_sur_le_site(self):
        from unittest import mock
        m = Mouvement.objects.get(numero=5)
        a = just.ajouter_lien(m, "https://documents.exemple.org/fichier/aaa/", "Document en ligne")
        b = just.ajouter_lien(m, "https://documents.exemple.org/fichier/bbb/", "Document en ligne")
        c = just.ajouter_lien(m, "https://documents.exemple.org/fichier/ccc/", "Document en ligne")
        reponses = {"aaa": (b"%PDF-1.4 facture", "Facture 12.pdf"), "bbb": (b"<!DOCTYPE html><html>connexion", "")}

        def faux(lien):
            if lien[-4:-1] not in reponses:
                raise ValueError("site du document injoignable depuis ComptaBB")
            return reponses[lien[-4:-1]]
        self.client.force_login(self.u)
        self.assertContains(self.client.get("/justificatifs/a-classer/"), "Documents encore en ligne (3)")
        with mock.patch.object(just, "telecharger", faux):
            r = self.client.post("/justificatifs/a-classer/", {"rapatrier": "1"}, follow=True)
        self.assertContains(r, "1 document(s) copié(s) sur le site ; 2 encore en ligne")
        self.assertContains(r, "page web (connexion demandée")
        a.refresh_from_db()
        self.assertEqual((a.lien, a.nom, a.taille), ("", "Facture 12.pdf", 16))
        self.assertEqual(just.chemin(a).read_bytes(), b"%PDF-1.4 facture")
        self.assertEqual(self.client.get(f"/justificatif/{a.pk}/").status_code, 200)            # servi par le site
        self.assertTrue(Modification.objects.filter(action="Document en ligne enregistré sur le site",
                                                    avant__contains="aaa").exists())
        b.refresh_from_db()
        c.refresh_from_db()
        self.assertTrue(b.lien and c.lien)                                                     # les échecs gardent leur lien

    def test_lien_colle_sur_le_mouvement(self):
        self.client.force_login(self.u)
        self.assertContains(self.client.get("/mouvement/5/"), "Joindre le lien")
        adresse = "https://documents.exemple.org/fichier/a2be4b4f-5493-4433-b40e-f6c0d81eaa83/"
        self.assertRedirects(self.client.post("/mouvement/5/justificatifs/", {"lien": adresse, "description": "Traiteur"}), "/mouvement/5/")
        j = Justificatif.objects.get(mouvement__numero=5)
        self.assertEqual((j.lien, j.nom, j.description), (adresse, "Document en ligne", "Traiteur"))
        self.client.post("/mouvement/5/justificatifs/", {"lien": adresse})                     # pas deux fois
        self.client.post("/mouvement/5/justificatifs/", {"lien": "javascript:alert(1)"})       # https seulement
        self.assertEqual(Justificatif.objects.count(), 1)

    def test_liens(self):
        deposes, refus = just.deposer("EXTRAIT.xlsx", self.extrait())
        self.assertEqual((len(deposes), refus), (3, []))
        self.assertEqual(just.deposer("EXTRAIT.xlsx", self.extrait())[0], [])       # pas de doublon
        lignes = {l["lien"][-4:-1]: l for l in just.a_classer()}
        self.assertEqual(lignes["aaa"]["mouvement"].numero, 5)
        self.assertTrue(lignes["aaa"]["sur"])
        self.assertEqual(lignes["aaa"]["affiche"], "2 - FETES - RALLYE · Cash")
        self.assertEqual(lignes["bbb"]["mouvement"].numero, 1)
        self.assertIsNone(lignes["ccc"]["mouvement"])
        self.client.force_login(self.u)
        page = self.client.get("/justificatifs/a-classer/")
        self.assertContains(page, "https://documents.exemple.org/fichier/aaa/")
        donnees = {"rattacher": "1", "nom": [l["nom"] for l in page.context["lignes"]]}
        for i, l in enumerate(page.context["lignes"]):
            if l["mouvement"]:
                donnees.update({f"mvt_{i}": str(l["mouvement"].numero), f"garder_{i}": "1"})
        self.client.post("/justificatifs/a-classer/", donnees)
        j = Justificatif.objects.get(mouvement__numero=5)
        self.assertEqual((j.lien, j.chemin, j.nom), ("https://documents.exemple.org/fichier/aaa/", None, "DEPENSES du 15/03/2026"))
        self.assertEqual(self.client.get(f"/justificatif/{j.pk}/")["Location"], j.lien)
        self.assertContains(self.client.get("/mouvement/5/"), "🔗")
        self.assertEqual([l["lien"][-4:-1] for l in just.a_classer()], ["ccc"])          # reste à classer
        just.ecarter(just.a_classer()[0]["nom"])
        self.assertEqual(just.a_classer(), [])
        self.assertEqual(just.deposer("EXTRAIT.xlsx", self.extrait())[0], ["lien 0"])  # seul « ccc » revient
        # export complet : les liens reviennent à la réinjection
        chemin = ec.exporter()
        Justificatif.objects.all().delete()
        ec.reinjecter(chemin)
        self.assertEqual(set(Justificatif.objects.values_list("lien", flat=True)),
                         {"https://documents.exemple.org/fichier/aaa/", "https://documents.exemple.org/fichier/bbb/"})

