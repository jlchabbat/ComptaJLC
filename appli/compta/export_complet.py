"""Export complet en une opération, et réinjection d'un export complet dans une base remise à zéro.

L'export (Exports/Export_complet_AAAA-MM-JJ_HHMMSS.zip) contient :
- comptabb.sqlite3 : sauvegarde de la base (restauration exacte, comptes utilisateurs compris) ;
- Parametres.xlsx : les paramètres (même classeur que la page Paramètres) ;
- Tiers.xlsx : les fiches tiers (même format que l'import Tiers.xlsx) ;
- Ecritures.xlsx : toutes les écritures (une ligne par ligne d'écriture, avec pointage, origine et auteur) ;
- Donnees.xlsx : comptes de tiers, exercices, pointages, relevés, budget, fiches bénévoles et leurs documents, justificatifs,
  historique ;
- Justificatifs/ : les scans et photos joints aux mouvements ;
- Etats_<exercice>.xlsx : états de chaque exercice (lecture seule) ;
- controle.json : nombres et totaux, pour vérifier une réinjection.
Seuls les comptes utilisateurs (identifiants, mots de passe) ne sont que dans comptabb.sqlite3.

La réinjection vide la comptabilité (les comptes utilisateurs sont gardés), recharge Parametres, Tiers, Ecritures et
Donnees, puis compare la base obtenue à controle.json. Tout ou rien : au moindre écart, rien n'est modifié.

Fichiers modifiés à la main (Excel) : décompresser l'export, modifier Parametres.xlsx, Tiers.xlsx, Ecritures.xlsx ou
Donnees.xlsx sans changer les en-têtes de colonnes, recompresser le dossier (clic droit › Envoyer vers › Dossier compressé),
puis Base de données › Remettre à zéro et recharger en cochant « Fichiers modifiés ». La base n'est alors pas comparée à
controle.json, mais aux contrôles de cohérence (Mvt équilibrés, débit OU crédit, codes existants) ; au moindre défaut,
rien n'est modifié.
"""

import datetime as dt
import hashlib
import io
import json
import zipfile
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.db import transaction
from django.db.models import Sum
from django.utils import timezone

from . import membres, parametres
from .models import (
    ZERO, Budget, CodeAnalytique, Compte, DocumentFiche, Exercice, Fiche, Journal, Justificatif, Ligne, LigneFiche, LigneReleve, LigneSchema, Membre,
    ModeFiche, ModeleOperation, Modification, Mouvement, MoyenPaiement, NatureFiche, ParametreReleve, Prefixe,
    Rapprochement, Reglage, TiersProvisoire, Traduction, TypeTiers, arrondi,
)

GARDER = 20                     # exports complets conservés dans Exports
FICHIERS_REINJECTES = ("Parametres.xlsx", "Tiers.xlsx", "Donnees.xlsx")       # obligatoires
FICHIERS_FACULTATIFS = ("Ecritures.xlsx", "controle.json")                    # absents des exports antérieurs ou retirés
ECRITURES = "Écritures"                                                       # feuille de Ecritures.xlsx (autrefois de Donnees.xlsx)


class ExportInvalide(ValueError):
    pass


# ---------------------------------------------------------------- conversions

def _heure(v):
    """Date-heure sans fuseau pour Excel (heure d'Israël), à la seconde."""
    return timezone.localtime(v).replace(tzinfo=None, microsecond=0) if v else None


def _lire_heure(v):
    if v in (None, ""):
        return None
    if isinstance(v, str):
        v = dt.datetime.fromisoformat(v)
    return timezone.make_aware(v) if timezone.is_naive(v) else v


def _jour(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return dt.date.fromisoformat(str(v)[:10])


def _montant(v):
    return None if v is None else float(v)


def _lire_montant(v):
    return ZERO if v in (None, "") else Decimal(str(v)).quantize(Decimal("0.01"))


def _texte(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def _oui(v):
    return "Oui" if v else "Non"


def _lire_oui(v):
    return parametres.Booleen().lire(v) if v not in (None, "") else False


def _entier(v):
    return None if v in (None, "") else int(float(v))


def _nom_utilisateur(u):
    return u.get_username() if u else ""


# ---------------------------------------------------------------- empreinte (contrôle)

COMPTES = [Compte, Journal, CodeAnalytique, Prefixe, Reglage, TypeTiers, MoyenPaiement, LigneSchema, ModeleOperation,
           NatureFiche, ModeFiche, Traduction, ParametreReleve, Exercice, Mouvement, Ligne, Rapprochement, LigneReleve,
           Budget, Membre, TiersProvisoire, Fiche, LigneFiche, Justificatif, DocumentFiche]


def empreinte():
    """Nombres d'éléments et totaux de la base : deux bases identiques ont la même empreinte."""
    e = {"nombres": {m._meta.model_name: m.objects.count() for m in COMPTES}}
    t = Ligne.objects.aggregate(d=Sum("debit"), c=Sum("credit"))
    soldes = Ligne.objects.values("compte").annotate(d=Sum("debit"), c=Sum("credit")).order_by("compte")
    texte = ";".join(f"{s['compte']}={arrondi(s['d']) - arrondi(s['c'])}" for s in soldes)
    e["totaux"] = {
        "debit": str(arrondi(t["d"])), "credit": str(arrondi(t["c"])),
        "soldes_par_compte": hashlib.sha256(texte.encode()).hexdigest()[:16],
        "lignes_lettrees": Ligne.objects.exclude(lettrage="").count(),
        "lignes_pointees": Ligne.objects.filter(rapprochement__isnull=False).count(),
        "releve": str(arrondi(LigneReleve.objects.aggregate(s=Sum("montant"))["s"])),
        "releve_pointe": LigneReleve.objects.filter(rapprochement__isnull=False).count(),
        "budget": str(arrondi(Budget.objects.aggregate(s=Sum("montant"))["s"])),
        "fiches": str(arrondi(LigneFiche.objects.aggregate(s=Sum("montant"))["s"])),
        "historique": Modification.objects.count(),
        "contenu": signature(),
    }
    return e


# colonnes qui portent un numéro interne : remplacé par le rang (l'ordre de création est conservé)
REFERENCES = {"Pointages": {"N°": "p"}, "Écritures": {"Pointage": "p"}, "Relevés": {"Pointage": "p"},
              "Tiers provisoires": {"N°": "t"}, "Fiches": {"N°": "f"}, "Lignes de fiches": {"Fiche": "f", "Tiers provisoire": "t"},
              "Documents fiches": {"Fiche": "f"}}


def _normal(v):
    if isinstance(v, dt.datetime):
        return (v + dt.timedelta(microseconds=500000)).replace(microsecond=0).isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        return f"{v:.2f}"
    return "" if v is None else str(v)


def signature():
    """Empreinte de tout le contenu : paramètres, tiers et données, cellule par cellule."""
    rangs = {"p": {pk: i for i, pk in enumerate(Rapprochement.objects.order_by("pk").values_list("pk", flat=True))},
             "t": {pk: i for i, pk in enumerate(TiersProvisoire.objects.order_by("pk").values_list("pk", flat=True))},
             "f": {pk: i for i, pk in enumerate(Fiche.objects.order_by("pk").values_list("pk", flat=True))}}
    h = hashlib.sha256()
    for f in parametres.FEUILLES:
        for obj in f.elements():
            h.update(repr([_normal(c.type.vers_excel(getattr(obj, c.champ))) for c in f.colonnes]).encode())
    for m in Membre.objects.order_by("compte"):
        h.update(repr([_normal(v) for v in _ligne_tiers(m)]).encode())
    for nom, lignes in _lignes_donnees():
        refs = {FEUILLES[nom].index(col): cle for col, cle in REFERENCES.get(nom, {}).items()}
        for l in lignes:
            l = [rangs[refs[i]].get(v) if i in refs else v for i, v in enumerate(l)]
            h.update(repr([nom] + [_normal(v) for v in l]).encode())
    return h.hexdigest()[:20]


def ecarts(attendu, obtenu):
    res = []
    for partie in ("nombres", "totaux"):
        for k, v in attendu.get(partie, {}).items():
            if obtenu[partie].get(k) != v:
                res.append(f"{k} : {v} attendu, {obtenu[partie].get(k)} obtenu")
    return res


# ---------------------------------------------------------------- Donnees.xlsx

FEUILLES = {
    "Comptes de tiers": ["Compte", "Libellé", "Axe 1", "Lettrable", "Actif"],
    "Exercices": ["Libellé", "Début", "Fin", "Clos", "Mvt à-nouveaux", "Résultat affecté", "Clôturé le", "Clôturé par",
                  "Archive"],
    "Écritures": ["Mvt", "Date", "Journal", "Pièce", "Origine", "Commentaire", "Créé le", "Créé par", "Ordre", "Compte",
                  "Libellé", "Débit", "Crédit", "Axe 2", "Lettrage", "Pointage", "Montant devise"],
    "Pointages": ["N°", "Journal", "Mode", "Créé le", "Créé par"],
    "Relevés": ["Journal", "Date", "Rang", "Référence", "Opération", "Montant", "Solde", "Ouverture", "Source",
                "Importé le", "Pointage"],
    "Budget": ["Exercice", "Nature", "Compte", "Axe 1", "Axe 2", "Montant"],
    "Tiers provisoires": ["N°", "Nom", "Prénom", "Remarque", "Compte", "Créé par", "Créé le"],
    "Fiches": ["N°", "Type", "Titre", "Axe 2", "Bénévoles", "Statut", "Créée le"],
    "Lignes de fiches": ["Fiche", "Sens", "Date", "Tiers", "Tiers provisoire", "Autre", "Personnes", "Nature", "Montant",
                         "Mode", "Sens du mode", "Justificatif", "Remarque", "Compte", "Axe 2", "Mvt", "Créé par",
                         "Créé le"],
    "Historique": ["Date", "Auteur", "Lot", "Action", "Objet", "Avant", "Après"],
    "Bénévoles": ["Compte", "Identifiant"],
    "Justificatifs": ["Mvt", "Fichier", "Nom", "Description", "Taille", "Ajouté le", "Ajouté par", "Empreinte", "Lien"],
    # reçus des bénévoles ; Ligne = rang de la ligne dans « Lignes de fiches » (1 = la première), vide = toute la fiche
    "Documents fiches": ["Fiche", "Ligne", "Fichier", "Nom", "Description", "Taille", "Ajouté le", "Ajouté par", "Empreinte"],
    "Axes de comptes": ["Compte", "Axe", "Ordre", "Valeur"],
}


def _lignes_donnees():
    tiers = set(Membre.objects.values_list("compte_id", flat=True))
    yield "Comptes de tiers", ([c.numero, c.libelle, c.anal1_id, _oui(c.lettrable), _oui(c.actif)]
                               for c in Compte.objects.filter(numero__in=tiers))
    yield "Exercices", ([e.libelle, e.debut, e.fin, _oui(e.clos), e.mouvement_an.numero if e.mouvement_an else None,
                         _montant(e.resultat), _heure(e.cloture_le), e.cloture_par, e.archive]
                        for e in Exercice.objects.select_related("mouvement_an"))
    yield "Écritures", ([m.numero, m.date, m.journal_id, m.piece, m.origine, m.commentaire, _heure(m.cree_le),
                         _nom_utilisateur(m.cree_par), l.ordre, l.compte_id, l.libelle, _montant(l.debit) or None,
                         _montant(l.credit) or None, l.anal2_id, l.lettrage, l.rapprochement_id,
                         _montant(l.montant_devise) if l.montant_devise is not None else None]
                        for m in Mouvement.objects.select_related("cree_par").prefetch_related("lignes").order_by("numero")
                        for l in m.lignes.all())
    yield "Pointages", ([r.pk, r.journal_id, r.mode, _heure(r.cree_le), _nom_utilisateur(r.cree_par)]
                        for r in Rapprochement.objects.select_related("cree_par").order_by("pk"))
    yield "Relevés", ([l.journal_id, l.date, l.rang, l.reference, l.operation, _montant(l.montant), _montant(l.solde),
                       _oui(l.ouverture), l.source, _heure(l.importe_le), l.rapprochement_id]
                      for l in LigneReleve.objects.order_by("journal", "date", "rang", "pk"))
    yield "Budget", ([b.exercice.libelle, b.nature, b.compte_id, b.anal1_id, b.anal2_id, _montant(b.montant)]
                     for b in Budget.objects.select_related("exercice"))
    yield "Tiers provisoires", ([t.pk, t.nom, t.prenom, t.remarque, t.compte_id, _nom_utilisateur(t.cree_par),
                                 _heure(t.cree_le)] for t in TiersProvisoire.objects.select_related("cree_par").order_by("pk"))
    yield "Fiches", ([f.pk, f.type, f.titre, f.anal2_id, ";".join(sorted(u.get_username() for u in f.benevoles.all())),
                      f.statut, _heure(f.cree_le)] for f in Fiche.objects.prefetch_related("benevoles").order_by("pk"))
    yield "Lignes de fiches", ([l.fiche_id, l.sens, l.date, l.tiers_id, l.provisoire_id, l.autre, l.personnes, l.nature.libelle,
                                _montant(l.montant), l.mode.libelle if l.mode else "", l.mode.sens if l.mode else "",
                                l.justificatif, l.remarque, l.compte_id, l.anal2_id,
                                l.mouvement.numero if l.mouvement else None, _nom_utilisateur(l.cree_par), _heure(l.cree_le)]
                               for l in LigneFiche.objects.select_related("nature", "mode", "mouvement", "cree_par")
                               .order_by("pk"))
    from .models import ValeurCompte
    yield "Axes de comptes", ([v.compte_id, v.axe.nom, v.axe.ordre, v.valeur]
                              for v in ValeurCompte.objects.select_related("axe").order_by("axe__ordre", "axe__nom", "compte"))
    from .justificatifs import chemin as chemin_justificatif
    yield "Justificatifs", ([j.mouvement.numero, j.chemin or "", j.nom, j.description, j.taille, _heure(j.ajoute_le), j.ajoute_par,
                             "lien" if j.lien else _empreinte_fichier(chemin_justificatif(j)), j.lien]
                            for j in Justificatif.objects.select_related("mouvement").order_by("mouvement__numero", "pk"))
    rang_ligne = {pk: i for i, pk in enumerate(LigneFiche.objects.order_by("pk").values_list("pk", flat=True), 1)}
    from .justificatifs import dossier as dossier_justificatifs
    yield "Documents fiches", ([d.fiche_id, rang_ligne.get(d.ligne_id), d.chemin, d.nom, d.description, d.taille, _heure(d.ajoute_le),
                                d.ajoute_par, _empreinte_fichier(dossier_justificatifs() / d.chemin)]
                               for d in DocumentFiche.objects.order_by("pk"))
    yield "Historique", ([_heure(m.date), m.auteur, m.lot, m.action, m.objet, m.avant, m.apres]
                         for m in Modification.objects.order_by("date", "pk"))
    yield "Bénévoles", ([m.compte_id, m.utilisateur.get_username()]
                        for m in Membre.objects.filter(utilisateur__isnull=False).select_related("utilisateur").order_by("compte"))


def _empreinte_fichier(chemin):
    return hashlib.sha256(chemin.read_bytes()).hexdigest()[:16] if chemin.exists() else "absent"


def classeur_donnees(feuilles=None):
    """Donnees.xlsx : toutes les feuilles sauf les écritures (feuilles=None), ou seulement celles demandées."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb._named_styles["Normal"].font = Font(name="Calibri", size=12)
    wb.remove(wb.active)
    for nom, lignes in _lignes_donnees():
        if (nom == ECRITURES) if feuilles is None else (nom not in feuilles):
            continue
        ws = wb.create_sheet(nom)
        ws.append(FEUILLES[nom])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF", size=12), PatternFill("solid", fgColor="1F3864")
        for l in lignes:
            ws.append(l)
        for i, e in enumerate(FEUILLES[nom], 1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(e) + 6))
            if e in ("Date", "Début", "Fin"):
                for cellule in ws[get_column_letter(i)][1:]:
                    cellule.number_format = "DD/MM/YYYY"
        ws.freeze_panes = "A2"
    return wb


def classeur_tiers():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tiers"
    ws.append(membres.ENTETES_MODELE)
    for m in Membre.objects.select_related("type").order_by("compte"):
        ws.append(_ligne_tiers(m))
    return wb


def _ligne_tiers(m):
    return [m.compte_id, m.type.libelle if m.type else "", m.nom, m.prenom, m.adresse, m.code_postal, m.ville,
            m.telephone, m.email, m.date_adhesion, m.statut, _montant(m.cotisation)]


def _octets(wb):
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


# ---------------------------------------------------------------- export

def dossier():
    from .dossiers import exports
    return exports()


def liste():
    return sorted(dossier().glob("Export_complet_*.zip"), key=lambda p: p.stat().st_mtime, reverse=True)


def exporter(auteur=""):
    """Crée Exports/Export_complet_<date>.zip ; renvoie son chemin."""
    from . import base_donnees
    from .export import classeur_exercice
    sauvegarde = base_donnees.sauvegarder("export_complet")
    chemin = dossier() / f"Export_complet_{dt.datetime.now():%Y-%m-%d_%H%M%S}.zip"
    controle = empreinte()
    controle.update(cree_le=f"{dt.datetime.now():%d/%m/%Y %H:%M}", auteur=auteur)
    with zipfile.ZipFile(chemin, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(sauvegarde, "comptabb.sqlite3")
        z.writestr("Parametres.xlsx", parametres.contenu_classeur())
        z.writestr("Tiers.xlsx", _octets(classeur_tiers()))
        z.writestr("Ecritures.xlsx", _octets(classeur_donnees([ECRITURES])))
        z.writestr("Donnees.xlsx", _octets(classeur_donnees()))
        for ex in Exercice.objects.all():
            nom = "".join(c if c.isalnum() else "_" for c in ex.libelle)
            z.writestr(f"Etats_{nom}.xlsx", _octets(classeur_exercice(ex)))
        from .justificatifs import chemin as chemin_justificatif
        for j in Justificatif.objects.exclude(chemin=None):
            if chemin_justificatif(j).exists():
                z.write(chemin_justificatif(j), f"Justificatifs/{j.chemin}")
        from .justificatifs import dossier as dossier_justificatifs
        for d in DocumentFiche.objects.all():                    # reçus des bénévoles (fiches non encore reportées)
            if (dossier_justificatifs() / d.chemin).exists():
                z.write(dossier_justificatifs() / d.chemin, f"Justificatifs/{d.chemin}")
        z.writestr("controle.json", json.dumps(controle, ensure_ascii=False, indent=1))
        z.writestr("LISEZMOI.txt", __doc__.strip() + "\n")
    for vieux in liste()[GARDER:]:
        vieux.unlink()
    Modification.objects.create(auteur=auteur, lot="Base de données", action="Export complet", objet=chemin.name,
                                apres=f"{controle['nombres']['mouvement']} Mvt, {controle['nombres']['ligne']} lignes")
    return chemin


# ---------------------------------------------------------------- réinjection

def vider():
    """Efface toute la comptabilité ; les comptes utilisateurs sont gardés."""
    Exercice.objects.update(mouvement_an=None)
    from .models import AxeCompte, ValeurCompte
    for m in (ValeurCompte, AxeCompte, Membre, Budget, LigneReleve, Rapprochement, ParametreReleve, Traduction, DocumentFiche, LigneFiche, Fiche, TiersProvisoire,
              NatureFiche, ModeFiche, Ligne, Mouvement, Modification, ModeleOperation, MoyenPaiement, LigneSchema, TypeTiers,
              Journal, Compte, Prefixe, CodeAnalytique, Exercice, Reglage):
        m.objects.all().delete()


FACULTATIVES = ("Justificatifs", "Documents fiches", "Axes de comptes")       # absentes des exports faits avant leur création
COLONNES_FACULTATIVES = {"Lien", "Montant devise"}         # idem pour les colonnes


def _rangees(wb, nom, fichier="Donnees.xlsx"):
    if nom not in wb.sheetnames:
        if nom in FACULTATIVES:
            return
        raise ExportInvalide(f"Feuille « {nom} » absente de {fichier}.")
    rangees = wb[nom].iter_rows(values_only=True)
    entetes = [str(e or "").strip() for e in next(rangees, [])]
    manque = [e for e in FEUILLES[nom] if e not in entetes and e not in COLONNES_FACULTATIVES]
    if manque:
        raise ExportInvalide(f"{fichier}, feuille {nom} : colonne(s) {', '.join(manque)} absente(s).")
    for r in rangees:
        if any(v not in (None, "") for v in r):
            yield dict(zip(entetes, r))


def _charger_comptes_tiers(wb):
    for r in _rangees(wb, "Comptes de tiers"):
        Compte.objects.create(numero=_texte(r["Compte"]), libelle=_texte(r["Libellé"]), anal1_id=_texte(r["Axe 1"]) or None,
                              lettrable=_lire_oui(r["Lettrable"]), actif=_lire_oui(r["Actif"]))


def _charger_donnees(wb, ecritures=None, pieces=None, ecrits=None):
    """ecritures : classeur Ecritures.xlsx (à défaut, la feuille Écritures de Donnees.xlsx des exports antérieurs) ;
    pieces : fichiers des justificatifs de l'export ; ecrits : fichiers écrits, pour tout remettre en cas d'échec."""
    utilisateurs = {u.get_username(): u for u in get_user_model().objects.all()}
    qui = lambda v: utilisateurs.get(_texte(v))  # noqa: E731
    an = {}
    for r in _rangees(wb, "Exercices"):
        e = Exercice.objects.create(libelle=_texte(r["Libellé"]), debut=_jour(r["Début"]), fin=_jour(r["Fin"]),
                                    clos=_lire_oui(r["Clos"]), resultat=None if r["Résultat affecté"] in (None, "")
                                    else _lire_montant(r["Résultat affecté"]), cloture_le=_lire_heure(r["Clôturé le"]),
                                    cloture_par=_texte(r["Clôturé par"]), archive=_texte(r["Archive"]))
        if r["Mvt à-nouveaux"] not in (None, ""):
            an[e.pk] = _entier(r["Mvt à-nouveaux"])
    pointages = {}
    for r in _rangees(wb, "Pointages"):
        p = Rapprochement.objects.create(journal_id=_texte(r["Journal"]), mode=_texte(r["Mode"]), cree_par=qui(r["Créé par"]))
        Rapprochement.objects.filter(pk=p.pk).update(cree_le=_lire_heure(r["Créé le"]) or p.cree_le)
        pointages[_entier(r["N°"])] = p
    mouvements, dates, lignes = {}, {}, []
    for r in _rangees(ecritures or wb, ECRITURES, "Ecritures.xlsx" if ecritures else "Donnees.xlsx"):
        n = _entier(r["Mvt"])
        if n not in mouvements:
            mouvements[n] = Mouvement(numero=n, date=_jour(r["Date"]), journal_id=_texte(r["Journal"]), piece=_entier(r["Pièce"]),
                                      origine=_texte(r["Origine"]), commentaire=_texte(r["Commentaire"]),
                                      cree_par=qui(r["Créé par"]))
            dates[n] = _lire_heure(r["Créé le"])
        lignes.append((n, Ligne(ordre=_entier(r["Ordre"]) or 0, compte_id=_texte(r["Compte"]), libelle=_texte(r["Libellé"]),
                                debit=_lire_montant(r["Débit"]), credit=_lire_montant(r["Crédit"]),
                                anal2_id=_texte(r["Axe 2"]) or None, lettrage=_texte(r["Lettrage"]),
                                rapprochement=pointages.get(_entier(r["Pointage"])),
                                montant_devise=None if r.get("Montant devise") in (None, "") else _lire_montant(r["Montant devise"]))))
    Mouvement.objects.bulk_create(mouvements.values())
    ids = dict(Mouvement.objects.values_list("numero", "pk"))
    for n, l in lignes:
        l.mouvement_id = ids[n]
    Ligne.objects.bulk_create([l for _, l in lignes])
    for n, d in dates.items():
        if d:
            Mouvement.objects.filter(numero=n).update(cree_le=d)
    for pk, n in an.items():
        Exercice.objects.filter(pk=pk).update(mouvement_an_id=ids[n])
    releve = []
    for r in _rangees(wb, "Relevés"):
        releve.append((LigneReleve(journal_id=_texte(r["Journal"]), date=_jour(r["Date"]), rang=1 if r["Rang"] in (None, "") else _entier(r["Rang"]),
                                   reference=_texte(r["Référence"]), operation=_texte(r["Opération"]),
                                   montant=_lire_montant(r["Montant"]),
                                   solde=None if r["Solde"] in (None, "") else _lire_montant(r["Solde"]),
                                   ouverture=_lire_oui(r["Ouverture"]), source=_texte(r["Source"]),
                                   rapprochement=pointages.get(_entier(r["Pointage"]))), _lire_heure(r["Importé le"])))
    for l, quand in releve:
        l.save()
        if quand:
            LigneReleve.objects.filter(pk=l.pk).update(importe_le=quand)
    exercices = {e.libelle: e for e in Exercice.objects.all()}
    for r in _rangees(wb, "Budget"):
        Budget.objects.create(exercice=exercices[_texte(r["Exercice"])], nature=_texte(r["Nature"]),
                              compte_id=_texte(r["Compte"]) or None, anal1_id=_texte(r["Axe 1"]) or None,
                              anal2_id=_texte(r["Axe 2"]) or None, montant=_lire_montant(r["Montant"]))
    provisoires = {}
    for r in _rangees(wb, "Tiers provisoires"):
        t = TiersProvisoire.objects.create(nom=_texte(r["Nom"]), prenom=_texte(r["Prénom"]), remarque=_texte(r["Remarque"]),
                                           compte_id=_texte(r["Compte"]) or None, cree_par=qui(r["Créé par"]))
        TiersProvisoire.objects.filter(pk=t.pk).update(cree_le=_lire_heure(r["Créé le"]) or t.cree_le)
        provisoires[_entier(r["N°"])] = t
    fiches = {}
    for r in _rangees(wb, "Fiches"):
        f = Fiche.objects.create(type=_texte(r["Type"]), titre=_texte(r["Titre"]), anal2_id=_texte(r["Axe 2"]) or None,
                                 statut=_texte(r["Statut"]))
        Fiche.objects.filter(pk=f.pk).update(cree_le=_lire_heure(r["Créée le"]) or f.cree_le)
        f.benevoles.set([u for u in (qui(n) for n in _texte(r["Bénévoles"]).split(";")) if u])
        fiches[_entier(r["N°"])] = f
    lignes_fiches = []
    for r in _rangees(wb, "Lignes de fiches"):
        f = fiches[_entier(r["Fiche"])]
        sens = _texte(r["Sens"])
        mode = None
        if _texte(r["Mode"]):
            mode = ModeFiche.objects.get(type_fiche=f.type, sens=_texte(r["Sens du mode"]), libelle=_texte(r["Mode"]))
        l = LigneFiche.objects.create(
            fiche=f, sens=sens, date=_jour(r["Date"]), tiers_id=_texte(r["Tiers"]) or None,
            provisoire=provisoires.get(_entier(r["Tiers provisoire"])), autre=_texte(r["Autre"]),
            personnes=_entier(r["Personnes"]), montant=_lire_montant(r["Montant"]), mode=mode,
            nature=NatureFiche.objects.get(type_fiche=f.type, sens=sens, libelle=_texte(r["Nature"])),
            justificatif=_texte(r["Justificatif"]), remarque=_texte(r["Remarque"]), compte_id=_texte(r["Compte"]) or None,
            anal2_id=_texte(r["Axe 2"]) or None, mouvement_id=ids.get(_entier(r["Mvt"])), cree_par=qui(r["Créé par"]))
        LigneFiche.objects.filter(pk=l.pk).update(cree_le=_lire_heure(r["Créé le"]) or l.cree_le)
        lignes_fiches.append(l)
    from .models import AxeCompte, ValeurCompte
    for r in _rangees(wb, "Axes de comptes"):
        axe, _ = AxeCompte.objects.get_or_create(nom=_texte(r["Axe"]), defaults={"ordre": _entier(r["Ordre"]) or 0})
        ValeurCompte.objects.create(compte_id=_texte(r["Compte"]), axe=axe, valeur=_texte(r["Valeur"]))
    from .justificatifs import dossier as dossier_justificatifs
    for r in _rangees(wb, "Justificatifs"):
        if _texte(r.get("Lien")):                           # document resté en ligne : seulement son lien
            j = Justificatif.objects.create(mouvement_id=ids[_entier(r["Mvt"])], lien=_texte(r["Lien"]), nom=_texte(r["Nom"]),
                                            description=_texte(r["Description"]), ajoute_par=_texte(r["Ajouté par"]))
            Justificatif.objects.filter(pk=j.pk).update(ajoute_le=_lire_heure(r["Ajouté le"]) or j.ajoute_le)
            continue
        relatif = _texte(r["Fichier"])
        contenu = (pieces or {}).get(relatif)
        if contenu is None:
            raise ExportInvalide(f"Justificatif {relatif} absent de l'export.")
        cible = dossier_justificatifs() / relatif
        if not cible.exists() or cible.read_bytes() != contenu:
            if ecrits is not None:
                ecrits.append((cible, cible.read_bytes() if cible.exists() else None))   # pour tout remettre en cas d'échec
            cible.parent.mkdir(parents=True, exist_ok=True)
            cible.write_bytes(contenu)
        j = Justificatif.objects.create(mouvement_id=ids[_entier(r["Mvt"])], chemin=relatif, nom=_texte(r["Nom"]),
                                        description=_texte(r["Description"]), taille=_entier(r["Taille"]) or 0,
                                        ajoute_par=_texte(r["Ajouté par"]))
        Justificatif.objects.filter(pk=j.pk).update(ajoute_le=_lire_heure(r["Ajouté le"]) or j.ajoute_le)
    for r in _rangees(wb, "Documents fiches"):
        relatif = _texte(r["Fichier"])
        contenu = (pieces or {}).get(relatif)
        if contenu is None:
            raise ExportInvalide(f"Document de fiche {relatif} absent de l'export.")
        cible = dossier_justificatifs() / relatif
        if not cible.exists() or cible.read_bytes() != contenu:
            if ecrits is not None:
                ecrits.append((cible, cible.read_bytes() if cible.exists() else None))
            cible.parent.mkdir(parents=True, exist_ok=True)
            cible.write_bytes(contenu)
        rang = _entier(r["Ligne"])
        d = DocumentFiche.objects.create(fiche=fiches[_entier(r["Fiche"])], ligne=lignes_fiches[rang - 1] if rang else None,
                                         chemin=relatif, nom=_texte(r["Nom"]), description=_texte(r["Description"]),
                                         taille=_entier(r["Taille"]) or 0, ajoute_par=_texte(r["Ajouté par"]))
        DocumentFiche.objects.filter(pk=d.pk).update(ajoute_le=_lire_heure(r["Ajouté le"]) or d.ajoute_le)
    historique = []
    for r in _rangees(wb, "Historique"):
        historique.append((Modification(auteur=_texte(r["Auteur"]), lot=_texte(r["Lot"]), action=_texte(r["Action"]),
                                        objet=_texte(r["Objet"]), avant=_texte(r["Avant"]), apres=_texte(r["Après"])),
                           _lire_heure(r["Date"])))
    for m, quand in historique:
        m.save()
        if quand:
            Modification.objects.filter(pk=m.pk).update(date=quand)
    if "Bénévoles" in wb.sheetnames:                          # absente des exports antérieurs
        for r in _rangees(wb, "Bénévoles"):
            Membre.objects.filter(compte_id=_texte(r["Compte"])).update(utilisateur=qui(r["Identifiant"]))


def lire_export(source):
    """Fichiers utiles d'un export complet (chemin ou octets d'un ZIP), repérés par leur nom : un export décompressé
    puis recompressé (fichiers rangés dans un dossier du ZIP) est accepté."""
    try:
        z = zipfile.ZipFile(source if not isinstance(source, bytes) else io.BytesIO(source))
    except zipfile.BadZipFile:
        raise ExportInvalide("Ce fichier n'est pas un export complet (ZIP illisible).") from None
    with z:
        par_nom = {}
        for n in z.namelist():
            base = n.replace("\\", "/").rsplit("/", 1)[-1]
            if base in FICHIERS_REINJECTES + FICHIERS_FACULTATIFS and not n.startswith("__MACOSX"):
                if base in par_nom:
                    raise ExportInvalide(f"{base} figure deux fois dans le ZIP ({par_nom[base]} et {n}) ; n'en laisser qu'un.")
                par_nom[base] = n
        manque = [n for n in FICHIERS_REINJECTES if n not in par_nom]
        if manque:
            raise ExportInvalide(f"Export complet incomplet : {', '.join(manque)} absent(s).")
        fichiers = {base: z.read(n) for base, n in par_nom.items()}
        fichiers["justificatifs"] = {}                      # Justificatifs/<année>/… , éventuellement dans un dossier du ZIP
        for n in z.namelist():
            chemin = n.replace("\\", "/")
            if n.endswith("/") or chemin.startswith("__MACOSX") or "Justificatifs/" not in chemin:
                continue
            fichiers["justificatifs"][chemin.split("Justificatifs/", 1)[1]] = z.read(n)
        return fichiers


def coherence():
    """Défauts qui interdisent de garder une base rechargée depuis des fichiers modifiés."""
    from django.db import IntegrityError, connection

    from . import controles
    defauts = []
    try:
        connection.check_constraints()
    except IntegrityError as e:
        defauts.append(f"code inexistant (compte, journal, code analytique, exercice…) : {e}")
    for r in controles.executer():
        if r.statut == "Anomalie" and r.regle in ("RG-01", "RG-03"):      # équilibre, débit OU crédit
            defauts.append(f"{r.regle} {r.libelle} : {r.valeur}" + (f" ({' ; '.join(map(str, r.details[:10]))})" if r.details else ""))
    return defauts


def reinjecter(source, auteur="", modifie=False):
    """Remet la comptabilité à zéro puis recharge l'export complet. Renvoie (message, sauvegarde préalable).

    Export intact : la base obtenue doit être identique à controle.json. Fichiers modifiés (modifie=True) : la base
    obtenue doit passer les contrôles de cohérence. Lève ExportInvalide sinon (rien n'est alors modifié)."""
    from . import base_donnees
    fichiers = lire_export(source)
    attendu = json.loads(fichiers["controle.json"]) if "controle.json" in fichiers else None
    avant = base_donnees.sauvegarder("avant-reinjection")
    ecrits = []
    try:
        message = _reinjecter(fichiers, attendu, auteur, ecrits, modifie)
    except Exception:
        for cible, ancien in reversed(ecrits):             # fichiers remis comme avant
            if ancien is None:
                cible.unlink(missing_ok=True)
            else:
                cible.write_bytes(ancien)
        raise
    return message, avant


def _reinjecter(fichiers, attendu, auteur, ecrits, modifie):
    with transaction.atomic():
        vider()
        r = parametres.importer(fichiers["Parametres.xlsx"], auteur=auteur, tracer=False)
        if r.erreurs:
            raise ExportInvalide("Parametres.xlsx : " + " · ".join(r.erreurs[:10]))
        wb = openpyxl.load_workbook(io.BytesIO(fichiers["Donnees.xlsx"]), data_only=True, read_only=True)
        ecritures = (openpyxl.load_workbook(io.BytesIO(fichiers["Ecritures.xlsx"]), data_only=True, read_only=True)
                     if "Ecritures.xlsx" in fichiers else None)
        try:
            _charger_comptes_tiers(wb)                       # comptes avant les fiches tiers
            t = membres.importer_tableau(membres.lire_tableau("Tiers.xlsx", fichiers["Tiers.xlsx"]))
            if t.erreurs:
                raise ExportInvalide("Tiers.xlsx : " + " · ".join(t.erreurs[:10]))
            _charger_donnees(wb, ecritures, fichiers["justificatifs"], ecrits)
        except ExportInvalide:
            raise
        except Exception as e:                               # cellule illisible, code inconnu, doublon…
            if not modifie:
                raise
            raise ExportInvalide(f"fichiers modifiés illisibles ({type(e).__name__} : {e}).") from e
        differences = ecarts(attendu, empreinte()) if attendu else ["controle.json absent"]
        if modifie:
            defauts = coherence()
            if defauts:
                raise ExportInvalide("Les fichiers modifiés ne sont pas cohérents : " + " · ".join(defauts[:10]))
        elif differences:
            raise ExportInvalide("La base rechargée diffère de l'export : " + " · ".join(differences[:10])
                                 + ". Si vous avez modifié les fichiers, cochez « Fichiers modifiés ».")
        n = {m: m.objects.count() for m in (Mouvement, Ligne, Membre, Justificatif)}
        Modification.objects.create(auteur=auteur, lot="Base de données", action="Réinjection d'un export complet",
                                    objet=f"export du {(attendu or {}).get('cree_le', '?')}"
                                          + (" (fichiers modifiés)" if modifie else ""),
                                    apres=f"{n[Mouvement]} Mvt, " + ("contrôles de cohérence OK" if modifie else "contrôle identique"))
    controle = ("contrôles de cohérence satisfaits (fichiers modifiés)" if modifie
                else "base identique à l'export")
    return (f"Export du {(attendu or {}).get('cree_le', '?')} réinjecté : {n[Mouvement]} mouvements, "
            f"{n[Ligne]} lignes, {n[Membre]} tiers, {n[Justificatif]} justificatif(s). Contrôle : {controle}.")
