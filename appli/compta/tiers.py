"""Tiers : fiches, situation (solde) et lettrage des comptes de tiers (clients, fournisseurs…)."""

import datetime as dt
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal

from django.db import transaction

from .models import ZERO, Compte, Ligne, Modification, Tiers, TypeTiers


def _cle_nom(nom):
    s = unicodedata.normalize("NFKD", nom.upper())
    return "".join(ch for ch in s if ch.isalpha() and ord(ch) < 128)[:5]


def compte_propose(nom, type_tiers=None):
    """Préfixe du type + 5 premières lettres du nom + rang : 411TAIEB001, 401PARTN001…"""
    t = type_tiers or TypeTiers.objects.order_by("pk").first()
    prefixe = (t.prefixe if t else "411") + _cle_nom(nom)
    rang = 1
    while Compte.objects.filter(numero=f"{prefixe}{rang:03d}").exists():
        rang += 1
    return f"{prefixe}{rang:03d}"


@transaction.atomic
def creer_tiers(type_tiers, nom, prenom="", **coordonnees):
    """Crée le compte (axe 1 repris d'un compte du même type) et la fiche du tiers."""
    libelle = f"{nom.strip().upper()} {prenom.strip().upper()}".strip()
    modele = Compte.objects.filter(numero__startswith=type_tiers.prefixe, anal1__isnull=False).first()
    compte = Compte.objects.create(numero=compte_propose(nom, type_tiers), libelle=libelle, lettrable=True,
                                   anal1=modele.anal1 if modele else None)
    Tiers.objects.create(compte=compte, type=type_tiers, nom=nom.strip().upper(), prenom=prenom.strip(),
                          **{k: (v or "").strip() for k, v in coordonnees.items()})
    return compte


def type_du_compte(numero):
    """Type de tiers d'après le préfixe du compte (le plus long qui convient)."""
    types = sorted(TypeTiers.objects.all(), key=lambda t: -len(t.prefixe))
    return next((t for t in types if numero.startswith(t.prefixe)), None)


def creer_manquants():
    """Une fiche pour chaque compte de tiers (clients, fournisseurs…) ; les comptes collectifs (tout en chiffres)
    sont exclus. Complète aussi le type des fiches qui n'en ont pas. Renvoie le nombre de fiches créées."""
    n = 0
    for t in TypeTiers.objects.all():
        for c in Compte.objects.filter(numero__startswith=t.prefixe, tiers__isnull=True):
            if c.numero.isdigit():                        # compte collectif (411000…) : pas de fiche
                continue
            Tiers.objects.create(compte=c, type=t, nom=c.libelle)
            n += 1
    for m in Tiers.objects.filter(type__isnull=True).select_related("compte"):
        m.type = type_du_compte(m.compte_id)
        if m.type:
            m.save(update_fields=["type"])
    return n


# ---------------------------------------------------------------- situation d'un compte de tiers

@dataclass
class Situation:
    facture: Decimal = ZERO
    regle: Decimal = ZERO
    solde: Decimal = ZERO
    impayes: list = field(default_factory=list)       # (ligne, restant dû)
    avance: Decimal = ZERO


def allouer(lignes):
    """Règlements (crédits) imputés sur les factures (débits) les plus anciennes ; renvoie [(ligne débit, restant)] et l'avance."""
    debits = [[l, l.debit] for l in lignes if l.debit > 0]
    credit = sum((l.credit for l in lignes), ZERO)
    for d in debits:
        pris = min(d[1], credit)
        d[1] -= pris
        credit -= pris
    return [(l, r) for l, r in debits if r > 0], credit


def situation(compte, jusquau=None):
    jusquau = jusquau or dt.date.today()
    ls = list(Ligne.objects.filter(compte=compte, mouvement__date__lte=jusquau).exclude(mouvement__origine="cloture")
              .select_related("mouvement").order_by("mouvement__date", "mouvement__numero", "ordre"))
    s = Situation(facture=sum((l.debit for l in ls), ZERO), regle=sum((l.credit for l in ls), ZERO))
    s.solde = s.facture - s.regle
    restants, s.avance = allouer([l for l in ls if not l.lettrage])
    s.impayes = restants
    return s


def montant(v):
    from .reglages import montant
    return montant(v)


# ---------------------------------------------------------------- lettrage

def code_suivant(compte):
    """AAA, AAB, … AAZ, ABA… : premier code de 3 lettres majuscules libre du compte."""
    pris = set(Ligne.objects.filter(compte=compte).exclude(lettrage="").values_list("lettrage", flat=True))
    for n in range(26 ** 3):
        code = "".join(chr(65 + n // 26 ** i % 26) for i in (2, 1, 0))
        if code not in pris:
            return code
    raise ValueError("Plus de code de lettrage libre sur ce compte.")


@transaction.atomic
def lettrer(compte, lignes, auteur=""):
    lignes = list(lignes)
    if len(lignes) < 2 or any(l.compte_id != compte.numero or l.lettrage for l in lignes):
        raise ValueError("Choisir au moins deux lignes non lettrées de ce compte.")
    d, c = sum((l.debit for l in lignes), ZERO), sum((l.credit for l in lignes), ZERO)
    if d != c:
        raise ValueError(f"Débit {montant(d)} ≠ crédit {montant(c)} : le lettrage doit être équilibré.")
    code = code_suivant(compte)
    Ligne.objects.filter(pk__in=[l.pk for l in lignes]).update(lettrage=code)
    if auteur:
        Modification.objects.create(auteur=auteur, lot="Tiers", action="Lettrage", objet=f"compte {compte.numero}",
                                    apres=f"{code} : {len(lignes)} lignes, {montant(d)}")
    return code


@transaction.atomic
def delettrer(compte, code, auteur=""):
    n = Ligne.objects.filter(compte=compte, lettrage=code).update(lettrage="")
    if n and auteur:
        Modification.objects.create(auteur=auteur, lot="Tiers", action="Délettrage", objet=f"compte {compte.numero}", avant=code)
    return n


def lettrage_automatique(compte, auteur=""):
    """1) facture et règlement de même montant dans un même mouvement ; 2) facture puis premier règlement
    ultérieur de même montant. Renvoie le nombre de lettrages."""
    ls = list(Ligne.objects.filter(compte=compte, lettrage="").exclude(mouvement__origine="cloture")
              .select_related("mouvement").order_by("mouvement__date", "mouvement__numero", "ordre"))
    libres = {l.pk: l for l in ls}
    paires = []
    for l in ls:
        if l.debit > 0 and l.pk in libres:
            meme = next((x for x in ls if x.pk in libres and x.mouvement_id == l.mouvement_id and x.credit == l.debit), None)
            if meme:
                paires.append((l, meme))
                del libres[l.pk], libres[meme.pk]
    for l in ls:
        if l.debit > 0 and l.pk in libres:
            suite = next((x for x in ls if x.pk in libres and x.credit == l.debit and x.mouvement.date >= l.mouvement.date), None)
            if suite:
                paires.append((l, suite))
                del libres[l.pk], libres[suite.pk]
    for a, b in paires:
        lettrer(compte, [a, b])
    if paires and auteur:
        Modification.objects.create(auteur=auteur, lot="Tiers", action="Lettrage automatique", objet=f"compte {compte.numero}",
                                    apres=f"{len(paires)} lettrage(s)")
    return len(paires)


# ---------------------------------------------------------------- import des tiers (Tiers.xlsx)

def cle(texte):
    return re.sub(r"[^a-z]", "", unicodedata.normalize("NFKD", texte or "").encode("ascii", "ignore").decode().lower())


COLONNES_TIERS = {
    "compte": ["compte", "numero", "ncompte", "numerodecompte", "comptetiers"],
    "type": ["type", "typedetiers", "categorie"],
    "nom": ["nom", "raisonsociale", "nomouraisonsociale", "societe", "nomraisonsociale"],
    "prenom": ["prenom"],
    "adresse": ["adresse", "rue", "adresse1"],
    "code_postal": ["codepostal", "cp", "zip"],
    "ville": ["ville", "localite", "commune"],
    "telephone": ["telephone", "tel", "portable", "mobile", "gsm"],
    "email": ["email", "mail", "courriel", "adressemail", "adresseemail"],
}
ENTETES_MODELE = ["Compte", "Type", "Nom", "Prénom", "Adresse", "Code postal", "Ville", "Téléphone", "E-mail"]


def _colonnes(entetes):
    res = {}
    for i, e in enumerate(entetes):
        k = cle(str(e or ""))
        for champ, noms in COLONNES_TIERS.items():
            if k in noms and champ not in res:
                res[champ] = i
    return res


@dataclass
class RapportImport:
    crees: list = field(default_factory=list)
    mis_a_jour: list = field(default_factory=list)
    erreurs: list = field(default_factory=list)


@transaction.atomic
def importer_tableau(rangees):
    """Importe des tiers (Tiers.xlsx ou CSV). Une ligne = un tiers ; en-têtes reconnus sous plusieurs formes.

    Compte présent : fiche de ce compte (compte créé s'il manque et que son préfixe correspond à un type).
    Compte absent : tiers retrouvé par type + nom + prénom, sinon créé avec un compte proposé.
    Une cellule vide ne remplace jamais une valeur déjà saisie ; rien n'est supprimé."""
    rapport, cols, defaut = RapportImport(), None, TypeTiers.objects.order_by("pk").first()
    for n, r in enumerate(rangees, 1):
        r = list(r)
        if cols is None:
            c = _colonnes(r)
            if "nom" in c or "compte" in c:
                cols = c
            continue
        val = lambda k: ("" if k not in cols or cols[k] >= len(r) or r[cols[k]] is None else r[cols[k]])  # noqa: E731
        texte = lambda k: str(val(k)).strip()  # noqa: E731
        numero, nom = texte("compte"), texte("nom")
        if not numero and not nom:
            continue
        numero = numero[:-2] if numero.endswith(".0") else numero          # nombre lu par Excel
        t = None
        if texte("type"):
            t = next((x for x in TypeTiers.objects.all() if cle(x.libelle) == cle(texte("type"))), None)
            t = t or (type_du_compte(numero) if numero else None)          # libellé différent : type d'après le compte
            if not t:
                rapport.erreurs.append(f"Ligne {n} : type « {texte('type')} » inconnu (Référentiels › Types de tiers).")
                continue
        fiche = None
        if numero:
            compte = Compte.objects.filter(numero=numero).first()
            t = t or type_du_compte(numero)
            if not t or not numero.startswith(t.prefixe):
                rapport.erreurs.append(f"Ligne {n} : le compte {numero} ne correspond à aucun type de tiers.")
                continue
            if not compte:
                if not nom:
                    rapport.erreurs.append(f"Ligne {n} : nom manquant pour créer le compte {numero}.")
                    continue
                modele = Compte.objects.filter(numero__startswith=t.prefixe, anal1__isnull=False).first()
                compte = Compte.objects.create(numero=numero, libelle=f"{nom.upper()} {texte('prenom').upper()}".strip(),
                                               lettrable=True, anal1=modele.anal1 if modele else None)
            fiche = Tiers.objects.filter(compte=compte).first()
            nouveau = fiche is None
            fiche = fiche or Tiers(compte=compte, nom=(nom or compte.libelle).upper())
        else:
            t = t or defaut
            if not t:
                rapport.erreurs.append(f"Ligne {n} : type de tiers manquant (Référentiels › Types de tiers).")
                continue
            fiche = Tiers.objects.filter(type=t, nom__iexact=nom, prenom__iexact=texte("prenom")).first()
            nouveau = fiche is None
            if nouveau:
                fiche = Tiers.objects.get(compte=creer_tiers(t, nom, texte("prenom")))
        fiche.type = fiche.type or t
        for champ in ("prenom", "adresse", "code_postal", "ville", "telephone", "email"):
            if texte(champ):
                setattr(fiche, champ, texte(champ))
        if nom:
            fiche.nom = nom.upper()
        if texte("email") and "@" not in texte("email"):
            rapport.erreurs.append(f"Ligne {n} : e-mail « {texte('email')} » ignoré.")
            fiche.email = ""
        fiche.save()
        (rapport.crees if nouveau else rapport.mis_a_jour).append(f"{fiche.compte_id} {fiche}")
    if cols is None:
        rapport.erreurs.append("En-têtes non reconnus : une colonne « Nom » ou « Compte » est nécessaire.")
    return rapport


def lire_tableau(nom_fichier, contenu):
    """Rangées d'un fichier .xlsx (première feuille contenant des en-têtes reconnus) ou .csv."""
    import csv
    import io
    if nom_fichier.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contenu), data_only=True, read_only=True)
        for ws in wb.worksheets:
            rangees = [list(r) for r in ws.iter_rows(values_only=True)]
            if any(_colonnes(r).keys() & {"nom", "compte"} for r in rangees[:20]):
                return rangees
        return []
    if nom_fichier.lower().endswith(".csv"):
        s = contenu.decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(s), delimiter=";" if s.count(";") >= s.count(",") else ","))
    raise ValueError("Format non reconnu : fichier .xlsx ou .csv attendu.")


def classeur_modele():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tiers"
    ws.append(ENTETES_MODELE)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    ws.append(["", "Client", "COHEN", "David", "12 rue Herzl", "4250000", "Netanya", "+972 50 000 0000", "david@example.org"])
    ws.append(["", "Fournisseur", "TRAITEUR EXEMPLE", "", "5 rue Allenby", "6100000", "Tel Aviv", "+972 3 000 0000",
               "contact@example.org"])
    for i, largeur in enumerate([14, 12, 24, 16, 28, 12, 16, 18, 26], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"
    aide = wb.create_sheet("Mode d'emploi")
    for ligne in ["Une ligne par tiers. Seule la colonne Nom est obligatoire (ou Compte pour mettre à jour un tiers existant).",
                  "Compte vide : le compte est créé d'après le type et le nom (411COHEN001, 401TRAIT001…).",
                  "Type : Client, Fournisseur ou un autre type des Référentiels ; vide = premier type (ou type du compte).",
                  "Un tiers déjà présent (même compte, ou même type + nom + prénom) est mis à jour ; une cellule vide ne remplace rien.",
                  "Supprimer les deux lignes d'exemple avant l'import."]:
        aide.append([ligne])
    aide.column_dimensions["A"].width = 120
    return wb
