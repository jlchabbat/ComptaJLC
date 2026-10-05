"""Reprise des données du classeur ComptaBB.xlsm (Lot 0 et suivants).

    python manage.py importer_classeur chemin/ComptaBB.xlsm [--remplacer]

Lit les tables du classeur (valeurs calculées par Excel au dernier
enregistrement) : T_Axe1, T_Axe2, T_Prefixes, T_PlanComptable, T_Journaux,
T_Ecritures, T_Journal, et les paramètres nommés P_*. Vérifie ensuite que les
totaux importés égalent ceux du classeur.
"""

import datetime as dt
from collections import defaultdict
from decimal import Decimal

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from compta.models import (
    ZERO, CodeAnalytique, Compte, Exercice, Journal, Ligne, Modification, Mouvement, Prefixe, Reglage, soldes,
)


def tables(wb):
    res = {}
    for ws in wb.worksheets:
        for t in ws.tables.values():
            lignes = list(ws[t.ref])
            entetes = [c.value for c in lignes[0]]
            res[t.name] = [dict(zip(entetes, (c.value for c in l))) for l in lignes[1:]]
    return res


def nom(wb, n):
    d = wb.defined_names.get(n)
    if d is None:
        return None
    for feuille, plage in d.destinations:
        return wb[feuille][plage.replace("$", "")].value


def texte(v):
    return "" if v is None else str(v).strip()


def montant(v):
    if v in (None, ""):
        return ZERO
    return Decimal(str(v)).quantize(Decimal("0.01"))


def jour(v):
    return v.date() if isinstance(v, dt.datetime) else v


class Command(BaseCommand):
    help = "Reprend les données du classeur ComptaBB.xlsm."

    def add_arguments(self, parser):
        parser.add_argument("classeur")
        parser.add_argument("--remplacer", action="store_true", help="Efface d'abord les données existantes.")

    @transaction.atomic
    def handle(self, classeur, remplacer=False, **options):
        if Mouvement.objects.exists() and not remplacer:
            raise CommandError("La base contient déjà des écritures : relancer avec --remplacer pour les effacer.")
        if remplacer:
            from compta.models import (Fiche, LigneFiche, LigneReleve, LigneSchema, ModeFiche, ModeleOperation, MoyenPaiement,
                                       NatureFiche, ParametreReleve, Rapprochement, TiersProvisoire, Traduction, TypeTiers)
            from compta.models import Budget, Membre
            for m in (Membre, Budget, LigneReleve, Rapprochement, ParametreReleve, Traduction, LigneFiche, Fiche, TiersProvisoire, NatureFiche, ModeFiche, Ligne, Mouvement, Modification, ModeleOperation, MoyenPaiement, LigneSchema, TypeTiers, Journal, Compte,
                      Prefixe, CodeAnalytique, Exercice, Reglage):
                m.objects.all().delete()
        wb = openpyxl.load_workbook(classeur, data_only=True)
        t = tables(wb)

        for axe, table in ((1, "T_Axe1"), (2, "T_Axe2")):
            for r in t[table]:
                if texte(r["Code"]):
                    statut = r.get("Actif")
                    CodeAnalytique.objects.create(code=texte(r["Code"]), axe=axe, libelle=texte(r["Libellé"]),
                                                  statut=int(statut) if isinstance(statut, (int, float)) else 1)
        for r in t["T_Prefixes"]:
            if texte(r["Préfixe"]):
                Prefixe.objects.create(prefixe=texte(r["Préfixe"]), axe=int(r["Axe"]), libelle=texte(r["Libellé"]))
        for r in t["T_PlanComptable"]:
            if texte(r["Compte"]):
                Compte.objects.create(numero=texte(r["Compte"]), libelle=texte(r["Libellé compte"]),
                                      anal1_id=texte(r["Axe 1 (Anal1)"]) or None, lettrable=texte(r.get("Lettrable")) == "oui")
        for r in t["T_Journaux"]:
            if texte(r["Code"]):
                numero = texte(r.get("N° de compte"))
                Journal.objects.create(code=texte(r["Code"]), intitule=texte(r["Intitulé"]), type=texte(r.get("Type")),
                                       compte_id=numero if numero and Compte.objects.filter(numero=numero).exists() else None,
                                       actif=r.get("Actif") in (1, None))

        # écritures : un mouvement par numéro de Mvt
        lignes = [r for r in t["T_Ecritures"] if r.get("Date")]
        par_mvt = defaultdict(list)
        for r in lignes:
            par_mvt[int(r["Mvt"])].append(r)
        for numero, ls in sorted(par_mvt.items()):
            premiere = ls[0]
            m = Mouvement.objects.create(numero=numero, date=jour(premiere["Date"]), journal_id=texte(premiere["Jnl"]),
                                         origine="import")
            Ligne.objects.bulk_create([
                Ligne(mouvement=m, ordre=i, compte_id=texte(r["Compte"]), libelle=texte(r["Libellé"]),
                      debit=montant(r["Débit"]), credit=montant(r["Crédit"]), anal2_id=(texte(r["Anal2"]) if texte(r["Compte"])[:1] in ("6", "7") else None), lettrage=texte(r.get("Let")))
                for i, r in enumerate(ls)])

        # exercices et réglages
        debut, fin = jour(nom(wb, "P_DebutExercice")), jour(nom(wb, "P_FinExercice"))
        cloture = jour(nom(wb, "P_DateCloture"))
        premiere_date = min(jour(r["Date"]) for r in lignes)
        if cloture and premiere_date <= cloture:
            Exercice.objects.create(libelle=f"Exercice court {premiere_date:%d/%m/%Y} – {cloture:%d/%m/%Y}", debut=premiere_date,
                                    fin=cloture, clos=True)
        if debut and fin:
            Exercice.objects.create(libelle=f"Exercice {debut.year}", debut=debut, fin=fin, clos=False)
        for cle, n, desc in (("compte_virement", "P_CompteVirement", "Virements internes et paiements carte Isracard (décision Q2)"),
                             ("compte_attente", "P_CompteAttente", "Opérations en attente d'éclaircissement (décision Q3)")):
            v = nom(wb, n)
            if v:
                Reglage.objects.create(cle=cle, valeur=texte(v), description=desc)

        for r in t.get("T_Journal", []):
            if texte(r.get("Action")):
                Modification.objects.create(auteur=texte(r["Auteur"]), lot=texte(r.get("Lot")), action=texte(r["Action"]),
                                            objet=texte(r["Objet"])[:200], avant=texte(r.get("Avant"))[:300],
                                            apres=texte(r.get("Après"))[:300])
        Modification.objects.create(auteur="Reprise", lot="Appli W0", action="Reprise",
                                    objet=f"classeur {classeur.rsplit('/', 1)[-1]}",
                                    apres=f"{len(par_mvt)} mouvements, {len(lignes)} lignes")

        from compta.saisie import initialiser_parametres
        initialiser_parametres()
        self.reprendre_releve(wb, t)

        # vérification
        d_x = sum(montant(r["Débit"]) for r in lignes)
        c_x = sum(montant(r["Crédit"]) for r in lignes)
        d, c, _ = soldes(Ligne.objects.all())
        if (d, c) != (d_x, c_x) or Ligne.objects.count() != len(lignes):
            raise CommandError(f"Totaux différents du classeur : {d}/{c} contre {d_x}/{c_x}")
        _, _, charges = soldes(Ligne.objects.filter(compte__numero__startswith="6"))
        _, _, produits = soldes(Ligne.objects.filter(compte__numero__startswith="7"))
        from compta.models import LigneReleve
        if LigneReleve.objects.exists():
            self.stdout.write(f"Relevé : {LigneReleve.objects.filter(ouverture=False).count()} lignes reprises")
        self.stdout.write(self.style.SUCCESS(
            f"{Mouvement.objects.count()} mouvements, {Ligne.objects.count()} lignes, débit = crédit = {d:,.2f}, "
            f"résultat {-(charges + produits):,.2f}"))

    def reprendre_releve(self, wb, t):
        """Relevé Banque 1 (T_Banque1), traductions (T_TradBanque) et paramètres de l'onglet Banque1, repérés par libellé."""
        from compta import releves
        from compta.models import ParametreReleve, Traduction
        for r in t.get("T_TradBanque", []):
            h, tr = texte(r.get("Opération (hébreu)")), texte(r.get("Traduction"))
            if h and tr and Traduction.cle_de(h):
                Traduction.objects.get_or_create(cle=Traduction.cle_de(h)[:120], defaults={"hebreu": h[:120], "traduction": tr[:120]})
        if "T_Banque1" not in t or "Banque1" not in wb.sheetnames:
            return
        params = {}
        for a, b in wb["Banque1"].iter_rows(min_col=1, max_col=2, values_only=True):
            if isinstance(a, str):
                params[a.split("(")[0].strip().lower()] = b
        journal = Journal.objects.filter(compte_id=texte(params.get("compte comptable de la banque"))).first()
        if not journal:
            return
        lignes = [{"date": jour(r["Date"]), "reference": texte(r["Référence"])[:40], "operation": texte(r["Opération (relevé)"])[:200],
                   "montant": montant(r["Montant"]), "solde": montant(r["Solde relevé"]) if r.get("Solde relevé") is not None else None}
                  for r in t["T_Banque1"] if r.get("Date") and r.get("Montant") not in (None, "")]
        releves.importer(journal, lignes, source="Reprise du classeur (T_Banque1)",
                         solde_ouverture=montant(params.get("solde d'ouverture du relevé")) if params.get("solde d'ouverture du relevé") else None)
        ParametreReleve.objects.update_or_create(journal=journal, defaults={
            "date_reprise": jour(params.get("date de reprise en compta")), "libelle": texte(wb["Banque1"]["A2"].value)[:200]})
        tol = params.get("tolérance de date pour le rapprochement")
        if isinstance(tol, (int, float)):
            Reglage.objects.update_or_create(cle="tolerance_rapprochement",
                                             defaults={"valeur": str(int(tol)), "description": "Tolérance de date du rapprochement (jours)"})
