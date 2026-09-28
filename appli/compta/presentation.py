"""Classeurs Excel « de présentation » (situation financière, états détaillés) : titre, sections, tableaux à bandes,
totaux, montants en ₪, mise en page A4 prête à imprimer."""

import io

from django.http import HttpResponse

BLEU, CLAIR, GRIS = "1F3864", "EEF3FB", "DDE1E7"
MONTANT = '#,##0.00 "₪";[Red]-#,##0.00 "₪";"–"'


class Presentation:
    def __init__(self, titre, sous_titre, largeurs, pied=""):
        import openpyxl
        from openpyxl.styles import Border, Font, PatternFill, Side
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = titre[:31].replace("/", "-")
        self.ws.sheet_view.showGridLines = False
        self.nb_colonnes = len(largeurs)
        for i, l in enumerate(largeurs, 1):
            self.ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = l
        self.ligne = 1
        self.police = {"titre": Font(bold=True, size=18, color=BLEU), "sous": Font(size=11, color="555555"),
                       "section": Font(bold=True, size=13, color=BLEU), "entete": Font(bold=True, color="FFFFFF"),
                       "gras": Font(bold=True), "sous_total": Font(bold=True, italic=True, color=BLEU)}
        self.fond = {"entete": PatternFill("solid", fgColor=BLEU), "clair": PatternFill("solid", fgColor=CLAIR),
                     "total": PatternFill("solid", fgColor=GRIS)}
        self.bas = Border(bottom=Side(style="thin", color=GRIS))
        self.trait = Border(bottom=Side(style="medium", color=BLEU))
        self.pied = pied

    def ecrire(self, valeurs, police=None, fond=None, montants=(), bordure=None, hauteur=None, formats=None):
        from openpyxl.styles import Alignment
        r = self.ligne
        for i, v in enumerate(valeurs, 1):
            c = self.ws.cell(r, i, v)
            if police:
                c.font = police
            if fond:
                c.fill = fond
            if bordure:
                c.border = bordure
            if i in montants:
                c.number_format = MONTANT
                c.alignment = Alignment(horizontal="right")
            if formats and i in formats:
                c.number_format = formats[i]
        if hauteur:
            self.ws.row_dimensions[r].height = hauteur
        self.ligne += 1
        return r

    def titre(self, titre, sous_titre):
        self.ecrire([titre], police=self.police["titre"], hauteur=30)
        for s in ([sous_titre] if isinstance(sous_titre, str) else sous_titre):
            self.ecrire([s], police=self.police["sous"])

    def section(self, titre):
        self.ligne += 1
        r = self.ecrire([titre], police=self.police["section"], hauteur=22)
        for i in range(1, self.nb_colonnes + 1):
            self.ws.cell(r, i).border = self.trait

    def note(self, texte):
        self.ecrire([texte], police=self.police["sous"])

    def tableau(self, entetes, lignes, montants=(), total=None, formats=None):
        self.ecrire(entetes, police=self.police["entete"], fond=self.fond["entete"])
        for k, l in enumerate(lignes):
            self.ecrire(l, fond=self.fond["clair"] if k % 2 else None, montants=montants, bordure=self.bas, formats=formats)
        if total:
            self.ecrire(total, police=self.police["gras"], fond=self.fond["total"], montants=montants)

    def sous_total(self, valeurs, montants=()):
        self.ecrire(valeurs, police=self.police["sous_total"], montants=montants, bordure=self.trait)

    def reponse(self, nom):
        from openpyxl.worksheet.page import PageMargins
        ws = self.ws
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
        ws.oddFooter.center.text = "Page &P / &N"
        ws.oddFooter.right.text = self.pied
        tampon = io.BytesIO()
        self.wb.save(tampon)
        r = HttpResponse(tampon.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        r["Content-Disposition"] = f'attachment; filename="{nom}"'
        return r
