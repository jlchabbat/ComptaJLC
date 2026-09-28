"""Suivi des membres (cahier des charges, Lot 2) : situation (solde), cotisations et lettrage des comptes de membres."""

import datetime as dt
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal

from django.db import transaction

from .models import ZERO, Compte, Ligne, Membre, Modification, Reglage, TypeTiers


def _cle_nom(nom):
    s = unicodedata.normalize("NFKD", nom.upper())
    return "".join(ch for ch in s if ch.isalpha() and ord(ch) < 128)[:5]


def compte_propose(nom, type_tiers=None):
    """Préfixe du type + 5 premières lettres du nom + rang : 411TAIEB001, 401PARTN001…"""
    t = type_tiers or TypeTiers.objects.filter(libelle="Membre").first()
    prefixe = (t.prefixe if t else "411") + _cle_nom(nom)
    rang = 1
    while Compte.objects.filter(numero=f"{prefixe}{rang:03d}").exists():
        rang += 1
    return f"{prefixe}{rang:03d}"


@transaction.atomic
def creer_tiers(type_tiers, nom, prenom="", **coordonnees):
    """Crée le compte (axe 1 repris d'un compte du même type) et la fiche du tiers."""
    libelle = f"{nom.strip().upper()} {prenom.strip().upper()}".strip()
    modele = Compte.objects.filter(numero__startswith=type_tiers.prefixe, anal1__isnull=False).first()
    compte = Compte.objects.create(numero=compte_propose(nom, type_tiers), libelle=libelle, lettrable=True,
                                   anal1=modele.anal1 if modele else None)
    Membre.objects.create(compte=compte, type=type_tiers, nom=nom.strip().upper(), prenom=prenom.strip(),
                          **{k: (v or "").strip() for k, v in coordonnees.items()})
    return compte


def type_du_compte(numero):
    """Type de tiers d'après le préfixe du compte (le plus long qui convient)."""
    types = sorted(TypeTiers.objects.all(), key=lambda t: -len(t.prefixe))
    return next((t for t in types if numero.startswith(t.prefixe)), None)


def creer_manquants():
    """Une fiche pour chaque compte de tiers (membres, fournisseurs…) ; les comptes collectifs de membres (tout en chiffres)
    sont exclus. Complète aussi le type des fiches qui n'en ont pas. Renvoie le nombre de fiches créées."""
    n = 0
    for t in TypeTiers.objects.all():
        for c in Compte.objects.filter(numero__startswith=t.prefixe, membre__isnull=True):
            if t.libelle == "Membre" and c.numero.isdigit():
                continue
            mots = c.libelle.split()
            if t.libelle == "Membre":
                Membre.objects.create(compte=c, type=t, nom=mots[0] if mots else c.libelle, prenom=" ".join(mots[1:]).title())
            else:
                Membre.objects.create(compte=c, type=t, nom=c.libelle)
            n += 1
    for m in Membre.objects.filter(type__isnull=True).select_related("compte"):
        m.type = type_du_compte(m.compte_id)
        if m.type:
            m.save(update_fields=["type"])
    return n


# ---------------------------------------------------------------- situation d'un compte de membre

@dataclass
class Situation:
    facture: Decimal = ZERO
    regle: Decimal = ZERO
    solde: Decimal = ZERO
    impayes: list = field(default_factory=list)       # (ligne, restant dû)
    avance: Decimal = ZERO


def allouer(lignes):
    """Règlements (crédits) imputés sur les factures (débits) les plus anciennes ; renvoie [(ligne débit, restant)] et l'avance."""
    debits = [[l, l.debit] for l in lignes if l.debit > 0]
    credit = sum((l.credit for l in lignes), ZERO)
    for d in debits:
        pris = min(d[1], credit)
        d[1] -= pris
        credit -= pris
    return [(l, r) for l, r in debits if r > 0], credit


def situation(compte, jusquau=None):
    jusquau = jusquau or dt.date.today()
    ls = list(Ligne.objects.filter(compte=compte, mouvement__date__lte=jusquau).exclude(mouvement__origine="cloture")
              .select_related("mouvement").order_by("mouvement__date", "mouvement__numero", "ordre"))
    s = Situation(facture=sum((l.debit for l in ls), ZERO), regle=sum((l.credit for l in ls), ZERO))
    s.solde = s.facture - s.regle
    restants, s.avance = allouer([l for l in ls if not l.lettrage])
    s.impayes = restants
    return s


def cotisations(exercice):
    """Cotisations de l'exercice par membre : attendue (membre actif), facturée, reçue, restant dû."""
    compte_cot = Reglage.lire("compte_cotisations")
    res = []
    for m in Membre.objects.filter(type__libelle="Membre").select_related("compte"):
        mvts = Ligne.objects.filter(compte=m.compte, debit__gt=0, mouvement__date__range=(exercice.debut, exercice.fin)).values("mouvement")
        cot = Ligne.objects.filter(mouvement__in=mvts, compte_id=compte_cot)
        facturee = sum((l.credit - l.debit for l in cot), ZERO)
        s = situation(m.compte, exercice.fin)
        du = sum((r for l, r in s.impayes if l.mouvement.date >= exercice.debut
                  and l.mouvement.lignes.filter(compte_id=compte_cot).exists()), ZERO)
        attendue = (m.cotisation or ZERO) if m.statut == "actif" else ZERO
        if attendue or facturee:
            res.append({"membre": m, "attendue": attendue, "facturee": facturee, "recue": facturee - du, "du": du})
    tot = {k: sum((r[k] for r in res), ZERO) for k in ("attendue", "facturee", "recue", "du")}
    tot["taux"] = round(tot["recue"] * 100 / tot["facturee"], 1) if tot["facturee"] else None
    tot["compte"] = compte_cot
    return res, tot


def montant(v):
    from .reglages import montant
    return montant(v)


# ---------------------------------------------------------------- lettrage

def code_suivant(compte):
    """A, B, … Z, AA, AB… : premier code libre du compte."""
    pris = set(Ligne.objects.filter(compte=compte).exclude(lettrage="").values_list("lettrage", flat=True))
    n = 0
    while True:
        code, k = "", n
        while True:
            code = chr(65 + k % 26) + code
            k = k // 26 - 1
            if k < 0:
                break
        if code not in pris:
            return code
        n += 1


@transaction.atomic
def lettrer(compte, lignes, auteur=""):
    lignes = list(lignes)
    if len(lignes) < 2 or any(l.compte_id != compte.numero or l.lettrage for l in lignes):
        raise ValueError("Choisir au moins deux lignes non lettrées de ce compte.")
    d, c = sum((l.debit for l in lignes), ZERO), sum((l.credit for l in lignes), ZERO)
    if d != c:
        raise ValueError(f"Débit {montant(d)} ≠ crédit {montant(c)} : le lettrage doit être équilibré.")
    code = code_suivant(compte)
    Ligne.objects.filter(pk__in=[l.pk for l in lignes]).update(lettrage=code)
    if auteur:
        Modification.objects.create(auteur=auteur, lot="Membres", action="Lettrage", objet=f"compte {compte.numero}",
                                    apres=f"{code} : {len(lignes)} lignes, {montant(d)}")
    return code


@transaction.atomic
def delettrer(compte, code, auteur=""):
    n = Ligne.objects.filter(compte=compte, lettrage=code).update(lettrage="")
    if n and auteur:
        Modification.objects.create(auteur=auteur, lot="Membres", action="Délettrage", objet=f"compte {compte.numero}", avant=code)
    return n


def lettrage_automatique(compte, auteur=""):
    """1) facture et règlement de même montant dans un même mouvement ; 2) facture puis premier règlement
    ultérieur de même montant. Renvoie le nombre de lettrages."""
    ls = list(Ligne.objects.filter(compte=compte, lettrage="").exclude(mouvement__origine="cloture")
              .select_related("mouvement").order_by("mouvement__date", "mouvement__numero", "ordre"))
    libres = {l.pk: l for l in ls}
    paires = []
    for l in ls:
        if l.debit > 0 and l.pk in libres:
            meme = next((x for x in ls if x.pk in libres and x.mouvement_id == l.mouvement_id and x.credit == l.debit), None)
            if meme:
                paires.append((l, meme))
                del libres[l.pk], libres[meme.pk]
    for l in ls:
        if l.debit > 0 and l.pk in libres:
            suite = next((x for x in ls if x.pk in libres and x.credit == l.debit and x.mouvement.date >= l.mouvement.date), None)
            if suite:
                paires.append((l, suite))
                del libres[l.pk], libres[suite.pk]
    for a, b in paires:
        lettrer(compte, [a, b])
    if paires and auteur:
        Modification.objects.create(auteur=auteur, lot="Membres", action="Lettrage automatique", objet=f"compte {compte.numero}",
                                    apres=f"{len(paires)} lettrage(s)")
    return len(paires)


# ---------------------------------------------------------------- import du modèle 07_membres.csv

def cle(texte):
    return re.sub(r"[^a-z]", "", unicodedata.normalize("NFKD", texte or "").encode("ascii", "ignore").decode().lower())


COLONNES_TIERS = {
    "compte": ["compte", "numero", "ncompte", "numerodecompte", "comptetiers"],
    "type": ["type", "typedetiers", "categorie"],
    "nom": ["nom", "raisonsociale", "nomouraisonsociale", "societe", "nomraisonsociale"],
    "prenom": ["prenom"],
    "adresse": ["adresse", "rue", "adresse1"],
    "code_postal": ["codepostal", "cp", "zip"],
    "ville": ["ville", "localite", "commune"],
    "telephone": ["telephone", "tel", "portable", "mobile", "gsm"],
    "email": ["email", "mail", "courriel", "adressemail", "adresseemail"],
    "date_adhesion": ["datedadhesion", "dateadhesion", "adhesion"],
    "statut": ["statut"],
    "cotisation": ["cotisation", "cotisationannuelle", "cotisationattendue"],
}
ENTETES_MODELE = ["Compte", "Type", "Nom", "Prénom", "Adresse", "Code postal", "Ville", "Téléphone", "E-mail",
                  "Date d'adhésion", "Statut", "Cotisation annuelle"]


def _colonnes(entetes):
    res = {}
    for i, e in enumerate(entetes):
        k = cle(str(e or ""))
        for champ, noms in COLONNES_TIERS.items():
            if k in noms and champ not in res:
                res[champ] = i
    return res


@dataclass
class RapportImport:
    crees: list = field(default_factory=list)
    mis_a_jour: list = field(default_factory=list)
    erreurs: list = field(default_factory=list)


@transaction.atomic
def importer_tableau(rangees):
    """Importe des tiers (Tiers.xlsx ou CSV). Une ligne = un tiers ; en-têtes reconnus sous plusieurs formes.

    Compte présent : fiche de ce compte (compte créé s'il manque et que son préfixe correspond à un type).
    Compte absent : tiers retrouvé par type + nom + prénom, sinon créé avec un compte proposé.
    Une cellule vide ne remplace jamais une valeur déjà saisie ; rien n'est supprimé."""
    from .releves import date as lire_date, nombre
    rapport, cols, defaut = RapportImport(), None, TypeTiers.objects.filter(libelle="Membre").first()
    for n, r in enumerate(rangees, 1):
        r = list(r)
        if cols is None:
            c = _colonnes(r)
            if "nom" in c or "compte" in c:
                cols = c
            continue
        val = lambda k: ("" if k not in cols or cols[k] >= len(r) or r[cols[k]] is None else r[cols[k]])  # noqa: E731
        texte = lambda k: str(val(k)).strip()  # noqa: E731
        numero, nom = texte("compte"), texte("nom")
        if not numero and not nom:
            continue
        numero = numero[:-2] if numero.endswith(".0") else numero          # nombre lu par Excel
        t = None
        if texte("type"):
            t = next((x for x in TypeTiers.objects.all() if cle(x.libelle) == cle(texte("type"))), None)
            if not t:
                rapport.erreurs.append(f"Ligne {n} : type « {texte('type')} » inconnu (Référentiels › Types de tiers).")
                continue
        fiche = None
        if numero:
            compte = Compte.objects.filter(numero=numero).first()
            t = t or type_du_compte(numero)
            if not t or not numero.startswith(t.prefixe):
                rapport.erreurs.append(f"Ligne {n} : le compte {numero} ne correspond à aucun type de tiers.")
                continue
            if not compte:
                if not nom:
                    rapport.erreurs.append(f"Ligne {n} : nom manquant pour créer le compte {numero}.")
                    continue
                modele = Compte.objects.filter(numero__startswith=t.prefixe, anal1__isnull=False).first()
                compte = Compte.objects.create(numero=numero, libelle=f"{nom.upper()} {texte('prenom').upper()}".strip(),
                                               lettrable=True, anal1=modele.anal1 if modele else None)
            fiche = Membre.objects.filter(compte=compte).first()
            nouveau = fiche is None
            fiche = fiche or Membre(compte=compte, nom=(nom or compte.libelle).upper())
        else:
            t = t or defaut
            fiche = Membre.objects.filter(type=t, nom__iexact=nom, prenom__iexact=texte("prenom")).first()
            nouveau = fiche is None
            if nouveau:
                fiche = Membre.objects.get(compte=creer_tiers(t, nom, texte("prenom")))
        fiche.type = fiche.type or t
        for champ in ("prenom", "adresse", "code_postal", "ville", "telephone", "email"):
            if texte(champ):
                setattr(fiche, champ, texte(champ))
        if nom:
            fiche.nom = nom.upper()
        if val("date_adhesion") != "":
            fiche.date_adhesion = lire_date(val("date_adhesion")) or fiche.date_adhesion
        if val("cotisation") != "":
            fiche.cotisation = nombre(val("cotisation"))
        if texte("statut") and cle(texte("statut")) in dict(Membre.STATUTS):
            fiche.statut = cle(texte("statut"))
        if texte("email") and "@" not in texte("email"):
            rapport.erreurs.append(f"Ligne {n} : e-mail « {texte('email')} » ignoré.")
            fiche.email = ""
        fiche.save()
        (rapport.crees if nouveau else rapport.mis_a_jour).append(f"{fiche.compte_id} {fiche}")
    if cols is None:
        rapport.erreurs.append("En-têtes non reconnus : une colonne « Nom » ou « Compte » est nécessaire.")
    return rapport


def lire_tableau(nom_fichier, contenu):
    """Rangées d'un fichier .xlsx (première feuille contenant des en-têtes reconnus) ou .csv."""
    import csv
    import io
    if nom_fichier.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contenu), data_only=True, read_only=True)
        for ws in wb.worksheets:
            rangees = [list(r) for r in ws.iter_rows(values_only=True)]
            if any(_colonnes(r).keys() & {"nom", "compte"} for r in rangees[:20]):
                return rangees
        return []
    if nom_fichier.lower().endswith(".csv"):
        s = contenu.decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(s), delimiter=";" if s.count(";") >= s.count(",") else ","))
    raise ValueError("Format non reconnu : fichier .xlsx ou .csv attendu.")


def classeur_modele():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tiers"
    ws.append(ENTETES_MODELE)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    ws.append(["", "Membre", "COHEN", "David", "12 rue Herzl", "4250000", "Netanya", "+972 50 000 0000", "david@example.org",
               "01/09/2020", "actif", 500])
    ws.append(["", "Fournisseur", "TRAITEUR EXEMPLE", "", "5 rue Allenby", "6100000", "Tel Aviv", "+972 3 000 0000",
               "contact@example.org", "", "", ""])
    for i, largeur in enumerate([14, 12, 24, 16, 28, 12, 16, 18, 26, 14, 12, 12], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"
    aide = wb.create_sheet("Mode d'emploi")
    for ligne in ["Une ligne par tiers. Seule la colonne Nom est obligatoire (ou Compte pour mettre à jour un tiers existant).",
                  "Compte vide : le compte est créé d'après le type et le nom (411COHEN001, 401TRAIT001…).",
                  "Type : Membre, Fournisseur ou un autre type des Référentiels ; vide = Membre (ou type du compte).",
                  "Un tiers déjà présent (même compte, ou même type + nom + prénom) est mis à jour ; une cellule vide ne remplace rien.",
                  "Statut (membres) : actif, honoraire ou démissionnaire. Dates : jj/mm/aaaa. Montants : 500 ou 500,00.",
                  "Supprimer les deux lignes d'exemple avant l'import."]:
        aide.append([ligne])
    aide.column_dimensions["A"].width = 120
    return wb
