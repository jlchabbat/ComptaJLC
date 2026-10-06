"""Export Excel figé en valeurs (archive de clôture, téléchargement des états)."""

import datetime as dt

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import etats
from .models import Ligne

ENTETE = Font(bold=True, color="FFFFFF")
FOND = PatternFill("solid", fgColor="1F3864")
MONTANT = '#,##0.00;-#,##0.00;""'


def feuille(wb, titre, entetes, lignes, montants=()):
    ws = wb.create_sheet(titre[:31])
    ws.append(entetes)
    for c in ws[1]:
        c.font, c.fill = ENTETE, FOND
    for l in lignes:
        ws.append([float(v) if hasattr(v, "quantize") else v for v in l])
    for i in montants:
        for c in ws.iter_cols(min_col=i, max_col=i, min_row=2):
            for x in c:
                x.number_format = MONTANT
    for i, e in enumerate(entetes, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(50, len(str(e)) + 4))
    ws.freeze_panes = "A2"
    return ws


def classeur_exercice(ex):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    prec = etats.precedent(ex)
    cr = etats.compte_de_resultat(ex, prec)
    lignes = []
    for s in cr["sections"]:
        lignes += [[l["numero"], l["libelle"], l["n"], l["n1"]] for l in s["lignes"]]
        lignes.append(["", f"Total {s['titre'].lower()}", s["total_n"], s["total_n1"]])
    lignes.append(["", "Résultat", cr["resultat_n"], cr["resultat_n1"]])
    feuille(wb, "Compte de résultat", ["Compte", "Libellé", ex.libelle, prec.libelle if prec else "N-1"], lignes, (3, 4))
    feuille(wb, "Résultat Anal", ["Code", "Libellé", "Produits", "Charges", "Résultat", "Résultat N-1"],
            [[r["code"], r["libelle"], r["produits"], r["charges"], r["resultat"], r["resultat_n1"]]
             for r in etats.par_anal(ex, prec)], (3, 4, 5, 6))
    b = etats.bilan(ex.fin, ex.debut)
    lignes = [["Actif", "", "", ""]]
    for _, lib, ls in b["actif"]:
        lignes += [["", l["numero"], l["libelle"], l["montant"]] for l in ls]
    lignes += [["Total actif", "", "", b["total_actif"]], ["Passif", "", "", ""]]
    for _, lib, ls in b["passif"]:
        lignes += [["", l["numero"], l["libelle"], l["montant"]] for l in ls]
    lignes += [["", "", "Résultat de l'exercice", b["resultat"]], ["", "", "Résultats antérieurs non reportés", b["resultats_anterieurs"]],
               ["Total passif", "", "", b["total_passif"]]]
    feuille(wb, "Bilan", ["", "Compte", "Libellé", f"Au {ex.fin:%d/%m/%Y}"], lignes, (4,))
    bal = etats.soldes_par_compte(Ligne.objects.filter(mouvement__date__range=(ex.debut, ex.fin)))
    feuille(wb, "Balance", ["Compte", "Libellé", "Solde"], [[k, lib, s] for k, (lib, s) in bal.items()], (3,))
    ecr = (Ligne.objects.filter(mouvement__date__range=(ex.debut, ex.fin))
           .select_related("mouvement", "compte", "compte__anal1").order_by("compte__numero", "mouvement__date", "mouvement__numero", "ordre"))
    feuille(wb, "Grand livre", ["Compte", "Intitulé", "Date", "Jnl", "Mvt", "Libellé", "Débit", "Crédit", "Anal", "Let"],
            [[l.compte_id, l.compte.libelle, l.mouvement.date, l.mouvement.journal_id, l.mouvement.numero,
              l.libelle, l.debit, l.credit, l.compte.anal1_id, l.lettrage] for l in ecr], (7, 8))
    bud = etats.budget(ex)
    if bud:
        feuille(wb, "Budget", ["Nature", "Cible", "Budget", "Réalisé", "Écart", "Écart %"],
                [[x["budget"].get_nature_display(), str(x["budget"].cible), x["budget"].montant, x["realise"], x["ecart"],
                  x["pourcent"]] for x in bud], (3, 4, 5))
    from .models import Modification
    hist = Modification.objects.filter(date__date__range=(ex.debut, ex.fin + dt.timedelta(days=90)))
    feuille(wb, "Historique", ["Date", "Auteur", "Lot", "Action", "Objet", "Avant", "Après"],
            [[m.date.replace(tzinfo=None), m.auteur, m.lot, m.action, m.objet, m.avant, m.apres] for m in hist])
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if hasattr(c.value, "year") and not isinstance(c.value, str):
                    c.number_format = "DD/MM/YYYY"
    return wb
