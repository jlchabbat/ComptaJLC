"""Classeur Parametres.xlsx : les paramètres de ComptaBB exportés, modifiés dans Excel puis réinjectés.

Une feuille par table de paramètres. À l'import :
- une ligne dont la clé existe met à jour l'élément ; une clé inconnue le crée ;
- une cellule vide ne remplace rien ; rien n'est jamais supprimé (mettre Actif = Non pour retirer) ;
- une seule erreur et rien n'est enregistré : le rapport liste toutes les lignes à corriger.

Les comptes de tiers (membres, fournisseurs…) n'y figurent pas : ils ont leur propre fichier Tiers.xlsx.
"""

import datetime as dt
import io
import re
from dataclasses import dataclass, field

from django.db import transaction

from .models import (
    CodeAnalytique, Compte, FICHE_TYPES, Journal, LigneSchema, Membre, ModeFiche, ModeleOperation, Modification,
    MoyenPaiement, NatureFiche, ParametreReleve, Prefixe, Reglage, SENS_FICHE, Traduction, TypeTiers,
)

NOM_FICHIER = "Parametres.xlsx"
OUI, NON = "Oui", "Non"


class Refus(ValueError):
    """Valeur de cellule refusée ; le message est affiché tel quel avec la feuille et la ligne."""


def cle(texte):
    """Minuscules sans accents ni ponctuation, chiffres gardés (« Axe 1 » ≠ « Axe 2 »)."""
    import unicodedata
    if isinstance(texte, float) and texte.is_integer():
        texte = int(texte)
    s = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


# ---------------------------------------------------------------- types de colonnes

class Texte:
    def __init__(self, maxi=None, majuscules=False):
        self.maxi, self.majuscules = maxi, majuscules

    def vers_excel(self, v):
        return v

    def lire(self, v):
        if isinstance(v, float) and v.is_integer():
            v = int(v)                                  # 512000 lu par Excel comme nombre
        s = str(v).strip()
        if self.majuscules:
            s = s.upper()
        if self.maxi and len(s) > self.maxi:
            raise Refus(f"« {s} » dépasse {self.maxi} caractères")
        return s


class Entier:
    def vers_excel(self, v):
        return v

    def lire(self, v):
        try:
            n = float(str(v).replace(",", ".").strip())
        except ValueError:
            raise Refus(f"« {v} » n'est pas un nombre entier") from None
        if not n.is_integer() or n < 0:
            raise Refus(f"« {v} » n'est pas un nombre entier positif")
        return int(n)


class Booleen:
    def vers_excel(self, v):
        return OUI if v else NON

    def lire(self, v):
        k = cle(v)
        if k in ("oui", "o", "vrai", "true", "x", "1"):
            return True
        if k in ("non", "n", "faux", "false", "0"):
            return False
        raise Refus(f"« {v} » : Oui ou Non attendu")


class Date:
    def vers_excel(self, v):
        return v

    def lire(self, v):
        if isinstance(v, dt.datetime):
            return v.date()
        if isinstance(v, dt.date):
            return v
        from .releves import date as lire_date
        d = lire_date(v)
        if not d:
            raise Refus(f"« {v} » n'est pas une date (jj/mm/aaaa)")
        return d


class Choix:
    """Valeur parmi une liste : l'export montre le libellé, l'import accepte le code ou le libellé."""

    def __init__(self, choix):
        self.choix = [(k, str(lib)) for k, lib in choix]

    def vers_excel(self, v):
        return dict(self.choix).get(v, v)

    def lire(self, v):
        k = cle(v)
        for code, lib in self.choix:
            if k and k in (cle(code), cle(lib)):
                return code
        raise Refus(f"« {v} » : valeurs possibles {', '.join(lib for _, lib in self.choix)}")


class Reference:
    """Clé d'une autre table (compte, journal, code analytique…)."""

    def __init__(self, modele, champ, nom, filtre=None):
        self.modele, self.champ, self.nom, self.filtre = modele, champ, nom, filtre or {}

    def vers_excel(self, obj):
        return None if obj is None else getattr(obj, self.champ)

    def lire(self, v):
        s = Texte().lire(v)
        qs = self.modele.objects.filter(**self.filtre)
        obj = qs.filter(**{self.champ: s}).first() or qs.filter(**{self.champ + "__iexact": s}).first()
        if obj is None:
            raise Refus(f"{self.nom} « {s} » inconnu(e)")
        return obj


@dataclass
class Colonne:
    entete: str
    champ: str
    type: object
    obligatoire: bool = False      # à la création
    largeur: int = 14
    aide: str = ""


@dataclass
class Feuille:
    nom: str
    modele: type
    cles: list                      # champs qui identifient une ligne
    colonnes: list
    aide: str
    requete: object = None          # éléments exportés (défaut : tous)
    fixe: dict = field(default_factory=dict)   # valeurs imposées (axe du code analytique…)
    preparer: object = None         # complète un nouvel objet (clé de traduction…)
    verifier: object = None         # contrôle de l'objet complet, lève Refus
    trouver: object = None          # recherche de l'élément existant (défaut : par les clés)

    def elements(self):
        return self.requete() if self.requete else self.modele.objects.all()


# ---------------------------------------------------------------- définition des feuilles

def _comptes_hors_tiers():
    tiers = Membre.objects.values_list("compte_id", flat=True)
    return Compte.objects.exclude(numero__in=tiers).select_related("anal1")


def _verifier_modele(m):
    if not LigneSchema.objects.filter(schema=m.schema).exists():
        raise Refus(f"schéma « {m.schema} » inconnu (schémas : {', '.join(sorted(set(LigneSchema.objects.values_list('schema', flat=True))))})")


def _verifier_mode(m):
    if m.genre == "TRESO" and not m.compte_id:
        raise Refus("un mode « banque, caisse ou carte » demande un compte de trésorerie")


def _preparer_traduction(t):
    t.cle = Traduction.cle_de(t.hebreu)
    if not t.cle:
        raise Refus("opération (hébreu) vide")


STATUTS_ANALYTIQUES = Choix(CodeAnalytique.STATUTS)
AXE1 = Reference(CodeAnalytique, "code", "code axe 1", {"axe": 1})
COMPTE = Reference(Compte, "numero", "compte")
JOURNAL = Reference(Journal, "code", "journal")

FEUILLES = [
    Feuille("Réglages", Reglage, ["cle"], [
        Colonne("Clé", "cle", Texte(40), True, 26),
        Colonne("Valeur", "valeur", Texte(200), True, 22),
        Colonne("Description", "description", Texte(200), False, 60),
    ], "Hypothèses nommées : compte de virement interne, compte d'attente, compte des cotisations…"),
    Feuille("Axe 1", CodeAnalytique, ["code"], [
        Colonne("Code", "code", Texte(20), True, 14),
        Colonne("Libellé", "libelle", Texte(100), False, 40),
        Colonne("Statut", "statut", STATUTS_ANALYTIQUES, False, 14, "Non affecté, En cours ou Terminé"),
    ], "Codes de nature (axe 1).", requete=lambda: CodeAnalytique.objects.filter(axe=1), fixe={"axe": 1}),
    Feuille("Axe 2", CodeAnalytique, ["code"], [
        Colonne("Code", "code", Texte(20), True, 14),
        Colonne("Libellé", "libelle", Texte(100), False, 40),
        Colonne("Statut", "statut", STATUTS_ANALYTIQUES, False, 14, "Non affecté, En cours ou Terminé"),
    ], "Codes d'événement ou de projet (axe 2).", requete=lambda: CodeAnalytique.objects.filter(axe=2), fixe={"axe": 2}),
    Feuille("Préfixes", Prefixe, ["prefixe"], [
        Colonne("Préfixe", "prefixe", Texte(10), True, 12),
        Colonne("Axe", "axe", Choix([(1, "1"), (2, "2")]), True, 8, "1 ou 2"),
        Colonne("Libellé", "libelle", Texte(60), False, 40),
    ], "Préfixes des codes analytiques proposés par l'application (ACT, MAN…)."),
    Feuille("Plan comptable", Compte, ["numero"], [
        Colonne("Compte", "numero", Texte(20), True, 14),
        Colonne("Libellé", "libelle", Texte(100), True, 42),
        Colonne("Axe 1", "anal1", AXE1, False, 12, "code de la feuille Axe 1"),
        Colonne("Lettrable", "lettrable", Booleen(), False, 11, "Oui ou Non"),
        Colonne("Actif", "actif", Booleen(), False, 9, "Oui ou Non"),
    ], "Comptes généraux. Les comptes de tiers (membres, fournisseurs…) sont dans Tiers.xlsx.",
        requete=_comptes_hors_tiers),
    Feuille("Journaux", Journal, ["code"], [
        Colonne("Code", "code", Texte(10), True, 10),
        Colonne("Intitulé", "intitule", Texte(60), True, 30),
        Colonne("Type", "type", Texte(10), False, 12),
        Colonne("Compte", "compte", COMPTE, False, 14, "compte de trésorerie (banque, caisse)"),
        Colonne("Actif", "actif", Booleen(), False, 9, "Oui ou Non"),
    ], "Journaux (banques, caisse, achats, ventes, opérations diverses)."),
    Feuille("Types de tiers", TypeTiers, ["libelle"], [
        Colonne("Libellé", "libelle", Texte(30), True, 20),
        Colonne("Préfixe de compte", "prefixe", Texte(10), True, 18),
    ], "Catégories de tiers et début de leurs numéros de compte (Membre 411, Fournisseur 401…)."),
    Feuille("Moyens de paiement", MoyenPaiement, ["libelle"], [
        Colonne("Libellé", "libelle", Texte(40), True, 26),
        Colonne("Journal", "journal", JOURNAL, False, 10, "vide = non réglé (facture seule)"),
        Colonne("Ordre", "ordre", Entier(), False, 8),
    ], "Moyens de paiement proposés à la saisie guidée."),
    Feuille("Schémas", LigneSchema, ["schema", "ligne"], [
        Colonne("Schéma", "schema", Texte(4, majuscules=True), True, 9),
        Colonne("Ligne", "ligne", Entier(), True, 8),
        Colonne("Mvt", "mvt", Entier(), False, 7),
        Colonne("Rôle", "role", Choix(LigneSchema.ROLES), True, 28),
        Colonne("Sens", "sens", Choix([("D", "Débit"), ("C", "Crédit")]), True, 9, "Débit ou Crédit"),
        Colonne("Seulement si réglé", "si_regle", Booleen(), False, 12, "Oui ou Non"),
        Colonne("Journal", "journal", Choix(LigneSchema.JOURNAUX), True, 44),
    ], "Lignes d'écritures générées par chaque schéma de la saisie guidée (à modifier avec prudence)."),
    Feuille("Modèles d'opération", ModeleOperation, ["type"], [
        Colonne("Type d'opération", "type", Texte(60), True, 34),
        Colonne("Schéma", "schema", Texte(4, majuscules=True), True, 9, "RT, DT, RS, DS, RM, RF, VI, CB"),
        Colonne("Compte", "compte", COMPTE, False, 12, "compte proposé ; vide = à choisir"),
        Colonne("Journal par défaut", "journal_defaut", JOURNAL, False, 12),
        Colonne("Tiers", "tiers", Reference(TypeTiers, "libelle", "type de tiers"), False, 14, "type de tiers ; vide = sans tiers"),
        Colonne("Paiement obligatoire", "paiement_obligatoire", Booleen(), False, 12, "Oui ou Non"),
        Colonne("Classe", "classe", Texte(1), False, 8, "6 ou 7"),
        Colonne("Libellé type", "libelle_type", Texte(40), True, 22),
        Colonne("Aide", "aide", Texte(200), False, 50),
        Colonne("Ordre", "ordre", Entier(), False, 8),
        Colonne("Actif", "actif", Booleen(), False, 9, "Oui ou Non"),
    ], "Modèles de la saisie guidée.", verifier=_verifier_modele),
    Feuille("Natures fiches", NatureFiche, ["type_fiche", "sens", "libelle"], [
        Colonne("Fiche", "type_fiche", Choix(FICHE_TYPES), True, 11, "Activité ou Gestion"),
        Colonne("Sens", "sens", Choix(SENS_FICHE), True, 10, "Recette ou Dépense"),
        Colonne("Libellé", "libelle", Texte(60), True, 30),
        Colonne("Compte", "compte", COMPTE, True, 12),
        Colonne("Libellé d'écriture", "libelle_ecriture", Texte(40), True, 24),
        Colonne("Ordre", "ordre", Entier(), False, 8),
    ], "Natures proposées sur les fiches bénévoles."),
    Feuille("Modes fiches", ModeFiche, ["type_fiche", "sens", "libelle"], [
        Colonne("Fiche", "type_fiche", Choix(FICHE_TYPES), True, 11, "Activité ou Gestion"),
        Colonne("Sens", "sens", Choix(SENS_FICHE + [("*", "Les deux")]), True, 10, "Recette, Dépense ou Les deux"),
        Colonne("Libellé", "libelle", Texte(40), True, 24),
        Colonne("Genre", "genre", Choix(ModeFiche.GENRES), True, 28),
        Colonne("Journal", "journal", JOURNAL, True, 10),
        Colonne("Compte", "compte", COMPTE, False, 12, "obligatoire pour banque, caisse ou carte"),
        Colonne("Ordre", "ordre", Entier(), False, 8),
    ], "Modes de paiement des fiches bénévoles.", verifier=_verifier_mode),
    Feuille("Relevés", ParametreReleve, ["journal"], [
        Colonne("Journal", "journal", JOURNAL, True, 10),
        Colonne("Date de reprise", "date_reprise", Date(), False, 16, "jj/mm/aaaa"),
        Colonne("Description du relevé", "libelle", Texte(200), False, 50),
    ], "Paramètres du rapprochement bancaire, par journal.", requete=lambda: ParametreReleve.objects.select_related("journal")),
    Feuille("Traductions", Traduction, ["hebreu"], [
        Colonne("Opération (hébreu)", "hebreu", Texte(120), True, 40),
        Colonne("Traduction", "traduction", Texte(120), True, 40),
    ], "Libellés des relevés bancaires et leur traduction.", preparer=_preparer_traduction,
        trouver=lambda v: Traduction.objects.filter(cle=Traduction.cle_de(v["hebreu"])).first()),
]


# ---------------------------------------------------------------- export

def classeur(seulement=None):
    """Classeur des paramètres actuels (openpyxl.Workbook) ; seulement = noms des feuilles à garder (un référentiel)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb._named_styles["Normal"].font = Font(name="Calibri", size=12)      # lisible sans zoom
    aide = wb.active
    aide.title = "Mode d'emploi"
    lignes = [
        ["Paramètres de ComptaBB", ""],
        [f"Exporté le {dt.datetime.now():%d/%m/%Y à %H:%M}", ""],
        ["", ""],
        ["1. Modifier les feuilles dans Excel (ajouter des lignes, corriger des libellés…). Ne pas renommer les feuilles ni les en-têtes.", ""],
        ["2. Enregistrer en .xlsx, puis dans ComptaBB : Administration › Paramètres (Excel) › Importer.", ""],
        ["3. Une ligne dont la clé (colonne grise) existe met à jour l'élément ; une clé nouvelle le crée.", ""],
        ["4. Une cellule vide ne change rien ; rien n'est jamais supprimé : pour retirer un élément, mettre Actif = Non.", ""],
        ["5. Au moindre problème, rien n'est enregistré et la liste des lignes à corriger s'affiche.", ""],
        ["   Une sauvegarde de la base est faite juste avant chaque import.", ""],
        ["", ""],
        ["Feuille", "Contenu"],
    ] + [[f.nom, f.aide] for f in FEUILLES if not seulement or f.nom in seulement]
    for l in lignes:
        aide.append(l)
    for rangee in aide.iter_rows():
        for c in rangee:
            c.font = Font(size=12)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    aide["A1"].font = Font(bold=True, size=16)
    for c in aide[11]:
        c.font = Font(bold=True, size=12)
    aide.column_dimensions["A"].width = 26
    aide.column_dimensions["B"].width = 95

    titre, cle_fond = PatternFill("solid", fgColor="1F3864"), PatternFill("solid", fgColor="D9D9D9")
    for f in FEUILLES:
        if seulement and f.nom not in seulement:
            continue
        ws = wb.create_sheet(f.nom)
        ws.append([c.entete for c in f.colonnes])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF", size=12), titre
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for obj in f.elements():
            ws.append([c.type.vers_excel(getattr(obj, c.champ)) for c in f.colonnes])
        for rangee in ws.iter_rows(min_row=2):
            for cellule in rangee:
                cellule.font = Font(size=12)
        ws.row_dimensions[1].height = 34
        for i, c in enumerate(f.colonnes, 1):
            lettre = get_column_letter(i)
            ws.column_dimensions[lettre].width = c.largeur + 3
            if c.champ in f.cles:
                for cellule in ws[lettre][1:]:
                    cellule.fill = cle_fond
            if isinstance(c.type, Date):
                for cellule in ws[lettre][1:]:
                    cellule.number_format = "DD/MM/YYYY"
            if c.aide:
                from openpyxl.comments import Comment
                ws.cell(1, i).comment = Comment(c.aide, "ComptaBB")
        ws.freeze_panes = "A2"
    return wb


def contenu_classeur(seulement=None):
    tampon = io.BytesIO()
    classeur(seulement).save(tampon)
    return tampon.getvalue()


# ---------------------------------------------------------------- import

@dataclass
class Rapport:
    crees: list = field(default_factory=list)
    modifies: list = field(default_factory=list)
    inchanges: int = 0
    erreurs: list = field(default_factory=list)
    feuilles: list = field(default_factory=list)

    @property
    def resume(self):
        return f"{len(self.crees)} créé(s), {len(self.modifies)} modifié(s), {self.inchanges} inchangé(s)"


class _Annulation(Exception):
    pass


def _vide(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _affiche(v):
    if v is None:
        return ""
    if isinstance(v, dt.date):
        return f"{v:%d/%m/%Y}"
    return str(v)


def _importer_feuille(f, rangees, rapport, auteur, tracer=True):
    entetes = [cle(e) for e in (rangees[0] if rangees else [])]
    positions = {}
    for c in f.colonnes:
        if cle(c.entete) in entetes:
            positions[c.champ] = entetes.index(cle(c.entete))
    manquantes = [c.entete for c in f.colonnes if c.champ in f.cles and c.champ not in positions]
    if manquantes:
        rapport.erreurs.append(f"{f.nom} : colonne(s) {', '.join(manquantes)} introuvable(s) en ligne 1.")
        return
    rapport.feuilles.append(f.nom)
    par_champ = {c.champ: c for c in f.colonnes}
    vues = set()
    for n, r in enumerate(rangees[1:], 2):
        brut = {ch: (r[i] if i < len(r) else None) for ch, i in positions.items()}
        if all(_vide(v) for v in brut.values()):
            continue
        ou = f"{f.nom}, ligne {n}"
        try:
            valeurs = {ch: par_champ[ch].type.lire(v) for ch, v in brut.items() if not _vide(v)}
        except Refus as e:
            rapport.erreurs.append(f"{ou} : {e}.")
            continue
        manque = [par_champ[ch].entete for ch in f.cles if ch not in valeurs]
        if manque:
            rapport.erreurs.append(f"{ou} : {', '.join(manque)} vide(s).")
            continue
        identite = tuple(valeurs[ch] for ch in f.cles)
        if identite in vues:
            rapport.erreurs.append(f"{ou} : doublon de {' / '.join(_affiche(v) for v in identite)}.")
            continue
        vues.add(identite)
        obj = f.trouver(valeurs) if f.trouver else f.modele.objects.filter(**{ch: valeurs[ch] for ch in f.cles}).first()
        nouveau = obj is None
        if not nouveau and any(getattr(obj, k) != v for k, v in f.fixe.items()):
            rapport.erreurs.append(f"{ou} : {identite[0]} existe déjà sur une autre feuille (axe {getattr(obj, 'axe', '?')}).")
            continue
        if nouveau:
            manque = [c.entete for c in f.colonnes if c.obligatoire and c.champ not in valeurs]
            if manque:
                rapport.erreurs.append(f"{ou} : pour créer cet élément, remplir {', '.join(manque)}.")
                continue
            obj = f.modele(**f.fixe)
        avant = {ch: getattr(obj, ch) for ch in valeurs} if not nouveau else {}
        for ch, v in valeurs.items():
            setattr(obj, ch, v)
        try:
            if f.preparer:
                f.preparer(obj)
            if f.verifier:
                f.verifier(obj)
            obj.full_clean()
        except Refus as e:
            rapport.erreurs.append(f"{ou} : {e}.")
            continue
        except Exception as e:                        # ValidationError de Django
            details = getattr(e, "message_dict", None)
            texte = "; ".join(f"{par_champ[k].entete if k in par_champ else k} : {' '.join(v)}" for k, v in details.items()) \
                if details else str(e)
            rapport.erreurs.append(f"{ou} : {texte}.")
            continue
        changes = [ch for ch in valeurs if not nouveau and avant[ch] != getattr(obj, ch)]
        nom = f"{f.nom} : {' / '.join(_affiche(par_champ[ch].type.vers_excel(valeurs[ch])) for ch in f.cles)}"
        if nouveau:
            obj.save()
            rapport.crees.append(nom)
            tracer and Modification.objects.create(auteur=auteur, lot="Paramètres", action="Création (import)", objet=nom[:200],
                                        apres=" · ".join(f"{par_champ[ch].entete} = {_affiche(par_champ[ch].type.vers_excel(v))}"
                                                         for ch, v in valeurs.items())[:300])
        elif changes:
            obj.save()
            rapport.modifies.append(nom)
            tracer and Modification.objects.create(
                auteur=auteur, lot="Paramètres", action="Modification (import)", objet=nom[:200],
                avant=" · ".join(f"{par_champ[ch].entete} = {_affiche(par_champ[ch].type.vers_excel(avant[ch]))}" for ch in changes)[:300],
                apres=" · ".join(f"{par_champ[ch].entete} = {_affiche(par_champ[ch].type.vers_excel(getattr(obj, ch)))}"
                                 for ch in changes)[:300])
        else:
            rapport.inchanges += 1


def importer(contenu, auteur="", tracer=True):
    """Importe un classeur Parametres.xlsx (octets). Tout ou rien : en cas d'erreur, la base reste inchangée.

    tracer=False : pas de ligne d'historique (réinjection d'un export complet, qui restaure l'historique d'origine)."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenu), data_only=True, read_only=True)
    except Exception:
        raise ValueError("Ce fichier n'est pas un classeur Excel .xlsx lisible.") from None
    noms = {cle(n): n for n in wb.sheetnames}
    rapport = Rapport()
    try:
        with transaction.atomic():
            for f in FEUILLES:                          # ordre : un code est créé avant d'être cité
                if cle(f.nom) in noms:
                    _importer_feuille(f, [list(r) for r in wb[noms[cle(f.nom)]].iter_rows(values_only=True)], rapport, auteur,
                                      tracer)
            if not rapport.feuilles and not rapport.erreurs:
                rapport.erreurs.append("Aucune feuille de paramètres reconnue (Réglages, Plan comptable, Journaux…).")
            if rapport.erreurs:
                raise _Annulation
            tracer and Modification.objects.create(auteur=auteur, lot="Paramètres", action="Import Parametres.xlsx",
                                        objet=", ".join(rapport.feuilles)[:200], apres=rapport.resume)
    except _Annulation:
        rapport.crees, rapport.modifies, rapport.inchanges = [], [], 0
    wb.close()
    return rapport
