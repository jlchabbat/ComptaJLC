import csv
import io
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from flask import (Blueprint, Response, flash, g, redirect, render_template,
                   request, url_for)
from sqlalchemy import func

from .models import (Axe, CodeAnalytique, Compte, Ecriture, Historique, Journal,
                     Ligne, LigneAnalytique, db)

bp = Blueprint("compta", __name__)


def fmt_montant(c):
    """Centimes -> « 1 234,50 »."""
    s = f"{abs(c or 0) / 100:,.2f}".replace(",", " ").replace(".", ",")
    return ("-" if (c or 0) < 0 else "") + s


def en_centimes(texte):
    texte = (texte or "").strip().replace(" ", "").replace(",", ".")
    if not texte:
        return 0
    try:
        return int((Decimal(texte) * 100).to_integral_value())
    except InvalidOperation:
        raise ValueError(f"Montant invalide : « {texte} »")


@bp.route("/")
def accueil():
    return render_template("accueil.html", nb=Ecriture.query.count())


# ---------- Plan comptable ----------

@bp.route("/plan", methods=["GET", "POST"])
def plan():
    if request.method == "POST":
        numero = request.form["numero"].strip()
        libelle = request.form["libelle"].strip()
        if not numero.isdigit() or not libelle:
            flash("Numéro (chiffres) et libellé obligatoires.", "erreur")
        elif db.session.get(Compte, numero):
            flash("Ce compte existe déjà.", "erreur")
        else:
            db.session.add(Compte(numero=numero, libelle=libelle))
            db.session.commit()
            flash("Compte ajouté.", "ok")
        return redirect(url_for("compta.plan"))
    return render_template("plan.html", comptes=Compte.query.order_by(Compte.numero).all())


@bp.post("/plan/<numero>/modifier")
def plan_modifier(numero):
    c = db.get_or_404(Compte, numero)
    libelle = request.form["libelle"].strip()
    if libelle:
        c.libelle = libelle
        db.session.commit()
        flash("Libellé modifié.", "ok")
    return redirect(url_for("compta.plan"))


@bp.post("/plan/<numero>/supprimer")
def plan_supprimer(numero):
    c = db.get_or_404(Compte, numero)
    if Ligne.query.filter_by(compte_numero=numero).first():
        flash("Compte utilisé par des écritures : suppression impossible.", "erreur")
    else:
        db.session.delete(c)
        db.session.commit()
        flash("Compte supprimé.", "ok")
    return redirect(url_for("compta.plan"))


@bp.post("/plan/importer")
def plan_importer():
    """Import CSV « numero;libelle » : ajoute les comptes absents, met à jour les libellés."""
    f = request.files.get("fichier")
    if not f:
        flash("Aucun fichier.", "erreur")
        return redirect(url_for("compta.plan"))
    texte = f.read().decode("utf-8-sig", errors="replace")
    delim = ";" if texte.count(";") >= texte.count(",") else ","
    ajoutes = maj = 0
    for row in csv.reader(io.StringIO(texte), delimiter=delim):
        if len(row) < 2 or not row[0].strip().isdigit():
            continue
        c = db.session.get(Compte, row[0].strip())
        if c:
            c.libelle = row[1].strip()
            maj += 1
        else:
            db.session.add(Compte(numero=row[0].strip(), libelle=row[1].strip()))
            ajoutes += 1
    db.session.commit()
    flash(f"Import terminé : {ajoutes} ajouté(s), {maj} mis à jour.", "ok")
    return redirect(url_for("compta.plan"))


# ---------- Axes et codes analytiques ----------

@bp.route("/axes", methods=["GET", "POST"])
def axes():
    if request.method == "POST":
        nom = request.form["nom"].strip()
        if nom and not Axe.query.filter_by(nom=nom).first():
            db.session.add(Axe(nom=nom))
            db.session.commit()
            flash("Axe ajouté.", "ok")
        else:
            flash("Nom vide ou déjà utilisé.", "erreur")
        return redirect(url_for("compta.axes"))
    return render_template("axes.html", axes=Axe.query.order_by(Axe.id).all())


@bp.post("/axes/<int:axe_id>/renommer")
def axe_renommer(axe_id):
    a = db.get_or_404(Axe, axe_id)
    nom = request.form["nom"].strip()
    if nom:
        a.nom = nom
        db.session.commit()
    return redirect(url_for("compta.axes"))


@bp.post("/axes/<int:axe_id>/codes")
def code_ajouter(axe_id):
    a = db.get_or_404(Axe, axe_id)
    code, libelle = request.form["code"].strip().upper(), request.form["libelle"].strip()
    if not code or not libelle:
        flash("Code et libellé obligatoires.", "erreur")
    elif CodeAnalytique.query.filter_by(axe_id=a.id, code=code).first():
        flash("Ce code existe déjà sur cet axe.", "erreur")
    else:
        db.session.add(CodeAnalytique(axe_id=a.id, code=code, libelle=libelle))
        db.session.commit()
        flash("Code ajouté.", "ok")
    return redirect(url_for("compta.axes"))


@bp.post("/codes/<int:code_id>/supprimer")
def code_supprimer(code_id):
    c = db.get_or_404(CodeAnalytique, code_id)
    if LigneAnalytique.query.filter_by(code_id=c.id).first():
        flash("Code utilisé par des écritures : suppression impossible.", "erreur")
    else:
        db.session.delete(c)
        db.session.commit()
        flash("Code supprimé.", "ok")
    return redirect(url_for("compta.axes"))


# ---------- Journaux ----------

@bp.route("/journaux", methods=["GET", "POST"])
def journaux():
    if request.method == "POST":
        code, libelle = request.form["code"].strip().upper(), request.form["libelle"].strip()
        if code and libelle and not db.session.get(Journal, code):
            db.session.add(Journal(code=code, libelle=libelle))
            db.session.commit()
            flash("Journal ajouté.", "ok")
        else:
            flash("Code vide ou déjà utilisé.", "erreur")
        return redirect(url_for("compta.journaux"))
    return render_template("journaux.html", journaux=Journal.query.order_by(Journal.code).all())


# ---------- Saisie, correction, suppression ----------

def instantane(e):
    """Copie lisible d'une écriture (stockée dans l'historique)."""
    return {"date": e.date.isoformat(), "journal": e.journal_code, "piece": e.piece or "",
            "libelle": e.libelle,
            "lignes": [{"compte": l.compte_numero, "debit": l.debit, "credit": l.credit,
                        "analytique": [f"{x.code.axe.nom} : {x.code.code}" for x in l.analytiques]}
                       for l in e.lignes]}


def journaliser(action, mvt, motif, avant, apres):
    db.session.add(Historique(utilisateur=g.utilisateur.nom, action=action, mvt=mvt, motif=motif,
                              avant=json.dumps(avant) if avant else None,
                              apres=json.dumps(apres) if apres else None))


def prochain_mvt():
    """Les numéros de mouvement supprimés ne sont jamais réutilisés (l'historique les cite)."""
    a = db.session.query(func.max(Ecriture.mvt)).scalar() or 0
    b = db.session.query(func.max(Historique.mvt)).scalar() or 0
    return max(a, b) + 1


def centimes_texte(c):
    return f"{c // 100},{c % 100:02d}"


def valeurs_ecriture(e):
    """Valeurs de départ du formulaire pour modifier une écriture."""
    v = {"date": e.date.isoformat(), "journal": e.journal_code, "piece": e.piece or "",
         "libelle": e.libelle}
    for i, l in enumerate(e.lignes):
        v[f"compte{i}"] = l.compte_numero
        v[f"debit{i}"] = centimes_texte(l.debit) if l.debit else ""
        v[f"credit{i}"] = centimes_texte(l.credit) if l.credit else ""
        for x in l.analytiques:
            v[f"axe{x.code.axe_id}_{i}"] = str(x.code_id)
    return v


def lire_formulaire(f, axes_, nb_lignes):
    """Valide le formulaire ; lève ValueError (message en français) sinon."""
    d = datetime.strptime(f["date"], "%Y-%m-%d").date()
    if not f["libelle"].strip():
        raise ValueError("Libellé obligatoire.")
    if not db.session.get(Journal, f["journal"]):
        raise ValueError("Journal inconnu.")
    lignes = []
    for i in range(nb_lignes):
        num = f.get(f"compte{i}", "").strip()
        deb, cre = en_centimes(f.get(f"debit{i}")), en_centimes(f.get(f"credit{i}"))
        if not num and not deb and not cre:
            continue
        if not db.session.get(Compte, num):
            raise ValueError(f"Ligne {i + 1} : compte « {num} » introuvable.")
        if (deb and cre) or not (deb or cre) or deb < 0 or cre < 0:
            raise ValueError(f"Ligne {i + 1} : un montant positif au débit OU au crédit.")
        codes = []
        for a in axes_:
            cid = f.get(f"axe{a.id}_{i}", "")
            if cid:
                c = db.session.get(CodeAnalytique, int(cid))
                if not c or c.axe_id != a.id:
                    raise ValueError(f"Ligne {i + 1} : code analytique invalide.")
                codes.append(c)
        lignes.append((num, deb, cre, codes))
    if len(lignes) < 2:
        raise ValueError("Au moins deux lignes.")
    if sum(l[1] for l in lignes) != sum(l[2] for l in lignes):
        raise ValueError("Écriture non équilibrée : total débit ≠ total crédit.")
    return d, f["journal"].strip(), f.get("piece", "").strip(), f["libelle"].strip(), lignes


def construire_lignes(lignes):
    out = []
    for num, deb, cre, codes in lignes:
        l = Ligne(compte_numero=num, debit=deb, credit=cre)
        l.analytiques = [LigneAnalytique(code_id=c.id) for c in codes]
        out.append(l)
    return out


def contexte_saisie(val, nb_lignes=6, ecriture=None):
    return dict(journaux=Journal.query.all(), comptes=Compte.query.order_by(Compte.numero).all(),
                axes=Axe.query.order_by(Axe.id).all(), aujourdhui=date.today().isoformat(),
                nb_lignes=nb_lignes, val=val, ecriture=ecriture)


@bp.route("/saisie", methods=["GET", "POST"])
def saisie():
    ctx = contexte_saisie(request.form if request.method == "POST" else {})
    if request.method == "GET":
        return render_template("saisie.html", **ctx)
    try:
        d, jnl, piece, libelle, lignes = lire_formulaire(request.form, ctx["axes"], ctx["nb_lignes"])
    except ValueError as e:
        flash(str(e), "erreur")
        return render_template("saisie.html", **ctx), 400
    e = Ecriture(mvt=prochain_mvt(), date=d, journal_code=jnl, piece=piece, libelle=libelle)
    e.lignes = construire_lignes(lignes)
    db.session.add(e)
    db.session.flush()
    journaliser("création", e.mvt, "", None, instantane(e))
    db.session.commit()
    flash(f"Écriture n° {e.mvt} enregistrée.", "ok")
    return redirect(url_for("compta.saisie"))


@bp.route("/ecriture/<int:mvt>/modifier", methods=["GET", "POST"])
def ecriture_modifier(mvt):
    e = db.get_or_404(Ecriture, mvt)
    nb = max(6, len(e.lignes) + 2)
    post = request.method == "POST"
    ctx = contexte_saisie(request.form if post else valeurs_ecriture(e), nb, ecriture=e)
    if not post:
        return render_template("saisie.html", **ctx)
    try:
        motif = request.form.get("motif", "").strip()
        if not motif:
            raise ValueError("Indiquez le motif de la modification.")
        d, jnl, piece, libelle, lignes = lire_formulaire(request.form, ctx["axes"], nb)
    except ValueError as err:
        flash(str(err), "erreur")
        return render_template("saisie.html", **ctx), 400
    avant = instantane(e)
    e.date, e.journal_code, e.piece, e.libelle = d, jnl, piece, libelle
    e.lignes = construire_lignes(lignes)
    db.session.flush()
    db.session.refresh(e)
    journaliser("modification", mvt, motif, avant, instantane(e))
    db.session.commit()
    flash(f"Écriture n° {mvt} modifiée.", "ok")
    return redirect(url_for("compta.journal"))


@bp.post("/ecriture/<int:mvt>/supprimer")
def ecriture_supprimer(mvt):
    e = db.get_or_404(Ecriture, mvt)
    motif = request.form.get("motif", "").strip()
    if not motif:
        flash("Indiquez le motif de la suppression.", "erreur")
        return redirect(url_for("compta.ecriture_modifier", mvt=mvt))
    journaliser("suppression", mvt, motif, instantane(e), None)
    db.session.delete(e)
    db.session.commit()
    flash(f"Écriture n° {mvt} supprimée (conservée dans l'historique).", "ok")
    return redirect(url_for("compta.journal"))


@bp.route("/historique")
def historique():
    rows = Historique.query.order_by(Historique.id.desc()).all()
    return render_template("historique.html", rows=[(h, json.loads(h.avant) if h.avant else None,
                                                     json.loads(h.apres) if h.apres else None) for h in rows])


@bp.app_template_filter("resume_ecriture")
def resume_ecriture(snap):
    if not snap:
        return ""
    out = [f"{snap['date']} · {snap['journal']} · {snap['piece']} · {snap['libelle']}"]
    for l in snap["lignes"]:
        sens = f"D {fmt_montant(l['debit'])}" if l["debit"] else f"C {fmt_montant(l['credit'])}"
        ana = " [" + " ; ".join(l["analytique"]) + "]" if l["analytique"] else ""
        out.append(f"  {l['compte']}  {sens}{ana}")
    return "\n".join(out)


# ---------- États ----------

@bp.route("/journal")
def journal():
    return render_template("journal.html", ecritures=Ecriture.query.order_by(Ecriture.date, Ecriture.mvt).all())


@bp.route("/balance")
def balance():
    rows = (db.session.query(Ligne.compte_numero, Compte.libelle, func.sum(Ligne.debit), func.sum(Ligne.credit))
            .join(Compte).group_by(Ligne.compte_numero).order_by(Ligne.compte_numero).all())
    return render_template("balance.html", rows=rows,
                           td=sum(r[2] for r in rows), tc=sum(r[3] for r in rows))


@bp.route("/grand-livre/<numero>")
def grand_livre(numero):
    c = db.get_or_404(Compte, numero)
    lignes = (Ligne.query.filter_by(compte_numero=numero).join(Ecriture)
              .order_by(Ecriture.date, Ecriture.mvt).all())
    return render_template("grand_livre.html", compte=c, lignes=lignes)


@bp.route("/analytique/<int:axe_id>")
def analytique(axe_id):
    a = db.get_or_404(Axe, axe_id)
    rows = (db.session.query(CodeAnalytique.code, CodeAnalytique.libelle,
                             func.sum(Ligne.debit), func.sum(Ligne.credit))
            .join(LigneAnalytique, LigneAnalytique.code_id == CodeAnalytique.id)
            .join(Ligne, Ligne.id == LigneAnalytique.ligne_id)
            .filter(CodeAnalytique.axe_id == a.id).group_by(CodeAnalytique.id)
            .order_by(CodeAnalytique.code).all())
    return render_template("analytique.html", axe=a, rows=rows)


# ---------- Exports Excel ----------

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _texte(ws, ligne, col, valeur):
    """Écrit un texte sans jamais l'interpréter comme une formule (=, +, -, @)."""
    c = ws.cell(row=ligne, column=col, value=valeur)
    if isinstance(valeur, str):
        c.data_type = "s"
    return c


def _reponse(wb, nom):
    from openpyxl.styles import Font
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                50, max(10, max(len(str(c.value or "")) for c in col) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype=XLSX,
                    headers={"Content-Disposition": f"attachment; filename={nom}"})


@bp.route("/export/journal.xlsx")
def export_journal():
    from openpyxl import Workbook
    axes_ = Axe.query.order_by(Axe.id).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    ent = ["Mvt", "Date", "Journal", "Pièce", "Libellé", "Compte", "Libellé du compte",
           "Débit", "Crédit"] + [a.nom for a in axes_]
    for k, h in enumerate(ent, 1):
        _texte(ws, 1, k, h)
    r = 2
    for e in Ecriture.query.order_by(Ecriture.date, Ecriture.mvt).all():
        for l in e.lignes:
            ws.cell(row=r, column=1, value=e.mvt)
            ws.cell(row=r, column=2, value=e.date).number_format = "DD/MM/YYYY"
            _texte(ws, r, 3, e.journal_code)
            _texte(ws, r, 4, e.piece or "")
            _texte(ws, r, 5, e.libelle)
            _texte(ws, r, 6, l.compte_numero)
            _texte(ws, r, 7, l.compte.libelle)
            ws.cell(row=r, column=8, value=l.debit / 100 if l.debit else None).number_format = "#,##0.00"
            ws.cell(row=r, column=9, value=l.credit / 100 if l.credit else None).number_format = "#,##0.00"
            par_axe = {x.code.axe_id: x.code for x in l.analytiques}
            for k, a in enumerate(axes_):
                c = par_axe.get(a.id)
                if c:
                    _texte(ws, r, 10 + k, f"{c.code} {c.libelle}")
            r += 1
    return _reponse(wb, "journal.xlsx")


@bp.route("/export/balance.xlsx")
def export_balance():
    from openpyxl import Workbook
    rows = (db.session.query(Ligne.compte_numero, Compte.libelle, func.sum(Ligne.debit), func.sum(Ligne.credit))
            .join(Compte).group_by(Ligne.compte_numero).order_by(Ligne.compte_numero).all())
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance"
    for k, h in enumerate(["Compte", "Libellé", "Débit", "Crédit", "Solde"], 1):
        _texte(ws, 1, k, h)
    for r, (n, lib, d, c) in enumerate(rows, 2):
        _texte(ws, r, 1, n)
        _texte(ws, r, 2, lib)
        ws.cell(row=r, column=3, value=d / 100).number_format = "#,##0.00"
        ws.cell(row=r, column=4, value=c / 100).number_format = "#,##0.00"
        ws.cell(row=r, column=5, value=f"=C{r}-D{r}").number_format = "#,##0.00"
    t = len(rows) + 2
    _texte(ws, t, 2, "Total")
    for col, L in ((3, "C"), (4, "D"), (5, "E")):
        ws.cell(row=t, column=col, value=f"=SUM({L}2:{L}{t - 1})").number_format = "#,##0.00"
    return _reponse(wb, "balance.xlsx")
